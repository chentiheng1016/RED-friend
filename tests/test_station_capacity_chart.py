"""針車/射出/包裝 三站每日產能折線圖（chart_station_capacity / station_capacity_report）。

涵蓋：站別別名（射出＝生管日報的「灌注」）、一次開檔算多站、跨月往前補到指定天數、
補不到時的降級與警告、折線圖 spec（三條線 + 跨月標籤）、[[MAIL_FILE:]] 附件標記。

隔離：openpyxl 合成兩份「生管日報」workbook（本月 2 天 + 上月 5 天），patch
_load_latest_sheet / _find_month_progress_file / _download_xlsx_bytes（不連 Drive）
與 chart_export.generate_chart（攔 spec、不畫圖/不送）。另有 setUpModule 全模組防
護網，漏 mock 而摸到 Google 的測試會在 tearDownModule 變紅（見下方註解）。
"""
import datetime as _dt
import io
import json
import os
import tempfile
import traceback
import unittest
from unittest import mock

import openpyxl

from agent_core import factory_production_report as fpr

# ── 全模組防護網：這支測試檔一律不准碰真的 Drive ──────────────────────────
# 歷史教訓（2026-08-12）：test_no_data_message 只 mock 了 _load_latest_sheet（本月），
# 沒擋「跨月往前補」那條路。有憑證的機器（~/RED 主 checkout）會真的去 Drive 抓 7MB
# 的上個月進度表回來，於是本月餵垃圾 bytes 也照樣畫出完整圖表、斷言 fail；CI 沒憑證
# 所以永遠是綠的 —— 一個只在本機紅、看起來像「測試順序問題」的假象。
#
# _prev_month_sheet 把所有例外吞成 warning（補資料失敗不該讓整張圖沒有），所以這裡
# 光 raise 沒用：還要記帳，由 tearDownModule 把「有人沒 mock 就摸到 Drive」變成紅燈，
# 而不是默默降級成綠的。unittest 沒有 conftest autouse fixture（見 CLAUDE.md）。
_DRIVE_HITS: list = []
_REAL_GET_SERVICE = None


def setUpModule():  # noqa: N802 (unittest 命名慣例)
    global _REAL_GET_SERVICE
    from agent_core import google_auth

    _REAL_GET_SERVICE = google_auth.get_service

    def _blocked(api_name, version):
        caller = traceback.extract_stack()[-2]
        _DRIVE_HITS.append(
            f"{api_name}/{version} @ {os.path.basename(caller.filename)}:{caller.lineno}")
        raise AssertionError(
            "測試不准連 Google：請 mock _find_month_progress_file / _download_xlsx_bytes")

    google_auth.get_service = _blocked


def tearDownModule():  # noqa: N802
    from agent_core import google_auth

    google_auth.get_service = _REAL_GET_SERVICE
    hits, _DRIVE_HITS[:] = list(_DRIVE_HITS), []
    if hits:
        raise AssertionError("有測試沒 mock 就摸到 Google API：" + "、".join(hits))


def _build_xlsx(days: dict) -> bytes:
    """days = {分頁名: [(客戶, 針車, 灌注, 包裝), ...]}；欄序含三站日計。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for day, rows in days.items():
        ws = wb.create_sheet(day)
        ws.append(["客戶", "指令", "雙數", "Stitching", "Injection", "Packing"])
        ws.append([None, None, None, "日計", "日計", "日計"])
        for cust, st, inj, pk in rows:
            ws.append([cust, "WO", 5000, st, inj, pk])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_CUR = _build_xlsx({
    "1": [("DECA.", 100, 200, 300), ("RICHTER", 10, 20, 30)],
    "2": [("DECA.", 400, 500, 600)],
})
_PREV = _build_xlsx({str(d): [("DECA.", d, d * 2, d * 3)] for d in (27, 28, 29, 30, 31)})


class ResolveStationTests(unittest.TestCase):
    def test_aliases_and_names(self):
        self.assertEqual(fpr._resolve_station("射出"), "Injection")     # 廠內口語
        self.assertEqual(fpr._resolve_station("灌注"), "Injection")     # 生管日報欄名
        self.assertEqual(fpr._resolve_station("針車"), "Stitching")
        self.assertEqual(fpr._resolve_station("packing"), "Packing")    # 大小寫不敏感
        self.assertIsNone(fpr._resolve_station("沒有這站"))
        self.assertIsNone(fpr._resolve_station(""))

    def test_unknown_station_falls_back_to_packing(self):
        # 舊行為：_compute_daily_output 認不得就當包裝（完工雙數）
        self.assertEqual(fpr._compute_daily_output(_CUR, station="沒有這站"),
                         [(1, 330.0), (2, 600.0)])


class ComputeMultiTests(unittest.TestCase):
    def test_three_stations_one_pass(self):
        out = fpr._compute_daily_output_multi(_CUR, ["Stitching", "Injection", "Packing"])
        self.assertEqual(out["Stitching"], [(1, 110.0), (2, 400.0)])
        self.assertEqual(out["Injection"], [(1, 220.0), (2, 500.0)])
        self.assertEqual(out["Packing"], [(1, 330.0), (2, 600.0)])

    def test_single_station_delegates_to_multi(self):
        self.assertEqual(fpr._compute_daily_output(_CUR, "針車"), [(1, 110.0), (2, 400.0)])

    def test_missing_station_column_is_empty_not_zero(self):
        """該站沒欄位 → 空 list（不是一排 0）：分得出「沒這站」與「這站沒產出」。"""
        only_pack = _build_xlsx({"1": [("DECA.", 1, 2, 3)]})
        wb = openpyxl.load_workbook(io.BytesIO(only_pack))
        ws = wb["1"]
        ws.cell(row=1, column=4).value = None      # 拿掉 Stitching 表頭錨點
        ws.cell(row=1, column=5).value = None      # 拿掉 Injection
        buf = io.BytesIO()
        wb.save(buf)
        out = fpr._compute_daily_output_multi(buf.getvalue(), ["Stitching", "Packing"])
        self.assertEqual(out["Stitching"], [])
        self.assertEqual(out["Packing"], [(1, 3.0)])

    def test_unparseable_bytes(self):
        out = fpr._compute_daily_output_multi(b"nope", ["Packing"])
        self.assertEqual(out, {"Packing": []})


class MonthOfNameTests(unittest.TestCase):
    def test_parses_leading_month(self):
        self.assertEqual(fpr._month_of_progress_name("08月份生產日報進度表08-04.xlsx"), 8)
        self.assertEqual(fpr._month_of_progress_name("1月份生產日報進度表1-24.xls"), 1)
        self.assertIsNone(fpr._month_of_progress_name("Production Daily Report"))
        self.assertIsNone(fpr._month_of_progress_name(""))


def _reset_prev_cache():
    """上個月進度表的 bytes 快取是 module-level 狀態，測試之間會互相汙染。

    unittest 沒有 conftest autouse fixture，隔離只能寫在 setUp（見 CLAUDE.md）。
    不清的話「找不到上個月的檔」那幾個 case 會拿到前一個 case 快取的真 bytes。
    """
    fpr._PREV_CACHE.update(month=None, bytes=None, name="", modified="", ts=0.0)


class SeriesTests(unittest.TestCase):
    _TODAY = _dt.date(2026, 8, 5)

    def setUp(self):
        _reset_prev_cache()
        self.addCleanup(_reset_prev_cache)

    def _series(self, days, *, prev=True, prev_raises=False):
        find = mock.patch.object(
            fpr, "_find_month_progress_file",
            return_value=(("prev-id", "07月份生產日報進度表07-31.xlsx") if prev
                          else (None, None)))
        dl = mock.patch.object(
            fpr, "_download_xlsx_bytes",
            side_effect=(RuntimeError("boom") if prev_raises
                         else lambda fid: (_PREV, {"modifiedTime": "2026-08-01T00:00:00Z"})))
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(_CUR, "08月份生產日報進度表08-02.xlsx",
                                             "2026-08-03")), find, dl:
            return fpr.station_capacity_series(days=days, today=self._TODAY)

    def test_current_month_only_when_enough_days(self):
        data = self._series(2)
        self.assertEqual([lab for lab, _v in data["points"]], ["8/1", "8/2"])
        self.assertEqual(len(data["sources"]), 1)
        self.assertEqual(data["warnings"], [])

    def test_tops_up_from_previous_month(self):
        data = self._series(5)
        # 本月只有 2 天 → 往前補上個月最後 3 天，標籤帶月份才不會 8/1 跟 7/31 混淆
        self.assertEqual([lab for lab, _v in data["points"]],
                         ["7/29", "7/30", "7/31", "8/1", "8/2"])
        self.assertEqual(data["points"][0][1]["針車"], 29.0)
        self.assertEqual(data["points"][0][1]["射出(灌注)"], 58.0)
        self.assertEqual(data["points"][-1][1]["包裝"], 600.0)
        self.assertEqual([n for n, _m in data["sources"]],
                         ["07月份生產日報進度表07-31.xlsx", "08月份生產日報進度表08-02.xlsx"])

    def test_caps_at_requested_days(self):
        data = self._series(4)
        self.assertEqual(len(data["points"]), 4)
        self.assertEqual(data["points"][0][0], "7/30")

    def test_prev_month_missing_degrades_with_warning(self):
        data = self._series(10, prev=False)
        self.assertEqual(len(data["points"]), 2)          # 只剩本月，但不是空的
        self.assertTrue(any("找不到 7 月份" in w for w in data["warnings"]))

    def test_prev_month_download_failure_degrades_with_warning(self):
        data = self._series(10, prev_raises=True)
        self.assertEqual(len(data["points"]), 2)
        self.assertTrue(any("讀上個月進度表失敗" in w for w in data["warnings"]))

    def test_explicit_file_id_does_not_cross_month(self):
        """指定 file_id＝「就看這一份」，不該偷偷混進別月的數字。"""
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(_CUR, "08月份生產日報進度表08-02.xlsx", "2026-08-03")), \
             mock.patch.object(fpr, "_find_month_progress_file") as find:
            data = fpr.station_capacity_series(days=30, file_id="abc", today=self._TODAY)
        find.assert_not_called()
        self.assertEqual(len(data["points"]), 2)


class StyleOutputRateTests(unittest.TestCase):
    """各型體近期實際包裝產出 → 週產（給「這些料還夠做幾週」換算）。"""

    _TODAY = _dt.date(2026, 8, 5)

    def setUp(self):
        _reset_prev_cache()
        self.addCleanup(_reset_prev_cache)

    @staticmethod
    def _model_xlsx(days: dict) -> bytes:
        """days = {分頁: [(客戶, 型體, 包裝日計)]}"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for day, rows in days.items():
            ws = wb.create_sheet(day)
            ws.append(["客戶", "指令", "型體", "雙數", "Packing"])
            ws.append([None, None, None, None, "日計"])
            for cust, model, pk in rows:
                ws.append([cust, "WO", model, 5000, pk])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _rate(self, days=4, prev=None):
        cur = self._model_xlsx({
            "1": [("DECA.", "DJS336189-02BEIGE", 100), ("DECA.", "DJS336195-01GY", 40)],
            "2": [("DECA.", "DJS336189-02GREEN", 60), ("RICHTER", "FE2303V", 500)],
        })
        find = mock.patch.object(fpr, "_find_month_progress_file",
                                 return_value=(("p", "07月份生產日報進度表07-31.xlsx")
                                               if prev else (None, None)))
        dl = mock.patch.object(fpr, "_download_xlsx_bytes",
                               return_value=(prev, {"modifiedTime": "2026-08-01T00:00:00Z"}))
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(cur, "08月份生產日報進度表08-02.xlsx",
                                             "2026-08-03")), find, dl:
            return fpr.style_output_rate(days=days, today=self._TODAY)

    def test_sums_packing_per_model(self):
        r = self._rate(days=2)
        self.assertEqual(r["days"], 2)
        self.assertEqual(r["pairs"]["DJS336189-02BEIGE"], 100.0)
        self.assertEqual(r["pairs"]["FE2303V"], 500.0)

    def test_colour_suffixes_roll_up_to_erp_style(self):
        r = self._rate(days=2)
        w = fpr.weekly_rates_for_styles(r, ["DJS336189-02", "DJS336195-01"])
        # BEIGE 100 + GREEN 60 = 160 雙 / 2 天 × 7 = 560
        self.assertAlmostEqual(w["DJS336189-02"], 560.0)
        self.assertAlmostEqual(w["DJS336195-01"], 140.0)

    def test_longest_prefix_wins(self):
        """舊款 DJS336195 不可以把 DJS336195-01 的產出整碗端走。"""
        r = self._rate(days=2)
        w = fpr.weekly_rates_for_styles(r, ["DJS336195", "DJS336195-01"])
        self.assertNotIn("DJS336195", w)
        self.assertAlmostEqual(w["DJS336195-01"], 140.0)

    def test_style_without_recent_output_absent(self):
        """近期沒生產就不給數字（呼叫端顯示「—」，不要拿猜的速度硬換算）。"""
        w = fpr.weekly_rates_for_styles(self._rate(days=2), ["NOSUCH-01"])
        self.assertEqual(w, {})

    def test_empty_rate_is_safe(self):
        self.assertEqual(fpr.weekly_rates_for_styles({"pairs": {}, "days": 0}, ["A"]), {})

    def test_tops_up_window_from_previous_month(self):
        prev = self._model_xlsx({str(d): [("DECA.", "DJS336189-02BEIGE", 10)]
                                 for d in (28, 29, 30, 31)})
        r = self._rate(days=4, prev=prev)
        self.assertEqual(r["days"], 4)              # 本月 2 天 + 上月尾巴 2 天
        self.assertEqual(r["pairs"]["DJS336189-02BEIGE"], 120.0)   # 本月 100 + 上月 10+10
        w = fpr.weekly_rates_for_styles(r, ["DJS336189-02"])
        self.assertAlmostEqual(w["DJS336189-02"], 315.0)           # (120+60)/4×7

    def test_prev_month_shared_cache_downloads_once(self):
        """折線圖與速率在同一輪簡報各需要上個月的檔 —— 只能下載一次（7MB/11 秒）。"""
        prev = self._model_xlsx({str(d): [("DECA.", "DJS336189-02BEIGE", 10)]
                                 for d in (28, 29, 30, 31)})
        cur = self._model_xlsx({"1": [("DECA.", "DJS336189-02BEIGE", 100)]})
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(cur, "08月份生產日報進度表08-02.xlsx", "2026-08-03")), \
             mock.patch.object(fpr, "_find_month_progress_file",
                               return_value=("p", "07月份生產日報進度表07-31.xlsx")), \
             mock.patch.object(fpr, "_download_xlsx_bytes",
                               return_value=(prev, {"modifiedTime": "2026-08-01T00:00:00Z"})) as dl:
            fpr.station_capacity_series(days=4, today=self._TODAY)
            fpr.style_output_rate(days=4, today=self._TODAY)
        self.assertEqual(dl.call_count, 1)


class FindMonthFileTests(unittest.TestCase):
    _FILES = [
        {"id": "a", "name": "07月份生產日報進度表07-31.xlsx", "modifiedTime": "2026-08-01T00:00:00Z"},
        {"id": "b", "name": "07月份生產日報進度表07-30.xlsx", "modifiedTime": "2025-08-01T00:00:00Z"},
        {"id": "c", "name": "08月份生產日報進度表08-04.xlsx", "modifiedTime": "2026-08-05T00:00:00Z"},
        {"id": "d", "name": "說明.txt", "modifiedTime": "2026-08-05T00:00:00Z"},
    ]

    def _svc(self):
        svc = mock.MagicMock()
        svc.files.return_value.list.return_value.execute.return_value = {"files": self._FILES}
        return svc

    def test_picks_requested_month_and_skips_stale_year(self):
        with mock.patch("agent_core.google_auth.get_service", return_value=self._svc()):
            fid, name = fpr._find_month_progress_file(7, min_modified="2026-06-01")
        self.assertEqual(fid, "a")                       # 去年那份 (b) 被 min_modified 擋掉
        self.assertTrue(name.startswith("07月份"))

    def test_no_match_returns_none(self):
        with mock.patch("agent_core.google_auth.get_service", return_value=self._svc()):
            self.assertEqual(fpr._find_month_progress_file(3), (None, None))

    def test_min_modified_can_filter_everything(self):
        with mock.patch("agent_core.google_auth.get_service", return_value=self._svc()):
            self.assertEqual(fpr._find_month_progress_file(7, min_modified="2026-09-01"),
                             (None, None))


class ReportTests(unittest.TestCase):
    def setUp(self):
        _reset_prev_cache()
        self.addCleanup(_reset_prev_cache)

    def _run(self, *, artifact=True, days=5):
        captured = {}
        tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
        tmp.write(b"png")
        tmp.close()
        self.addCleanup(os.unlink, tmp.name)

        def fake_gen(spec_json, filename="", deliver=True, chat_id=""):
            captured["spec"] = json.loads(spec_json)
            captured["filename"] = filename
            captured["deliver"] = deliver
            from agent_core.tool_result import ToolResult
            return ToolResult.success("ok", artifacts=([tmp.name] if artifact else []))

        with mock.patch("agent_core.chart_export.generate_chart", side_effect=fake_gen), \
             mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(_CUR, "08月份生產日報進度表08-02.xlsx", "2026-08-03")), \
             mock.patch.object(fpr, "_find_month_progress_file",
                               return_value=("prev-id", "07月份生產日報進度表07-31.xlsx")), \
             mock.patch.object(fpr, "_download_xlsx_bytes",
                               return_value=(_PREV, {"modifiedTime": "2026-08-01T00:00:00Z"})):
            out = fpr.station_capacity_report(days=days, today=_dt.date(2026, 8, 5))
        return out, captured, tmp.name

    def test_line_spec_has_three_series(self):
        _out, cap, _p = self._run()
        spec = cap["spec"]
        self.assertEqual(spec["type"], "line")
        self.assertEqual([s["name"] for s in spec["series"]], ["針車", "射出(灌注)", "包裝"])
        self.assertEqual(spec["categories"], ["7/29", "7/30", "7/31", "8/1", "8/2"])
        self.assertEqual(spec["series"][0]["values"], [29.0, 30.0, 31.0, 110.0, 400.0])

    def test_defaults_to_no_telegram_push(self):
        """排程寄信是走附件；預設不推大王的 Telegram（否則每天兩次噪音）。"""
        _out, cap, _p = self._run()
        self.assertFalse(cap["deliver"])

    def test_text_has_table_average_and_mail_marker(self):
        out, _cap, path = self._run()
        self.assertIn("針車", out)
        self.assertIn("平均", out)
        self.assertIn("近 5 天平均日產", out)
        self.assertIn(f"[[MAIL_FILE:{path}]]", out)
        self.assertIn("射出即生管日報的『灌注』站", out)

    def test_missing_artifact_says_so_instead_of_marker(self):
        out, _cap, _p = self._run(artifact=False)
        self.assertNotIn("[[MAIL_FILE:", out)
        self.assertIn("折線圖產檔失敗", out)

    def test_no_data_message(self):
        """本月這份算不出天數、上個月也補不到 → 明說算不出，不要生一張空圖。

        只 mock 本月是不夠的：本月湊不滿天數時會往前補上個月，那條路沒擋掉的話，
        有憑證的機器會真的去 Drive 抓上個月的檔，畫出一張「全是上個月」的圖。
        """
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(b"nope", "x.xlsx", "2026-08-03")), \
             mock.patch.object(fpr, "_find_month_progress_file",
                               return_value=(None, None)) as find, \
             mock.patch.object(fpr, "_download_xlsx_bytes") as dl:
            out = fpr.station_capacity_report(days=5, today=_dt.date(2026, 8, 5))
        self.assertIn("算不出每日產量", out)
        find.assert_called_once()      # 有去找上個月
        dl.assert_not_called()         # 找不到就不該下載（也就沒有 7MB 的真檔混進來）

    def test_unreadable_sheet_says_so(self):
        """檔打不開 ≠ 月初還沒填 —— 前者要找人修檔，後者等就好，訊息得分得出來。"""
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(b"nope", "x.xlsx", "2026-08-03")), \
             mock.patch.object(fpr, "_find_month_progress_file",
                               return_value=(None, None)), \
             mock.patch.object(fpr, "_download_xlsx_bytes"):
            out = fpr.station_capacity_report(days=5, today=_dt.date(2026, 8, 5))
        self.assertIn("算不出每日產量", out)
        self.assertIn("「x.xlsx」打不開", out)
        # 找不到上個月那份的原因以前也是看不到的，一起帶出來
        self.assertIn("找不到 7 月份", out)

    def test_openable_but_empty_sheet_is_not_reported_as_unreadable(self):
        """打得開、只是還沒有任何一天的分頁（月初）→ 不可以說「打不開」。"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb.create_sheet("說明")          # 有分頁，但沒有一個是日期分頁
        buf = io.BytesIO()
        wb.save(buf)

        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(buf.getvalue(), "08月份生產日報進度表08-01.xlsx",
                                             "2026-08-01")), \
             mock.patch.object(fpr, "_find_month_progress_file",
                               return_value=(None, None)), \
             mock.patch.object(fpr, "_download_xlsx_bytes"):
            out = fpr.station_capacity_report(days=5, today=_dt.date(2026, 8, 5))
        self.assertIn("算不出每日產量", out)
        self.assertNotIn("打不開", out)

    def test_sheet_not_found_returns_message(self):
        with mock.patch.object(fpr, "_load_latest_sheet",
                               side_effect=RuntimeError("Drive 上找不到生產日報進度表")):
            out = fpr.station_capacity_report(days=5)
        self.assertIn("找不到生產日報進度表", out)


if __name__ == "__main__":
    unittest.main()
