"""chart_daily_output（每日產量趨勢/產能高峰）與 chart_delivery_risk（達交風險）的測試。

涵蓋：_compute_daily_output 逐日加總某站產出 + 找峰、_compute_delivery_risk 篩未出貨
有欠數的 PO 並依希望出貨日排序、端到端 chart spec（峰日標紅、達交逾期紅/7日內橘/其餘綠）、
SAFE tier、無資料/全數出貨的友善回覆。

隔離：openpyxl 合成多日『生管日報』workbook（達交日期相對 today 計算→測試與執行日期無關）；
patch _load_latest_sheet（不連 Drive）與 chart_export.generate_chart（攔 spec、不畫圖/不送）。
"""
import datetime as _dt
import io
import json
import unittest
from unittest import mock

import openpyxl

from agent_core import factory_production_report as fpr

_TODAY = _dt.date.today()
_OVERDUE = (_TODAY - _dt.timedelta(days=2)).isoformat()
_SOON = (_TODAY + _dt.timedelta(days=3)).isoformat()
_LATER = (_TODAY + _dt.timedelta(days=30)).isoformat()


def _build_multiday_xlsx() -> bytes:
    """3 個日分頁(13/14/15)，欄＝客戶/指令/雙數/希望出貨日/實際出貨日/累計/欠數/Packing。
    14 號為產能高峰；最新日(15)帶達交資訊（相對 today）。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    days = {
        "13": [("DECA.26", 100), ("RICHTER", 80)],
        "14": [("DECA.26", 300), ("RICHTER", 200)],            # 峰：總 500
        "15": [("DECA.26", 180), ("RICHTER", 120),
               ("NEWWAVE", 90), ("JALAS", 60)],
    }
    risk = {                                                    # 最新日的 (want, actual, rem)
        "DECA.26": (_LATER, "", 4980),                          # 綠：>7 天
        "RICHTER": (_OVERDUE, "", 3350),                        # 紅：逾期
        "NEWWAVE": (_SOON, "", 1000),                           # 橘：7 日內
        "JALAS": (_SOON, "2026-06-12", 0),                      # 排除：已出貨 + 欠 0
    }
    for day, custs in days.items():
        ws = wb.create_sheet(day)
        ws.append(["客戶", "指令", "雙數", "希望出貨日", "實際出貨日", None, None, "Packing"])
        ws.append([None, None, None, None, None, "累計", "欠數", "日計"])
        for cust, pack in custs:
            want, actual, rem = risk.get(cust, ("", "", 0))
            ws.append([cust, "WO", 5000, want, actual, 5000 - rem, rem, pack])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ComputeDailyOutputTests(unittest.TestCase):
    def setUp(self):
        self.xlsx = _build_multiday_xlsx()

    def test_daily_totals_and_peak(self):
        out = fpr._compute_daily_output(self.xlsx)   # 預設 包裝(Packing)
        self.assertEqual(out, [(13, 180.0), (14, 500.0), (15, 450.0)])
        peak = max(out, key=lambda dv: dv[1])
        self.assertEqual(peak[0], 14)

    def test_unparseable_returns_empty(self):
        self.assertEqual(fpr._compute_daily_output(b"nope"), [])


class ComputeDeliveryRiskTests(unittest.TestCase):
    def setUp(self):
        self.xlsx = _build_multiday_xlsx()

    def test_excludes_shipped_and_zero_remaining_and_sorts_by_want(self):
        pos = fpr._compute_delivery_risk(self.xlsx)
        custs = [p["customer"] for p in pos]
        self.assertNotIn("JALAS", custs)                       # 已出貨/欠 0 → 排除
        self.assertEqual(custs, ["RICHTER", "NEWWAVE", "DECA.26"])  # 依希望出貨日升冪
        self.assertEqual(pos[0]["remaining"], 3350)

    def test_top_n_limit(self):
        self.assertEqual(len(fpr._compute_delivery_risk(self.xlsx, top_n=1)), 1)


class ChartDailyOutputTests(unittest.TestCase):
    def _run(self):
        captured = {}

        def fake_gen(spec_json, filename="", deliver=True, chat_id=""):
            captured["spec"] = json.loads(spec_json)
            from agent_core.tool_result import ToolResult
            return ToolResult.success("ok", data={"path": "/x.png"}, artifacts=["/x.png"])

        with mock.patch("agent_core.chart_export.generate_chart", side_effect=fake_gen), \
             mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(_build_multiday_xlsx(), "x.xlsx", "2026-06-16")):
            fpr.chart_daily_output(deliver=False)
        return captured["spec"]

    def test_bar_spec_peak_highlighted(self):
        spec = self._run()
        self.assertEqual(spec["type"], "bar")
        self.assertEqual(spec["categories"], ["13", "14", "15"])
        colors = spec["series"][0]["colors"]
        # 峰值(14號, index 1)為紅，其餘藍
        self.assertEqual(colors[1], "#E74C3C")
        self.assertEqual(colors[0], "#3498DB")
        self.assertIn("產能高峰：14 號", spec["subtitle"])


class ChartDeliveryRiskTests(unittest.TestCase):
    def _run(self):
        captured = {}

        def fake_gen(spec_json, filename="", deliver=True, chat_id=""):
            captured["spec"] = json.loads(spec_json)
            from agent_core.tool_result import ToolResult
            return ToolResult.success("ok", data={"path": "/x.png"}, artifacts=["/x.png"])

        with mock.patch("agent_core.chart_export.generate_chart", side_effect=fake_gen), \
             mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(_build_multiday_xlsx(), "x.xlsx", "2026-06-16")):
            out = fpr.chart_delivery_risk(deliver=False)
        return out, captured.get("spec")

    def test_barh_customer_aggregation_order_and_colors(self):
        _out, spec = self._run()
        self.assertEqual(spec["type"], "barh")
        # 依客戶欠數由多到少：DECA.26(4980,綠/later) → RICHTER(3350,紅/逾期) → NEWWAVE(1000,橘/7日內)
        self.assertEqual(spec["categories"], ["DECA.26", "RICHTER", "NEWWAVE"])
        self.assertEqual(spec["series"][0]["colors"], ["#2ECC71", "#E74C3C", "#E67E22"])
        self.assertIn("逾期", spec["series"][0]["annotations"][1])   # RICHTER 那條標逾期
        self.assertIn("共 3 筆未出貨", spec["subtitle"])              # 真實總數（非截斷）

    def test_aggregates_multiple_pos_and_flags_overdue(self):
        # 同一客戶多張單 → 彙總欠數/單數，只要有一張逾期整個客戶就標紅
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("15")
        ws.append(["客戶", "指令", "雙數", "希望出貨日", "實際出貨日", None, None, "Packing"])
        ws.append([None, None, None, None, None, "累計", "欠數", "日計"])
        ws.append(["DECA.26", "WO1", 5000, _LATER, "", 4000, 1000, 0])     # 未逾期
        ws.append(["DECA.26", "WO2", 5000, _OVERDUE, "", 3000, 2000, 0])   # 逾期
        buf = io.BytesIO()
        wb.save(buf)
        captured = {}

        def fake_gen(spec_json, filename="", deliver=True, chat_id=""):
            captured["spec"] = json.loads(spec_json)
            from agent_core.tool_result import ToolResult
            return ToolResult.success("ok", data={"path": "/x.png"}, artifacts=["/x.png"])

        with mock.patch("agent_core.chart_export.generate_chart", side_effect=fake_gen), \
             mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(buf.getvalue(), "x.xlsx", "2026-06-16")):
            fpr.chart_delivery_risk(deliver=False)
        spec = captured["spec"]
        self.assertEqual(spec["categories"], ["DECA.26"])
        self.assertEqual(spec["series"][0]["values"], [3000.0])          # 1000+2000
        self.assertEqual(spec["series"][0]["colors"], ["#E74C3C"])       # 有逾期→紅
        self.assertIn("2 單", spec["series"][0]["annotations"][0])
        self.assertIn("1 逾期", spec["series"][0]["annotations"][0])

    def test_no_risk_returns_friendly_message(self):
        # 全部已出貨/欠 0 → 友善回覆，不畫圖
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("15")
        ws.append(["客戶", "指令", "雙數", "希望出貨日", "實際出貨日", None, None, "Packing"])
        ws.append([None, None, None, None, None, "累計", "欠數", "日計"])
        ws.append(["DECA.26", "WO", 5000, _SOON, "2026-06-12", 5000, 0, 100])  # 已出貨
        buf = io.BytesIO()
        wb.save(buf)
        with mock.patch.object(fpr, "_load_latest_sheet",
                               return_value=(buf.getvalue(), "x.xlsx", "2026-06-16")):
            out = fpr.chart_delivery_risk(deliver=False)
        self.assertIn("沒有", str(out))


class TierAndRegistrationTests(unittest.TestCase):
    def test_safe_tier(self):
        from agent_core.tg_auth import is_sensitive
        from agent_core.tool_tiers import get_tier, TIER_SAFE
        for name in ("chart_daily_output", "chart_delivery_risk"):
            self.assertFalse(is_sensitive(name), name)
            self.assertEqual(get_tier(name), TIER_SAFE, name)

    def test_registered(self):
        from agent_core import tool_registry_catalog as cat
        names = {getattr(f, "__name__", "") for f in cat.BASE_BUILTIN_TOOLS}
        self.assertIn("chart_daily_output", names)
        self.assertIn("chart_delivery_risk", names)


if __name__ == "__main__":
    unittest.main()
