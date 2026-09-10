"""doc_export 多格式匯出 tool 的測試。

涵蓋：寬鬆輸入正規化、三種 renderer 結構正確、數字型別保留、格式別名、
錯誤路徑、Telegram 交付 wiring。

隔離：用 tempdir monkeypatch doc_export.EXPORTS_DIR（不碰真的 var/data/exports），
全部 deliver=False 或 patch 掉 telegram_send_file，不真的連 Telegram。
不 hardcode 任何 /Users/... 路徑。
"""
import io
import json
import os
import tempfile
import unittest
from unittest import mock

from agent_core import doc_export
from agent_core.doc_export import export_report

_RICH = {
    "title": "化學防具 & 緩衝物料 庫存盤點報告",
    "subtitle": "資料截止：2026-06-02",
    "blocks": [
        {"type": "heading", "text": "摘要"},
        {"type": "paragraph", "text": "盤點日截止於 5 月中旬及 3 月中旬，現場數據預估已變動。"},
        {"type": "list", "title": "建議", "items": ["請倉庫主管提供 6 月月結。", "由小紅匯入系統。"]},
        {"type": "table", "title": "依據資料",
         "columns": ["#", "檔案", "數量", "修改日期"],
         "rows": [
             [1, "5-15Jalas buffer stocks_.xlsx", 120, "2026-05-18"],
             [2, "260525 化學防具種類 & 庫存數.xls", 88, "2026-06-02"],
         ]},
    ],
}


class DocExportBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="red_export_test_")
        self._patch = mock.patch.object(doc_export, "EXPORTS_DIR", self._tmp)
        self._patch.start()

    def tearDown(self):
        self._patch.stop()
        import shutil
        shutil.rmtree(self._tmp, ignore_errors=True)

    def _dumps(self, obj):
        return json.dumps(obj, ensure_ascii=False)


class TestNormalization(DocExportBase):
    def test_parse_formats_aliases(self):
        self.assertEqual(doc_export._parse_formats("all"), ["excel", "word", "pdf"])
        self.assertEqual(doc_export._parse_formats("excel,pdf"), ["excel", "pdf"])
        self.assertEqual(doc_export._parse_formats("試算表"), ["excel"])
        self.assertEqual(doc_export._parse_formats("word docx"), ["word"])  # dedup
        self.assertEqual(doc_export._parse_formats(""), ["excel"])
        self.assertEqual(doc_export._parse_formats("garbage"), [])

    def test_list_of_dicts_becomes_single_table(self):
        spec = doc_export._normalize_spec([
            {"品名": "A", "數量": 10},
            {"品名": "B", "數量": 20, "備註": "急"},  # 多一欄 → 聯集
        ])
        self.assertEqual(len(spec["blocks"]), 1)
        tbl = spec["blocks"][0]
        self.assertEqual(tbl["type"], "table")
        self.assertEqual(tbl["columns"], ["品名", "數量", "備註"])
        self.assertEqual(tbl["rows"][0], ["A", 10, ""])      # 數字保留、缺欄補空
        self.assertEqual(tbl["rows"][1], ["B", 20, "急"])

    def test_lenient_block_types(self):
        spec = doc_export._normalize_spec({
            "title": "T",
            "blocks": [
                {"type": "title", "text": "標題級"},          # heading 別名
                "純字串也當段落",                               # 裸 str → paragraph
                {"items": ["a", "b"]},                          # 無 type，有 items → list
                {"type": "kv", "pairs": {"k1": "v1", "k2": 2}},  # kv 別名 + dict pairs
            ],
        })
        types = [b["type"] for b in spec["blocks"]]
        self.assertEqual(types, ["heading", "paragraph", "list", "keyvalue"])
        self.assertEqual(spec["blocks"][3]["pairs"], [["k1", "v1"], ["k2", "2"]])

    def test_empty_content_yields_no_blocks(self):
        self.assertEqual(doc_export._normalize_spec([])["blocks"], [])
        self.assertEqual(doc_export._normalize_spec({})["blocks"], [])

    def test_cellify_preserves_numbers_only(self):
        self.assertEqual(doc_export._cellify(120), 120)
        self.assertEqual(doc_export._cellify(1.5), 1.5)
        self.assertEqual(doc_export._cellify("120"), "120")   # 字串數字維持字串
        self.assertEqual(doc_export._cellify(True), "是")
        self.assertEqual(doc_export._cellify(None), "")


class TestExcel(DocExportBase):
    def test_excel_structure_and_native_numbers(self):
        import openpyxl
        res = export_report(self._dumps(_RICH), formats="excel", deliver=False)
        self.assertTrue(res.ok, res)
        xlsx = res.artifacts[0]
        wb = openpyxl.load_workbook(xlsx)
        # 報告分頁 + 表格分頁
        self.assertIn("依據資料", wb.sheetnames)
        report = wb[wb.sheetnames[0]]
        self.assertEqual(report["A1"].value, _RICH["title"])
        self.assertEqual(report["A1"].fill.fgColor.rgb[-6:], doc_export._DARK)
        # 表格：表頭深色、凍結、autofilter、數字是原生 int
        tbl = wb["依據資料"]
        self.assertEqual([tbl.cell(1, c).value for c in (1, 2, 3, 4)],
                         ["#", "檔案", "數量", "修改日期"])
        self.assertEqual(tbl.cell(1, 1).fill.fgColor.rgb[-6:], doc_export._DARK)
        self.assertEqual(tbl.freeze_panes, "A2")
        self.assertTrue(tbl.auto_filter.ref)
        # 數量欄第一筆是 int 120（不是字串），且右對齊
        qty = tbl.cell(2, 3)
        self.assertEqual(qty.value, 120)
        self.assertIsInstance(qty.value, int)
        self.assertEqual(qty.alignment.horizontal, "right")

    def test_sheet_name_dedup_and_sanitize(self):
        import openpyxl
        spec = {"blocks": [
            {"type": "table", "title": "A/B:C*?", "columns": ["x"], "rows": [[1]]},
            {"type": "table", "title": "A/B:C*?", "columns": ["y"], "rows": [[2]]},
        ]}
        res = export_report(self._dumps(spec), formats="excel", deliver=False)
        wb = openpyxl.load_workbook(res.artifacts[0])
        # 禁用字元被清掉，且兩個同名表得到不同分頁名
        self.assertEqual(len(wb.sheetnames), 2)
        self.assertNotEqual(wb.sheetnames[0], wb.sheetnames[1])
        for name in wb.sheetnames:
            self.assertFalse(set(name) & set('/\\?*[]:'))


class TestWord(DocExportBase):
    def test_word_structure(self):
        from docx import Document
        from docx.oxml.ns import qn
        res = export_report(self._dumps(_RICH), formats="word", deliver=False)
        self.assertTrue(res.ok, res)
        d = Document(res.artifacts[0])
        heads = [p.text for p in d.paragraphs
                 if p.style.name.startswith(("Title", "Heading"))]
        self.assertIn(_RICH["title"], heads)
        self.assertIn("摘要", heads)
        bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
        self.assertEqual(len(bullets), 2)
        t = d.tables[0]
        self.assertEqual([c.text for c in t.rows[0].cells], ["#", "檔案", "數量", "修改日期"])
        self.assertEqual(len(t.rows), 3)                       # header + 2
        self.assertEqual(t.rows[1].cells[2].text, "120")       # 數字 stringify
        shd = t.rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:shd"))
        self.assertEqual(shd.get(qn("w:fill")), doc_export._DARK)


class TestPdf(DocExportBase):
    def test_pdf_builds_and_embeds_font_when_available(self):
        res = export_report(self._dumps(_RICH), formats="pdf", deliver=False)
        self.assertTrue(res.ok, res)
        pdf = res.artifacts[0]
        with open(pdf, "rb") as f:
            data = f.read()
        self.assertTrue(data.startswith(b"%PDF-"))
        font = doc_export._ensure_cjk_font()
        if font.startswith("RED-CJK"):
            # 系統有可嵌入字型時，PDF 必含嵌入的 TrueType（FontFile2），
            # 否則手機 / 輕量檢視器會看不到中文字。
            self.assertIn(b"FontFile2", data,
                          "嵌入字型可用卻沒嵌入 TrueType 子集")
            self.assertGreater(len(data), 20000)


class TestErrors(DocExportBase):
    def test_bad_json(self):
        res = export_report("{not json", formats="excel")
        self.assertFalse(res.ok)
        self.assertEqual(res.error_code, "invalid_input")

    def test_empty_content(self):
        res = export_report("[]", formats="excel")
        self.assertFalse(res.ok)
        self.assertEqual(res.error_code, "invalid_input")

    def test_unknown_format(self):
        res = export_report(self._dumps(_RICH), formats="png")
        self.assertFalse(res.ok)
        self.assertEqual(res.error_code, "invalid_input")


class TestDelivery(DocExportBase):
    def test_deliver_true_sends_each_file(self):
        from agent_core.tool_result import ToolResult
        sent = []

        def fake_send(path, caption="", chat_id=""):
            sent.append((path, caption, chat_id))
            return ToolResult.success(f"已送 {os.path.basename(path)}")

        with mock.patch("agent_core.telegram.telegram_send_file", fake_send):
            res = export_report(self._dumps(_RICH), formats="excel,pdf",
                                deliver=True, chat_id="123")
        self.assertTrue(res.ok, res)
        self.assertEqual(len(sent), 2)                          # 兩個檔各送一次
        self.assertTrue(all(s[2] == "123" for s in sent))       # chat_id 有帶過去
        self.assertEqual({os.path.splitext(s[0])[1] for s in sent}, {".xlsx", ".pdf"})
        self.assertIn("已傳到 Telegram", res)

    def test_deliver_false_does_not_send(self):
        with mock.patch("agent_core.telegram.telegram_send_file") as m:
            res = export_report(self._dumps(_RICH), formats="excel", deliver=False)
        m.assert_not_called()
        self.assertTrue(res.ok)
        self.assertTrue(os.path.isfile(res.artifacts[0]))

    def test_delivery_failure_is_best_effort(self):
        # 傳送丟例外不該讓整個 tool 失敗（檔案已產出）
        def boom(*a, **k):
            raise RuntimeError("network down")

        with mock.patch("agent_core.telegram.telegram_send_file", boom):
            res = export_report(self._dumps(_RICH), formats="excel", deliver=True)
        self.assertTrue(res.ok)                                 # 仍算成功
        self.assertTrue(os.path.isfile(res.artifacts[0]))
        self.assertIn("傳送失敗", res)


class TestArtifacts(DocExportBase):
    def test_all_three_formats_produced(self):
        res = export_report(self._dumps(_RICH), formats="all", filename="樣本",
                            deliver=False)
        self.assertTrue(res.ok, res)
        exts = sorted(os.path.splitext(p)[1] for p in res.artifacts)
        self.assertEqual(exts, [".docx", ".pdf", ".xlsx"])
        for p in res.artifacts:
            self.assertTrue(os.path.isfile(p) and os.path.getsize(p) > 0)
            self.assertTrue(os.path.basename(p).startswith("樣本_"))


class TestImageCells(DocExportBase):
    """表格儲存格放圖片本體（2026-08-12 UserAng 案：規格單 PDF → tracking log）。"""

    def setUp(self):
        super().setUp()
        from PIL import Image
        self.img_dir = os.path.join(self._tmp, "imgs")
        os.makedirs(self.img_dir, exist_ok=True)
        self.photo = os.path.join(self.img_dir, "2001_p01_01.png")
        Image.new("RGB", (600, 400), (120, 160, 200)).save(self.photo)
        self.tall = os.path.join(self.img_dir, "2004_p01_01.png")
        Image.new("RGB", (200, 800), (200, 120, 60)).save(self.tall)
        # 大王路徑（沒有部門色）走 telegram._validate_send_path 的目錄閘，
        # tempdir 不在它的允許根裡 → 這裡放行掉，白名單本身另外測。
        self._roots = mock.patch.object(doc_export, "_image_allowed_roots",
                                        return_value=None)
        self._roots.start()
        self._ok = mock.patch("agent_core.telegram._validate_send_path",
                              side_effect=lambda p: (True, p))
        self._ok.start()

    def tearDown(self):
        self._ok.stop()
        self._roots.stop()
        super().tearDown()

    def _table(self, cell, **extra):
        return {"title": "tracking log",
                "blocks": [dict({"type": "table", "title": "samples",
                                 "columns": ["article", "remarks / Bemerkungen"],
                                 "rows": [["2001A", cell]]}, **extra)]}

    def _anchors(self, xlsx_path):
        import openpyxl
        ws = openpyxl.load_workbook(xlsx_path)["samples"]
        return ws, [(im.anchor._from.col, im.anchor._from.row,
                     round(im.anchor.ext.width / 9525),
                     round(im.anchor.ext.height / 9525)) for im in ws._images]

    def test_dict_image_cell_is_embedded_in_excel(self):
        res = export_report(self._dumps(self._table({"image": self.photo})),
                            formats="excel", deliver=False)
        self.assertTrue(res.ok, res)
        self.assertIn("已嵌入 1 張圖片", res.summary)
        self.assertEqual(res.data["images_embedded"], 1)
        ws, anchors = self._anchors(res.artifacts[0])
        # 錨在 B2（0-indexed col 1 / row 1）、等比縮到預設列高 90px
        self.assertEqual(anchors, [(1, 1, 135, 90)])
        self.assertIsNone(ws.cell(2, 2).value)          # 格子本身留空
        self.assertGreater(ws.row_dimensions[2].height, 60)   # 列高撐開裝得下
        self.assertGreater(ws.column_dimensions["B"].width, 15)

    def test_marker_string_form_also_works(self):
        res = export_report(self._dumps(self._table(f"[[IMG:{self.photo}]]")),
                            formats="excel", deliver=False)
        _, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(len(anchors), 1)

    def test_aspect_ratio_preserved_for_tall_image(self):
        res = export_report(self._dumps(self._table({"image": self.tall})),
                            formats="excel", deliver=False)
        _, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(anchors[0][2:], (22, 90))      # 200x800 → 22x90

    def test_image_height_override(self):
        res = export_report(self._dumps(self._table({"image": self.photo},
                                                    image_height=150)),
                            formats="excel", deliver=False)
        _, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(anchors[0][3], 150)

    def test_bad_path_degrades_to_text_and_is_reported(self):
        res = export_report(self._dumps(self._table({"image": "/nope/gone.png"})),
                            formats="excel", deliver=False)
        self.assertTrue(res.ok, res)
        ws, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(anchors, [])
        self.assertIn("gone.png", ws.cell(2, 2).value)
        # 🛑 不能靜默：LLM 要知道哪幾格沒圖，才不會對員工宣稱「圖都放好了」
        self.assertIn("沒放進去", res.summary)
        self.assertTrue(res.warnings)

    def test_failure_note_does_not_leak_absolute_path(self):
        secret = os.path.join(self._tmp, "私密資料夾", "gone.png")
        res = export_report(self._dumps(self._table({"image": secret})),
                            formats="excel", deliver=False)
        ws, _ = self._anchors(res.artifacts[0])
        self.assertIn("gone.png", ws.cell(2, 2).value)
        self.assertNotIn("私密資料夾", ws.cell(2, 2).value)
        self.assertNotIn("私密資料夾", res.summary)

    def test_non_image_extension_rejected(self):
        script = os.path.join(self.img_dir, "evil.sh")
        with open(script, "w") as f:
            f.write("echo hi")
        res = export_report(self._dumps(self._table({"image": script})),
                            formats="excel", deliver=False)
        _, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(anchors, [])
        self.assertIn("不支援的圖片格式", res.summary)

    def test_corrupt_image_rejected(self):
        fake = os.path.join(self.img_dir, "fake.png")
        with open(fake, "wb") as f:
            f.write(b"not really a png")
        res = export_report(self._dumps(self._table({"image": fake})),
                            formats="excel", deliver=False)
        _, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(anchors, [])
        self.assertIn("讀不出來", res.summary)

    def test_per_doc_cap_downgrades_extras_to_text(self):
        content = {"blocks": [{"type": "table", "title": "samples",
                               "columns": ["a", "img"],
                               "rows": [["1", {"image": self.photo}],
                                        ["2", {"image": self.photo}]]}]}
        with mock.patch.object(doc_export, "_IMAGE_MAX_PER_DOC", 1):
            res = export_report(self._dumps(content), formats="excel", deliver=False)
        ws, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(len(anchors), 1)
        self.assertEqual(ws.cell(3, 2).value, os.path.basename(self.photo))
        self.assertIn("上限", res.summary)

    def test_word_and_pdf_embed_too(self):
        import zipfile
        res = export_report(self._dumps(self._table({"image": self.photo})),
                            formats="all", deliver=False)
        docx = [p for p in res.artifacts if p.endswith(".docx")][0]
        media = [n for n in zipfile.ZipFile(docx).namelist() if "media/" in n]
        self.assertEqual(len(media), 1, media)
        pdf = [p for p in res.artifacts if p.endswith(".pdf")][0]
        self.assertGreater(os.path.getsize(pdf), 0)
        # 產出的檔會傳給員工/客戶：本機絕對路徑不該印在裡面
        with open(docx, "rb") as f:
            self.assertNotIn(self._tmp.encode(), f.read())

    def test_plain_dict_cell_still_becomes_json_text(self):
        """回歸：只有帶 image/圖 之類 key 的 dict 才算圖，一般 dict 照舊。"""
        res = export_report(self._dumps(self._table({"品名": "A", "數量": 3})),
                            formats="excel", deliver=False)
        ws, anchors = self._anchors(res.artifacts[0])
        self.assertEqual(anchors, [])
        self.assertIn("品名", ws.cell(2, 2).value)


class TestImageAllowedRoots(DocExportBase):
    """嵌圖來源白名單：部門員工收窄成 per-color 分艙，大王走敏感路徑黑名單。"""

    def test_owner_has_no_root_whitelist(self):
        with mock.patch.object(doc_export, "_dept_scope_color", return_value=""):
            self.assertIsNone(doc_export._image_allowed_roots())

    def test_dept_color_roots_are_per_color(self):
        with mock.patch.object(doc_export, "_dept_scope_color", return_value="orange"), \
             mock.patch("agent_core.logging_and_paths.DATA_DIR", self._tmp), \
             mock.patch("agent_core.logging_and_paths.GENERATED_IMAGES_DIR", self._tmp):
            roots = doc_export._image_allowed_roots()
        joined = " ".join(roots)
        self.assertIn(os.path.join("doc_images", "dept", "orange"), joined)
        self.assertNotIn(os.path.join("doc_images", "dept", "green"), joined)

    def test_image_outside_dept_roots_is_rejected(self):
        from PIL import Image
        outsider = os.path.join(self._tmp, "someone_else.png")
        Image.new("RGB", (200, 200), (0, 0, 0)).save(outsider)
        allowed = os.path.join(self._tmp, "allowed")
        os.makedirs(allowed, exist_ok=True)
        insider = os.path.join(allowed, "mine.png")
        Image.new("RGB", (200, 200), (0, 0, 0)).save(insider)
        with mock.patch.object(doc_export, "_image_allowed_roots",
                               return_value=(os.path.realpath(allowed),)):
            self.assertFalse(doc_export._resolve_image(outsider).ok)
            self.assertTrue(doc_export._resolve_image(insider).ok)


if __name__ == "__main__":
    unittest.main()
