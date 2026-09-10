"""factory_bom：型體→style 映射、BOM parser(DECA/JALAS/Supremo)、RICHTER 商品號解析、get_model_bom 工具。"""
import io
import unittest
from unittest import mock

import openpyxl

from agent_core import factory_bom as fb


def _deca_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append([None] * 12)
    ws.append([None, "Shoe Code", "DSM Code", "dim", "Supplier", None, "Net", "loss", "Gross", "Price"])
    ws.append([None, "UPPER SUMUP", None, None, None, None, None, None, None, None])  # 區段→跳
    ws.append([None, "010 - Toe cap", "9810042942", '54"', "三芳", "m", 0.028, 0.015, 0.0286, 17.99])
    ws.append([None, "020 - Vamp", "9810031446", '57"', "Italy", "m", 0.061, 0.015, 0.0624, 8.44])
    ws.append([None, "021 - Vamp", "Processing cost", '57"', "fude", "m", 0.06, 0.015, 0.06, 1.25])  # 加工費→跳
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _jalas_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Preliminary BOM"
    ws.append(["Model 1055"] + [None] * 8)
    ws.append(["Material Description", "Material Code", "Vendor Name", "Vendor Code",
               "Consumption", "Waste %", "Total Consumption", "Unit", "Unit Price"])
    ws.append(["PU split Rubber", "1230041", "Gruppo Mastrotto", "206742", 0.437, 44, 0.629, "SQF", 1.84])
    ws.append(["DRI-LEX 867", "12014010", "Toung Far", "215310", 0.159, 49, 0.237, "M2", 3.44])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class ModelToStyleTests(unittest.TestCase):
    def test_deca_jalas_unknown(self):
        self.assertEqual(fb._model_to_style("DJS336189-02"), ("DECATHLON", "336189"))
        self.assertEqual(fb._model_to_style("JA1055 BLACK"), ("JALAS", "1055"))
        self.assertEqual(fb._model_to_style("EG2305V-02"), ("LURCHI", "EG2305V"))
        self.assertEqual(fb._model_to_style("EH2311V-06"), ("LURCHI", "EH2311V"))
        self.assertEqual(fb._model_to_style("XX9999"), ("", ""))   # 真無 BOM 來源
        self.assertEqual(fb._model_to_style(""), ("", ""))


class ParseTests(unittest.TestCase):
    def test_deca_skips_section_and_processing(self):
        mats = fb._parse_deca_bom(_deca_xlsx())
        self.assertEqual(len(mats), 2)  # UPPER SUMUP + Processing cost 都跳掉
        self.assertEqual(mats[0]["code"], "9810042942")
        self.assertEqual(mats[0]["supplier"], "三芳")
        self.assertEqual(mats[0]["usage"], 0.0286)        # col8 毛用量
        self.assertEqual(mats[0]["unit_price"], 17.99)    # col9 單價

    def test_jalas_parses_by_header_keywords(self):
        mats = fb._parse_jalas_bom(_jalas_xlsx())
        self.assertEqual(len(mats), 2)
        self.assertEqual(mats[0]["code"], "1230041")
        self.assertEqual(mats[0]["supplier"], "Gruppo Mastrotto")
        self.assertEqual(mats[0]["usage"], 0.629)         # Total Consumption


class ToolTests(unittest.TestCase):
    def test_unsupported_customer_message(self):
        # 非 DECA/JALAS/LURCHI 會退到生管日報解析；mock 成查無，避免測試打 Drive
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("", "")):
            out = fb.get_model_bom("XX9999")
        self.assertIn("沒有 BOM 來源", out)
        self.assertIn("DECATHLON", out)
        self.assertIn("RICHTER", out)

    def test_formats_materials_and_supplier_summary(self):
        mats = [
            {"part": "010 Toe cap", "code": "9810042942", "supplier": "三芳",
             "unit": "m", "usage": 0.0286, "unit_price": 17.99},
            {"part": "020 Vamp", "code": "9810031446", "supplier": "三芳",
             "unit": "m", "usage": 0.06, "unit_price": 8.44},
        ]
        with mock.patch.object(fb, "_load_model_bom",
                               return_value=("DECATHLON", "336189", "src.xlsm", mats)):
            out = fb.get_model_bom("DJS336189")
        self.assertIn("DECATHLON 型體 336189", out)
        self.assertIn("9810042942", out)
        self.assertIn("三芳：2 項料", out)   # 供應商彙總（追採購用）

    def test_no_materials_message(self):
        # 找到來源檔但解不出料（格式不符）→ 與「沒做 CBD(src='')」區分
        with mock.patch.object(fb, "_load_model_bom",
                               return_value=("JALAS", "9999", "rfq.xlsx", [])):
            out = fb.get_model_bom("JA9999")
        self.assertIn("抓不到料表", out)


class RegistrationTests(unittest.TestCase):
    def test_registered(self):
        from agent_core import tool_registry_catalog as cat
        names = {getattr(t, "__name__", "") for t in cat.BASE_BUILTIN_TOOLS}
        self.assertIn("get_model_bom", names)
        self.assertIn("check_material_readiness", names)


class ReadinessTests(unittest.TestCase):
    def _mats(self):
        return [
            {"part": "toe", "code": "98", "supplier": "Coats Phong Phú", "unit": "m", "usage": 0.1, "unit_price": 1},
            {"part": "vamp", "code": "98", "supplier": "Coats Phong Phú", "unit": "m", "usage": 0.1, "unit_price": 1},
            {"part": "lining", "code": "99", "supplier": "三芳", "unit": "m", "usage": 0.1, "unit_price": 1},
        ]

    def test_readiness_shows_po_and_flags_misses(self):
        def fake_timeline(dept, query="", days=0, limit=20):
            if "Coats" in query:
                return ("🧵 採購\n  [2026-06-03] FUCHUN to COATS => JF0P26060002\n"
                        "       ↳ 確認兩筆訂單")
            return "查無 採購 相關信件"
        with mock.patch.object(fb, "_load_model_bom",
                               return_value=("DECATHLON", "336189", "src", self._mats())), \
                mock.patch("agent_core.lake_dept_timeline.read_dept_email_timeline",
                           side_effect=fake_timeline):
            out = fb.check_material_readiness("DJS336189")
        self.assertIn("採購備料進度", out)
        self.assertIn("Coats", out)
        self.assertIn("JF0P26060002", out)   # 採購信抽到的 PO
        self.assertIn("人工確認", out)         # 三芳 對不上 → 標人工確認
        self.assertIn("三芳", out)

    def test_readiness_unsupported_model(self):
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("", "")):
            out = fb.check_material_readiness("XX9999")
        self.assertIn("沒有 BOM 來源", out)

    def test_latest_procurement_parses_date_and_summary(self):
        def fake_timeline(dept, query="", days=0, limit=20):
            return ("🧵 採購\n  [2025-12-09] GRUPPO => PO#JA251209-04\n"
                    "       ↳ 確認採購單，2026 Q4 出貨")
        with mock.patch("agent_core.lake_dept_timeline.read_dept_email_timeline",
                        side_effect=fake_timeline):
            r = fb._latest_procurement("Gruppo Mastrotto", 240)
        self.assertIn("2025-12-09", r)
        self.assertIn("PO#JA251209-04", r)


def _lurchi_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "#28"
    for _ in range(7):
        ws.append([None] * 14)
    h = [None] * 14
    h[2], h[3], h[4], h[5], h[6], h[12] = (
        "parts", "material specification", "Supplier", "Pairs/SF/Y", "measurement", "unit price")
    ws.append(h)

    def row(part, mat, sup, usage):
        r = [None] * 14
        if part:
            r[0], r[2] = 1, part
        r[3], r[5] = mat, usage
        if sup:
            r[4] = sup
        return r

    ws.append(row("Strap 1", "PU經編絨", "", 142.4))
    ws.append(row("", "PU經編絨", "", 142.4))
    ws.append(row("", "紙襯 paperlining", "南寶", 142.4))
    ws.append(row("", "貼合費", "", 142.4))
    ws.append(row("Strap 2", "PU經編絨", "", 100.0))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class SupremoParseTests(unittest.TestCase):
    def test_dedup_skip_fee_and_supplier(self):
        mats = fb._parse_supremo_cbd(_lurchi_xlsx())
        parts = [m["part"] for m in mats]
        self.assertEqual(len(mats), 3)  # Strap1 PU(去重) + Strap1 紙襯 + Strap2 PU
        self.assertFalse(any("貼合費" in p for p in parts))
        self.assertTrue(any("Strap 1: PU經編絨" in p for p in parts))
        self.assertTrue(any("Strap 2: PU經編絨" in p for p in parts))
        npm = next(m for m in mats if "紙襯" in m["part"])
        self.assertEqual(npm["supplier"], "南寶")


class RichterResolveTests(unittest.TestCase):
    """RICHTER：型體碼不帶 CBD 號，靠生管日報 cust_style 換成商品號再找 CBD。"""

    _SCHED = [
        {"customer": "RICHTER", "model": "FE2303V-02 AT", "cust_style": "5001L-4691-7201"},
        {"customer": "RICHTER", "model": "FE2303V4-01MYSTIC", "cust_style": "5001L-4292-7800"},
        {"customer": "LURCHI", "model": "EG2305V-02", "cust_style": "63L1033038-AT"},
    ]

    def test_model_to_article_via_schedule(self):
        with mock.patch("agent_core.production_schedule._load_schedule",
                        return_value=(self._SCHED, "報表.xlsx")):
            # 精確型體 → 自己的商品號（去 L、去顏色尾碼）
            self.assertEqual(fb._resolve_from_schedule("FE2303V-02 AT"), ("RICHTER", "5001-4691"))
            # 不帶顏色、用核心前綴比對；且不會誤配到 FE2303V-02
            self.assertEqual(fb._resolve_from_schedule("FE2303V4"), ("RICHTER", "5001-4292"))
            # 查無 → 空（不亂猜、不跨形體）
            self.assertEqual(fb._resolve_from_schedule("ZZ0000"), ("", ""))

    def test_richter_bom_via_article(self):
        mats = [{"part": "Strap 1: PU經編絨", "code": "", "supplier": "",
                 "unit": "Y", "usage": 173.4, "unit_price": 7.88}]
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("RICHTER", "5001-4691")), \
                mock.patch.object(fb, "_find_supremo_cbd_by_article",
                                  return_value=("CBD 5001-4691 (2026) NYLON 70D.xlsx", mats)):
            out = fb.get_model_bom("FE2303V-02")
        self.assertIn("RICHTER 型體 5001-4691", out)
        self.assertIn("Strap 1: PU經編絨", out)

    def test_richter_no_cbd_is_normal(self):
        # 形體沒做 CBD（找不到）→ 明確訊息、非錯誤
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("RICHTER", "5001-4693")), \
                mock.patch.object(fb, "_find_supremo_cbd_by_article", return_value=("", [])):
            out = fb.get_model_bom("FE2303V5-01")
        self.assertIn("找不到對應的 CBD", out)
        self.assertIn("並非每個形體", out)


class BomProbeTests(unittest.TestCase):
    """bom_probe：給生產排程落後清單一鍵標示『有沒有料表』的輕量探針。"""

    def test_lurchi_existence_only(self):
        # DECA/JALAS/LURCHI 走檔名搜尋(不下載)；找到=yes、無=no
        with mock.patch.object(fb, "_find_bom_file", return_value={"name": "CBD EH2311V.xlsx", "id": "1"}):
            self.assertEqual(fb.bom_probe("EH2311V-04"), ("yes", "CBD EH2311V.xlsx", 0))
        with mock.patch.object(fb, "_find_bom_file", return_value=None):
            self.assertEqual(fb.bom_probe("EH2311V-04"), ("no", "", 0))

    def test_richter_parse_confirm(self):
        # RICHTER 要解析確認（規格表與 CBD 同名）；命中帶料數
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("RICHTER", "5001-4292")), \
                mock.patch.object(fb, "_find_supremo_cbd_by_article",
                                  return_value=("CBD 5001-4292.xlsx", [{"x": 1}] * 79)):
            self.assertEqual(fb.bom_probe("FE2303V4"), ("yes", "CBD 5001-4292.xlsx", 79))
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("RICHTER", "5001-4691")), \
                mock.patch.object(fb, "_find_supremo_cbd_by_article", return_value=("", [])):
            self.assertEqual(fb.bom_probe("FE2303V-02"), ("no", "", 0))

    def test_unsupported(self):
        with mock.patch.object(fb, "_resolve_from_schedule", return_value=("", "")):
            self.assertEqual(fb.bom_probe("ZZ0000"), ("unsupported", "", 0))


if __name__ == "__main__":
    unittest.main()
