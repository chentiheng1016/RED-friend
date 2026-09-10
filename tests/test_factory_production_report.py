"""結構化生產日報進度表解析器的測試（合成 fixture，免網路）。

驗證 agent_core/factory_production_report.py 能從『以日為分頁、寬交叉表』的
xlsx 正確抽出 每日×各客戶 日產量 + 截至最新日的包裝累計/未完。
"""
import datetime
import io
import unittest
from unittest import mock

import openpyxl

from agent_core.factory_production_report import (
    _detect_layout,
    _iter_day_sheets,
    _parse_ship_date,
    _ship_date_sort_key,
    _summarize,
    iter_production_rows,
)


def _write_template_sheet(ws, data_rows):
    """把一個分頁寫成福群進度表版型：表頭在第 3-4 列、資料從第 5 列。

    data_rows: list[dict]，鍵見下方 colmap。
    """
    # 第 3 列：主表頭（站名用穩定英文錨點）
    ws.cell(row=3, column=3, value="客戶型體")   # Supremo 業務訂單型體欄
    ws.cell(row=3, column=4, value="客戶")
    ws.cell(row=3, column=6, value="指 令")
    ws.cell(row=3, column=7, value="型體")
    ws.cell(row=3, column=11, value="雙數")
    ws.cell(row=3, column=19, value="Stitching")
    ws.cell(row=3, column=22, value="Molding")
    ws.cell(row=3, column=23, value="Insock")
    ws.cell(row=3, column=24, value="Injection")
    ws.cell(row=3, column=29, value="Packing")
    ws.cell(row=3, column=32, value="客戶\n希望\n出貨日")
    ws.cell(row=3, column=34, value="實際\n出貨日")
    # 第 4 列：累計/欠數/日計 子標
    ws.cell(row=4, column=17, value="累 計\nLŨY KẾ\nsum")
    ws.cell(row=4, column=18, value="欠 數\nremaining")
    ws.cell(row=4, column=19, value="日計\nS/LTR.NG\nday")
    ws.cell(row=4, column=20, value="累 計\nsum")
    ws.cell(row=4, column=21, value="欠 數\nremaining")
    ws.cell(row=4, column=22, value="日計\nday")
    ws.cell(row=4, column=23, value="日計\nday")
    ws.cell(row=4, column=24, value="日計\nday")
    ws.cell(row=4, column=26, value="累 計\nsum")          # Packing 累計
    ws.cell(row=4, column=27, value="累 計\n不良 NG sum")   # Packing NG 累計（要被跳過）
    ws.cell(row=4, column=28, value="欠 數\nremaining")     # Packing 欠數
    ws.cell(row=4, column=29, value="日計\nday")            # Packing 日計
    # 資料列
    r = 5
    for d in data_rows:
        ws.cell(row=r, column=3, value=d.get("cust_style", ""))
        ws.cell(row=r, column=4, value=d["customer"])
        ws.cell(row=r, column=6, value=d["work_order"])
        ws.cell(row=r, column=7, value=d.get("model", ""))
        ws.cell(row=r, column=11, value=d.get("pairs", 0))
        ws.cell(row=r, column=19, value=d.get("stitch_day", 0))
        ws.cell(row=r, column=22, value=d.get("mold_day", 0))
        ws.cell(row=r, column=23, value=d.get("insock_day", 0))
        ws.cell(row=r, column=24, value=d.get("inject_day", 0))
        ws.cell(row=r, column=26, value=d.get("pack_cum", 0))
        ws.cell(row=r, column=28, value=d.get("pack_rem", 0))
        ws.cell(row=r, column=29, value=d.get("pack_day", 0))
        ws.cell(row=r, column=32, value=d.get("want_ship"))
        ws.cell(row=r, column=34, value=d.get("actual_ship"))
        r += 1


def _build_fixture():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    s1 = wb.create_sheet("01")
    s2 = wb.create_sheet("02")
    # 日 01
    _write_template_sheet(s1, [
        dict(customer="DECA.26", work_order="JFC25001", model="MH100",
             cust_style="74L1303003", pairs=224, stitch_day=100, insock_day=50,
             inject_day=80, pack_day=70,
             pack_cum=224, pack_rem=0, want_ship=datetime.datetime(2026, 6, 20), actual_ship="OK"),
        dict(customer="DECA.26", work_order="JFC25002", model="MH100",
             pairs=300, stitch_day=20, inject_day=30, pack_day=40,
             pack_cum=120, pack_rem=180, want_ship=datetime.datetime(2026, 6, 25), actual_ship=None),
        dict(customer="JALAS", work_order="JFC25003", model="EJ2401",
             pairs=96, inject_day=289, pack_day=0,
             pack_cum=96, pack_rem=0, want_ship=datetime.datetime(2026, 6, 18), actual_ship="OK"),
    ])
    # 日 02（最新一天 → 累計/欠數從這頁取）
    _write_template_sheet(s2, [
        dict(customer="DECA.26", work_order="JFC25001", model="MH100",
             pairs=224, stitch_day=10, inject_day=5, pack_day=15,
             pack_cum=224, pack_rem=0, want_ship=datetime.datetime(2026, 6, 20), actual_ship="OK"),
        dict(customer="DECA.26", work_order="JFC25002", model="MH100",
             pairs=300, stitch_day=0, inject_day=0, pack_day=0,
             pack_cum=260, pack_rem=40, want_ship=datetime.datetime(2026, 6, 25), actual_ship=None),
        dict(customer="JALAS", work_order="JFC25003", model="EJ2401",
             pairs=96, inject_day=223, pack_day=1320,
             pack_cum=96, pack_rem=0, want_ship=datetime.datetime(2026, 6, 18), actual_ship="OK"),
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DetectLayoutTests(unittest.TestCase):
    def test_colmap_anchored_by_header_text(self):
        wb = openpyxl.load_workbook(io.BytesIO(_build_fixture()), data_only=True, read_only=True)
        rows = list(wb["01"].iter_rows(values_only=True))
        data_start, (colmap, warns) = _detect_layout(rows)
        self.assertEqual(data_start, 4)  # 表頭第 3 列(0-based idx2)+2 → 資料 idx4 = 第 5 列
        self.assertEqual(colmap["cust_style"], 2)  # 客戶型體：1-based col3 → 0-based 2
        self.assertEqual(colmap["customer"], 3)    # 1-based col4 → 0-based 3
        self.assertEqual(colmap["work_order"], 5)
        self.assertEqual(colmap["pairs"], 10)
        self.assertEqual(colmap["Stitching_day"], 18)
        self.assertEqual(colmap["Injection_day"], 23)
        self.assertEqual(colmap["Packing_day"], 28)
        # 包裝累計要抓 col26(idx25) 而非 NG 的 col27(idx26)
        self.assertEqual(colmap["pack_cum"], 25)
        self.assertEqual(colmap["pack_rem"], 27)
        self.assertEqual(warns, [])


class DetectLayoutPackWarningTests(unittest.TestCase):
    """Packing 日計欄在、但 累計/欠數 子標解不到 → 必須 append 警告（整欄會靜默變 0，
    production_alert 會假全綠）。"""

    @staticmethod
    def _rows_without_pack_subheaders():
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("01")
        ws.cell(row=3, column=4, value="客戶")
        ws.cell(row=3, column=6, value="指 令")
        ws.cell(row=3, column=11, value="雙數")
        ws.cell(row=3, column=29, value="Packing")
        ws.cell(row=4, column=29, value="日計\nday")
        # 刻意不放 累計/欠數 子標（模擬版型位移）
        ws.cell(row=5, column=4, value="DECA.26")
        ws.cell(row=5, column=6, value="JFC1")
        ws.cell(row=5, column=11, value=100)
        buf = io.BytesIO()
        wb.save(buf)
        wb2 = openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True, read_only=True)
        return list(wb2["01"].iter_rows(values_only=True))

    def test_missing_pack_cum_rem_warned(self):
        data_start, (colmap, warns) = _detect_layout(self._rows_without_pack_subheaders())
        self.assertIsNotNone(data_start)
        self.assertNotIn("pack_cum", colmap)
        self.assertNotIn("pack_rem", colmap)
        self.assertTrue(any("包裝累計" in w for w in warns))
        self.assertTrue(any("欠數" in w for w in warns))

    def test_good_fixture_has_no_pack_warnings(self):
        wb = openpyxl.load_workbook(io.BytesIO(_build_fixture()), data_only=True, read_only=True)
        rows = list(wb["01"].iter_rows(values_only=True))
        _start, (_colmap, warns) = _detect_layout(rows)
        self.assertEqual(warns, [])


class ParseShipDateTests(unittest.TestCase):
    """無年日期取「離今天最近」的年份（一律補今年會讓 1 月讀到的 12/28 不再逾期）。"""

    def test_full_date_passthrough(self):
        self.assertEqual(_parse_ship_date("2026-06-20"), datetime.date(2026, 6, 20))
        self.assertEqual(_parse_ship_date("2026/06/20"), datetime.date(2026, 6, 20))

    def test_yearless_near_new_year_resolves_to_prior_year(self):
        today = datetime.date(2026, 1, 5)
        self.assertEqual(_parse_ship_date("12/28", today=today), datetime.date(2025, 12, 28))

    def test_yearless_late_december_resolves_to_next_year(self):
        today = datetime.date(2026, 12, 20)
        self.assertEqual(_parse_ship_date("1/15", today=today), datetime.date(2027, 1, 15))

    def test_yearless_mid_year_stays_current(self):
        today = datetime.date(2026, 7, 10)
        self.assertEqual(_parse_ship_date("6/20", today=today), datetime.date(2026, 6, 20))

    def test_unparseable_returns_none(self):
        self.assertIsNone(_parse_ship_date("排程中"))
        self.assertIsNone(_parse_ship_date(""))

    def test_sort_key_orders_by_real_date_not_string(self):
        # '6/5'（無年）字串排序會排在 '2026-05-30' 後面 → 要按真日期排
        wants = ["2026-07-01", "6/5", "2026-05-30", "排程中"]
        ordered = sorted(wants, key=_ship_date_sort_key)
        self.assertEqual(ordered[-1], "排程中")               # 解析不出 → 最後
        d = [_parse_ship_date(w) for w in ordered[:-1]]
        self.assertEqual(d, sorted(d))                        # 前段依真日期升冪


class SummariseTests(unittest.TestCase):
    def setUp(self):
        self.data = _build_fixture()

    def test_daily_per_customer_sums(self):
        out = _summarize(self.data)
        self.assertIn("1號", out)
        self.assertIn("2號", out)
        # 日01 DECA：針車100+20=120，灌注80+30=110，包裝70+40=110，中底50
        self.assertIn("針車120", out)
        self.assertIn("灌注110", out)
        self.assertIn("包裝110", out)
        # 日01 JALAS 灌注289
        self.assertIn("灌注289", out)
        # 日02 JALAS 包裝1320
        self.assertIn("包裝1320", out)

    def test_cumulative_from_latest_day(self):
        out = _summarize(self.data)
        self.assertIn("截至 2 號", out)
        # 本季含今年已出貨：DECA 本季在表 524（已出貨 224 ＋ 在產 300）、累計 484、未完 40
        self.assertRegex(out, r"DECA\.26: 本季在表 524 雙（已出貨 224｜在產 300）｜已包裝累計 484｜未完 40")

    def test_this_year_shipped_counted_not_excluded(self):
        # 關鍵回歸：今年已出貨的**單**仍算本季在表（不再被「已出貨就排除」逐列誤殺）。
        out = _summarize(self.data)
        # DECA 的 JFC25001(224 雙)今年已出貨 → 仍計進本季在表 524，不是只剩在產 300
        self.assertRegex(out, r"DECA\.26: 本季在表 524 雙（已出貨 224｜在產 300）")
        # fixture 全是今年單 → 不該有「往年舊單」區、也不再有「訂單」誤導標籤
        self.assertNotIn("往年舊單", out)
        self.assertNotIn("訂單 ", out)

    def test_fully_shipped_customer_leaves_progress_table(self):
        """大王規則「已經出完貨的就不需要回報」：整季出完的客戶不佔進度表，數字不消失。"""
        out = _summarize(self.data)
        prog = out.split("【本季已出完貨")[0]
        # JALAS 96 雙全數出貨、未完 0 → 不再佔「本季各客戶 生產進度」一列
        self.assertNotIn("JALAS: 本季在表 96 雙（已出貨", prog)
        # …改成一行結案摘要（別人單獨問 JALAS 時仍查得到，不是憑空消失）
        self.assertIn("【本季已出完貨（在產 0、未完 0）—— 每日回報不需列入】共 1 個客戶 96 雙", out)
        self.assertIn("JALAS: 本季在表 96 雙 已全數出貨（已包裝累計 96）", out)
        # 還有未完的 DECA 不受影響，照樣在進度表
        self.assertIn("DECA.26: 本季在表 524", prog)

    def test_packed_but_not_shipped_stays_in_progress_table(self):
        """未完 0 但『在產 > 0』＝包裝完還沒出貨 → 還要追，不能當出完貨踢掉。"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        _write_template_sheet(wb.create_sheet("01"), [
            dict(customer="RICHTER", work_order="JFC26501", model="FE23",
                 pairs=400, pack_cum=400, pack_rem=0,
                 want_ship=datetime.datetime(2026, 8, 20), actual_ship=None),
        ])
        buf = io.BytesIO()
        wb.save(buf)
        out = _summarize(buf.getvalue())
        self.assertRegex(out, r"RICHTER: 本季在表 400 雙（已出貨 0｜在產 400）")
        self.assertNotIn("已出完貨", out)

    def test_missing_remaining_column_disables_done_split(self):
        """欠數欄解不到時 pack_rem 全 0 是假象 → 整段分流停用，寧可多念也不靜默漏報。"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("01")
        _write_template_sheet(ws, [
            dict(customer="JALAS", work_order="JFC26601", model="EJ24",
                 pairs=96, pack_cum=96, pack_rem=0,
                 want_ship=datetime.datetime(2026, 6, 18), actual_ship="OK"),
        ])
        # 抹掉 Packing「欠 數」子標＝版型位移。要走 .value 設值：openpyxl 的
        # cell(row, col, value=None) 是「不給值」語意，不會清掉既有內容。
        ws.cell(row=4, column=28).value = None
        buf = io.BytesIO()
        wb.save(buf)
        out = _summarize(buf.getvalue())
        self.assertIn("找不到『欠數』欄", out)
        self.assertRegex(out, r"JALAS: 本季在表 96 雙（已出貨 96｜在產 0）")
        self.assertNotIn("已出完貨", out)

    def test_overdue_shipment_listed(self):
        out = _summarize(self.data)
        # DECA JFC25002 未完 40、未標實際出貨 → 出現在應出貨清單
        self.assertIn("JFC25002", out)
        self.assertIn("2026-06-25", out)
        # 已出貨(actual='OK')的 JFC25001 不該出現在「未出貨」清單尾段
        tail = out.split("近期應出貨")[-1]
        self.assertNotIn("JFC25001", tail)

    def test_day_filter(self):
        out = _summarize(self.data, day="2")
        self.assertIn("2號", out)
        self.assertNotIn("1號  ", out)
        # 單日視圖不出累計區
        self.assertNotIn("截至", out)

    def test_customer_filter(self):
        out = _summarize(self.data, customer="JALAS")
        self.assertIn("JALAS", out)
        self.assertNotIn("DECA", out)

    def test_non_day_sheet_workbook_graceful(self):
        wb = openpyxl.Workbook()
        wb.active.title = "封面"
        wb.active["A1"] = "不是進度表"
        buf = io.BytesIO()
        wb.save(buf)
        out = _summarize(buf.getvalue())
        self.assertIn("找不到", out)


class IterDaySheetsTests(unittest.TestCase):
    def test_only_numeric_sheets_sorted(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for n in ["02", "01", "封面", "14", "說明"]:
            wb.create_sheet(n)
        days = _iter_day_sheets(wb)
        self.assertEqual(days, [(1, "01"), (2, "02"), (14, "14")])


class IterProductionRowsTests(unittest.TestCase):
    """新公開入口 iter_production_rows()：數據倉 loader 與 _summarize 共用的逐列產出。"""

    def setUp(self):
        self.data = _build_fixture()

    def test_yields_all_data_rows(self):
        # 2 天 × 3 列 = 6
        self.assertEqual(len(list(iter_production_rows(self.data))), 6)

    def test_row_shape_and_values(self):
        rows = list(iter_production_rows(self.data))
        r = next(x for x in rows if x["prod_day"] == 1 and x["work_order"] == "JFC25001")
        self.assertEqual(r["customer"], "DECA.26")
        self.assertEqual(r["model"], "MH100")
        self.assertEqual(r["cust_style"], "74L1303003")  # 客戶型體(Supremo)同列抽出
        self.assertEqual(r["pairs"], 224.0)
        self.assertEqual(r["stitching_day"], 100.0)
        self.assertEqual(r["insock_day"], 50.0)
        self.assertEqual(r["injection_day"], 80.0)
        self.assertEqual(r["packing_day"], 70.0)
        self.assertEqual(r["molding_day"], 0.0)  # PU 灌注廠，成型常 0
        self.assertEqual(r["pack_cum"], 224.0)
        self.assertEqual(r["pack_rem"], 0.0)

    def test_is_latest_day_flag(self):
        rows = list(iter_production_rows(self.data))
        self.assertTrue(all(not x["is_latest_day"] for x in rows if x["prod_day"] == 1))
        self.assertTrue(all(x["is_latest_day"] for x in rows if x["prod_day"] == 2))

    def test_raw_ship_dates_preserved(self):
        # want_ship 原樣保留為 datetime（落庫時才轉），未出貨 actual_ship 為 None
        r = next(x for x in iter_production_rows(self.data)
                 if x["prod_day"] == 2 and x["work_order"] == "JFC25002")
        self.assertIsInstance(r["want_ship"], datetime.datetime)
        self.assertIsNone(r["actual_ship"])

    def test_blank_customer_rows_skipped(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        s = wb.create_sheet("01")
        _write_template_sheet(s, [
            dict(customer="DECA.26", work_order="JFC1", pairs=10, pack_day=5),
            dict(customer="", work_order="", pairs=0),        # 空客戶 → 跳過
            dict(customer="  ", work_order="JFC2", pairs=20),  # 純空白 → 跳過
        ])
        buf = io.BytesIO()
        wb.save(buf)
        rows = list(iter_production_rows(buf.getvalue()))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["customer"], "DECA.26")

    def test_non_day_sheet_yields_nothing_with_warning(self):
        wb = openpyxl.Workbook()
        wb.active.title = "封面"
        wb.active["A1"] = "不是進度表"
        buf = io.BytesIO()
        wb.save(buf)
        warns = set()
        self.assertEqual(list(iter_production_rows(buf.getvalue(), warnings=warns)), [])
        self.assertTrue(warns)  # 有記下原因，不靜默

    def test_bad_bytes_yields_nothing_with_warning(self):
        warns = set()
        self.assertEqual(list(iter_production_rows(b"not an xlsx", warnings=warns)), [])
        self.assertTrue(any("解析 xlsx 失敗" in w for w in warns))

    def test_generator_is_source_of_truth_for_summary(self):
        # 不變量：generator 算出的 day02 JALAS 包裝日計，與 _summarize 文字『包裝1320』同源
        jalas_d2 = next(x for x in iter_production_rows(self.data)
                        if x["prod_day"] == 2 and x["customer"] == "JALAS")
        self.assertEqual(jalas_d2["packing_day"], 1320.0)
        self.assertIn(f"包裝{int(jalas_d2['packing_day'])}", _summarize(self.data))


def _build_year_fixture():
    """單日 fixture：本季(今年)在產+已出貨 + 往年(去年)結案舊單，驗『依年份』分流。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    s = wb.create_sheet("01")
    _write_template_sheet(s, [
        dict(customer="DECA.26", work_order="JFC26001", model="MH1",
             pairs=300, pack_cum=100, pack_rem=200,
             want_ship=datetime.datetime(2026, 9, 20), actual_ship=None),    # 今年在產
        dict(customer="DECA.26", work_order="JFC26002", model="MH1",
             pairs=500, pack_cum=500, pack_rem=0,
             want_ship=datetime.datetime(2026, 9, 25), actual_ship="OK"),     # 今年已出貨 → 仍算本季
        dict(customer="NISSHIN", work_order="JFC25900", model="DX1",
             pairs=620, pack_cum=620, pack_rem=0,
             want_ship=datetime.datetime(2025, 11, 30), actual_ship="OK"),    # 去年舊單 → 不計本季
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class YearSegregationTests(unittest.TestCase):
    def test_this_year_shipped_included_prior_year_segregated(self):
        out = _summarize(_build_year_fixture())
        # DECA 含今年已出貨(500)：本季在表 800（已出貨 500｜在產 300）
        self.assertRegex(out, r"DECA\.26: 本季在表 800 雙（已出貨 500｜在產 300）")
        # NISSHIN 是去年單 → 不在本季生產進度區
        prod_section = out.split("往年舊單")[0]
        self.assertNotIn("NISSHIN", prod_section)
        # …而是落在「往年舊單」區
        self.assertIn("往年舊單", out)
        self.assertRegex(out, r"NISSHIN: 1 單 620 雙")


def _build_daily_fixture():
    """兩天 fixture 給每日回報：昨天(2 號)有產出的、已出完貨的、還有未完的都各來一家。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rows_day = [
        # 昨天有產出、還有未完 → 一定要列
        dict(customer="DECA.26", work_order="JFC26001-1-1", model="MH100",
             pairs=800, stitch_day=244, inject_day=900, pack_day=320,
             pack_cum=500, pack_rem=300,
             want_ship=datetime.datetime(2026, 8, 20), actual_ship=None),
        # 昨天沒動、但還有未完（＝上線前被 LLM 漏掉的 RICHTER 那種）→ 也要列
        dict(customer="RICHTER", work_order="JFC26050-1-1", model="FE2303",
             pairs=600, pack_cum=560, pack_rem=40,
             want_ship=datetime.datetime(2026, 5, 29), actual_ship=None),
        # 本季已出完貨 → 不列、只計進「另有 N 家」
        dict(customer="NEW WAVE", work_order="JFC26099-1-1", model="DV2501",
             pairs=1440, pack_cum=1440, pack_rem=0,
             want_ship=datetime.datetime(2026, 5, 1), actual_ship="OK"),
    ]
    for name in ("01", "02"):
        _write_template_sheet(wb.create_sheet(name), [
            d if name == "02" else {**d, "stitch_day": 0, "inject_day": 0, "pack_day": 0}
            for d in rows_day
        ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DailyProductionReportTests(unittest.TestCase):
    """每日生產數量回報（確定性版）——同一份資料每次都要產出同一張表。"""

    def setUp(self):
        from agent_core import factory_production_report as fpr
        self.fpr = fpr
        self.data = _build_daily_fixture()
        # ERP 段獨立驗，這裡固定回空 → 表格本身的行為不受鏡像在不在影響。
        self.erp = mock.patch.object(fpr, "_erp_order_facts", return_value={})
        self.erp.start()
        self.addCleanup(self.erp.stop)

    def _run(self, now):
        return self.fpr._build_daily_production_report(
            now, self.data, "08月份生產日報進度表08-02.xlsx", "2026-08-03")

    def test_table_lists_active_customers_and_hides_shipped_out(self):
        out = self._run(datetime.datetime(2026, 8, 3, 10, 0))
        self.assertIn("2026-08-02（日）每日生產數量回報", out)
        # 昨天有產出的排前面、數字照表念
        self.assertIn("| **DECA.26** | 900 雙 | 320 雙 | 244 雙 | 500 雙 | 300 雙 |", out)
        # 昨天沒動但還有未完的仍在表上（漏這種才是漏報）
        self.assertIn("| **RICHTER** | － | － | － | 560 雙 | 40 雙 |", out)
        # 已出完貨的不列名，只報家數
        self.assertNotIn("NEW WAVE", out)
        self.assertIn("另有 1 家本季已出完貨、共 1,440 雙", out)

    def test_same_input_renders_identical_report(self):
        """改確定性的重點：同一份資料重跑不會產出不同的表（去重才會生效）。"""
        now = datetime.datetime(2026, 8, 3, 10, 0)
        self.assertEqual(self._run(now), self._run(now))

    def test_quiet_before_chase_hour_when_sheet_lags(self):
        # 進度表只到 2 號，昨天是 4 號 → 早上先安靜等下一輪
        out = self._run(datetime.datetime(2026, 8, 5, 9, 30))
        self.assertEqual(out, "(無新發現)")

    def test_chases_after_chase_hour_on_workday(self):
        out = self._run(datetime.datetime(2026, 8, 5, 11, 30))
        self.assertIn("仍停在 8/2", out)
        self.assertIn("建議催收", out)

    def test_quiet_when_yesterday_was_sunday(self):
        # 2026-08-10 是週一 → 昨天是週日，工廠不生產，過 11 點也不催
        self.assertEqual(datetime.date(2026, 8, 9).weekday(), 6)
        out = self._run(datetime.datetime(2026, 8, 10, 11, 30))
        self.assertEqual(out, "(無新發現)")


class DailyReportFetchFailureTests(unittest.TestCase):
    """讀不到 Drive 不寄給收報表的人 —— 走排程健康告警那條線（#370）。

    2026-08-10 實際踩到：Drive 偶發 403 User rate limit exceeded，09:13 就把整段
    HttpError（含 API URL、Details JSON）寄給四個人，而 30 分鐘後那輪其實就恢復了。
    """

    def setUp(self):
        from agent_core import factory_production_report as fpr
        self.fpr = fpr
        # 真實踩到的那串（googleapiclient HttpError 的 str()）
        self.err = RuntimeError(
            '<HttpError 403 when requesting https://www.googleapis.com/drive/v3/files?'
            'q=name+contains+%27x%27&alt=json returned "User rate limit exceeded.". '
            'Details: "[{\'message\': \'User rate limit exceeded.\'}]">')

    def _run(self, hour):
        with mock.patch.object(self.fpr, "_load_latest_sheet", side_effect=self.err):
            return self.fpr._daily_production_report(datetime.datetime(2026, 8, 10, hour, 13))

    def test_transient_failure_before_chase_hour_is_silent(self):
        # 下一輪 30 分鐘後就會再跑，暫時性失敗不該佔用四個人的信箱
        self.assertEqual(self._run(9), "(無新發現)")

    def test_still_failing_after_chase_hour_raises_for_the_alerting_path(self):
        # 整個上午都讀不到 → raise，dispatcher 記 last_error，告警通知維運
        with self.assertRaises(RuntimeError) as ctx:
            self._run(11)
        msg = str(ctx.exception)
        self.assertIn("讀不到生管的生產日報進度表", msg)
        self.assertIn("403 User rate limit exceeded.", msg)
        # 摘要過的訊息不該再帶整段 API URL
        self.assertNotIn("googleapis.com", msg)

    def test_short_error_falls_back_to_type_and_message(self):
        out = self.fpr._short_error(ValueError("欄位對不上"))
        self.assertEqual(out, "ValueError: 欄位對不上")


class DailyReportErpSectionTests(unittest.TestCase):
    """ERP 交期核對段：兩邊對不上要標出來，查不到也不能讓整封信掛掉。"""

    def setUp(self):
        from agent_core import factory_production_report as fpr
        self.fpr = fpr
        self.focus = [("2026-05-29", "RICHTER", "JFC26050-1-1", "FE2303", 40.0)]

    def test_flags_erp_closed_but_report_still_owing(self):
        facts = {"JFC26050": {"交期": "2026-05-29", "數量": 600, "狀態": "完工"}}
        with mock.patch.object(self.fpr, "_erp_order_facts", return_value=facts):
            out = self.fpr._erp_delivery_check_section(self.focus)
        self.assertIn("| JFC26050-1-1 | RICHTER |", out)
        self.assertIn("ERP 狀態已是「完工」，日報卻還有 40 雙未完", out)

    def test_flags_date_mismatch(self):
        facts = {"JFC26050": {"交期": "2026-06-30", "數量": 600, "狀態": "生效"}}
        with mock.patch.object(self.fpr, "_erp_order_facts", return_value=facts):
            out = self.fpr._erp_delivery_check_section(self.focus)
        self.assertIn("日報希望出貨日 2026-05-29、ERP 客戶交期 2026-06-30", out)

    def test_erp_failure_does_not_break_the_mail(self):
        with mock.patch.object(self.fpr, "_erp_order_facts",
                                        side_effect=RuntimeError("mirror gone")):
            out = self.fpr._erp_delivery_check_section(self.focus)
        self.assertIn("查 ERP 鏡像失敗", out)
        self.assertIn("本段略過", out)


class CustomerOrderPosTests(unittest.TestCase):
    """read_customer_order_pos 的 PO 分析表解析（純 bytes，免 Drive）。"""

    def test_po_qty_from_xlsx_bytes_groups_by_po(self):
        from agent_core.factory_production_report import _po_qty_from_xlsx_bytes
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "工作表2"
        ws.append(["Bestellung", "Datum", "Groupname", "Artikel", "*",
                   "Groupname", "Farbe", "Verlauf", "Menge"])
        for r in [("20262602036", "x", "g", "63L1073044", "", "g", "atlantic", "25-30", 146),
                  ("20262602036", "x", "g", "63L1073044", "", "g", "fuchsia", "25-30", 154),
                  ("20262602057", "x", "g", "63L1033038", "", "g", "black", "26-30", 8058)]:
            ws.append(list(r))
        buf = io.BytesIO()
        wb.save(buf)
        qty = _po_qty_from_xlsx_bytes(buf.getvalue())
        self.assertEqual(qty.get("20262602036"), 300.0)   # 146 + 154
        self.assertEqual(qty.get("20262602057"), 8058.0)
        self.assertEqual(len(qty), 2)


if __name__ == "__main__":
    unittest.main()
