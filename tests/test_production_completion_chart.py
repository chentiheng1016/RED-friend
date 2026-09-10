"""chart_production_completion（各客戶訂單完工達成率橫條圖）的測試。

涵蓋：_compute_completion 解析合成生管表算出達成率/排序/客戶篩選、達成率分色門檻、
端到端建出正確的 barh chart spec（紅橘綠分色 + 條尾標註 + 100% 目標線）、SAFE tier
（免 +確認）、無資料時回可讀提示。

隔離：用 openpyxl 在記憶體合成一張「生管日報」版型的 workbook；patch Drive 取檔
（_find_latest_progress_file / _download_xlsx_bytes）與 chart_export.generate_chart
（攔截 spec、不真畫圖/不連 Telegram）。不 hardcode /Users/... 路徑、不連網。
"""
import io
import json
import unittest
from unittest import mock

import openpyxl

from agent_core import factory_production_report as fpr


def _build_progress_xlsx() -> bytes:
    """合成一張最小但合版型的『生管日報進度表』：分頁名 '14'，
    欄＝客戶/指令/雙數/累計/欠數/Packing(日計)。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "14"
    ws.append(["客戶", "指令", "雙數", None, None, "Packing"])   # header
    ws.append([None, None, None, "累計", "欠數", "日計"])          # sub-header
    ws.append(["DECA.26", "WO1", 12000, 7020, 4980, 100])
    ws.append(["JALAS", "WO2", 7020, 7020, 0, 50])
    ws.append(["RICHTER", "WO3", 5000, 1650, 3350, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ComputeCompletionTests(unittest.TestCase):
    def setUp(self):
        self.xlsx = _build_progress_xlsx()

    def test_completion_math_and_sort(self):
        rows = fpr._compute_completion(self.xlsx)
        # 依達成率升冪（落後的在前）
        self.assertEqual([r["customer"] for r in rows], ["RICHTER", "DECA.26", "JALAS"])
        by = {r["customer"]: r for r in rows}
        self.assertAlmostEqual(by["DECA.26"]["pct"], 58.5)
        self.assertAlmostEqual(by["RICHTER"]["pct"], 33.0)
        self.assertAlmostEqual(by["JALAS"]["pct"], 100.0)
        self.assertEqual(by["DECA.26"]["remaining"], 4980)
        self.assertEqual(by["JALAS"]["remaining"], 0)

    def test_customer_filter(self):
        rows = fpr._compute_completion(self.xlsx, customer="deca")
        self.assertEqual([r["customer"] for r in rows], ["DECA.26"])

    def test_empty_or_unparseable_returns_list(self):
        self.assertEqual(fpr._compute_completion(b"not a real xlsx"), [])
        # workbook 沒有「日計分頁」→ []
        wb = openpyxl.Workbook()
        wb.active.title = "封面"
        buf = io.BytesIO()
        wb.save(buf)
        self.assertEqual(fpr._compute_completion(buf.getvalue()), [])


class CompletionColorTests(unittest.TestCase):
    def test_thresholds(self):
        self.assertEqual(fpr._completion_color(0), "#E74C3C")     # 紅 <50
        self.assertEqual(fpr._completion_color(49.9), "#E74C3C")
        self.assertEqual(fpr._completion_color(50), "#E67E22")    # 橘 50–80
        self.assertEqual(fpr._completion_color(79.9), "#E67E22")
        self.assertEqual(fpr._completion_color(80), "#2ECC71")    # 綠 ≥80
        self.assertEqual(fpr._completion_color(100), "#2ECC71")


def _reset_ps_cache():
    """_load_latest_sheet(留空) 走 production_schedule 的 5 分鐘共用快取 ——
    測試間要重置，否則上一個測試的 workbook 會被下一個測試撿到。"""
    from agent_core import production_schedule as ps
    ps._CACHE.update(bytes=None, name="", modified="", ts=0.0, key="", schedule=None)


class ChartProductionCompletionTests(unittest.TestCase):
    def setUp(self):
        self.xlsx = _build_progress_xlsx()
        _reset_ps_cache()

    def tearDown(self):
        _reset_ps_cache()

    def _run(self, deliver=False):
        captured = {}

        def fake_gen(spec_json, filename="", deliver=True, chat_id=""):
            captured["spec"] = json.loads(spec_json)
            captured["filename"] = filename
            captured["deliver"] = deliver
            from agent_core.tool_result import ToolResult
            return ToolResult.success("✅ 已畫出圖表", data={"path": "/x.png"}, artifacts=["/x.png"])

        meta = {"name": "6月份生產日報進度表06-14.xlsx", "modifiedTime": "2026-06-14T03:00:00Z"}
        with mock.patch("agent_core.chart_export.generate_chart", side_effect=fake_gen), \
             mock.patch.object(fpr, "_find_latest_progress_file",
                               return_value=("FID", meta["name"])), \
             mock.patch.object(fpr, "_download_xlsx_bytes", return_value=(self.xlsx, meta)):
            out = fpr.chart_production_completion(deliver=deliver)
        return out, captured

    def test_builds_barh_spec(self):
        out, cap = self._run()
        spec = cap["spec"]
        self.assertEqual(spec["type"], "barh")
        self.assertEqual(spec["reference_line"], 100)
        self.assertEqual(spec["categories"], ["RICHTER", "DECA.26", "JALAS"])
        s0 = spec["series"][0]
        self.assertEqual(s0["values"], [33.0, 58.5, 100.0])
        self.assertEqual(s0["colors"], ["#E74C3C", "#E67E22", "#2ECC71"])  # 紅/橘/綠

    def test_annotations_have_counts_and_status(self):
        _out, cap = self._run()
        anns = cap["spec"]["series"][0]["annotations"]
        self.assertIn("33% (1,650/5,000)", anns[0])
        self.assertIn("欠 3,350", anns[0])
        self.assertIn("100% (7,020/7,020)", anns[2])
        self.assertIn("已結案", anns[2])

    def test_deliver_flag_passthrough(self):
        _out, cap = self._run(deliver=False)
        self.assertFalse(cap["deliver"])

    def test_no_data_returns_readable_message(self):
        wb = openpyxl.Workbook()
        wb.active.title = "封面"   # 無日計分頁
        buf = io.BytesIO()
        wb.save(buf)
        meta = {"name": "x.xlsx", "modifiedTime": "2026-06-14T00:00:00Z"}
        with mock.patch.object(fpr, "_find_latest_progress_file", return_value=("FID", "x.xlsx")), \
             mock.patch.object(fpr, "_download_xlsx_bytes", return_value=(buf.getvalue(), meta)):
            out = fpr.chart_production_completion(deliver=False)
        self.assertIn("算不出", str(out))


class TierTests(unittest.TestCase):
    def test_safe_tier_no_confirmation(self):
        from agent_core.tg_auth import is_sensitive
        from agent_core.tool_tiers import get_tier, TIER_SAFE
        self.assertFalse(is_sensitive("chart_production_completion"))
        self.assertEqual(get_tier("chart_production_completion"), TIER_SAFE)

    def test_registered(self):
        from agent_core import tool_registry_catalog as cat
        names = {getattr(f, "__name__", "") for f in cat.BASE_BUILTIN_TOOLS}
        self.assertIn("chart_production_completion", names)


if __name__ == "__main__":
    unittest.main()
