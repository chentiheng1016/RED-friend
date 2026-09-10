from __future__ import annotations

import csv
import os
import tempfile
import unittest
from unittest import mock


def _make_bom_xlsx(path: str, sku: str, materials: list[tuple[str, str, str, str]]) -> None:
    """Build a minimal Pricing BOM xlsx that mirrors the real layout.

    materials = list of (description, material_code, vendor_name, unit_price)
    Header row is at R15 to match the production templates.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sku
    ws.cell(row=1, column=1, value="Pricing BOM")
    ws.cell(row=1, column=2, value="2026-04-03")
    ws.cell(row=3, column=1, value=f"2PEAK {sku} - upper lasted")
    headers = [
        "Material Description", "Material Code", "Vendor Name", "Vendor Code",
        "Consumption", "Waste %", "Total Consumption", "Unit",
        "Unit Price €", "EURO",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=15, column=col, value=h)
    for offset, (desc, code, vendor, price) in enumerate(materials):
        r = 16 + offset
        ws.cell(row=r, column=1, value=desc)
        ws.cell(row=r, column=2, value=code)
        ws.cell(row=r, column=3, value=vendor)
        ws.cell(row=r, column=4, value="VC001")
        ws.cell(row=r, column=5, value=0.1)
        ws.cell(row=r, column=6, value=25)
        ws.cell(row=r, column=7, value=0.125)
        ws.cell(row=r, column=8, value="M2")
        ws.cell(row=r, column=9, value=float(price))
        ws.cell(row=r, column=10, value=float(price) * 0.125)
    wb.save(path)


class BomHistoryTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmpdir.cleanup)
        self.tmpdir = self._tmpdir.name

    def _patch_bom_dir(self):
        from agent_core.agents.orange_sales import bom
        bom_dir = os.path.join(self.tmpdir, "bom_history")
        os.makedirs(bom_dir, exist_ok=True)
        return [
            mock.patch.object(bom, "BOM_HISTORY_DIR", bom_dir),
            mock.patch.object(bom, "_BOM_CSV", os.path.join(bom_dir, "auto_extracted.csv")),
            mock.patch.object(bom, "_BOM_STATE", os.path.join(bom_dir, "ingested_files.json")),
        ]

    def test_classifier_distinguishes_membrane_from_repellent(self):
        """The Jalas hallucination root cause — these three must never collapse."""
        from agent_core.agents.orange_sales.bom import _classify_material

        self.assertEqual(_classify_material("Sympatex W55 PERF"), "防水膜")
        self.assertEqual(_classify_material("Gore-Tex Extended Comfort"), "防水膜")
        self.assertEqual(_classify_material("CASPER 005 GA IDROREPELLEN+TAC"), "撥水")
        self.assertEqual(_classify_material("DWR FINISHING BLACK"), "撥水")
        # Lining must not be mistaken for membrane even when it contains MTP3D.
        self.assertEqual(_classify_material("DRI-LEX 867 PERF MTP3D/24DT BL"), "襯裡")
        # Microfiber (Huafon family) — the actual material on Jalas BOMs.
        self.assertEqual(_classify_material("MICROFIBER SUEDE 1.8 C0 S2 BIO"), "微纖維")

    def test_filename_alias_maps_2peak_to_jalas(self):
        from agent_core.agents.orange_sales.bom import _infer_customer_from_filename

        self.assertEqual(_infer_customer_from_filename("Pricing BOM 2peak 1155.xlsx"), "Jalas")
        self.assertEqual(_infer_customer_from_filename("ejendals_q2.xlsx"), "Jalas")
        self.assertEqual(_infer_customer_from_filename("Lurchi sample BOM.xlsx"), "Lurchi")
        self.assertEqual(_infer_customer_from_filename("random_file.xlsx"), "")

    def test_parse_bom_xlsx_extracts_full_row(self):
        from agent_core.agents.orange_sales.bom import parse_bom_xlsx

        fp = os.path.join(self.tmpdir, "Pricing BOM 2peak 1155.xlsx")
        _make_bom_xlsx(fp, "1155", [
            ("CASPER 005 GA IDROREPELLEN+TAC", "12006077", "Toung Far", "1.85"),
            ("MICROFIBER SUEDE 1.8 C0", "13080819", "Huafon", "14.69"),
        ])
        rows = parse_bom_xlsx(fp)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["customer"], "Jalas")
        self.assertEqual(rows[0]["factory_code"], "2PEAK")
        self.assertEqual(rows[0]["sku"], "1155")
        self.assertEqual(rows[0]["category"], "撥水")
        self.assertEqual(rows[0]["unit_price_eur"], 1.85)
        self.assertEqual(rows[1]["category"], "微纖維")

    def test_query_bom_for_missing_membrane_returns_negative_evidence(self):
        """The Jalas case — querying category=防水膜 with customer=Jalas
        must surface the LACK of evidence, not silence."""
        from agent_core.agents.orange_sales import bom

        with self._patch_bom_dir()[0], self._patch_bom_dir()[1], self._patch_bom_dir()[2]:
            fp = os.path.join(self.tmpdir, "Pricing BOM 2peak 1155.xlsx")
            _make_bom_xlsx(fp, "1155", [
                ("CASPER 005 GA IDROREPELLEN+TAC", "12006077", "Toung Far", "1.85"),
            ])
            out = bom.build_bom_history(folder=self.tmpdir)
            self.assertIn("重建完成", out)

            result = bom.query_bom(customer="Jalas", category="防水膜")
            self.assertIn("查到 0 筆", result)
            self.assertIn("沒有任何屬於『防水膜』類", result)
            self.assertIn("『沒有用』的直接證據", result)

    def test_query_bom_returns_actual_repellent_rows_for_jalas(self):
        from agent_core.agents.orange_sales import bom

        with self._patch_bom_dir()[0], self._patch_bom_dir()[1], self._patch_bom_dir()[2]:
            fp = os.path.join(self.tmpdir, "Pricing BOM 2peak 1155.xlsx")
            _make_bom_xlsx(fp, "1155", [
                ("CASPER 005 GA IDROREPELLEN+TAC", "12006077", "Toung Far", "1.85"),
                ("MICROFIBER SUEDE 1.8 C0", "13080819", "Huafon", "14.69"),
            ])
            bom.build_bom_history(folder=self.tmpdir)

            result = bom.query_bom(customer="Jalas", category="撥水")
            self.assertIn("查到 1 筆", result)
            self.assertIn("IDROREPELLEN", result)
            self.assertIn("Pricing BOM 2peak 1155.xlsx", result)

    def test_sync_bom_from_drive_downloads_and_rebuilds(self):
        """sync_bom_from_drive() must download every BOM hit then chain into
        build_bom_history so query_bom sees the new rows in the same call."""
        from agent_core.agents.orange_sales import bom

        # Prebuild one xlsx in a temp dir, then mock the Drive service to
        # "return" that file as if it lived on Drive.
        source = os.path.join(self.tmpdir, "source.xlsx")
        _make_bom_xlsx(source, "1155", [
            ("Sympatex W55 PERF", "MEMB001", "Sympatex GmbH", "12.50"),
        ])
        with open(source, "rb") as f:
            xlsx_bytes = f.read()

        fake_files = [{
            "id": "fake-id-1",
            "name": "Pricing BOM 2peak demo.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "modifiedTime": "2026-05-19T10:00:00.000Z",
        }]

        class _FakeFiles:
            def list(self, **kwargs):
                class _Req:
                    def execute(self_inner):
                        return {"files": fake_files}
                return _Req()

            def get_media(self, fileId, supportsAllDrives=False):
                class _Req:
                    def execute(self_inner):
                        return xlsx_bytes
                return _Req()

        class _FakeService:
            def files(self):
                return _FakeFiles()

        patches = self._patch_bom_dir()
        cache_dir = os.path.join(self.tmpdir, "bom_history", "_drive_cache")
        with patches[0], patches[1], patches[2], \
             mock.patch.object(bom, "_BOM_DRIVE_CACHE", cache_dir), \
             mock.patch(
                 "agent_core.google_auth.get_service",
                 return_value=_FakeService(),
             ):
            out = bom.sync_bom_from_drive()

            self.assertIn("下載 1", out)
            self.assertIn("重建完成", out)

            # The membrane row must now be queryable.
            result = bom.query_bom(category="防水膜")
            self.assertIn("Sympatex W55", result)

    def test_query_bom_without_csv_directs_to_build(self):
        from agent_core.agents.orange_sales import bom

        with self._patch_bom_dir()[0], self._patch_bom_dir()[1], self._patch_bom_dir()[2]:
            out = bom.query_bom(customer="Jalas")
            self.assertIn("尚未建立", out)
            self.assertIn("build_bom_history", out)


if __name__ == "__main__":
    unittest.main()
