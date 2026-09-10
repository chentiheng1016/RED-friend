"""doc_images：從規格單 PDF 與客人 Excel 母表挖產品圖。

涵蓋：PDF 內嵌圖抽取、logo/icon 濾除、跨頁重複濾除、向量頁 render fallback、
頁範圍/張數上限；Excel 儲存格內嵌圖 + **錨點列標籤**（把圖對回款號的唯一依據）；
員工通道工具的路徑閘（只讀 Telegram 上傳目錄）。

隔離：PDF 與輸出都在 tempdir；RED_TELEGRAM_UPLOAD_DIR / DOC_IMAGES_DIR 在
setUp 換掉、tearDown 還原（不碰主 checkout 的 var/）。不 hardcode /Users 路徑。
"""
import os
import shutil
import tempfile
import unittest
from unittest import mock

from agent_core import doc_images


def _png(path: str, size: tuple[int, int], color: tuple[int, int, int]) -> str:
    from PIL import Image
    Image.new("RGB", size, color).save(path)
    return path


def _make_spec_pdf(dirpath: str, name: str = "2001.pdf") -> str:
    """兩頁的假規格單：
      p1 = 頁首 logo + 產品照 + 小 icon
      p2 = 同一個頁首 logo（跨頁重複）+ 純向量繪圖（沒有內嵌圖）
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    photo = _png(os.path.join(dirpath, "_photo.png"), (600, 400), (120, 160, 200))
    logo = _png(os.path.join(dirpath, "_logo.png"), (200, 160), (250, 250, 250))
    icon = _png(os.path.join(dirpath, "_icon.png"), (40, 40), (10, 200, 10))

    pdf_path = os.path.join(dirpath, name)
    c = canvas.Canvas(pdf_path, pagesize=A4)
    w, h = A4
    c.drawImage(logo, 40, h - 80, width=140, height=42)
    c.drawImage(photo, 60, h - 400, width=300, height=200)
    c.drawImage(icon, 400, h - 120, width=28, height=28)
    c.showPage()
    c.drawImage(logo, 40, h - 80, width=140, height=42)
    c.rect(120, h - 380, 260, 180)
    c.showPage()
    c.save()
    return pdf_path


def _make_master_xlsx(dirpath: str, name: str = "master.xlsx",
                      gap_row: bool = False) -> str:
    """假的客人母表：A1 頁首 logo、E1 小 icon、每列一張產品照貼在 D 欄。

    ``gap_row=True`` 時中間插一列「有資料但 remarks 放文字、沒有圖」的品項
    —— 2026-08-17 UserAng 母表第 11 列（``Davos | 7904 | Unspecified``）就是
    這型，抽圖工具看不到它，整列在下游消失。
    """
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage

    logo = _png(os.path.join(dirpath, "_logo.png"), (240, 60), (250, 250, 250))
    icon = _png(os.path.join(dirpath, "_icon.png"), (24, 24), (10, 200, 10))
    photos = [_png(os.path.join(dirpath, f"_ph{i}.png"), (500, 340), c)
              for i, c in enumerate([(150, 140, 190), (110, 120, 90), (80, 90, 110)], 1)]

    rows = [("2001A", "Magic forest viola", ""),
            ("2001B", "2001-2672-6302", "")]
    if gap_row:
        rows.append(("7904", "Unspecified", "Davos winter style"))   # 這列沒有圖
    rows.append(("2004-Motif girl", "Lotus (15-1906 TPX)", ""))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AW27"
    ws.append(["group", "article number", "color number", "remarks", "Target date"])
    photo_anchors: list[str] = []
    for offset, (art, color, remarks) in enumerate(rows, start=2):
        ws.append(["Freestyle", art, color, remarks, "2026/9/5"])
        if not remarks:
            photo_anchors.append(f"D{offset}")
    for anchor, src, size in ([("A1", logo, (120, 30)), ("E1", icon, (24, 24))]
                              + [(a, p, (110, 75))
                                 for a, p in zip(photo_anchors, photos)]):
        img = XLImage(src)
        img.width, img.height = size
        img.anchor = anchor
        ws.add_image(img)
    wb.create_sheet("notes")["A1"] = "沒有圖的分頁"
    path = os.path.join(dirpath, name)
    wb.save(path)
    return path


def _make_incell_xlsx(dirpath: str, name: str = "incell.xlsx",
                      float_at: str = "") -> str:
    """「置於儲存格」（rich value）版的母表 —— openpyxl 的 ws._images 看不到這種。

    結構照 2026-08-12 UserAng 那份真檔複製：xl/media/ 有圖、xl/drawings/ 空，
    儲存格帶 vm="N" 指向 metadata.xml → richData 的對應鏈。用 zip 後製，因為
    openpyxl 不會寫這種格式（它只寫 drawing anchor）。
    """
    import shutil
    import zipfile

    import openpyxl

    base = os.path.join(dirpath, "_base.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AW27"
    ws.append(["group", "article number", "color number", "remarks", "Target date"])
    for art, color in (("2001A", "Magic forest viola"),
                       ("2001B", "2001-2672-6302"),
                       ("2004", "atlantic - neongreen")):
        ws.append(["Freestyle", art, color, None, "2026/9/5"])
    if float_at:
        # ⚠️ 浮動圖必須在 zip 後製**之前**加：openpyxl 存檔會整份重寫、
        # 把它不認得的 richData part 丟掉（真實副作用：別拿 openpyxl 改寫
        # 客人的母表，會弄丟他們「置於儲存格」的圖）。
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(_png(os.path.join(dirpath, "_float.png"), (300, 200), (9, 9, 9)))
        img.anchor = float_at
        ws.add_image(img)
    wb.save(base)

    photos = [_png(os.path.join(dirpath, f"_rv{i}.png"), (500, 340), c)
              for i, c in enumerate([(150, 140, 190), (110, 120, 90)], 1)]
    n = len(photos)
    # rv[i] 第一個 <v> = richValueRel 序號；rel 順序 = media 順序
    rd_rv = "".join(f"<rv s=\"0\"><v>{i}</v><v>5</v></rv>" for i in range(n))
    rels = "".join(
        f'<Relationship Id="rId{i + 1}" Type="http://schemas.openxmlformats.org/'
        f'officeDocument/2006/relationships/image" Target="../media/rv{i + 1}.png"/>'
        for i in range(n))
    rel_list = "".join(f'<rel r:id="rId{i + 1}"/>' for i in range(n))
    fut = "".join('<bk><extLst><ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}">'
                  f'<xlrd:rvb i="{i}"/></ext></extLst></bk>' for i in range(n))
    vmeta = "".join(f'<bk><rc t="1" v="{i}"/></bk>' for i in range(n))
    parts = {
        "xl/metadata.xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
            ' xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<metadataTypes count="1"><metadataType name="XLRICHVALUE"'
            ' minSupportedVersion="120000"/></metadataTypes>'
            f'<futureMetadata name="XLRICHVALUE" count="{n}">{fut}</futureMetadata>'
            f'<valueMetadata count="{n}">{vmeta}</valueMetadata></metadata>',
        "xl/richData/rdrichvalue.xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"'
            f' count="{n}">{rd_rv}</rvData>',
        "xl/richData/richValueRel.xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel"'
            ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'{rel_list}</richValueRels>',
        "xl/richData/_rels/richValueRel.xml.rels":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'{rels}</Relationships>',
    }

    out = os.path.join(dirpath, name)
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.namelist():
            data = src.read(item)
            if item == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                # ⚠️ openpyxl 對 None 的儲存格**整格不寫**（row 裡直接跳過 D2），
                # 所以是「插入」不是「取代」—— 插在同列 E 之前。
                for row, vm in ((2, 1), (3, 2)):
                    cell = f'<c r="D{row}" t="e" vm="{vm}"><v>#VALUE!</v></c>'
                    anchor = f'<c r="E{row}"'
                    assert anchor in xml, f"fixture 壞了：找不到 {anchor}"
                    xml = xml.replace(anchor, cell + anchor, 1)
                data = xml.encode("utf-8")
            dst.writestr(item, data)
        for i, p in enumerate(photos, start=1):
            with open(p, "rb") as f:
                dst.writestr(f"xl/media/rv{i}.png", f.read())
        for part, text in parts.items():
            dst.writestr(part, text)
    shutil.os.remove(base)
    return out


class PdfImagesBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="red_pdfimg_test_")
        self.src = os.path.join(self._tmp, "src")
        self.out = os.path.join(self._tmp, "out")
        os.makedirs(self.src, exist_ok=True)
        self.pdf = _make_spec_pdf(self.src)

    def tearDown(self):
        shutil.rmtree(self._tmp, ignore_errors=True)


class TestExtract(PdfImagesBase):
    def test_keeps_product_photo_drops_logo_and_icon(self):
        items = doc_images.extract_pdf_images(self.pdf, self.out)
        embedded = [i for i in items if i["source"] == "embedded"]
        # 內嵌圖只該留產品照：icon 太小被濾、logo 兩頁都有被當頁首濾掉
        self.assertEqual(len(embedded), 1, embedded)
        self.assertEqual((embedded[0]["width"], embedded[0]["height"]), (600, 400))
        self.assertEqual(embedded[0]["page"], 1)
        self.assertTrue(os.path.isfile(embedded[0]["path"]))
        self.assertTrue(embedded[0]["path"].endswith(".png"))

    def test_vector_only_page_falls_back_to_render(self):
        items = doc_images.extract_pdf_images(self.pdf, self.out)
        rendered = [i for i in items if i["source"] == "rendered"]
        self.assertEqual([i["page"] for i in rendered], [2])
        self.assertTrue(os.path.getsize(rendered[0]["path"]) > 0)

    def test_fallback_render_can_be_disabled(self):
        items = doc_images.extract_pdf_images(self.pdf, self.out, fallback_render=False)
        self.assertEqual([i["source"] for i in items], ["embedded"])

    def test_min_px_filter_can_let_small_images_through(self):
        items = doc_images.extract_pdf_images(self.pdf, self.out, min_px=10,
                                              fallback_render=False)
        sizes = {(i["width"], i["height"]) for i in items}
        self.assertIn((40, 40), sizes)      # icon 這次留下來
        self.assertNotIn((200, 160), sizes)  # logo 仍因跨頁重複被濾

    def test_pages_filter(self):
        items = doc_images.extract_pdf_images(self.pdf, self.out, pages="1")
        self.assertEqual({i["page"] for i in items}, {1})

    def test_max_images_cap(self):
        items = doc_images.extract_pdf_images(self.pdf, self.out, min_px=10,
                                              max_images=1)
        self.assertEqual(len(items), 1)

    def test_output_dir_created_and_named_after_pdf(self):
        nested = os.path.join(self.out, "deep", "deeper")
        items = doc_images.extract_pdf_images(self.pdf, nested)
        self.assertTrue(items)
        for it in items:
            self.assertEqual(os.path.dirname(it["path"]), nested)
            self.assertTrue(os.path.basename(it["path"]).startswith("2001_p"))


class TestExtractXlsx(PdfImagesBase):
    """客人母表（圖貼在儲存格裡）→ 圖 + 該列款號。"""

    def setUp(self):
        super().setUp()
        self.xlsx = _make_master_xlsx(self.src)

    def test_maps_each_image_to_its_row_and_label(self):
        items = doc_images.extract_xlsx_images(self.xlsx, self.out)
        by_cell = {i["cell"]: i for i in items}
        # 三張產品照各自對到自己那列的款號 —— 這是把舊表的圖搬進新表的唯一依據
        self.assertIn("2001A", by_cell["D2"]["label"])
        self.assertIn("2001B", by_cell["D3"]["label"])
        self.assertIn("2004-Motif girl", by_cell["D4"]["label"])
        self.assertEqual(by_cell["D2"]["row"], 2)
        self.assertEqual(by_cell["D2"]["sheet"], "AW27")
        self.assertEqual((by_cell["D2"]["width"], by_cell["D2"]["height"]), (500, 340))
        for it in items:
            self.assertTrue(os.path.isfile(it["path"]))

    def test_row_label_skips_the_image_column_itself(self):
        items = doc_images.extract_xlsx_images(self.xlsx, self.out)
        d2 = next(i for i in items if i["cell"] == "D2")
        self.assertNotIn("remarks", d2["label"])   # D 欄是圖欄，標籤不該取它

    def test_tiny_icon_filtered_but_logo_reported_with_its_cell(self):
        items = doc_images.extract_xlsx_images(self.xlsx, self.out)
        cells = {i["cell"] for i in items}
        self.assertNotIn("E1", cells)              # 24px icon 被濾掉
        # logo 留著但錨點是 A1（表頭列）——照實回報，由 LLM 判斷要不要用
        self.assertIn("A1", cells)

    def test_sheet_filter(self):
        self.assertEqual(doc_images.extract_xlsx_images(self.xlsx, self.out,
                                                        sheet="notes"), [])
        self.assertTrue(doc_images.extract_xlsx_images(self.xlsx, self.out,
                                                       sheet="AW27"))

    def test_max_images_cap(self):
        items = doc_images.extract_xlsx_images(self.xlsx, self.out, max_images=2)
        self.assertEqual(len(items), 2)

    def test_sorted_by_row(self):
        items = doc_images.extract_xlsx_images(self.xlsx, self.out)
        rows = [i["row"] for i in items]
        self.assertEqual(rows, sorted(rows))

    def test_in_cell_rich_value_images_are_found(self):
        """「置於儲存格」的圖：openpyxl 的 ws._images 看不到，要走 richData 那條。

        2026-08-12 實測踩到：客人真的 master 表 18 張圖全是這型，只靠
        drawing anchor 會回 0 張（合成 fixture 全過、真檔全滅）。
        """
        import openpyxl
        path = _make_incell_xlsx(self.src)
        # 先確認 openpyxl 真的看不到（不然這條測試就白測了）
        wb = openpyxl.load_workbook(path)
        self.assertEqual(list(getattr(wb["AW27"], "_images", [])), [])
        wb.close()

        items = doc_images.extract_xlsx_images(path, self.out)
        self.assertEqual(len(items), 2, items)
        by_cell = {i["cell"]: i for i in items}
        self.assertEqual(sorted(by_cell), ["D2", "D3"])
        self.assertIn("2001A", by_cell["D2"]["label"])
        self.assertIn("2001B", by_cell["D3"]["label"])
        self.assertEqual((by_cell["D2"]["width"], by_cell["D2"]["height"]), (500, 340))
        self.assertEqual(by_cell["D2"]["source"], "xlsx-incell")
        for it in items:
            self.assertTrue(os.path.isfile(it["path"]))

    def test_in_cell_and_drawing_images_coexist(self):
        """同一份表兩種圖都有時要一起回報、依列排序。"""
        path = _make_incell_xlsx(self.src, name="mixed.xlsx", float_at="D4")
        items = doc_images.extract_xlsx_images(path, self.out)
        self.assertEqual([i["cell"] for i in items], ["D2", "D3", "D4"])
        self.assertEqual([i["source"] for i in items],
                         ["xlsx-incell", "xlsx-incell", "xlsx"])

    def test_broken_rich_data_degrades_instead_of_raising(self):
        """richData 壞掉/缺件時整組放棄，不能讓整顆工具炸掉。"""
        import zipfile
        path = _make_incell_xlsx(self.src, name="broken.xlsx")
        broken = os.path.join(self.src, "broken2.xlsx")
        with zipfile.ZipFile(path) as src, \
                zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.namelist():
                data = src.read(item)
                if item == "xl/richData/rdrichvalue.xml":
                    data = b"<not-xml"
                dst.writestr(item, data)
        self.assertEqual(doc_images.extract_xlsx_images(broken, self.out), [])

    def test_workbook_without_images_returns_empty(self):
        import openpyxl
        plain = os.path.join(self.src, "plain.xlsx")
        wb = openpyxl.Workbook()
        wb.active["A1"] = "no drawings"
        wb.save(plain)
        self.assertEqual(doc_images.extract_xlsx_images(plain, self.out), [])


class TestUploadedTool(PdfImagesBase):
    """員工通道版：路徑限縮在 Telegram 上傳目錄、產出分艙 per-color。"""

    def setUp(self):
        super().setUp()
        self.upload_root = os.path.join(self._tmp, "uploads")
        os.makedirs(self.upload_root, exist_ok=True)
        self._env = os.environ.get("RED_TELEGRAM_UPLOAD_DIR")
        os.environ["RED_TELEGRAM_UPLOAD_DIR"] = self.upload_root
        self._patch = mock.patch.object(doc_images, "DOC_IMAGES_DIR",
                                        os.path.join(self._tmp, "doc_images"))
        self._patch.start()
        self.uploaded = os.path.join(self.upload_root, "2026-08-12", "2001.pdf")
        os.makedirs(os.path.dirname(self.uploaded), exist_ok=True)
        shutil.copy(self.pdf, self.uploaded)
        self.uploaded_xlsx = os.path.join(os.path.dirname(self.uploaded), "master.xlsx")
        shutil.copy(_make_master_xlsx(self.src), self.uploaded_xlsx)

    def tearDown(self):
        self._patch.stop()
        if self._env is None:
            os.environ.pop("RED_TELEGRAM_UPLOAD_DIR", None)
        else:
            os.environ["RED_TELEGRAM_UPLOAD_DIR"] = self._env
        super().tearDown()

    def test_extracts_and_reports_paths(self):
        out = doc_images.extract_uploaded_pdf_images(self.uploaded)
        self.assertIn("抽出", out)
        paths = [ln.split("（")[0].strip(" •") for ln in out.splitlines()
                 if ln.strip().startswith("•")]
        self.assertTrue(paths)
        # macOS 的 /var 是 symlink，output_root() 會 realpath —— 比對前先對齊
        root = os.path.realpath(os.path.join(self._tmp, "doc_images"))
        for p in paths:
            self.assertTrue(os.path.isfile(p), p)
            self.assertTrue(p.startswith(root), p)
        self.assertIn('{"image"', out)   # 要教 LLM 怎麼把圖填進 export_report

    def test_rejects_path_outside_upload_root(self):
        out = doc_images.extract_uploaded_pdf_images(self.pdf)  # 在 tmp/src，不在上傳區
        self.assertTrue(out.startswith("錯誤："), out)
        self.assertIn("上傳目錄", out)

    def test_rejects_traversal_out_of_upload_root(self):
        sneaky = os.path.join(self.upload_root, "..", "src", "2001.pdf")
        out = doc_images.extract_uploaded_pdf_images(sneaky)
        self.assertTrue(out.startswith("錯誤："), out)

    def test_rejects_non_pdf(self):
        other = os.path.join(self.upload_root, "note.txt")
        with open(other, "w") as f:
            f.write("x")
        out = doc_images.extract_uploaded_pdf_images(other)
        self.assertIn("只吃 PDF", out)

    def test_missing_file_is_reported_not_raised(self):
        out = doc_images.extract_uploaded_pdf_images(
            os.path.join(self.upload_root, "nope.pdf"))
        self.assertIn("找不到", out)

    def test_dept_color_partitions_output(self):
        with mock.patch("agent_core.dept_tool_scope.dept_scope_color",
                        return_value="orange"):
            out = doc_images.extract_uploaded_pdf_images(self.uploaded)
        self.assertIn(os.path.join("doc_images", "dept", "orange"), out)

    def test_empty_pdf_says_so_instead_of_faking_paths(self):
        blank = os.path.join(self.upload_root, "blank.pdf")
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(blank)
        c.drawString(100, 700, "just text")
        c.showPage()
        c.save()
        out = doc_images.extract_uploaded_pdf_images(blank, pages="1")
        # 純文字頁：render fallback 會裁成整頁，抽得到才正常；重點是不能謊報
        self.assertTrue(out.startswith("✅") or out.startswith("⚠️"), out)

    # ── 客人 Excel 母表 ────────────────────────────────────────────
    def test_excel_tool_reports_row_labels(self):
        out = doc_images.extract_uploaded_excel_images(self.uploaded_xlsx)
        self.assertIn("抽出", out)
        self.assertIn("AW27!D2", out)
        self.assertIn("2001A", out)         # 沒有列標籤就對不回款號
        self.assertIn("照順序硬配", out)     # 要提醒 LLM 別亂配
        self.assertIn('{"image"', out)

    # 2026-08-17 UserAng 案：抽圖工具只回報「有圖的」儲存格，母表裡 remarks
    # 放文字沒放圖的那列在下游整列消失（19 列產出 18 列）。現在要一起報出來。
    def test_excel_tool_reports_rows_that_have_data_but_no_image(self):
        gapped = os.path.join(self.upload_root, "gapped.xlsx")
        shutil.copy(_make_master_xlsx(self.src, name="g.xlsx", gap_row=True), gapped)
        out = doc_images.extract_uploaded_excel_images(gapped)
        self.assertIn("有資料但沒有圖", out)
        self.assertIn("第 4 列", out)          # 中間那列（表頭 1、資料 2~5）
        self.assertIn("7904", out)             # 要講是哪一款，不能只說「有一列」
        self.assertIn("Unspecified", out)
        self.assertIn("read_uploaded_table", out)   # 指路到「讀整張表」那顆

    def test_excel_tool_without_gaps_does_not_invent_missing_rows(self):
        out = doc_images.extract_uploaded_excel_images(self.uploaded_xlsx)
        self.assertNotIn("有資料但沒有圖", out)

    def test_rows_without_images_ignores_header_and_footer(self):
        """表頭/表尾本來就沒有圖 —— 只看第一張到最後一張圖之間的列，否則洗版。"""
        gapped = _make_master_xlsx(self.src, name="g2.xlsx", gap_row=True)
        items = doc_images.extract_xlsx_images(gapped, self.out)
        gaps = doc_images.rows_without_images(gapped, items)
        self.assertEqual([(s, r) for s, r, _label in gaps], [("AW27", 4)])

    def test_rows_without_images_empty_when_sheet_has_no_images(self):
        self.assertEqual(doc_images.rows_without_images(self.uploaded_xlsx, []), [])

    def test_excel_tool_rejects_pdf_and_xls(self):
        self.assertIn("只吃 Excel",
                      doc_images.extract_uploaded_excel_images(self.uploaded))
        legacy = os.path.join(self.upload_root, "old.xls")
        with open(legacy, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0")
        self.assertIn("只吃 Excel",
                      doc_images.extract_uploaded_excel_images(legacy))

    def test_excel_tool_rejects_path_outside_upload_root(self):
        outside = _make_master_xlsx(self.src, name="outside.xlsx")
        out = doc_images.extract_uploaded_excel_images(outside)
        self.assertTrue(out.startswith("錯誤："), out)
        self.assertIn("上傳目錄", out)

    def test_excel_tool_says_so_when_sheet_has_no_images(self):
        out = doc_images.extract_uploaded_excel_images(self.uploaded_xlsx,
                                                       sheet="notes")
        self.assertIn("沒抽到", out)
        self.assertIn("不要自己編圖檔路徑", out)

    def test_excel_tool_partitions_output_by_color(self):
        with mock.patch("agent_core.dept_tool_scope.dept_scope_color",
                        return_value="orange"):
            out = doc_images.extract_uploaded_excel_images(self.uploaded_xlsx)
        self.assertIn(os.path.join("doc_images", "dept", "orange"), out)


class TestEmployeeFlowIntegration(PdfImagesBase):
    """UserAng 案的整條路：員工上傳規格單 PDF／客人母表 → 抽圖 → 圖嵌進新的
    tracking log Excel。

    這條是兩個模組的接縫：doc_images 產出的路徑必須落在 doc_export 嵌圖白名單
    的**同一個** per-color 目錄裡，任一邊改了 base dir 就會靜默變成「圖都沒進去」。
    """

    def setUp(self):
        super().setUp()
        from agent_core import doc_export
        self.doc_export = doc_export
        self.data_dir = os.path.join(self._tmp, "var_data")
        self.upload_root = os.path.join(self._tmp, "uploads", "2026-08-12")
        os.makedirs(self.upload_root, exist_ok=True)
        self.uploaded = os.path.join(self.upload_root, "2001.pdf")
        shutil.copy(self.pdf, self.uploaded)
        self.uploaded_xlsx = os.path.join(self.upload_root, "master.xlsx")
        shutil.copy(_make_master_xlsx(self.src), self.uploaded_xlsx)

        self._env = os.environ.get("RED_TELEGRAM_UPLOAD_DIR")
        os.environ["RED_TELEGRAM_UPLOAD_DIR"] = os.path.dirname(self.upload_root)
        self._patches = [
            mock.patch.object(doc_images, "DOC_IMAGES_DIR",
                              os.path.join(self.data_dir, "doc_images")),
            mock.patch("agent_core.logging_and_paths.DATA_DIR", self.data_dir),
            mock.patch.object(doc_export, "EXPORTS_DIR",
                              os.path.join(self._tmp, "exports")),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in reversed(self._patches):
            p.stop()
        if self._env is None:
            os.environ.pop("RED_TELEGRAM_UPLOAD_DIR", None)
        else:
            os.environ["RED_TELEGRAM_UPLOAD_DIR"] = self._env
        super().tearDown()

    def _as_orange(self):
        from agent_core.agents.middleware import AgentRequest, agent_request_context
        from agent_core.agents.permission_matrix import Agent
        return agent_request_context(AgentRequest(
            caller=Agent.ORANGE, target=Agent.ORANGE,
            intent="telegram.freeform.test", payload={}))

    def test_extracted_image_survives_into_the_excel(self):
        import json
        import openpyxl
        with self._as_orange():
            out = doc_images.extract_uploaded_pdf_images(self.uploaded, pages="1")
            path = [ln for ln in out.splitlines() if ln.strip().startswith("•")][0]
            path = path.split("（")[0].strip(" •")
            self.assertIn(os.path.join("doc_images", "dept", "orange"), path)

            content = {"title": "AW27 tracking log",
                       "blocks": [{"type": "table", "title": "samples",
                                   "columns": ["article", "remarks"],
                                   "rows": [["2001A", {"image": path}]]}]}
            res = self.doc_export.export_report(json.dumps(content, ensure_ascii=False),
                                                formats="excel", deliver=False)
        self.assertTrue(res.ok, res)
        self.assertEqual(res.data["images_embedded"], 1, res.summary)
        self.assertIn(os.path.join("dept", "orange"), res.artifacts[0])
        ws = openpyxl.load_workbook(res.artifacts[0])["samples"]
        self.assertEqual(len(ws._images), 1)

    def test_master_xlsx_images_are_reused_in_the_new_excel(self):
        """客人母表的圖 → 依「該列款號」對回新表 → 嵌進新表的對應列。"""
        import json
        import openpyxl
        with self._as_orange():
            out = doc_images.extract_uploaded_excel_images(self.uploaded_xlsx,
                                                           sheet="AW27")
            # 依款號配對（不是照順序）——這是工具存在的理由
            by_article = {}
            for ln in out.splitlines():
                if not ln.strip().startswith("•"):
                    continue
                path = ln.split("（")[0].strip(" •")
                for art in ("2001A", "2001B", "2004-Motif girl"):
                    if f"| {art} |" in ln or f"：Freestyle | {art}" in ln:
                        by_article[art] = path
            self.assertEqual(set(by_article), {"2001A", "2001B", "2004-Motif girl"},
                             out)

            rows = [[art, {"image": by_article[art]}]
                    for art in ("2004-Motif girl", "2001A")]   # 新表順序不同
            content = {"title": "AW27 tracking log v2",
                       "blocks": [{"type": "table", "title": "samples",
                                   "columns": ["article", "remarks"], "rows": rows}]}
            res = self.doc_export.export_report(json.dumps(content, ensure_ascii=False),
                                                formats="excel", deliver=False)
        self.assertTrue(res.ok, res)
        self.assertEqual(res.data["images_embedded"], 2, res.summary)
        ws = openpyxl.load_workbook(res.artifacts[0])["samples"]
        self.assertEqual(len(ws._images), 2)
        self.assertEqual([ws.cell(r, 1).value for r in (2, 3)],
                         ["2004-Motif girl", "2001A"])

    def test_other_colors_image_is_rejected(self):
        """orange 的 session 不能把 green 目錄裡的圖嵌進自己的檔案。"""
        from PIL import Image
        green_dir = os.path.join(self.data_dir, "doc_images", "dept", "green")
        os.makedirs(green_dir, exist_ok=True)
        theirs = os.path.join(green_dir, "secret.png")
        Image.new("RGB", (300, 200), (0, 0, 0)).save(theirs)
        with self._as_orange():
            ref = self.doc_export._resolve_image(theirs)
        self.assertFalse(ref.ok)
        self.assertIn("不在允許的目錄", ref.note)


class TestToolWiring(unittest.TestCase):
    TOOLS = ("extract_uploaded_pdf_images", "extract_uploaded_excel_images")

    def test_tools_are_safe_tier_and_in_common_dept_scope(self):
        from agent_core.dept_tool_scope import _COMMON_TOOLS
        from agent_core.tool_tiers import TIER_SAFE, get_tier
        # 員工白名單只收 SAFE；這兩件事任一漏掉，工具在員工 session 就靜默消失
        for name in self.TOOLS:
            self.assertIn(name, _COMMON_TOOLS)
            self.assertEqual(get_tier(name), TIER_SAFE, name)

    def test_tools_are_registered_in_catalog(self):
        from agent_core import tool_registry_catalog as cat
        names = {getattr(f, "__name__", "") for f in cat.BASE_BUILTIN_TOOLS}
        for name in self.TOOLS:
            self.assertIn(name, names)

    def test_real_output_root_is_inside_doc_export_whitelist(self):
        """兩邊的 base dir 必須真的對得上。

        整合測試把 DOC_IMAGES_DIR 與 DATA_DIR 各自 patch 成一致值，會掩蓋
        「其中一邊改了 base dir」的漂移 —— 這條用真常數比對，不碰檔案系統。
        """
        from agent_core import doc_export
        from agent_core.agents.middleware import AgentRequest, agent_request_context
        from agent_core.agents.permission_matrix import Agent
        req = AgentRequest(caller=Agent.ORANGE, target=Agent.ORANGE,
                           intent="test.whitelist", payload={})
        with agent_request_context(req):
            roots = doc_export._image_allowed_roots()
        self.assertIn(doc_images.output_root("orange"), roots)


if __name__ == "__main__":
    unittest.main()
