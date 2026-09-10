"""factory_warehouse 建倉 + text-to-SQL 查詢層測試（合成 fixture，免網路／Gemini）。"""
import datetime
import io
import json
import os
import tempfile
import types
import unittest
from unittest import mock

import openpyxl

from agent_core import factory_warehouse as fw


def _sheet(ws, rows):
    """寫成福群進度表版型（表頭 3-4 列、資料第 5 列起）—— 與 _detect_layout 錨點對齊。"""
    ws.cell(row=3, column=4, value="客戶")
    ws.cell(row=3, column=6, value="指令")
    ws.cell(row=3, column=7, value="型體")
    ws.cell(row=3, column=11, value="雙數")
    ws.cell(row=3, column=19, value="Stitching")
    ws.cell(row=3, column=24, value="Injection")
    ws.cell(row=3, column=29, value="Packing")
    ws.cell(row=3, column=32, value="希望\n出貨日")
    ws.cell(row=3, column=34, value="實際\n出貨日")
    ws.cell(row=4, column=19, value="日計\nday")
    ws.cell(row=4, column=24, value="日計\nday")
    ws.cell(row=4, column=26, value="累計\nsum")
    ws.cell(row=4, column=28, value="欠數\nremaining")
    ws.cell(row=4, column=29, value="日計\nday")
    r = 5
    for d in rows:
        ws.cell(row=r, column=4, value=d["customer"])
        ws.cell(row=r, column=6, value=d["wo"])
        ws.cell(row=r, column=7, value=d.get("model", ""))
        ws.cell(row=r, column=11, value=d.get("pairs", 0))
        ws.cell(row=r, column=19, value=d.get("stitch", 0))
        ws.cell(row=r, column=24, value=d.get("inject", 0))
        ws.cell(row=r, column=26, value=d.get("pack_cum", 0))
        ws.cell(row=r, column=28, value=d.get("pack_rem", 0))
        ws.cell(row=r, column=29, value=d.get("pack_day", 0))
        ws.cell(row=r, column=32, value=d.get("want_ship"))
        ws.cell(row=r, column=34, value=d.get("actual_ship"))
        r += 1


def _fixture_bytes():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb.create_sheet("01"), [
        dict(customer="DECA.26", wo="JFC1", model="MH100", pairs=200,
             stitch=100, inject=80, pack_day=70, pack_cum=120, pack_rem=80,
             want_ship=datetime.datetime(2026, 6, 20), actual_ship=None),
        dict(customer="JALAS", wo="JFC2", model="EJ24", pairs=100,
             inject=50, pack_day=100, pack_cum=100, pack_rem=0,
             want_ship=datetime.datetime(2026, 6, 18), actual_ship="OK"),
    ])
    _sheet(wb.create_sheet("02"), [   # 最新一天 → fact_order_state 從這天取
        dict(customer="DECA.26", wo="JFC1", model="MH100", pairs=200,
             stitch=10, inject=5, pack_day=80, pack_cum=200, pack_rem=0,
             want_ship=datetime.datetime(2026, 6, 20), actual_ship=None),   # 完工待出
        dict(customer="JALAS", wo="JFC2", model="EJ24", pairs=100,
             inject=0, pack_day=0, pack_cum=100, pack_rem=0,
             want_ship=datetime.datetime(2026, 6, 18), actual_ship="OK"),   # 已出貨
        dict(customer="LURCHI", wo="JFC3", model="LX9", pairs=300,
             inject=200, pack_day=50, pack_cum=50, pack_rem=250,
             want_ship=datetime.datetime(2026, 6, 30), actual_ship=None),   # 未完
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class BuildWarehouseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "wh.duckdb")
        self.summary = fw.build_warehouse_from_bytes(
            _fixture_bytes(), db_path=self.db,
            report_file_id="FID", report_modified="2026-06-17")

    def tearDown(self):
        for p in (self.db, self.db + ".tmp", self.db + ".wal", self.db + ".tmp.wal"):
            if os.path.exists(p):
                os.remove(p)

    def _ro(self):
        import duckdb
        return duckdb.connect(self.db, read_only=True)

    def test_daily_rows_loaded(self):
        self.assertEqual(self.summary["daily_rows"], 5)   # day01:2 + day02:3
        con = self._ro()
        try:
            self.assertEqual(con.execute("SELECT count(*) FROM fact_production_daily").fetchone()[0], 5)
        finally:
            con.close()

    def test_order_state_completion_and_status(self):
        con = self._ro()
        try:
            rows = {r[0]: r for r in con.execute(
                "SELECT work_order, completion_pct, status, pack_rem_latest "
                "FROM fact_order_state").fetchall()}
        finally:
            con.close()
        self.assertEqual(self.summary["orders"], 3)
        self.assertAlmostEqual(rows["JFC1"][1], 100.0, places=1)
        self.assertEqual(rows["JFC1"][2], "completed")   # rem=0、未標出貨
        self.assertEqual(rows["JFC2"][2], "shipped")     # actual='OK'
        self.assertEqual(rows["JFC3"][2], "open")        # rem=250
        self.assertAlmostEqual(rows["JFC3"][1], 50 / 300 * 100, places=1)
        self.assertEqual(rows["JFC3"][3], 250)

    def test_date_stored_as_iso_text(self):
        con = self._ro()
        try:
            want = con.execute(
                "SELECT want_ship FROM fact_order_state WHERE work_order='JFC3'").fetchone()[0]
        finally:
            con.close()
        self.assertEqual(want, "2026-06-30")   # datetime → 'YYYY-MM-DD' 文字

    def test_atomic_rename_no_tmp_left(self):
        self.assertTrue(os.path.exists(self.db))
        self.assertFalse(os.path.exists(self.db + ".tmp"))

    def test_rebuild_is_idempotent(self):
        again = fw.build_warehouse_from_bytes(
            _fixture_bytes(), db_path=self.db, report_file_id="FID", report_modified="2026-06-17")
        self.assertEqual(again["daily_rows"], 5)   # 全量重建，不累加
        con = self._ro()
        try:
            self.assertEqual(con.execute("SELECT count(*) FROM fact_production_daily").fetchone()[0], 5)
        finally:
            con.close()


class QueryToolTests(unittest.TestCase):
    """text-to-SQL 工具的驗證／執行層（不碰 Gemini，除了一個 mock 端到端）。"""

    def setUp(self):
        import skills.factory_warehouse as sk
        self.sk = sk
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "wh.duckdb")
        fw.build_warehouse_from_bytes(
            _fixture_bytes(), db_path=self.db, report_file_id="FID", report_modified="2026-06-17")
        p = mock.patch.object(sk, "_db_path", lambda: self.db)
        p.start()
        self.addCleanup(p.stop)

    def tearDown(self):
        for p in (self.db, self.db + ".tmp", self.db + ".wal", self.db + ".tmp.wal"):
            if os.path.exists(p):
                os.remove(p)

    def test_select_ok(self):
        out = self.sk.run_warehouse_sql("SELECT count(*) AS n FROM fact_production_daily")
        self.assertIn("5", out)

    def test_open_orders_query(self):
        out = self.sk.run_warehouse_sql(
            "SELECT customer_raw, pack_rem_latest FROM fact_order_state WHERE status='open'")
        self.assertIn("LURCHI", out)
        self.assertIn("250", out)

    def test_rejects_write_statements(self):
        for bad in ("DELETE FROM fact_order_state",
                    "INSERT INTO fact_order_state(work_order) VALUES ('X')",
                    "DROP TABLE fact_order_state",
                    "UPDATE fact_order_state SET status='x'"):
            self.assertIn("拒", self.sk.run_warehouse_sql(bad), msg=bad)

    def test_rejects_file_access(self):
        self.assertIn("拒", self.sk.run_warehouse_sql("SELECT * FROM read_csv('/etc/passwd')"))
        self.assertIn("拒", self.sk.run_warehouse_sql("COPY fact_order_state TO '/tmp/x.csv'"))

    def test_rejects_multistatement_and_non_select(self):
        self.assertIn("拒", self.sk.run_warehouse_sql("SELECT 1; DROP TABLE fact_order_state"))
        self.assertIn("拒", self.sk.run_warehouse_sql("EXPLAIN SELECT 1"))

    def test_readonly_engine_blocks_write_even_if_validation_bypassed(self):
        # 直接走 _run_ro（繞過 _validate_select）→ 引擎層 read_only 仍擋寫
        with self.assertRaises(Exception):
            self.sk._run_ro("INSERT INTO fact_order_state(work_order) VALUES ('X')")

    def test_external_file_access_disabled_at_engine(self):
        # 即使關鍵字驗證被繞過，enable_external_access=false 仍擋讀本機檔
        with self.assertRaises(Exception):
            self.sk._run_ro("SELECT * FROM read_csv_auto('/etc/hosts')")

    def test_result_row_cap(self):
        # 用 range 產 100 列，確認封頂訊息
        out = self.sk.run_warehouse_sql("SELECT * FROM range(100)")
        self.assertIn(f"只顯示前 {self.sk._MAX_ROWS}", out)

    def test_query_factory_warehouse_end_to_end_mocked_gemini(self):
        fake = types.SimpleNamespace(
            text="SELECT customer_raw, completion_pct FROM fact_order_state ORDER BY completion_pct")
        with mock.patch("agent_core.gemini_client._gemini_generate", return_value=fake):
            out = self.sk.query_factory_warehouse("各客戶完工率由低到高")
        self.assertIn("SQL：", out)        # 附上實際 SQL
        self.assertIn("LURCHI", out)        # 完工率最低排最前

    def test_query_rejects_unsafe_generated_sql(self):
        fake = types.SimpleNamespace(text="DROP TABLE fact_order_state")
        with mock.patch("agent_core.gemini_client._gemini_generate", return_value=fake):
            out = self.sk.query_factory_warehouse("刪掉資料")
        self.assertIn("被拒", out)


def _email_parquet(path):
    """合成內部郵件 parquet（欄位對齊 data_lake_internal/emails.parquet）。"""
    import pandas as pd

    def ent(**kw):
        base = {"customers": [], "suppliers": [], "po_numbers": [], "products": [],
                "promised_dates": []}
        base.update(kw)
        return json.dumps(base, ensure_ascii=False)

    rows = [
        # T1：客戶寫 'Decathlon' → 該靠別名接到倉的 'DECA.26'
        dict(thread_id="T1", date="2026-06-01", last_message_date="2026-06-01", sender="a@x.com",
             primary_dept="sales", direction="inbound", state="進行中", summary="Decathlon 訂單確認",
             message_count=2, topic_tags='["PO確認"]',
             entities_json=ent(customers=["Decathlon"], po_numbers=["65L1"], promised_dates=["6/20 出貨"])),
        # T2：LURCHI，2 個 PO
        dict(thread_id="T2", date="2026-06-02", last_message_date="2026-06-03", sender="b@x.com",
             primary_dept="ship", direction="outbound", state="已完成", summary="LURCHI 出貨",
             message_count=1, topic_tags='["交期延誤"]',
             entities_json=ent(customers=["LURCHI"], suppliers=["FU CHUN"], po_numbers=["P1", "P2"])),
        # T3：FU CHUN（非倉客戶）→ 不該連 customer
        dict(thread_id="T3", date="2026-06-04", last_message_date="2026-06-04", sender="c@x.com",
             primary_dept="finance", direction="inbound", state="僅參考", summary="供應商對帳",
             message_count=1, topic_tags='["對帳"]',
             entities_json=ent(customers=["FU CHUN"])),
    ]
    pd.DataFrame(rows).to_parquet(path, index=False)


class EmailLakePhase2Tests(unittest.TestCase):
    """Phase 2：郵件 lake → fact_email_thread + bridge_email_po/customer + dim_customer。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "wh.duckdb")
        self.parquet = os.path.join(self.tmp, "emails.parquet")
        _email_parquet(self.parquet)
        self.summary = fw.build_warehouse_from_bytes(
            _fixture_bytes(), db_path=self.db, report_file_id="FID",
            report_modified="2026-06-17", email_parquet=self.parquet)

    def tearDown(self):
        import glob
        for p in glob.glob(os.path.join(self.tmp, "*")):
            try:
                os.remove(p)
            except OSError:
                pass

    def _ro(self):
        import duckdb
        return duckdb.connect(self.db, read_only=True)

    def test_threads_loaded(self):
        self.assertEqual(self.summary["email_threads"], 3)
        con = self._ro()
        try:
            self.assertEqual(con.execute("SELECT count(*) FROM fact_email_thread").fetchone()[0], 3)
        finally:
            con.close()

    def test_po_bridge_exploded(self):
        con = self._ro()
        try:
            n = con.execute("SELECT count(*) FROM bridge_email_po").fetchone()[0]
        finally:
            con.close()
        self.assertEqual(n, 3)   # T1:1 + T2:2

    def test_customer_bridge_alias_and_noise_rejected(self):
        con = self._ro()
        try:
            rev = {r[0]: r[1] for r in con.execute(
                "SELECT thread_id, customer_id FROM bridge_email_customer").fetchall()}
        finally:
            con.close()
        self.assertEqual(rev.get("T1"), "DECA.26")   # 'Decathlon' → 別名接上
        self.assertEqual(rev.get("T2"), "LURCHI")
        self.assertNotIn("T3", rev)                  # 'FU CHUN' 非客戶 → 不連

    def test_dim_customer_has_alias(self):
        con = self._ro()
        try:
            al = con.execute(
                "SELECT aliases FROM dim_customer WHERE customer_id='DECA.26'").fetchone()[0]
        finally:
            con.close()
        self.assertIn("DECATHLON", al.upper())

    def test_cross_source_customer_join(self):
        # 跨源：DECA.26 同時有生產訂單與郵件 —— Phase 2 的可行連結
        con = self._ro()
        try:
            row = con.execute(
                "SELECT d.customer_id, "
                "(SELECT count(*) FROM fact_order_state o WHERE o.customer_raw=d.customer_id), "
                "(SELECT count(*) FROM bridge_email_customer b WHERE b.customer_id=d.customer_id) "
                "FROM dim_customer d WHERE d.customer_id='DECA.26'").fetchone()
        finally:
            con.close()
        self.assertEqual(row[0], "DECA.26")
        self.assertGreaterEqual(row[1], 1)   # 有生產訂單
        self.assertEqual(row[2], 1)           # 有 1 封郵件（T1）

    def test_order_xref_customer_crosswalk(self):
        # Phase 3c：dim_order_xref 以客戶串 —— DECA.26 → 指令 JFC1（production）+ PO 65L1（email T1）
        con = self._ro()
        try:
            ids = {(r[0], r[1]) for r in con.execute(
                "SELECT id_type, identifier FROM dim_order_xref WHERE customer_id='DECA.26'").fetchall()}
            po_rows = con.execute(
                "SELECT identifier FROM dim_order_xref WHERE id_type='po' AND customer_id='LURCHI'").fetchall()
        finally:
            con.close()
        self.assertIn(("work_order", "JFC1"), ids)   # 生產指令 → 客戶（1.0）
        self.assertIn(("po", "65L1"), ids)            # 郵件 PO → 客戶（thread 級 0.7）
        self.assertEqual({r[0] for r in po_rows}, {"P1", "P2"})   # LURCHI 的兩個 PO（T2）

    def test_no_parquet_skips_email_but_tables_exist(self):
        db2 = os.path.join(self.tmp, "wh2.duckdb")
        s = fw.build_warehouse_from_bytes(_fixture_bytes(), db_path=db2, email_parquet=None)
        self.assertNotIn("email_threads", s)   # 沒攝取 → summary 無此鍵
        import duckdb
        con = duckdb.connect(db2, read_only=True)
        try:
            self.assertEqual(con.execute("SELECT count(*) FROM fact_email_thread").fetchone()[0], 0)
        finally:
            con.close()


class ParseDocTitleTests(unittest.TestCase):
    """Phase 3a：純檔名解析（零 OCR）。"""

    def test_invoice_with_lot(self):
        p = fw._parse_doc_title("INVOICE -- LOT 309-2025 .pdf")
        self.assertEqual(p["doc_type"], "invoice")
        self.assertEqual(p["lot_number"], "309-2025")

    def test_remittance_with_amount(self):
        p = fw._parse_doc_title("USD 9575--LOT 203-2025 匯款單.jpg")
        self.assertEqual(p["doc_type"], "remittance")
        self.assertEqual(p["lot_number"], "203-2025")
        self.assertEqual(p["currency"], "USD")
        self.assertEqual(p["amount"], 9575.0)

    def test_courier_awb_no_lot(self):
        p = fw._parse_doc_title("越輝快遞單 AWB 15329035.png")
        self.assertEqual(p["doc_type"], "courier")
        self.assertEqual(p["awb"], "15329035")
        self.assertEqual(p["lot_number"], "")

    def test_coo_shipping_insurance(self):
        self.assertEqual(fw._parse_doc_title("COO -- LOT 309-2025.pdf")["doc_type"], "coo")
        self.assertEqual(
            fw._parse_doc_title("Shipping document-- LOT 303-2025 (MASTROTTO).pdf")["doc_type"], "shipping")
        self.assertEqual(fw._parse_doc_title("保單-- LOT 206-2025.jpg")["doc_type"], "insurance")

    def test_unmatched_is_other(self):
        p = fw._parse_doc_title("會議記錄.pdf")
        self.assertEqual(p["doc_type"], "other")
        self.assertEqual(p["lot_number"], "")


class ShipmentDocPhase3aTests(unittest.TestCase):
    DOCS = [
        {"file_id": "f1", "title": "INVOICE -- LOT 309-2025 .pdf",
         "mime": "application/pdf", "modified": "2025-09-01", "drive_id": "D1"},
        {"file_id": "f2", "title": "Shipping document-- LOT 309-2025.pdf",
         "mime": "application/pdf", "modified": "2025-09-01", "drive_id": "D1"},
        {"file_id": "f3", "title": "USD 9575--LOT 203-2025 匯款單.jpg",
         "mime": "image/jpeg", "modified": "2025-07-01", "drive_id": "D1"},
        {"file_id": "f4", "title": "越輝快遞單 AWB 15329035.png",
         "mime": "image/png", "modified": "2025-06-01", "drive_id": "D1"},
        {"file_id": "f5", "title": "會議記錄.pdf",
         "mime": "application/pdf", "modified": "2025-05-01", "drive_id": "D1"},
    ]

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "wh.duckdb")
        self.summary = fw.build_warehouse_from_bytes(
            _fixture_bytes(), db_path=self.db, shipment_docs=self.DOCS)

    def tearDown(self):
        import glob
        for p in glob.glob(os.path.join(self.tmp, "*")):
            try:
                os.remove(p)
            except OSError:
                pass

    def _ro(self):
        import duckdb
        return duckdb.connect(self.db, read_only=True)

    def test_docs_loaded(self):
        self.assertEqual(self.summary["shipment_docs"], 5)
        con = self._ro()
        try:
            self.assertEqual(con.execute("SELECT count(*) FROM fact_shipment_doc").fetchone()[0], 5)
        finally:
            con.close()

    def test_parsed_fields(self):
        con = self._ro()
        try:
            rows = {r[0]: r for r in con.execute(
                "SELECT file_id, doc_type, lot_number, awb, currency, amount "
                "FROM fact_shipment_doc").fetchall()}
        finally:
            con.close()
        self.assertEqual(rows["f1"][1:3], ("invoice", "309-2025"))
        self.assertEqual(rows["f3"][1:6], ("remittance", "203-2025", "", "USD", 9575.0))
        self.assertEqual(rows["f4"][1], "courier")
        self.assertEqual(rows["f4"][3], "15329035")
        self.assertEqual(rows["f5"][1], "other")

    def test_lot_doc_inventory_query(self):
        # 「LOT 309-2025 有哪些單據」 → invoice + shipping
        con = self._ro()
        try:
            types_ = {r[0] for r in con.execute(
                "SELECT doc_type FROM fact_shipment_doc WHERE lot_number='309-2025'").fetchall()}
        finally:
            con.close()
        self.assertEqual(types_, {"invoice", "shipping"})

    def test_no_docs_skips_but_table_exists(self):
        db2 = os.path.join(self.tmp, "wh2.duckdb")
        s = fw.build_warehouse_from_bytes(_fixture_bytes(), db_path=db2)
        self.assertNotIn("shipment_docs", s)
        import duckdb
        con = duckdb.connect(db2, read_only=True)
        try:
            self.assertEqual(con.execute("SELECT count(*) FROM fact_shipment_doc").fetchone()[0], 0)
        finally:
            con.close()


class PaymentExtractPhase3bTests(unittest.TestCase):
    """Phase 3b：單據內容抽結構化金額 → 持久化 store → fact_payment（注入式，免網路/Gemini）。"""

    SHIP = [
        {"file_id": "p1", "title": "INVOICE -- LOT 309-2025.pdf",
         "mime": "application/pdf", "modified": "2025-09-02", "drive_id": "D"},
        {"file_id": "p2", "title": "USD 9575--LOT 203-2025 匯款單.pdf",
         "mime": "application/pdf", "modified": "2025-09-01", "drive_id": "D"},
        {"file_id": "p3", "title": "出貨單 -- LOT 309-2025.pdf",   # shipping，非付款類→排除
         "mime": "application/pdf", "modified": "2025-08-01", "drive_id": "D"},
        {"file_id": "p4", "title": "匯款單 LOT 5.jpg",             # 圖片→現在納入（ocrmac OCR）
         "mime": "image/jpeg", "modified": "2025-07-01", "drive_id": "D"},
        {"file_id": "p5", "title": "對帳資料夾",                     # folder→排除（抽不出文字）
         "mime": "application/vnd.google-apps.folder", "modified": "2025-07-01", "drive_id": "D"},
    ]
    FAKE = {"amount": 9575.0, "currency": "USD", "doc_date": "2025-09-01",
            "counterparty": "佳桀", "invoice_no": "INV1"}

    def setUp(self):
        from agent_core import factory_warehouse_extract as ex
        self.ex = ex
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "wh.duckdb")
        self.store = os.path.join(self.tmp, "payments.jsonl")
        fw.build_warehouse_from_bytes(_fixture_bytes(), db_path=self.db, shipment_docs=self.SHIP)
        p = mock.patch.object(ex, "payments_store_path", lambda: self.store)
        p.start()
        self.addCleanup(p.stop)

    def tearDown(self):
        import glob
        for f in glob.glob(os.path.join(self.tmp, "*")):
            try:
                os.remove(f)
            except OSError:
                pass

    def _run(self, **kw):
        return self.ex.extract_payments_batch(
            db_path=self.db, read_text=lambda fid: "invoice text 9575 USD",
            structure=lambda t: dict(self.FAKE), **kw)

    def test_money_docs_include_images_exclude_folders(self):
        r = self._run(max_docs=10)
        # p1 invoice + p2 remittance(pdf) + p4 remittance(圖片，現在納入)；
        # p3 出貨單(非付款類)、p5 folder(抽不出文字) 排除。
        self.assertEqual(r["extracted"], 3)
        self.assertEqual(r["candidates_remaining"], 0)

    def test_incremental_skips_done(self):
        self._run(max_docs=10)
        self.assertEqual(self._run(max_docs=10)["extracted"], 0)

    def test_max_docs_cap(self):
        self.assertEqual(self._run(max_docs=1)["extracted"], 1)

    def test_time_budget_zero_stops_before_any_doc(self):
        # 時間預算逐份開工前檢查：0 秒 → 一份都不開工、標記 time_exhausted
        r = self._run(max_docs=10, time_budget_s=0)
        self.assertEqual(r["extracted"], 0)
        self.assertTrue(r["time_exhausted"])
        self.assertEqual(r["candidates_remaining"], 3)

    def test_no_time_budget_not_marked_exhausted(self):
        self.assertFalse(self._run(max_docs=10)["time_exhausted"])

    def test_records_persist_incrementally(self):
        # 逐筆 append：第 N+1 份開工前，前 N 份必須已落盤（看門狗 os._exit 砍掉
        # 整批結束才寫的話，被砍的輪全部白抽、下輪重付 —— 2026-07-24 事故的損失面）
        seen = []

        def read_text(fid):
            if seen:
                with open(self.store, encoding="utf-8") as fh:
                    self.assertEqual(len(fh.readlines()), len(seen))
            seen.append(fid)
            return "invoice text 9575 USD"

        r = self.ex.extract_payments_batch(
            db_path=self.db, read_text=read_text,
            structure=lambda t: dict(self.FAKE), max_docs=10)
        self.assertEqual(r["extracted"], 3)

    def test_loaded_into_fact_payment(self):
        self._run(max_docs=10)
        db2 = os.path.join(self.tmp, "wh2.duckdb")
        s = fw.build_warehouse_from_bytes(_fixture_bytes(), db_path=db2, payments_store=self.store)
        self.assertEqual(s["payment_rows"], 3)   # p1 + p2 + p4（圖片）
        import duckdb
        con = duckdb.connect(db2, read_only=True)
        try:
            row = con.execute(
                "SELECT amount, currency, amount_source FROM fact_payment "
                "WHERE file_id='p2'").fetchone()
        finally:
            con.close()
        self.assertEqual(row, (9575.0, "USD", "content"))   # 模型抽到 → content

    def test_structure_payment_parses_json(self):
        fake = types.SimpleNamespace(
            text='{"amount":"9575","currency":"USD","doc_date":"2025-09-01",'
                 '"counterparty":"佳桀","invoice_no":"INV1"}')
        with mock.patch("agent_core.gemini_client._gemini_generate", return_value=fake):
            d = self.ex._structure_payment("some text")
        self.assertEqual(d["amount"], 9575.0)   # 字串 → float
        self.assertEqual(d["currency"], "USD")

    def test_structure_payment_json_array_does_not_crash(self):
        # Gemini 偶爾回陣列 [...]（live 抓到的真實案例）→ 當抽不到、不炸
        fake = types.SimpleNamespace(text='[{"amount": 1}]')
        with mock.patch("agent_core.gemini_client._gemini_generate", return_value=fake):
            d = self.ex._structure_payment("text")
        self.assertIsNone(d["amount"])
        self.assertEqual(d["currency"], "")

    def test_fact_payment_empty_without_store(self):
        db3 = os.path.join(self.tmp, "wh3.duckdb")
        s = fw.build_warehouse_from_bytes(_fixture_bytes(), db_path=db3)
        self.assertEqual(s["payment_rows"], 0)


class PaymentLoadEnrichTests(unittest.TestCase):
    """_load_payments 咽喉：正規化 + 檔名回填 + 去重（手寫 store，免抽取/網路）。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db = os.path.join(self.tmp, "wh.duckdb")
        self.store = os.path.join(self.tmp, "payments.jsonl")

    def tearDown(self):
        import glob
        for f in glob.glob(os.path.join(self.tmp, "*")):
            try:
                os.remove(f)
            except OSError:
                pass

    def _write_store(self, recs):
        with open(self.store, "w", encoding="utf-8") as fh:
            for r in recs:
                fh.write(json.dumps(r, ensure_ascii=False) + "\n")

    def _rows(self, cols, where=""):
        import duckdb
        fw.build_warehouse_from_bytes(
            _fixture_bytes(), db_path=self.db, payments_store=self.store)
        con = duckdb.connect(self.db, read_only=True)
        try:
            return con.execute(
                f"SELECT {cols} FROM fact_payment {where}").fetchall()
        finally:
            con.close()

    def test_filename_fallback_amount_and_date(self):
        # 內容抽不到金額/日期，但檔名有 → 回填、標 filename
        self._write_store([{
            "file_id": "f1", "doc_type": "remittance",
            "title": "佳桀-匯款-MASTROTTO EUR16,006.89-LOT 310-2026--20260327.pdf",
            "amount": None, "currency": "", "doc_date": "", "counterparty": "",
            "invoice_no": "", "extracted_at": "2026-03-28T00:00:00",
        }])
        row = self._rows(
            "amount, currency, doc_date, amount_source, counterparty",
            "WHERE file_id='f1'")[0]
        self.assertAlmostEqual(row[0], 16006.89, places=2)   # 千分位逗號正確
        self.assertEqual(row[1], "EUR")
        self.assertEqual(row[2], "2026-03-27")
        self.assertEqual(row[3], "filename")
        self.assertEqual(row[4], "MASTROTTO")                # 檔名回填對方

    def test_currency_and_counterparty_normalized(self):
        # US$ → USD；大小寫變體 → 同 counterparty_norm
        self._write_store([
            {"file_id": "a", "doc_type": "invoice", "title": "x", "amount": 100,
             "currency": "US$", "counterparty": "Jai Jye Corporation"},
            {"file_id": "b", "doc_type": "invoice", "title": "y", "amount": 200,
             "currency": "USD", "counterparty": "JAI JYE CORPORATION"},
        ])
        rows = self._rows("currency, counterparty_norm", "ORDER BY file_id")
        self.assertEqual(rows[0][0], "USD")                  # US$ 正規化
        self.assertEqual(rows[1][0], "USD")
        self.assertEqual(rows[0][1], rows[1][1])             # 兩種大小寫併成同一鍵
        self.assertEqual(rows[0][1], "JAI JYE CORPORATION")

    def test_dedup_by_file_id_keeps_last(self):
        # 同 file_id append 兩次（store 是 append-only）→ 只留最後一筆，免 SUM 重計
        self._write_store([
            {"file_id": "dup", "doc_type": "remittance", "title": "t", "amount": 111,
             "currency": "USD", "extracted_at": "2026-01-01T00:00:00"},
            {"file_id": "dup", "doc_type": "remittance", "title": "t", "amount": 222,
             "currency": "USD", "extracted_at": "2026-02-01T00:00:00"},
        ])
        rows = self._rows("amount", "WHERE file_id='dup'")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 222.0)                  # 最後一筆勝出

    def test_content_amount_kept_over_filename(self):
        # 內容有金額時不被檔名覆蓋，標 content
        self._write_store([{
            "file_id": "c", "doc_type": "remittance",
            "title": "佳桀-匯款-COATS USD626.82-20260616.pdf",
            "amount": 999.0, "currency": "USD",
        }])
        row = self._rows("amount, amount_source", "WHERE file_id='c'")[0]
        self.assertEqual(row[0], 999.0)                      # 內容值保留
        self.assertEqual(row[1], "content")


class PaymentOcrFallbackTests(unittest.TestCase):
    """_read_doc_text：掃描 PDF（文字層空/薄）fallback 本地 ocrmac；⚠ 不支援/太大不重抽。

    只測編排（mock read_drive_file + _ocr_scanned_pdf），不需 pypdfium2/Drive/ocrmac。
    """

    def setUp(self):
        from agent_core import factory_warehouse_extract as ex
        self.ex = ex

    def _run(self, read_ret, ocr_ret):
        with mock.patch("agent_core.ingest.drive_search.read_drive_file", return_value=read_ret), \
             mock.patch.object(self.ex, "_ocr_scanned_pdf", return_value=ocr_ret) as m_ocr:
            out = self.ex._read_doc_text("fid")
        return out, m_ocr

    def test_good_text_no_ocr(self):
        out, m = self._run(
            "📄 inv.pdf（2026-06-16｜application/pdf）\nINVOICE USD 1234.56 內容夠長夠長夠長",
            "SHOULD-NOT-USE")
        self.assertIn("1234.56", out)
        m.assert_not_called()             # 文字夠 → 不打 OCR

    def test_empty_scanned_pdf_triggers_ocr(self):
        out, m = self._run(
            "（「scan.pdf」抽出來是空的 — 可能是空白檔或無文字內容）",
            "USD 626.82 COATS 2026-06-16")
        self.assertIn("626.82", out)
        m.assert_called_once()

    def test_thin_pdf_body_triggers_ocr(self):
        out, m = self._run("📄 r.pdf（2026｜application/pdf）\n1", "EUR 16006.89")
        self.assertIn("16006.89", out)
        m.assert_called_once()            # 📄 但本文過短 → 也補 OCR

    def test_unsupported_or_toobig_not_reocr(self):
        out, m = self._run("⚠️ 「big.pdf」太大（50.0 MB，上限 30 MB）…", "X")
        m.assert_not_called()             # ⚠ → 不重抽（避免重新下載）
        self.assertTrue(out.startswith("⚠"))

    def test_ocr_empty_falls_back_to_original(self):
        out, m = self._run("（「scan.pdf」抽出來是空的）", "")
        self.assertIn("空的", out)         # OCR 也抽不到 → 回原文字、不炸


if __name__ == "__main__":
    unittest.main()
