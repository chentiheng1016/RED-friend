"""樣品單線稿驅動生圖：挑圖、每日配額、部門分艙與交付。

2026-08-03 UserC 案（Richter 8105-3272）：開發部員工丟樣品單 xlsx 要「自動生成
鞋圖」，小紅答「系統尚無此功能」。四層真因，本檔覆蓋其中三層：
  1. green 白名單沒有 generate_from_order/parse_sample_order（TG 白名單測試）
  2. tier=CONFIRM 進不了員工 session（tier + filter_tools_for_color 測試）
  3. 生成圖走 telegram_send_photo → 傳到大王手機，員工收不到（交付測試）
  4. _pick_shoe_sketch 挑錯圖：emf→png 是整頁 A4 白畫布，面積壓過真線稿；
     且 0.05 置信地板把真線稿的 footwear=0.045 抹平（挑圖測試）

注意（unittest discover）：conftest fixture 不生效，隔離全在 setUp/tearDown；
mock.patch 一律在主執行緒。配額/輸出目錄一律指到 tmp，不碰主 checkout 的 var/。
"""
import json
import os
import sys
import tempfile
import unittest
from unittest import mock

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)


def _draw(path: str, size: tuple[int, int], ink: tuple[int, int, int, int]) -> str:
    """畫一張白底圖，ink=(x0,y0,x1,y1) 區域塗黑（模擬線稿墨跡）。"""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, (255, 255, 255))
    ImageDraw.Draw(im).rectangle(ink, outline=(0, 0, 0), fill=(30, 30, 30))
    im.save(path)
    return path


class PickShoeSketchTests(unittest.TestCase):
    """挑線稿：先去白邊、鞋類訊號優先於面積。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = self.tmp.name
        # emf→soffice 轉出來的整頁 A4：內容只有中間一小塊（保養符號那頁）
        self.page = _draw(os.path.join(d, "page.png"), (794, 1123),
                          (320, 500, 461, 565))
        # 真正的手繪線稿：墨跡幾乎填滿畫面
        self.sketch = _draw(os.path.join(d, "sketch.png"), (442, 225),
                            (10, 5, 432, 220))
        # 大底線稿：夠大但太扁（寬高比排除）
        self.sole = _draw(os.path.join(d, "sole.png"), (900, 150),
                          (5, 5, 895, 145))

    def tearDown(self):
        self.tmp.cleanup()

    def test_autocrop_strips_a4_padding(self):
        from skills import sample_order as so
        from PIL import Image
        cropped = so._autocrop_white(self.page)
        self.assertNotEqual(cropped, self.page)
        self.assertEqual(Image.open(cropped).size, (142, 66))

    def test_no_shoe_signal_falls_back_to_cropped_area(self):
        """Vision 全無訊號時比的是**去白邊後**的面積 —— 大白紙不再自動獲勝。"""
        from skills import sample_order as so
        with mock.patch.object(so, "_shoe_confidence", return_value=0.0):
            picked = so._pick_shoe_sketch([self.page, self.sketch, self.sole])
        self.assertIsNotNone(picked)
        self.assertIn("sketch", os.path.basename(picked))
        # 回傳的是去白邊後的圖（線稿填滿畫面，圖生圖鎖形狀更準）
        self.assertTrue(os.path.basename(picked).startswith("crop_"))

    def test_weak_shoe_signal_beats_bigger_blank_page(self):
        """真線稿的 footwear 置信只有 0.045 —— 舊版 0.05 地板會把它抹平成跟白紙同分。"""
        from skills import sample_order as so
        big = _draw(os.path.join(self.tmp.name, "big.png"), (1200, 900),
                    (10, 10, 1190, 890))   # 面積遠大於 sketch、但沒有鞋訊號

        def _conf(path: str) -> float:
            return 0.045 if "sketch" in os.path.basename(path) else 0.0

        with mock.patch.object(so, "_shoe_confidence", side_effect=_conf):
            picked = so._pick_shoe_sketch([big, self.sketch])
        self.assertIn("sketch", os.path.basename(picked))

    def test_all_candidates_filtered_returns_none(self):
        from skills import sample_order as so
        icon = _draw(os.path.join(self.tmp.name, "icon.png"), (60, 60), (5, 5, 55, 55))
        with mock.patch.object(so, "_shoe_confidence", return_value=0.0):
            self.assertIsNone(so._pick_shoe_sketch([icon, self.sole]))


class RenderQuotaTests(unittest.TestCase):
    """每日生圖張數硬上限（取代 CONFIRM 閘擋意外花費）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._state = mock.patch("agent_core.logging_and_paths.STATE_DIR", self.tmp.name)
        self._state.start()

    def tearDown(self):
        self._state.stop()
        self.tmp.cleanup()

    def test_counts_per_scope_and_blocks_at_limit(self):
        from skills import sample_order as so
        with mock.patch.dict(os.environ, {"RED_ORDER_RENDER_DAILY_MAX": "2"}):
            self.assertEqual(so._consume_render_quota("green")[:2], (True, 1))
            self.assertEqual(so._consume_render_quota("green")[:2], (True, 2))
            allowed, used, limit = so._consume_render_quota("green")
            self.assertFalse(allowed)
            self.assertEqual((used, limit), (2, 2))
            # 別色有自己的額度，不被 green 佔用
            self.assertEqual(so._consume_render_quota("owner")[:2], (True, 1))

    def test_zero_limit_means_unlimited(self):
        from skills import sample_order as so
        with mock.patch.dict(os.environ, {"RED_ORDER_RENDER_DAILY_MAX": "0"}):
            for _ in range(5):
                self.assertTrue(so._consume_render_quota("green")[0])

    def test_yesterdays_count_does_not_carry_over(self):
        from skills import sample_order as so
        path = os.path.join(self.tmp.name, "order_render_quota.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump({"2020-01-01": {"green": 99}}, f)
        with mock.patch.dict(os.environ, {"RED_ORDER_RENDER_DAILY_MAX": "2"}):
            self.assertEqual(so._consume_render_quota("green")[:2], (True, 1))

    def test_corrupt_state_file_fails_open(self):
        from skills import sample_order as so
        path = os.path.join(self.tmp.name, "order_render_quota.json")
        with open(path, "w", encoding="utf-8") as f:
            f.write("{not json")
        with mock.patch.dict(os.environ, {"RED_ORDER_RENDER_DAILY_MAX": "2"}):
            self.assertTrue(so._consume_render_quota("green")[0])


class GenerateFromOrderDeliveryTests(unittest.TestCase):
    """部門 context → 分艙目錄 + [[TG_PHOTO:]] 回覆附圖；大王 context → 主動推送。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.gen = os.path.join(self.tmp.name, "generated_images")
        self.state = os.path.join(self.tmp.name, "state")
        os.makedirs(self.gen, exist_ok=True)
        os.makedirs(self.state, exist_ok=True)
        self._p = [
            mock.patch("agent_core.logging_and_paths.GENERATED_IMAGES_DIR", self.gen),
            mock.patch("agent_core.logging_and_paths.STATE_DIR", self.state),
        ]
        for p in self._p:
            p.start()

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["MODELSPECIFICATION SS2026"])
        ws.append(["style number richter", "8105-3272"])
        self.xlsx = os.path.join(self.tmp.name, "Richter_8105-3272-.xlsx")
        wb.save(self.xlsx)
        self.sketch = _draw(os.path.join(self.tmp.name, "sk.png"), (442, 225),
                            (10, 5, 432, 220))

    def tearDown(self):
        for p in self._p:
            p.stop()
        self.tmp.cleanup()

    def _fake_gemini(self):
        """同一顆 fake 同時服務 _spec_for_render（.text）與生圖（.candidates）。"""
        part = mock.MagicMock()
        part.inline_data.data = b"\x89PNG fake"
        cand = mock.MagicMock()
        cand.content.parts = [part]
        resp = mock.MagicMock(text="鞋身 pro action 深藍、魔鬼氈 青綠", candidates=[cand])
        fake = mock.MagicMock()
        fake.models.generate_content.return_value = resp
        return mock.patch("agent_core.gemini_client._get_gemini_client", return_value=fake)

    def _run(self, color: str, channel: str = ""):
        import contextlib

        from agent_core.channel_context import channel_context
        from skills import sample_order as so
        ctx = channel_context(channel) if channel else contextlib.nullcontext()
        with mock.patch.object(so, "_extract_order_images", return_value=[self.sketch]), \
                mock.patch.object(so, "_dept_scope_color", return_value=color), \
                mock.patch("agent_core.telegram.telegram_send_photo") as send, \
                self._fake_gemini(), ctx:
            out = so.generate_from_order(self.xlsx)
        return out, send

    def test_dept_context_writes_to_color_subdir_and_returns_marker(self):
        out, send = self._run("green")
        expected_dir = os.path.join(self.gen, "dept", "green")
        self.assertIn("[[TG_PHOTO:", out)
        self.assertIn(expected_dir, out)
        self.assertTrue(os.listdir(expected_dir))
        # 員工的圖絕不能走出站推送（那條 chat_id 閘只認大王）
        send.assert_not_called()

    def test_owner_context_pushes_and_has_no_marker(self):
        out, send = self._run("")
        self.assertNotIn("[[TG_PHOTO:", out)
        self.assertTrue(any(f.startswith("order_") for f in os.listdir(self.gen)))
        send.assert_called_once()

    def test_owner_in_telegram_chat_gets_marker_not_push(self):
        """2026-09-02 green 聊天室 5004 案：大王在 Telegram 對話發起 → 圖跟著
        回覆走（送回當前對話），不走只認紅 bot 主對話的出站推送。"""
        out, send = self._run("", channel="telegram")
        self.assertIn("[[TG_PHOTO:", out)
        # 大王的圖仍在 generated_images 頂層（owner 標記白名單=整個目錄，
        # 見 ReplyPhotoPerColorRootTests.test_owner_unchanged_sees_whole_generated_dir）
        self.assertTrue(any(f.startswith("order_") for f in os.listdir(self.gen)))
        self.assertNotIn(os.path.join(self.gen, "dept"), out)
        send.assert_not_called()

    def test_extra_note_material_lands_texture_ref(self):
        # 2026-09-02 PSS 6 案：「皮料改羊巴戈」只活在 extra_note、parts 材質欄
        # 照單上仍是 PU → 特寫從沒附上，模型只能憑字面想像質感。
        from PIL import Image
        from agent_core.logging_and_paths import DATA_DIR
        from skills import sample_order as so
        lib = os.path.join(DATA_DIR, "material_swatches")
        os.makedirs(lib, exist_ok=True)
        p = os.path.join(lib, "pro_action.png")
        Image.new("RGB", (32, 32), (40, 40, 40)).save(p)
        try:
            with mock.patch.object(so, "_extract_order_images", return_value=[self.sketch]), \
                    mock.patch.object(so, "_dept_scope_color", return_value=""), \
                    mock.patch("agent_core.telegram.telegram_send_photo"), \
                    self._fake_gemini():
                out = so.generate_from_order(self.xlsx, extra_note="皮料改羊巴戈")
            self.assertIn("材質特寫已附", out)
            self.assertIn("羊巴戈", out)
        finally:
            os.remove(p)

    def test_quota_exhausted_refuses_before_paying(self):
        from skills import sample_order as so
        with mock.patch.dict(os.environ, {"RED_ORDER_RENDER_DAILY_MAX": "1"}):
            self._run("green")
            with mock.patch.object(so, "_extract_order_images", return_value=[self.sketch]), \
                    mock.patch.object(so, "_dept_scope_color", return_value="green"), \
                    self._fake_gemini() as client:
                out = so.generate_from_order(self.xlsx)
        self.assertIn("上限", out)
        client.assert_not_called()   # 撞上限就完全不打 Gemini


class RenderGuardTests(unittest.TestCase):
    def test_color_text_must_appear_in_order_text(self):
        # 2026-09-02 PSS 6 案：完稿在場時 LLM 腦補「大底=黑色」，經內建色名表
        # 以 builtin 之姿穿過 #458 的 source 濾網 → 黑大底蓋掉客人畫的白/灰大底。
        from skills.sample_order import _color_text_in_order
        order = "Husky2.0\nRichtex, PU\nPSS 6 Zip inside\ndenim\nPANTONE 19-4025 TPX"
        self.assertFalse(_color_text_in_order("黑色", order))
        self.assertTrue(_color_text_in_order("denim", order))
        self.assertTrue(_color_text_in_order("Richtex ,PU", order))       # 空白差異不影響
        self.assertTrue(_color_text_in_order("深藍 Pantone 19-4025 TPX", order))  # 帶數字色號在原文
        self.assertFalse(_color_text_in_order("", order))

    def test_composition_flaws_detects_pair_and_fails_open(self):
        # 2026-09-02 PSS 6 案：一雙斜角產品照——OCR/ΔE 都攔不到構圖跑掉。
        import tempfile
        from PIL import Image
        from skills import sample_order as so
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.png")
            Image.new("RGB", (32, 32), (200, 200, 200)).save(p)
            with mock.patch("agent_core.gemini_client.generate_content_tracked",
                            return_value=mock.MagicMock(text='{"shoes": 2, "side_view": false}')):
                flaws = so._composition_flaws(p)
            self.assertEqual(len(flaws), 2)
            self.assertIn("2 隻鞋", flaws[0])
            with mock.patch("agent_core.gemini_client.generate_content_tracked",
                            return_value=mock.MagicMock(text='{"shoes": 1, "side_view": true}')):
                self.assertEqual(so._composition_flaws(p), [])
            with mock.patch("agent_core.gemini_client.generate_content_tracked",
                            side_effect=RuntimeError("api down")):
                self.assertEqual(so._composition_flaws(p), [])  # 驗證器壞了不擋交件

    def test_foreign_color_flagged_when_artwork_present(self):
        # 2026-09-02 5004 案：完稿棕褐印花被渲染成深綠鞋筒——ΔE 只驗「該有的
        # 在不在」，外來色靠同一通 flash 連完稿一起看。
        import tempfile
        from PIL import Image
        from skills import sample_order as so
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "r.png")
            art = os.path.join(d, "art.png")
            Image.new("RGB", (32, 32), (40, 60, 45)).save(p)
            Image.new("RGB", (32, 32), (150, 108, 76)).save(art)
            resp = '{"shoes": 1, "side_view": true, "foreign_color": "鞋筒變成深綠色"}'
            with mock.patch("agent_core.gemini_client.generate_content_tracked",
                            return_value=mock.MagicMock(text=resp)) as g:
                flaws = so._composition_flaws(p, artwork_path=art)
            self.assertEqual(len(flaws), 1)
            self.assertIn("深綠", flaws[0])
            self.assertEqual(len(g.call_args.kwargs["contents"]), 3)  # 成品+完稿+prompt
            # 「無」類字樣視同沒有外來色
            with mock.patch("agent_core.gemini_client.generate_content_tracked",
                            return_value=mock.MagicMock(
                                text='{"shoes": 1, "side_view": true, "foreign_color": "無"}')):
                self.assertEqual(so._composition_flaws(p, artwork_path=art), [])

class ArtColorTargetsTests(unittest.TestCase):
    """完稿 ΔE 目標＝主色 top4＋高彩度點綴色，同色（ΔE≤12）不重複收。

    2026-09-03 PSS 2 案：後套寶藍只佔完稿 2.3% 進不了主色 top4、色名又解不出
    hex → 全管線沒人驗這個部位，渲染成灰藍（ΔE 36.7）照樣交件。
    圖固定 ≤256px：dominant/accent 兩路都不觸發縮圖重採樣，色值可精確斷言。
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (250, 210), (255, 255, 255))
        d = ImageDraw.Draw(im)
        d.rectangle((25, 42, 125, 147), fill=(43, 48, 66))     # 深藍鞋身 ~20%（主色、低彩度）
        d.rectangle((150, 105, 225, 178), fill=(236, 122, 45))  # 橘拼接 ~11%（主色、高彩度）
        d.rectangle((200, 53, 232, 78), fill=(25, 63, 143))     # 寶藍後套 ~1.6%（點綴）
        self.art = os.path.join(self.tmp.name, "art.png")
        im.save(self.art)

    def tearDown(self):
        self.tmp.cleanup()

    def test_accent_appended_dominant_saturated_not_duplicated(self):
        from skills.sample_order import _art_color_targets
        got = _art_color_targets(self.art, [])
        by_part = {t["part"]: t["hex"] for t in got}
        self.assertIn("193f8f", by_part.values())               # 後套寶藍進了驗證
        accents = [p for p in by_part if p.startswith("完稿點綴色")]
        self.assertEqual(len(accents), 1)                       # 橘塊已是主色，不再另收
        self.assertIn("ec7a2d", by_part.values())

    def test_part_swatch_covering_accent_suppresses_duplicate(self):
        from skills.sample_order import _art_color_targets
        got = _art_color_targets(self.art, [{"part": "後套", "hex": "193f8f"}])
        self.assertFalse([t for t in got if t["part"].startswith("完稿點綴色")])


class TextMaskAccentTests(unittest.TestCase):
    """2026-09-08 review：完稿標註墨水不得成為點綴色 ΔE 目標（畫過色的完稿
    刻意不洗標號，抽色前先把 OCR 文字框塗白）；鞋身明暗叢集不得吃光點綴名額。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (250, 210), (255, 255, 255))
        d = ImageDraw.Draw(im)
        d.rectangle((25, 42, 125, 147), fill=(43, 48, 66))     # 深藍鞋身（低彩度）
        d.rectangle((150, 20, 210, 45), fill=(210, 30, 30))    # 紅色手寫標註 ~2.9%
        self.art = os.path.join(self.tmp.name, "art.png")
        im.save(self.art)

    def tearDown(self):
        self.tmp.cleanup()

    def test_annotation_ink_masked_out_of_accent_targets(self):
        from skills.sample_order import _art_color_targets
        # 對照組：OCR 讀不到字 → 標註紅會被當點綴色（這正是要堵的洞）
        with mock.patch("agent_core.vision_ocr.ocr_text_boxes", return_value=[]):
            raw = _art_color_targets(self.art, [])
        self.assertTrue([t for t in raw if t["part"].startswith("完稿點綴色")])
        # Vision 正規化座標（原點左下）：紅框 (150,20)-(210,45)
        box = (150 / 250, (210 - 45) / 210, 60 / 250, 25 / 210)
        with mock.patch("agent_core.vision_ocr.ocr_text_boxes",
                        return_value=[("材質不對", box)]):
            got = _art_color_targets(self.art, [])
        self.assertFalse([t for t in got if t["part"].startswith("完稿點綴色")])

    def test_ocr_empty_falls_back_to_original_path(self):
        from skills.sample_order import _mask_text_for_accents
        with mock.patch("agent_core.vision_ocr.ocr_text_boxes", return_value=[]):
            self.assertEqual(_mask_text_for_accents(self.art), self.art)

    def test_saturated_body_shades_do_not_eat_accent_slots(self):
        """先截 2 再去重會讓彩色鞋身的明暗叢集吃光名額、真點綴色落空——
        改成抓 6 個候選、對主色去重後才截。"""
        from PIL import Image, ImageDraw
        from skills.sample_order import _art_color_targets
        im = Image.new("RGB", (250, 210), (255, 255, 255))
        d = ImageDraw.Draw(im)
        d.rectangle((20, 30, 140, 180), fill=(34, 81, 186))    # 鞋身亮面（高彩藍）
        d.rectangle((70, 30, 140, 180), fill=(25, 59, 136))    # 鞋身暗面
        d.rectangle((200, 40, 232, 65), fill=(200, 30, 40))    # 紅點綴 ~1.5%
        p = os.path.join(self.tmp.name, "body.png")
        im.save(p)
        with mock.patch("agent_core.vision_ocr.ocr_text_boxes", return_value=[]):
            got = _art_color_targets(p, [])
        accents = [t["hex"] for t in got if t["part"].startswith("完稿點綴色")]
        self.assertEqual(len(accents), 1)                      # 只剩真點綴色
        r, g = int(accents[0][:2], 16), int(accents[0][2:4], 16)
        self.assertGreater(r, g + 80)                          # 而且是紅色系


class DeptScopeSampleOrderTests(unittest.TestCase):
    """green 拿得到樣品單兩顆；會花錢的那顆不隨 QUERY_MATRIX 繼承出去。"""

    def test_green_home_has_both_tools(self):
        from agent_core.dept_tool_scope import allowed_tool_names_for_color
        allowed = allowed_tool_names_for_color("green")
        self.assertIn("parse_sample_order", allowed)
        self.assertIn("generate_from_order", allowed)

    def test_inheriting_colors_get_parse_but_not_render(self):
        from agent_core.agents.permission_matrix import Agent, QUERY_MATRIX
        from agent_core.dept_tool_scope import allowed_tool_names_for_color
        inheritors = [a.value for a in Agent
                      if Agent.GREEN in QUERY_MATRIX.get(a, ()) and a is not Agent.RED]
        self.assertTrue(inheritors, "QUERY_MATRIX 應該有色可查 green")
        for color in inheritors:
            allowed = allowed_tool_names_for_color(color)
            self.assertIn("parse_sample_order", allowed, color)
            self.assertNotIn("generate_from_order", allowed, color)

    def test_render_tool_is_safe_tier_so_it_survives_the_filter(self):
        from agent_core.dept_tool_scope import filter_tools_for_color

        def generate_from_order():  # noqa: D401 - 只借名字給過濾器
            return ""

        def parse_sample_order():
            return ""

        kept, removed = filter_tools_for_color(
            [generate_from_order, parse_sample_order], "green")
        self.assertEqual([f.__name__ for f in kept],
                         ["generate_from_order", "parse_sample_order"])
        self.assertEqual(removed, [])

    def test_green_addendum_tells_llm_it_can_render(self):
        from agent_core.dept_tool_scope import dept_scope_addendum
        green = dept_scope_addendum("green")
        self.assertIn("generate_from_order", green)
        self.assertIn("TG_PHOTO", green)
        # 別色不該被塞開發部的生圖指示
        self.assertNotIn("generate_from_order", dept_scope_addendum("yellow"))


class ReplyPhotoPerColorRootTests(unittest.TestCase):
    """生成圖的回覆附圖白名單按色分艙（比照 [[TG_FILE:]]）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.gen = os.path.join(self.tmp.name, "generated_images")
        self.data = os.path.join(self.tmp.name, "data")
        self.green_dir = os.path.join(self.gen, "dept", "green")
        os.makedirs(self.green_dir, exist_ok=True)
        os.makedirs(os.path.join(self.data, "product_photos", "fetched"), exist_ok=True)
        self._p = [
            mock.patch("agent_core.logging_and_paths.GENERATED_IMAGES_DIR", self.gen),
            mock.patch("agent_core.logging_and_paths.DATA_DIR", self.data),
        ]
        for p in self._p:
            p.start()
        self.green_png = self._img(os.path.join(self.green_dir, "order_1.png"))
        self.owner_png = self._img(os.path.join(self.gen, "order_owner.png"))
        self.fetched_png = self._img(
            os.path.join(self.data, "product_photos", "fetched", "shoe.jpg"))

    def tearDown(self):
        for p in self._p:
            p.stop()
        self.tmp.cleanup()

    def _img(self, path: str) -> str:
        with open(path, "wb") as f:
            f.write(b"x" * 10)
        return os.path.realpath(path)

    def test_employee_gets_own_color_render_only(self):
        from agent_core.daemon_telegram import _extract_reply_photos
        _clean, photos = _extract_reply_photos(
            f"[[TG_PHOTO:{self.green_png}]]\n[[TG_PHOTO:{self.owner_png}]]", "green")
        self.assertEqual(photos, [self.green_png])

    def test_other_color_cannot_reach_green_render(self):
        from agent_core.daemon_telegram import _extract_reply_photos
        _clean, photos = _extract_reply_photos(f"[[TG_PHOTO:{self.green_png}]]", "indigo")
        self.assertEqual(photos, [])

    def test_fetched_photos_stay_shared_across_colors(self):
        from agent_core.daemon_telegram import _extract_reply_photos
        _clean, photos = _extract_reply_photos(f"[[TG_PHOTO:{self.fetched_png}]]", "indigo")
        self.assertEqual(photos, [self.fetched_png])

    def test_owner_unchanged_sees_whole_generated_dir(self):
        from agent_core.daemon_telegram import _extract_reply_photos
        _clean, photos = _extract_reply_photos(
            f"[[TG_PHOTO:{self.owner_png}]]\n[[TG_PHOTO:{self.green_png}]]", "")
        self.assertEqual(photos, [self.owner_png, self.green_png])


# ────────────────────────────────────────────────────────────────────
# 2026-08-04 UserC「不要再在生成的鞋圖上標數字」
#   #344 只在 prompt 講「不要出現數字」，當天 13 張成品實測仍有 3 張把
#   紅 1/綠 2/紅 5 畫上鞋面。治本＝輸入線稿先洗掉彩色標號，成品再 OCR 把關。
# ────────────────────────────────────────────────────────────────────
def _draw_annotated(path: str, size=(400, 200), annot_boxes=((20, 20, 34, 40),),
                    annot_rgb=(220, 20, 20)) -> str:
    """白底 + 黑線稿框 + 若干彩色小方塊（模擬紅/綠部位標號）。"""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.rectangle((5, 5, size[0] - 6, size[1] - 6), outline=(0, 0, 0), width=3)
    d.line((5, size[1] // 2, size[0] - 6, size[1] // 2), fill=(0, 0, 0), width=2)
    for box in annot_boxes:
        d.rectangle(box, fill=annot_rgb)
    im.save(path)
    return path


class ScrubColorAnnotationsTests(unittest.TestCase):
    """線稿上的彩色標號塗白；黑線稿完好；彩圖與純黑白都原樣放行。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _sat_pixels(self, path: str) -> int:
        import numpy as np
        from PIL import Image
        with Image.open(path) as raw:
            a = np.asarray(raw.convert("RGB")).astype(int)
        return int(((a.max(axis=2) - a.min(axis=2)) > 25).sum())

    def _dark_pixels(self, path: str) -> int:
        import numpy as np
        from PIL import Image
        with Image.open(path) as raw:
            a = np.asarray(raw.convert("RGB")).astype(int)
        return int((a.max(axis=2) < 120).sum())

    def test_color_marks_erased_line_art_kept(self):
        from skills import sample_order as so
        src = _draw_annotated(os.path.join(self.tmp.name, "sk.png"),
                              annot_boxes=((20, 20, 34, 40), (300, 150, 314, 170)))
        out = so._scrub_color_annotations(src)
        self.assertNotEqual(out, src)                 # 有洗 → 走新檔，不改原圖
        self.assertEqual(self._sat_pixels(out), 0)    # 彩色標號全消
        # 黑線稿保留（膨脹 1px 會蹭掉極少量，給 3% 容差）
        self.assertGreater(self._dark_pixels(out), self._dark_pixels(src) * 0.97)

    def test_pure_black_and_white_sketch_untouched(self):
        from skills import sample_order as so
        src = _draw_annotated(os.path.join(self.tmp.name, "bw.png"), annot_boxes=())
        self.assertEqual(so._scrub_color_annotations(src), src)

    def test_full_color_image_left_alone(self):
        """整張是彩圖（不是黑白線稿＋標註）→ 洗了會毀圖，原樣放行。"""
        from PIL import Image
        from skills import sample_order as so
        src = os.path.join(self.tmp.name, "photo.png")
        Image.new("RGB", (200, 200), (200, 40, 40)).save(src)
        self.assertEqual(so._scrub_color_annotations(src), src)

    def test_dark_painted_art_not_scrubbed(self):
        """全黑完稿彩色佔比低於 _ANNOT_MAX_FRACTION，但帶色的是客人畫的印花
        ——洗了就是毀客稿（PSS 7 案），深色填色判定要擋下來。"""
        from skills import sample_order as so
        src = _draw_dark_art(os.path.join(self.tmp.name, "dark_art.png"), (700, 590))
        self.assertEqual(so._scrub_color_annotations(src), src)
        self.assertGreater(self._sat_pixels(src), 0)   # 印花還在

    def test_unreadable_file_falls_back_to_original(self):
        from skills import sample_order as so
        bad = os.path.join(self.tmp.name, "broken.png")
        with open(bad, "wb") as f:
            f.write(b"not a png")
        self.assertEqual(so._scrub_color_annotations(bad), bad)   # 洗不了也不擋生圖

    def test_wired_into_generate_from_order(self):
        """挑完線稿一定要先洗再餵模型——漏接這一步等於整個修正沒生效。"""
        import inspect
        from skills import sample_order as so
        src = inspect.getsource(so.generate_from_order)
        self.assertIn("_scrub_color_annotations(sketch)", src)


class RenderedDigitLabelsTests(unittest.TestCase):
    """成品圖 OCR 只挑純數字 token；品牌字樣不算，OCR 不可用不擋。"""

    def _run(self, ocr_return):
        from skills import sample_order as so
        with mock.patch("agent_core.vision_ocr.ocr_image", return_value=ocr_return):
            return so._rendered_digit_labels("/tmp/whatever.png")

    def test_digits_flagged_brand_text_ignored(self):
        # 2026-08-04 order_20260804_102606.png 的實際 OCR 結果
        self.assertEqual(self._run({"text": "2 1 1 1 Richter 1 1 2 1"}),
                         ["2", "1", "1", "1", "1", "1", "2", "1"])

    def test_clean_render_reads_only_brand(self):
        self.assertEqual(self._run({"text": "Richter R"}), [])

    def test_ocr_unavailable_does_not_block(self):
        self.assertEqual(self._run(None), [])

    def test_ocr_exception_does_not_block(self):
        from skills import sample_order as so
        with mock.patch("agent_core.vision_ocr.ocr_image", side_effect=RuntimeError("boom")):
            self.assertEqual(so._rendered_digit_labels("/tmp/x.png"), [])


class RenderRetryOnDigitsTests(unittest.TestCase):
    """OCR 讀到數字就重生一張；重試照樣扣配額；全失敗要照實講。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.gen = os.path.join(self.tmp.name, "generated_images")
        self.state = os.path.join(self.tmp.name, "state")
        os.makedirs(self.gen, exist_ok=True)
        os.makedirs(self.state, exist_ok=True)
        self._p = [
            mock.patch("agent_core.logging_and_paths.GENERATED_IMAGES_DIR", self.gen),
            mock.patch("agent_core.logging_and_paths.STATE_DIR", self.state),
        ]
        for p in self._p:
            p.start()
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["style number richter", "8105-3272"])
        self.xlsx = os.path.join(self.tmp.name, "o.xlsx")
        wb.save(self.xlsx)
        self.sketch = _draw_annotated(os.path.join(self.tmp.name, "sk.png"), (442, 225))

    def tearDown(self):
        for p in self._p:
            p.stop()
        self.tmp.cleanup()

    def _fake_gemini(self):
        part = mock.MagicMock()
        part.inline_data.data = b"\x89PNG fake"
        cand = mock.MagicMock()
        cand.content.parts = [part]
        resp = mock.MagicMock(text="鞋身 深藍、拼接 青綠", candidates=[cand])
        fake = mock.MagicMock()
        fake.models.generate_content.return_value = resp
        return mock.patch("agent_core.gemini_client._get_gemini_client", return_value=fake)

    def _run(self, digit_sequence, env=None):
        """digit_sequence：每次 OCR 依序回傳的數字清單（末項會重複沿用）。"""
        from skills import sample_order as so
        seq = list(digit_sequence)

        def _ocr(_path):
            return seq.pop(0) if len(seq) > 1 else seq[0]

        with mock.patch.dict(os.environ, env or {}), \
                mock.patch.object(so, "_extract_order_images", return_value=[self.sketch]), \
                mock.patch.object(so, "_dept_scope_color", return_value="green"), \
                mock.patch.object(so, "_rendered_digit_labels", side_effect=_ocr), \
                self._fake_gemini() as client:
            out = so.generate_from_order(self.xlsx)
        # 打了幾次生圖 = 總呼叫數扣掉 _spec_for_render 那一發
        renders = client.return_value.models.generate_content.call_count - 1
        return out, renders

    def test_clean_first_try_renders_once(self):
        out, renders = self._run([[]])
        self.assertEqual(renders, 1)
        self.assertNotIn("殘留部位標號", out)

    def test_dirty_render_is_retried_until_clean(self):
        out, renders = self._run([["1", "2"], []])
        self.assertEqual(renders, 2)
        self.assertNotIn("殘留部位標號", out)   # 交出去的是乾淨那張
        self.assertIn("[[TG_PHOTO:", out)

    def test_retry_prompt_names_the_miss(self):
        from skills import sample_order as so
        self.assertNotIn("上一版", so._render_prompt("spec", retry=False))
        self.assertIn("上一版", so._render_prompt("spec", retry=True))

    def test_spec_prompt_forbids_part_numbering(self):
        """規格會原封不動進圖生圖 prompt——留著「部位 1」模型就會把數字畫上去。"""
        from skills import sample_order as so
        with mock.patch("agent_core.gemini_client.generate_content_tracked",
                        return_value=mock.MagicMock(text="鞋身深藍")) as g:
            so._spec_for_render("樣品單內容", "7202")
        self.assertIn("不可出現編號", g.call_args.kwargs["contents"][0])

    def test_all_attempts_dirty_says_so_instead_of_pretending(self):
        out, renders = self._run([["1"]], env={"RED_ORDER_RENDER_MAX_ATTEMPTS": "3"})
        self.assertEqual(renders, 3)
        self.assertIn("殘留部位標號", out)
        self.assertIn("[[TG_PHOTO:", out)      # 還是交件，只是講清楚

    def test_retries_consume_quota_and_stop_when_exhausted(self):
        """重試也是真的在花錢：不扣配額的話，日花費上限會被偷偷放大 N 倍。"""
        out, renders = self._run([["1"]], env={"RED_ORDER_RENDER_DAILY_MAX": "2",
                                               "RED_ORDER_RENDER_MAX_ATTEMPTS": "5"})
        self.assertEqual(renders, 2)           # 配額 2 張就停手，不是硬跑 5 次
        self.assertIn("殘留部位標號", out)
        with open(os.path.join(self.state, "order_render_quota.json"), encoding="utf-8") as f:
            self.assertEqual(list(json.load(f).values())[0]["green"], 2)


def _draw_colored(path: str, size: tuple[int, int]) -> str:
    """畫一張「上色完稿」樣式圖：白底 + 分離的彩色色塊。

    色塊刻意不填滿內容 bbox——真完稿是不規則鞋形，去白邊後 bbox 內仍有大量
    白底（PSS 8 實測 47%），這正是「設計稿 vs 實照」的判別訊號。
    """
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, (255, 255, 255))
    d = ImageDraw.Draw(im)
    w, h = size
    d.rectangle((int(w * .1), int(h * .2), int(w * .5), int(h * .7)), fill=(121, 100, 81))
    d.rectangle((int(w * .6), int(h * .5), int(w * .9), int(h * .85)), fill=(236, 122, 45))
    im.save(path)
    return path


def _draw_dark_art(path: str, size: tuple[int, int]) -> str:
    """畫一張「全黑款上色完稿」樣式圖：白底 + L 形大面積深色填色 + 一小塊彩色印花。

    彩色佔比刻意壓在 _ART_COLOR_MIN 之下（PSS 7 實測 0.145）——判別只能靠
    深色填色佔比；L 形（像鞋側影）讓去白邊後 bbox 內仍留白底。
    """
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, (255, 255, 255))
    d = ImageDraw.Draw(im)
    w, h = size
    d.rectangle((int(w * .1), int(h * .15), int(w * .5), int(h * .75)), fill=(20, 20, 22))
    d.rectangle((int(w * .1), int(h * .55), int(w * .8), int(h * .75)), fill=(20, 20, 22))
    d.rectangle((int(w * .55), int(h * .25), int(w * .65), int(h * .35)), fill=(200, 80, 200))
    im.save(path)
    return path


def _draw_photo(path: str, size: tuple[int, int]) -> str:
    """畫一張「實照」樣式圖：整張填滿彩色、零白底（studio 照片的樣子）。"""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, (150, 120, 90))
    ImageDraw.Draw(im).rectangle((10, 10, size[0] - 10, size[1] - 10), fill=(170, 100, 60))
    im.save(path)
    return path


class SketchTierTests(unittest.TestCase):
    """挑圖三級優先（2026-09-01 PSS 8 案）：上色完稿 > 黑白線稿 > 照片類。

    實照的 Vision 置信天生比設計稿高（實測 0.41 vs 0.16），純比分數必挑到
    exsample 實照——但客人要的設計在完稿上。這組測試釘死優先序。
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = self.tmp.name
        self.art = _draw_colored(os.path.join(d, "art.png"), (700, 590))
        self.lineart = _draw(os.path.join(d, "line.png"), (700, 590), (20, 20, 680, 570))
        self.photo = _draw_photo(os.path.join(d, "photo.png"), (600, 540))

    def tearDown(self):
        self.tmp.cleanup()

    def _pick(self, pngs, confs):
        from skills import sample_order as so
        with mock.patch.object(so, "_shoe_confidence",
                               side_effect=lambda p: confs.get(os.path.basename(p).replace("crop_", ""), 0.0)):
            return so._pick_shoe_sketch(pngs)

    def test_colored_art_beats_higher_confidence_photo(self):
        pick = self._pick([self.photo, self.art],
                          {"art.png": 0.16, "photo.png": 0.41})
        self.assertEqual(os.path.basename(pick).replace("crop_", ""), "art.png")

    def test_line_art_beats_photo_when_no_colored_art(self):
        pick = self._pick([self.photo, self.lineart],
                          {"line.png": 0.08, "photo.png": 0.41})
        self.assertEqual(os.path.basename(pick).replace("crop_", ""), "line.png")

    def test_colored_art_beats_line_art(self):
        pick = self._pick([self.lineart, self.art],
                          {"art.png": 0.16, "line.png": 0.30})
        self.assertEqual(os.path.basename(pick).replace("crop_", ""), "art.png")

    def test_dark_painted_art_beats_line_art(self):
        """PSS 7 型：全黑完稿彩色佔比不夠，靠深色填色佔比仍要贏黑白線稿。"""
        dark = _draw_dark_art(os.path.join(self.tmp.name, "dark.png"), (700, 590))
        pick = self._pick([self.lineart, dark],
                          {"dark.png": 0.14, "line.png": 0.30})
        self.assertEqual(os.path.basename(pick).replace("crop_", ""), "dark.png")


class XyCutTests(unittest.TestCase):
    """XY-cut 白縫切割：logo/標題/主圖分家（PDF 頁面渲染的前處理）。"""

    def test_two_blocks_split_by_gap(self):
        import numpy as np
        from skills.sample_order import _xy_cut_boxes
        ink = np.zeros((400, 600), dtype=bool)
        ink[20:120, 20:200] = True       # 標題塊
        ink[200:380, 100:560] = True     # 主圖塊
        boxes = sorted(_xy_cut_boxes(ink))
        self.assertEqual(len(boxes), 2)
        self.assertEqual(boxes[0], (20, 20, 200, 120))
        self.assertEqual(boxes[1], (100, 200, 560, 380))

    def test_single_block_returns_trimmed_bbox(self):
        import numpy as np
        from skills.sample_order import _xy_cut_boxes
        ink = np.zeros((300, 300), dtype=bool)
        ink[50:250, 60:260] = True
        self.assertEqual(_xy_cut_boxes(ink), [(60, 50, 260, 250)])


class PdfExtractTests(unittest.TestCase):
    """PDF 樣品單抽圖：整頁渲染＋切塊（向量線稿只能這樣拿）＋內嵌 raster。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def test_pdf_pages_yield_segmented_candidates(self):
        from PIL import Image, ImageDraw
        from skills.sample_order import _extract_order_images
        page = Image.new("RGB", (800, 1000), (255, 255, 255))
        d = ImageDraw.Draw(page)
        d.rectangle((40, 40, 300, 120), fill=(30, 30, 30))       # 標題塊
        d.rectangle((100, 300, 700, 800), fill=(121, 100, 81))   # 主圖塊
        pdf = os.path.join(self.tmp.name, "order.pdf")
        page.save(pdf, format="PDF")
        pngs = _extract_order_images(pdf)
        self.assertGreaterEqual(len(pngs), 2)                    # 至少切出兩個內容塊
        segs = [p for p in pngs if "seg" in os.path.basename(p)]
        self.assertGreaterEqual(len(segs), 2)
        sizes = []
        for p in segs:
            with Image.open(p) as im:
                sizes.append(im.size)
        self.assertTrue(any(w > 1000 for w, _ in sizes))         # 主圖塊（scale=2.0 渲染）

    def test_non_pdf_non_xlsx_returns_empty(self):
        from skills.sample_order import _extract_order_images
        self.assertEqual(_extract_order_images("/tmp/whatever.docx"), [])


class BasePhotoLookupTests(unittest.TestCase):
    """_find_base_photo：只查本機產品照索引、檔名前綴優先；索引沒建絕不炸。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.makedirs(os.path.join(self.tmp.name, "product_photos"), exist_ok=True)
        with open(os.path.join(self.tmp.name, "product_photos", "meta.json"), "w",
                  encoding="utf-8") as f:
            json.dump([
                {"id": "b", "file": "舊款5008-4292參考.jpg", "folder": "R-JAJLCG"},
                {"id": "a", "file": "5008-4292-6301.jpg", "folder": "Richter/FC-Husky 2.0/22w"},
                {"id": "c", "file": "CG2302V-03.jpg", "folder": "R-JAJLCG"},
                {"id": "d", "file": "PSS7.jpg", "folder": "Richter/雪靴/20260827寄雪靴開發樣品"},
            ], f, ensure_ascii=False)
        self._p = mock.patch("agent_core.logging_and_paths.DATA_DIR", self.tmp.name)
        self._p.start()

    def tearDown(self):
        self._p.stop()
        self.tmp.cleanup()

    def test_prefix_match_beats_substring(self):
        from skills import sample_order as so
        self.assertEqual(so._find_base_photo("5008-4292")["id"], "a")

    def test_substring_fallback_when_no_prefix(self):
        from skills import sample_order as so
        self.assertEqual(so._find_base_photo("4292-6301")["id"], "a")

    def test_style_with_space_matches_compact_filename(self):
        # 2026-09-02 PSS 7 案：客人單上寫「PSS 7」、寄樣照檔名是「PSS7.jpg」
        # ——比對雙方都去空白才對得上。
        from skills import sample_order as so
        self.assertEqual(so._find_base_photo("PSS 7")["id"], "d")

    def test_unknown_or_short_style_returns_none(self):
        from skills import sample_order as so
        self.assertIsNone(so._find_base_photo("9999-0000"))
        self.assertIsNone(so._find_base_photo("22"))       # 太短＝亂中一氣，不查
        self.assertIsNone(so._find_base_photo(""))

    def test_missing_index_returns_none(self):
        from skills import sample_order as so
        with mock.patch("agent_core.logging_and_paths.DATA_DIR",
                        os.path.join(self.tmp.name, "nonexistent")):
            self.assertIsNone(so._find_base_photo("5008-4292"))


class RenderColorAnchorTests(unittest.TestCase):
    """Phase A 色錨（2026-09-01）：規格結構化→色票當第二張輸入圖→ΔE 驗證→偏色重生。

    釘死三件事：① Pantone 色號查表值蓋掉 LLM 的 hex 推估（顏色不隨模型心情漂）
    ② 偏色會帶著「哪個部位、應為何色」的糾正語重生 ③ 重生到底仍偏色要照實講。
    """

    _SPEC_JSON = json.dumps({
        "描述": "鞋身麂皮經典藍、大底橡膠",
        "款號": "8105-3272",
        "部位": [
            {"部位": "鞋身", "材質": "麂皮", "顏色": "PANTONE 19-4052 TCX", "hex": "#123456"},
            {"部位": "大底", "材質": "橡膠", "顏色": "說不上來的顏色", "hex": "#a1b2c3"},
        ],
    }, ensure_ascii=False)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.gen = os.path.join(self.tmp.name, "generated_images")
        self.state = os.path.join(self.tmp.name, "state")
        os.makedirs(self.gen, exist_ok=True)
        os.makedirs(self.state, exist_ok=True)
        self._p = [
            mock.patch("agent_core.logging_and_paths.GENERATED_IMAGES_DIR", self.gen),
            mock.patch("agent_core.logging_and_paths.STATE_DIR", self.state),
        ]
        for p in self._p:
            p.start()
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(["style number richter", "8105-3272"])
        # 完稿鎖色接地濾網（PSS 6 案）：色票的顏色敘述必須真的寫在單上——
        # 模擬客人把 Pantone 碼寫進規格；「說不上來的顏色」刻意不寫（腦補色）。
        wb.active.append(["upper suede", "PANTONE 19-4052 TCX"])
        self.xlsx = os.path.join(self.tmp.name, "o.xlsx")
        wb.save(self.xlsx)
        self.sketch = _draw_annotated(os.path.join(self.tmp.name, "sk.png"), (442, 225))

    def tearDown(self):
        for p in self._p:
            p.stop()
        self.tmp.cleanup()

    def _fake_gemini(self, spec_text: str):
        """第一通回規格（spec_text）、之後每通回帶假圖的渲染回應。"""
        img_part = mock.MagicMock()
        img_part.inline_data.data = b"\x89PNG fake"
        cand = mock.MagicMock()
        cand.content.parts = [img_part]
        spec_resp = mock.MagicMock(text=spec_text, candidates=[])
        render_resp = mock.MagicMock(text="", candidates=[cand])
        fake = mock.MagicMock()
        fake.models.generate_content.side_effect = (
            lambda *a, **k: spec_resp
            if fake.models.generate_content.call_count == 1 else render_resp)
        return mock.patch("agent_core.gemini_client._get_gemini_client", return_value=fake)

    def _run(self, miss_sequence, spec_text=None, env=None, base_hit=None, base_bytes=b"",
             sketch=None, tool_kwargs=None):
        """miss_sequence：find_color_misses 每次依序回的清單（末項重複沿用）。

        _find_base_photo 預設打成 None——測試不可讀真索引（DATA_DIR 內容不可控）；
        要測形體參考接線就給 base_hit（+base_bytes 模擬 Drive 下載）。
        """
        from skills import sample_order as so
        seq = list(miss_sequence)
        self._fm_targets = []

        def _misses(_path, targets, _de):
            self._fm_targets = targets
            return seq.pop(0) if len(seq) > 1 else seq[0]

        with mock.patch.dict(os.environ, env or {}), \
                mock.patch.object(so, "_extract_order_images",
                                  return_value=[sketch or self.sketch]), \
                mock.patch.object(so, "_dept_scope_color", return_value="green"), \
                mock.patch.object(so, "_rendered_digit_labels", return_value=[]), \
                mock.patch.object(so, "_find_base_photo", return_value=base_hit), \
                mock.patch.object(so, "_download_drive_image",
                                  return_value=base_bytes) as dl, \
                mock.patch("agent_core.color_anchor.find_color_misses",
                           side_effect=_misses) as fm, \
                self._fake_gemini(spec_text or self._SPEC_JSON) as client:
            out = so.generate_from_order(self.xlsx, **(tool_kwargs or {}))
        calls = client.return_value.models.generate_content.call_args_list
        self._dl = dl
        return out, calls, fm

    def _prev_render(self, subdir=("dept", "green"), name="order_prev.png"):
        from PIL import Image
        d = os.path.join(self.gen, *subdir)
        os.makedirs(d, exist_ok=True)
        p = os.path.join(d, name)
        Image.new("RGB", (100, 80), (40, 45, 60)).save(p)
        return p

    def test_edit_mode_uses_previous_render_as_base(self):
        """修圖模式（2026-09-07 大王條件二）：上一版成品當第一張輸入圖、只修指正
        部位、其餘鎖定不動；不附同款實照；摘要標明修圖模式。"""
        prev = self._prev_render()
        out, calls, _ = self._run(
            [[]], base_hit={"id": "F1", "file": "DAVOS-PSS2.jpg"},
            tool_kwargs={"previous_render_path": prev,
                         "extra_note": "下面那條魔鬼氈材質不對，應該跟上面一樣"})
        contents = calls[1].kwargs["contents"]
        self.assertEqual(len(contents), 4)         # prev + 線稿 + 色票 + prompt（無實照）
        prompt = contents[-1]
        self.assertIn("上一版的渲染成品", prompt[:40])
        self.assertIn("只修改", prompt)
        self.assertIn("保持完全一致", prompt)
        self.assertIn("下面那條魔鬼氈", prompt)
        self._dl.assert_not_called()               # 修圖模式不附同款實照
        self.assertIn("🖌️ 修圖模式", out)

    def test_edit_mode_requires_note_and_scoped_path(self):
        """修圖模式沒帶指正原話→拒絕；路徑超出部門生圖範圍→拒絕（都不打 LLM）。"""
        prev = self._prev_render()
        out, calls, _ = self._run([[]], tool_kwargs={"previous_render_path": prev})
        self.assertIn("❌", out)
        self.assertEqual(calls, [])                # 連規格 LLM 都沒打、不扣配額
        outside = self._prev_render(subdir=(), name="order_outside.png")  # gen 根目錄
        out2, calls2, _ = self._run(
            [[]], tool_kwargs={"previous_render_path": outside, "extra_note": "改鞋頭"})
        self.assertIn("❌", out2)                   # green context 只能碰 dept/green
        self.assertEqual(calls2, [])

    def test_edit_mode_verification_defers_to_user(self):
        """review V1：修圖是使用者主導的改動（可能刻意偏離完稿）——ΔE 與完稿
        外來色/正側面比對分不出「要求的改動」和「偏色」，修圖模式停用，
        只留數字 OCR 與「單隻」檢查；視角以上一版為準。"""
        from skills import sample_order as so
        prev = self._prev_render(name="order_editv.png")
        art = _draw_colored(os.path.join(self.tmp.name, "art_ev.png"), (700, 590))
        with mock.patch.object(so, "_extract_order_images", return_value=[art]), \
                mock.patch.object(so, "_dept_scope_color", return_value="green"), \
                mock.patch.object(so, "_rendered_digit_labels", return_value=[]), \
                mock.patch.object(so, "_find_base_photo", return_value=None), \
                mock.patch.object(so, "_composition_flaws",
                                  return_value=["視角不是正側面"]) as cf, \
                mock.patch("agent_core.color_anchor.find_color_misses") as fm, \
                self._fake_gemini(self._SPEC_JSON) as client:
            out = so.generate_from_order(self.xlsx, extra_note="下面魔鬼氈改布料",
                                         previous_render_path=prev)
        fm.assert_not_called()                                 # ΔE 停用
        self.assertEqual(cf.call_args.kwargs["artwork_path"], "")  # 不與完稿比對
        calls = client.return_value.models.generate_content.call_args_list
        self.assertEqual(len(calls), 2)                        # 視角 flaw 被濾掉→不重試
        self.assertNotIn("視角不是正側面", out)
        self.assertNotIn("色準驗證", out)                       # ΔE 那行整個不出現
        self.assertNotIn("色準未達標", out)
        self.assertIn("由你驗收", out)                          # 改用誠實聲明

    def test_edit_mode_prompt_drops_full_repaint_clauses(self):
        """review：修圖模式的 prompt 不得殘留「整張重上色」指令——完稿整套配色
        與「正側面」硬鎖都要讓位給「與第一張圖一致」。"""
        art = _draw_colored(os.path.join(self.tmp.name, "art_ep.png"), (700, 590))
        prev = self._prev_render(name="order_editp.png")
        _out, calls, _ = self._run(
            [[]], sketch=art,
            tool_kwargs={"previous_render_path": prev, "extra_note": "後套改寶藍"})
        prompt = calls[1].kwargs["contents"][-1]
        self.assertNotIn("整體配色以完稿本身的顏色為準", prompt)
        self.assertNotIn("正側面視角", prompt)
        self.assertIn("維持第一張圖的視角", prompt)
        self.assertIn("未被點名的部位一律維持第一張圖原樣", prompt)

    def test_long_extra_note_truncation_is_reported(self):
        """review V2：>1000 字的修正指示不再被靜默腰斬——截斷要在摘要明講。"""
        _out, calls, _ = self._run([[]], tool_kwargs={"extra_note": "A" * 1100})
        prompt = calls[1].kwargs["contents"][-1]
        self.assertIn("A" * 1000, prompt)
        self.assertNotIn("A" * 1001, prompt)
        self.assertIn("超過 1000 字", _out)

    def test_edit_mode_rejects_unreadable_previous_render(self):
        """review V3：壞檔/半寫入 png 過去會拖到扣完配額才籠統爆掉——
        驗證時就 load() 探測、給明確訊息且不打任何 LLM。"""
        prev = self._prev_render(name="order_corrupt.png")
        with open(prev, "wb") as f:
            f.write(b"\x89PNG fake")
        out, calls, _ = self._run([[]], tool_kwargs={"previous_render_path": prev,
                                                     "extra_note": "改鞋頭"})
        self.assertIn("讀不開", out)
        self.assertEqual(calls, [])

    def test_extra_note_gets_viewpoint_guard(self):
        """一般模式帶修正指示 → 尾端補視角護欄（2026-09-07 大王條件一：聊天層曾
        自加「45 度立體展示視角」進 extra_note，把正側面鎖定蓋掉生出組圖）。"""
        _out, calls, _ = self._run([[]], tool_kwargs={"extra_note": "鞋頭改平滑皮料"})
        self.assertIn("視角一律維持與樣品單設計稿相同", calls[1].kwargs["contents"][-1])

    def test_swatch_becomes_second_input_image(self):
        out, calls, _ = self._run([[]])
        self.assertEqual(len(calls), 2)                    # 規格 1 + 渲染 1
        contents = calls[1].kwargs["contents"]
        self.assertEqual(len(contents), 3)                 # 線稿 + 色票 + prompt
        self.assertIn("指定色票", contents[2])
        self.assertIn("鞋身", contents[2])

    def test_pantone_table_beats_llm_hex_guess(self):
        out, _, _ = self._run([[]])
        self.assertIn("0f4c81", out)                       # 19-4052 查表值
        self.assertNotIn("123456", out)                    # LLM 推估被蓋掉
        self.assertIn("Pantone 對照", out)
        self.assertIn("a1b2c3", out)                       # 查不到的色才收 LLM 推估
        self.assertIn("LLM 推估", out)

    def test_color_miss_retries_with_named_correction(self):
        miss = [{"part": "鞋身", "hex": "0f4c81", "delta_e": 45.0}]
        out, calls, _ = self._run([miss, []])
        self.assertEqual(len(calls), 3)                    # 規格 1 + 渲染 2
        self.assertIn("偏離指定色票", calls[2].kwargs["contents"][2])
        self.assertIn("鞋身", calls[2].kwargs["contents"][2])
        self.assertIn("色準驗證", out)                      # 最後那張過了就說過了
        self.assertNotIn("色準未達標", out)

    def test_persistent_miss_reported_honestly(self):
        miss = [{"part": "鞋身", "hex": "0f4c81", "delta_e": 45.0}]
        out, calls, _ = self._run([miss], env={"RED_ORDER_RENDER_MAX_ATTEMPTS": "3"})
        self.assertEqual(len(calls), 4)                    # 規格 1 + 渲染 3（打滿）
        self.assertIn("色準未達標", out)
        self.assertIn("實體色卡", out)                      # 誠實界線要講
        self.assertIn("[[TG_PHOTO:", out)                  # 還是交件

    def test_env_switch_disables_color_check(self):
        out, calls, fm = self._run([[]], env={"RED_ORDER_COLOR_CHECK": "0"})
        self.assertEqual(len(calls), 2)
        fm.assert_not_called()
        self.assertNotIn("色準驗證", out)
        self.assertIn("0f4c81", out)                       # 色錨/色票照常運作

    def test_non_json_spec_falls_back_to_single_image(self):
        out, calls, fm = self._run([[]], spec_text="鞋身 深藍、拼接 青綠")
        self.assertEqual(len(calls[1].kwargs["contents"]), 2)   # 線稿 + prompt（無色票）
        fm.assert_not_called()                             # 無色錨就無從驗
        self.assertNotIn("🎨 色錨", out)

    def test_base_photo_attached_when_index_has_same_style(self):
        """Phase C：索引裡有同款實照 → 附為形體參考、prompt 講明只抄形體不抄配色。"""
        import io as _io
        from PIL import Image
        buf = _io.BytesIO()
        Image.new("RGB", (80, 60), (90, 90, 90)).save(buf, format="PNG")
        out, calls, _ = self._run(
            [[]], base_hit={"id": "F1", "file": "8105-3272-6301.jpg"},
            base_bytes=buf.getvalue())
        self._dl.assert_called_once_with("F1")
        contents = calls[1].kwargs["contents"]
        self.assertEqual(len(contents), 4)                 # 線稿 + 同款實照 + 色票 + prompt
        self.assertIn("同款（款號 8105-3272）", contents[3])
        self.assertIn("大底齒紋", contents[3])
        self.assertIn("不要照抄實照", contents[3])
        self.assertIn("📸 形體參考已附", out)
        self.assertIn("8105-3272-6301.jpg", out)

    def test_colored_art_suppresses_base_photo(self):
        """完稿在場不附同款實照（2026-09-03 PSS 2 案：索引命中的 8/27 寄樣舊配色
        被 nano-banana 整套照抄，正確色票＋ΔE 糾正語三 attempt 全拉不回）。"""
        import io as _io
        from PIL import Image
        buf = _io.BytesIO()
        Image.new("RGB", (80, 60), (70, 110, 120)).save(buf, format="PNG")
        art = _draw_colored(os.path.join(self.tmp.name, "art_base.png"), (700, 590))
        out, calls, _ = self._run([[]], base_hit={"id": "F1", "file": "DAVOS-PSS2.jpg"},
                                  base_bytes=buf.getvalue(), sketch=art)
        self._dl.assert_not_called()                       # 連 Drive 下載都不發生
        self.assertNotIn("同款", calls[1].kwargs["contents"][-1])
        self.assertNotIn("📸", out)

    def test_swatch_prompt_forbids_literal_color_names(self):
        """客戶專有色名字面常是錯的（lagoon=寶藍非「環礁湖」藍綠）→ 色票指令明令
        以色塊為準、不得按字面聯想。"""
        _out, calls, _ = self._run([[]])
        self.assertIn("不得按色名的字面意義自行聯想", calls[1].kwargs["contents"][-1])

    def test_single_shoe_side_view_leads_the_prompt(self):
        """「單隻、同視角」要在 prompt 開頭定調（2026-09-03 PSS 2 案：埋句尾時
        nano-banana 三連發四視角組圖，構圖驗證打回也拉不回）。"""
        _out, calls, _ = self._run([[]])
        prompt = calls[1].kwargs["contents"][-1]
        self.assertIn("畫面裡只有這一隻鞋", prompt[:160])
        self.assertIn("不要畫成多視角組圖", prompt)

    def test_colored_art_enforces_paired_part_consistency(self):
        """完稿上畫法相同的成對部件（兩條魔鬼氈）→ 同材質同色，材質詞不得落到
        完稿沒畫出該質感的部位（PSS 2 案：Web.tape 被猜到下條魔鬼氈）。"""
        art = _draw_colored(os.path.join(self.tmp.name, "art_pair.png"), (700, 590))
        _out, calls, _ = self._run([[]], sketch=art)
        prompt = calls[1].kwargs["contents"][-1]
        self.assertIn("畫法相同的重複部件", prompt)
        self.assertIn("完稿沒有畫出的部位不得自行分配材質", prompt)

    def test_spec_parse_prompt_bans_color_paraphrase_and_material_guessing(self):
        """規格解析 LLM：色名照抄不意譯（lagoon≠「環礁湖藍綠」）、沒點名部位的
        材質不得猜給魔鬼氈/鞋身。"""
        _out, calls, _ = self._run([[]])
        spec_prompt = calls[0].kwargs["contents"][0]
        self.assertIn("不得意譯成中文色系", spec_prompt)
        self.assertIn("不得自行猜給魔鬼氈", spec_prompt)

    def test_base_photo_download_failure_degrades_silently(self):
        """Drive 下載失敗＝加分項拿不到：不擋生圖、不多算輸入圖、不假稱有附。"""
        out, calls, _ = self._run([[]], base_hit={"id": "F1", "file": "x.jpg"},
                                  base_bytes=b"")
        self.assertEqual(len(calls[1].kwargs["contents"]), 3)   # 線稿 + 色票 + prompt
        self.assertNotIn("📸", out)

    def test_colored_art_sketch_locks_palette(self):
        """PSS 8 型：上色完稿當 sketch → prompt 鎖完稿配色、完稿主色進 ΔE 驗證。"""
        art = _draw_colored(os.path.join(self.tmp.name, "art_sk.png"), (700, 590))
        out, calls, fm = self._run([[]], sketch=art)
        prompt = calls[1].kwargs["contents"][-1]
        self.assertIn("已上色的鞋款設計完稿", prompt)
        self.assertIn("完稿本身的顏色", prompt)
        fm.assert_called_once()
        art_parts = [t for t in self._fm_targets if str(t["part"]).startswith("完稿主色")]
        self.assertTrue(art_parts)                          # 完稿主色進了驗證
        self.assertGreater(len(self._fm_targets), len(art_parts))   # 規格色錨也還在
        self.assertIn("🖼️ 客人上色完稿", out)

    def test_colored_sketch_prompt_maps_texture_to_material(self):
        """PSS 8 三輪：織紋蔓延到平滑皮料部位 → prompt 明定「紋理分佈=材質分佈」。"""
        art = _draw_colored(os.path.join(self.tmp.name, "art_tex.png"), (700, 590))
        _out, calls, _ = self._run([[]], sketch=art)
        prompt = calls[1].kwargs["contents"][-1]
        self.assertIn("平滑皮料", prompt)
        self.assertIn("紋理不得蔓延", prompt)

    def test_extra_note_reaches_render_prompt_and_summary(self):
        """使用者對前一版的修正原話 → 直接進渲染指令＋輸出摘要。"""
        from skills import sample_order as so
        note = "鞋面上沒有紋路的是皮料，有紋路的是布料"
        seq = [[]]
        with mock.patch.object(so, "_extract_order_images", return_value=[self.sketch]), \
                mock.patch.object(so, "_dept_scope_color", return_value="green"), \
                mock.patch.object(so, "_rendered_digit_labels", return_value=[]), \
                mock.patch.object(so, "_find_base_photo", return_value=None), \
                mock.patch("agent_core.color_anchor.find_color_misses",
                           side_effect=lambda *a: seq[0]), \
                self._fake_gemini(self._SPEC_JSON) as client:
            out = so.generate_from_order(self.xlsx, extra_note=note)
        calls = client.return_value.models.generate_content.call_args_list
        self.assertIn(note, calls[1].kwargs["contents"][-1])
        self.assertIn("使用者對前一版的修正", calls[1].kwargs["contents"][-1])
        self.assertIn("✏️ 已套用修正指示", out)

    def test_colored_sketch_drops_llm_guessed_swatch(self):
        """完稿在場時 LLM 推估色不進色票/驗證（PSS 8 實測：LLM 猜的淺米色票
        蓋過完稿棕灰主色）——模型猜的永遠不能贏過客人自己畫的。"""
        art = _draw_colored(os.path.join(self.tmp.name, "art_sk3.png"), (700, 590))
        out, calls, fm = self._run([[]], sketch=art)
        hexes = [t["hex"] for t in self._fm_targets]
        self.assertNotIn("a1b2c3", hexes)              # LLM 推估（大底）被剔除
        self.assertIn("0f4c81", hexes)                 # Pantone 查表（鞋身）保留
        self.assertNotIn("a1b2c3", out)
        self.assertNotIn("LLM 推估", out)

    def test_dark_painted_art_locks_palette_too(self):
        """PSS 7 型：全黑完稿（彩色佔比不夠、靠深色填色認）也要進完稿鎖色——
        LLM 推估色不進色票/驗證、完稿主色進 ΔE、prompt 鎖完稿配色。"""
        art = _draw_dark_art(os.path.join(self.tmp.name, "dark_art.png"), (700, 590))
        out, calls, fm = self._run([[]], sketch=art)
        self.assertIn("完稿本身的顏色", calls[1].kwargs["contents"][-1])
        self.assertIn("🖼️ 客人上色完稿", out)
        fm.assert_called_once()
        self.assertNotIn("a1b2c3", [t["hex"] for t in self._fm_targets])   # LLM 推估被剔除
        self.assertTrue(any(str(t["part"]).startswith("完稿主色") for t in self._fm_targets))

    def test_colored_sketch_verifies_even_without_spec_targets(self):
        """規格 LLM 掉鏈（非 JSON）但完稿有色 → 仍用完稿主色驗色，不是整段裸奔。"""
        art = _draw_colored(os.path.join(self.tmp.name, "art_sk2.png"), (700, 590))
        out, calls, fm = self._run([[]], spec_text="鞋身 深藍、拼接 青綠", sketch=art)
        fm.assert_called_once()
        self.assertTrue(all(str(t["part"]).startswith("完稿主色") for t in self._fm_targets))
        self.assertIn("已上色的鞋款設計完稿", calls[1].kwargs["contents"][-1])

    def test_material_swatch_joins_inputs_when_library_has_it(self):
        """Phase B：庫裡有麂皮特寫 → 多一張輸入圖、prompt 講明只取紋理不取顏色。"""
        from PIL import Image
        lib_root = os.path.join(self.tmp.name, "data")
        os.makedirs(os.path.join(lib_root, "material_swatches"), exist_ok=True)
        Image.new("RGB", (64, 64), (120, 90, 60)).save(
            os.path.join(lib_root, "material_swatches", "suede.png"))
        with mock.patch("agent_core.logging_and_paths.DATA_DIR", lib_root):
            out, calls, _ = self._run([[]])
        contents = calls[1].kwargs["contents"]
        self.assertEqual(len(contents), 4)                 # 線稿 + 材質特寫 + 色票 + prompt
        self.assertIn("材質特寫", contents[3])
        self.assertIn("只取紋理", contents[3])
        self.assertIn("色票為準", contents[3])
        self.assertIn("🧵 材質特寫已附", out)
        self.assertIn("麂皮→鞋身", out)


if __name__ == "__main__":
    unittest.main()
