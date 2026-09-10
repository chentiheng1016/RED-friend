"""Phase 2：客人樣品單(Excel/PDF)結構化解析——LLM 抽出每款的款號/顏色/逐部位材料規格。

樣品單兩型都吃：表格型(型號|顏色|部位|材料|材料顏色)、文字描述型(款號+逐部位文字)。
用 LLM 讀全份提取（不硬編表格位置）。抽出的款號可對接 search_product_photos 找舊款、餵生成。
"""
import io
import json
import os
import subprocess
import tempfile


# launchd daemon 的 PATH 沒有 /Applications 下的 app bundle，逐一探測
# （比照 drive_sync._SOFFICE_CANDIDATES）。
_SOFFICE_CANDIDATES = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
)


def _find_soffice() -> str:
    import shutil
    found = shutil.which("soffice")
    if found:
        return found
    for candidate in _SOFFICE_CANDIDATES:
        if os.path.exists(candidate):
            return candidate
    return ""


def _soffice_convert(src: str, fmt: str, outdir: str, timeout: int = 120) -> str:
    """soffice headless 轉檔 → 回輸出檔路徑；任何失敗回 ""（絕不 raise）。

    兩個非做不可的細節（2026-08-04 實測踩到，drive_sync 早就解過）：
      1. **每次轉檔用獨立 profile**（`-env:UserInstallation`）：共用 profile 有
         單一 instance 鎖，GUI 開著、或夜跑 drive_sync 正在轉檔時，headless 會
         整整卡到逾時。實測有一顆殘留 soffice 就讓樣品單抽圖全滅。
      2. **逾時要吃掉**：`subprocess.run(timeout=...)` 逾時是 raise
         TimeoutExpired，`check=False` 不擋這個 —— 原本會直接從工具裡噴出去
         變成 tool call 例外，而不是一句「線稿抽不到」。
    另外 soffice 有「exit 0 但沒產出檔」的靜默失敗模式，所以認產出檔不認 exit code。
    """
    soffice = _find_soffice()
    if not soffice:
        return ""
    stem = os.path.splitext(os.path.basename(src))[0]
    out = os.path.join(outdir, f"{stem}.{fmt}")
    profile = tempfile.mkdtemp(prefix="so_profile_")
    try:
        subprocess.run(
            [soffice, "--headless",
             f"-env:UserInstallation=file://{profile}",
             "--convert-to", fmt, "--outdir", outdir, src],
            capture_output=True, timeout=timeout, check=False,
        )
    except (subprocess.TimeoutExpired, OSError):
        return ""
    finally:
        import shutil
        shutil.rmtree(profile, ignore_errors=True)
    return out if os.path.exists(out) else ""


def _read_order_text(path: str) -> str:
    """樣品單 Excel(.xlsx/.xls)/PDF → 純文字（逐 sheet、逐列）。"""
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    ext = path.lower().rsplit(".", 1)[-1]

    if ext == "xls":
        # 舊 .xls：LibreOffice headless 轉 xlsx（比照 drive_sync 的退路）
        with tempfile.TemporaryDirectory() as td:
            conv = _soffice_convert(path, "xlsx", td)
            if conv:
                return _read_xlsx(conv)
        raise ValueError("xls 轉檔失敗（需 LibreOffice；或另一個 LibreOffice 佔用中）")
    if ext in ("xlsx", "xlsm"):
        return _read_xlsx(path)
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    raise ValueError(f"不支援的樣品單格式: .{ext}")


def _read_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def parse_sample_order(order_file_path: str) -> str:
    """讀客人樣品單(Excel/PDF) → 用 LLM 抽出每個款號的改款規格（款號/顏色/逐部位材料）。

    吃表格型與文字描述型兩種樣品單。抽出的款號可接 search_product_photos 找舊款、餵 generate_product_concept
    改款生成。⚠️ 只做解析、不生圖（生圖另用 generate_from_sample/generate_product_concept，避免意外花費）。

    Args:
        order_file_path: 樣品單本機路徑（Telegram 上傳的 .xlsx/.xls/.pdf）。
    """
    try:
        text = _read_order_text(order_file_path)
    except Exception as e:  # noqa: BLE001
        return f"❌ 讀樣品單失敗：{type(e).__name__}: {str(e)[:120]}"
    if not text.strip():
        return "❌ 樣品單無可讀文字（可能是掃描 PDF，需先 OCR）"

    try:
        from agent_core.gemini_client import GEMINI_MODEL, generate_content_tracked
        prompt = (
            "這是製鞋廠客人的樣品單。請抽出每個款號的改款規格，輸出繁體中文 JSON 陣列，"
            "每個元素格式：{\"款號\":\"型號/Pattern\",\"顏色\":\"Color Way\","
            "\"部位規格\":\"逐部位的材料與顏色簡述（一句話）\"}。"
            "同一份可能有多個款號。只輸出 JSON、不要多餘說明、不要 markdown code fence。\n\n樣品單內容：\n"
            + text[:9000]
        )
        resp = generate_content_tracked(model=GEMINI_MODEL, contents=[prompt],
                                        caller="sample_order.parse_sample_order")
        raw = (resp.text or "").strip()
    except Exception as e:  # noqa: BLE001
        return f"❌ LLM 解析失敗：{type(e).__name__}: {str(e)[:120]}"

    # 客戶辨識吃「檔名 + 單內文」——樣品單常只有款號沒印客戶名，光靠 LLM 會臆測品牌
    ident_text = os.path.basename(order_file_path) + "\n" + text[:9000]

    # 嘗試解析 JSON、給乾淨摘要
    cleaned = raw.strip().removeprefix("```json").removeprefix("```").removesuffix("```").strip()
    try:
        items = json.loads(cleaned)
    except Exception:  # noqa: BLE001
        return (f"✅ 樣品單解析（原始）：\n{raw[:1500]}\n"
                + _customer_line(ident_text, []))

    styles = [str(it.get("款號", "")).strip() for it in items if str(it.get("款號", "")).strip()]
    lines = [f"✅ 樣品單解析出 {len(items)} 個款號：", _customer_line(ident_text, styles)]
    for it in items[:20]:
        lines.append(f"  • 款號 {it.get('款號', '?')}｜{it.get('顏色', '')}｜{it.get('部位規格', '')[:70]}")
    lines.append("　（要「依這張單生鞋圖」→ 直接用 generate_from_order 走線稿驅動；"
                 "款號也可接 search_product_photos 找舊款）")
    return "\n".join(lines)


def _customer_line(sheet_text: str, styles: list) -> str:
    """辨識客戶並組輸出行；辨識失敗要明講，禁止下游 LLM 自行臆測品牌。"""
    try:
        from agent_core.customer_identify import identify_customer
        ident = identify_customer(sheet_text, styles)
    except Exception as e:  # noqa: BLE001
        return f"🏷️ 客戶：辨識程序失敗（{type(e).__name__}），請勿臆測客戶名稱"
    if ident["customer"]:
        line = f"🏷️ 客戶：{ident['customer']}（依{ident['source']}研判）"
        if ident["evidence"]:
            line += "，佐證：" + "；".join(ident["evidence"])
        return line
    return "🏷️ 客戶：樣品單未標註、email 紀錄查無此款號——無法辨識，請勿臆測客戶名稱"


def _xy_cut_boxes(ink, x0: int = 0, y0: int = 0, min_gap: int = 12,
                  min_size: int = 60, out: list | None = None) -> list:
    """遞迴白縫切割（XY-cut）：ink=bool 墨跡遮罩 → [(x0,y0,x1,y1)] 內容區塊。

    為什麼需要（2026-09-01 PSS 8 案）：PDF 樣品單的線稿/上色完稿是**向量**，
    只能整頁渲染取得——但整頁還裹著品牌 logo、標題文字、exsample 實照，直接
    餵圖生圖會把 logo 文字畫進成品、或渲染到錯的那隻鞋。XY-cut 沿最大白色
    縫隙遞迴切割，把每個內容塊分離成獨立候選，再交給挑圖邏輯。
    """
    import numpy as np
    if out is None:
        out = []
    h, w = ink.shape
    if h < min_size or w < min_size:
        return out
    rows, cols = ink.any(axis=1), ink.any(axis=0)
    if not rows.any():
        return out
    r0, r1 = int(np.argmax(rows)), int(len(rows) - np.argmax(rows[::-1]))
    c0, c1 = int(np.argmax(cols)), int(len(cols) - np.argmax(cols[::-1]))
    ink = ink[r0:r1, c0:c1]
    x0, y0 = x0 + c0, y0 + r0
    h, w = ink.shape
    if h < min_size or w < min_size:
        return out
    for axis in (0, 1):     # 先試橫切、再試直切；沿「最大」白縫切最穩
        prof = ink.any(axis=1 - axis)
        gaps, run, start = [], 0, 0
        for i, v in enumerate(prof):
            if not v:
                run += 1
                if run == 1:
                    start = i
            else:
                if run >= min_gap:
                    gaps.append((run, start))
                run = 0
        if gaps:
            glen, gstart = max(gaps)
            mid = gstart + glen // 2
            if axis == 0:
                _xy_cut_boxes(ink[:mid], x0, y0, min_gap, min_size, out)
                _xy_cut_boxes(ink[mid:], x0, y0 + mid, min_gap, min_size, out)
            else:
                _xy_cut_boxes(ink[:, :mid], x0, y0, min_gap, min_size, out)
                _xy_cut_boxes(ink[:, mid:], x0 + mid, y0, min_gap, min_size, out)
            return out
    out.append((x0, y0, x0 + w, y0 + h))
    return out


_PDF_MAX_PAGES = 6      # 樣品單 PDF 通常 1-2 頁；多頁型錄不整本掃


def _extract_pdf_images(order_file_path: str) -> list:
    """PDF 樣品單抽圖：整頁渲染→XY-cut 切塊（向量線稿/完稿只能這樣拿）＋內嵌 raster。

    實測 PSS 8（Richter）：線稿與上色完稿都是向量、pypdf 抽不到，只有 exsample
    實照是內嵌 raster——所以頁面渲染是主路徑、內嵌圖是補充候選。
    """
    import numpy as np
    from PIL import Image
    outdir = tempfile.mkdtemp(prefix="order_img_")
    pngs = []
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(order_file_path)
        for pi in range(min(len(pdf), _PDF_MAX_PAGES)):
            try:
                page = pdf[pi].render(scale=2.0).to_pil().convert("RGB")
            except Exception:  # noqa: BLE001 — 單頁渲染失敗跳過
                continue
            a = np.asarray(page).astype(int)
            ink = (a.max(axis=2) - a.min(axis=2) > 25) | (a.mean(axis=2) < 235)
            for bi, (x0, y0, x1, y1) in enumerate(_xy_cut_boxes(ink)):
                if (x1 - x0) * (y1 - y0) < _SKETCH_MIN_AREA:
                    continue
                p = os.path.join(outdir, f"page{pi}_seg{bi}.png")
                page.crop((x0, y0, x1, y1)).save(p)
                pngs.append(p)
    except Exception:  # noqa: BLE001 — pypdfium2 不可用/壞檔 → 只靠內嵌圖
        pass
    try:
        from pypdf import PdfReader
        for pi, page in enumerate(PdfReader(order_file_path).pages[:_PDF_MAX_PAGES]):
            for ii, im in enumerate(page.images):
                try:
                    p = os.path.join(outdir, f"emb{pi}_{ii}.png")
                    im.image.convert("RGB").save(p)
                    pngs.append(p)
                except Exception:  # noqa: BLE001
                    continue
    except Exception:  # noqa: BLE001
        pass
    return pngs


def _extract_order_images(order_file_path: str) -> list:
    """從樣品單抽候選鞋圖：xlsx 抽嵌入圖（emf/wmf 走 soffice）；PDF 走渲染+切塊。"""
    import zipfile
    low = order_file_path.lower()
    if low.endswith(".pdf"):
        return _extract_pdf_images(order_file_path)
    if not low.endswith((".xlsx", ".xlsm")):
        return []  # 其餘格式無圖可抽（openpyxl 會漏 emf，xlsx 直接讀 zip）
    outdir = tempfile.mkdtemp(prefix="order_img_")
    pngs = []
    try:
        z = zipfile.ZipFile(order_file_path)
    except Exception:  # noqa: BLE001
        return []
    for name in z.namelist():
        if "/media/" not in name:
            continue
        base = os.path.basename(name)
        raw = os.path.join(outdir, base)
        with open(raw, "wb") as f:
            f.write(z.read(name))
        low = base.lower()
        if low.endswith((".emf", ".wmf")):
            conv = _soffice_convert(raw, "png", outdir)
            if conv:
                pngs.append(conv)
        elif low.endswith((".png", ".jpg", ".jpeg")):
            pngs.append(raw)
    return pngs


_SHOE_KEYWORDS = ("shoe", "sneaker", "sandal", "boot", "footwear", "clog", "loafer")
_SKETCH_MIN_AREA = 30000        # 去白邊後的最小面積（材質圖標/logo 都遠小於此）
_SKETCH_RATIO_RANGE = (0.4, 2.3)


def _autocrop_white(path: str) -> str:
    """去掉四周白邊，回新檔路徑（無白邊/失敗 → 回原路徑）。

    為什麼非做不可（2026-08-03 UserC Richter 8105-3272 案）：emf/wmf 走 soffice
    轉檔是**整頁 A4 畫布**輸出（794×1123、真正內容只佔中間一小塊），面積一進
    排序，這種「大白紙」必然壓過真正的線稿 —— 該案實測「鞋類保養符號那頁」
    （內容其實只有 141×65）贏過真手繪線稿（422×216），生出來的圖必錯。
    """
    try:
        from PIL import Image, ImageChops
        with Image.open(path) as raw:
            im = raw.convert("RGB")
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bbox = ImageChops.difference(im, bg).getbbox()
        if not bbox or bbox == (0, 0, im.width, im.height):
            return path
        out = os.path.join(os.path.dirname(path), "crop_" + os.path.basename(path))
        im.crop(bbox).save(out)
        return out
    except Exception:  # noqa: BLE001
        return path


def _shoe_confidence(path: str) -> float:
    """macOS Vision 影像分類的鞋類置信度；Vision 不可用 → 0.0（＝無訊號）。"""
    try:
        import Vision
        from Foundation import NSData
        data = NSData.dataWithContentsOfFile_(path)
        handler = Vision.VNImageRequestHandler.alloc().initWithData_options_(data, None)
        req = Vision.VNClassifyImageRequest.alloc().init()
        handler.performRequests_error_([req], None)
        conf = 0.0
        for o in (req.results() or [])[:12]:
            if any(k in o.identifier().lower() for k in _SHOE_KEYWORDS):
                conf = max(conf, float(o.confidence()))
        return conf
    except Exception:  # noqa: BLE001
        return 0.0


_ART_COLOR_MIN = 0.15    # 彩色佔比 ≥ 此 = 上色圖（對齊 _ANNOT_MAX_FRACTION 的語意）
_ART_WHITE_MIN = 0.20    # 白底佔比 ≥ 此 = 畫在白底上的「設計稿」而非實照
                         # （PSS 8 實測：上色完稿 白 48.6%、exsample 實照 白 0.0%）
_ART_DARK_MIN = 0.25     # 深色填色佔比 ≥ 此 = 大面積塗色的設計稿。全黑/深色鞋款的
                         # 完稿彩色佔比不夠（PSS 7 實測：全黑完稿 colored=0.145 差
                         # 0.005 落榜、dark=0.473；線稿只有細線 dark=0.077）


def _color_stats(path: str) -> tuple[float, float, float]:
    """(彩色佔比, 白底佔比, 深色佔比)；讀不了回 (0.0, 0.0, 0.0)。"""
    try:
        import numpy as np
        from PIL import Image
        with Image.open(path) as raw:
            im = raw.convert("RGB")
        im.thumbnail((256, 256))
        a = np.asarray(im).astype(int)
        colored = float(((a.max(axis=2) - a.min(axis=2)) > 25).mean())
        gray = a.mean(axis=2)
        return colored, float((gray > 235).mean()), float((gray < 160).mean())
    except Exception:  # noqa: BLE001
        return 0.0, 0.0, 0.0


def _looks_painted(colored: float, white: float, dark: float) -> bool:
    """「客人上過色/塗色的設計完稿」判定：彩色佔比夠，或白底上有大面積深色填色。

    只看彩色佔比會漏掉全黑/深色鞋款的完稿（2026-09-02 PSS 7 案：全黑完稿只有
    印花帶色、colored=0.145 差 0.005 落榜 → 被當黑白線稿 → 完稿鎖色三道保護全
    失效，重生時 LLM 把「大底黑」寫進色票、蓋掉客人畫的白中底＋銀蔥大底）。
    深色路徑要求白底：沒有白底的深色圖是照片/滿版圖，不是畫在白紙上的設計稿。
    """
    return colored >= _ART_COLOR_MIN or (white >= _ART_WHITE_MIN and dark >= _ART_DARK_MIN)


def _pick_shoe_sketch(pngs: list):
    """挑「客人的設計圖」：去白邊後按【上色完稿 > 黑白線稿 > 照片類】三級優先。

    級內仍是鞋類置信優先、面積次之。三級的道理（2026-09-01 PSS 8 案實測）：
    PDF 樣品單常同頁放「上色完稿＋exsample 參考實照」，實照的 Vision 置信天生
    比設計稿高（0.41 vs 0.16），純比分數必挑到實照——但客人要的設計（雙魔鬼氈
    vs 實照的綁帶）在完稿上。設計稿 vs 實照用「白底佔比」分：設計稿畫在白底上。

    回傳的是**去白邊後**的圖（填滿畫面，圖生圖鎖形狀更準）。
    ⚠️ 置信度刻意不設地板：Vision 對線稿的絕對分數本來就低（Richter 案實測
    真線稿 footwear=0.045），舊版 0.05 地板會把唯一的真訊號抹平成跟白紙同分，
    排序就退化成純比面積 —— 那正是挑錯圖的主因。
    """
    from PIL import Image
    cands: list[tuple[int, float, int, str]] = []
    for p in pngs:
        cropped = _autocrop_white(p)
        try:
            with Image.open(cropped) as im:
                w, h = im.size
        except Exception:  # noqa: BLE001
            continue
        if not h or w * h < _SKETCH_MIN_AREA:  # 太小 = 材質圖標 / logo
            continue
        lo, hi = _SKETCH_RATIO_RANGE
        if not (lo <= w / h <= hi):  # 太扁 = 大底線稿 / 標題列
            continue
        conf = _shoe_confidence(cropped)
        colored, white, dark = _color_stats(cropped)
        if conf <= 0:
            tier = 3                                             # 無訊號（含 Vision 不可用）
        elif white >= _ART_WHITE_MIN and _looks_painted(colored, white, dark):
            tier = 0                                             # 上色完稿（含全黑款）：設計+配色都在
        elif colored < _ART_COLOR_MIN:
            tier = 1                                             # 黑白線稿：客人的設計
        else:
            tier = 2                                             # 照片類（exsample 等參考圖）
        cands.append((tier, conf, w * h, cropped))
    if not cands:
        return None
    best_tier = min(c[0] for c in cands)
    pool = [c for c in cands if c[0] == best_tier]
    return max(pool, key=lambda c: (c[1] + 0.01) * (c[2] ** 0.5))[3]


# ── 部位標號：從輸入線稿洗掉，別只靠 prompt 拜託模型不要畫 ────────────
_ANNOT_SAT_MIN = 25          # RGB(max-min) 高於此 = 彩色標註；黑白線稿恆為 0
_ANNOT_MAX_FRACTION = 0.15   # 彩色佔比超過此 = 這張本來就是彩圖，別動它


def _scrub_color_annotations(path: str) -> str:
    """把線稿上的彩色部位標號（紅 1 / 綠 2 …）與引線塗白，回新檔路徑。

    為什麼不能只靠 prompt（2026-08-04 UserC 回報）：#344 已經在 prompt 明講
    「數字與引線完全不要出現」，實測當天 13 張成品裡仍有 3 張把標號畫上鞋面
    （紅 1 / 綠 2 / 紅 5）。圖生圖對否定指令本來就不可靠——來源圖裡有的東西
    它就是會抄。治本是讓**輸入圖裡根本沒有那些數字**。

    靠顏色分離、不靠 OCR：客人線稿是純黑白工程線圖、只有標號帶顏色，實測
    Richter 8105-3272 線稿的飽和度是乾淨雙峰（黑線 9054 px 裡只有 2 px
    飽和度 >60；標號 289 px 全在另一端，門檻 18~32 之間結果完全相同）。
    OCR 在這種尺寸（422×216）讀不到那些小數字（實測只認得 Richter），
    所以顏色是這一端唯一可靠的訊號。
    ⚠️ 標號若是黑色的就洗不掉——那條靠 _rendered_digit_labels 事後把關。
    ⚠️ 深色塗色的完稿不洗（2026-09-02 PSS 7 案）：全黑完稿彩色佔比 0.145 低於
    _ANNOT_MAX_FRACTION，上面帶色的是客人畫的印花/車縫線，洗了就是毀客稿。
    """
    try:
        import numpy as np
        from PIL import Image, ImageFilter
        _colored, white, dark = _color_stats(path)
        if white >= _ART_WHITE_MIN and dark >= _ART_DARK_MIN:
            return path
        with Image.open(path) as raw:
            im = raw.convert("RGB")
        arr = np.asarray(im).astype(np.int16)
        sat = arr.max(axis=2) - arr.min(axis=2)
        # 膨脹 1px：反鋸齒邊緣飽和度不足門檻，不吃掉會留一圈彩色殘影。
        mask_img = Image.fromarray(((sat > _ANNOT_SAT_MIN) * 255).astype("uint8"), "L")
        mask = np.asarray(mask_img.filter(ImageFilter.MaxFilter(3))) > 0
        frac = float(mask.mean())
        if frac <= 0 or frac > _ANNOT_MAX_FRACTION:
            return path       # 純黑白（沒東西可洗）／整張彩圖（洗了會毀圖）
        out_arr = np.asarray(im).copy()
        out_arr[mask] = 255
        out = os.path.join(os.path.dirname(path), "clean_" + os.path.basename(path))
        Image.fromarray(out_arr).save(out)
        return out
    except Exception:  # noqa: BLE001 — 洗不了就用原圖，別擋生圖
        return path


def _rendered_digit_labels(png_path: str) -> list:
    """OCR 成品圖，回「純數字」token 清單（品牌字樣 Richter/R 不算）。

    第二道保險：線稿洗乾淨後模型仍可能自己加標號。實測 macOS Vision 對這種
    成品照分得很開——2026-08-04 那 13 張裡，髒的 3 張全抓到、乾淨的 10 張
    只讀到 'Richter'/'R'，零誤判。Vision 不可用（Cloud Run / 沒裝 ocrmac）
    → 回 [] ＝不擋，維持舊行為。
    """
    try:
        from agent_core import vision_ocr
        res = vision_ocr.ocr_image(png_path)
    except Exception:  # noqa: BLE001
        return []
    if not res:
        return []
    return [t for t in str(res.get("text") or "").split() if t.isdigit()]


def _norm_for_grounding(text: str) -> str:
    return "".join(str(text or "").split()).lower()


def _color_text_in_order(color_text: str, order_text: str) -> bool:
    """完稿在場時色票的入場券：顏色敘述必須真的出現在樣品單原文（去空白、不分大小寫）。

    2026-09-02 PSS 6 案：規格沒寫大底顏色，LLM 腦補「大底=黑色」，「黑色」又剛好
    解得動內建色名表（source=builtin）→ 穿過 #458 的 source!="llm" 濾網進色票，
    ΔE 驗證反過來把客人完稿的白/灰大底當偏色在糾正。Pantone 碼/色彙詞是照抄的、
    自然在原文裡；腦補的不在。LLM 有時在照抄的色號前後加註（「深藍 Pantone 19-4025
    TPX」）——帶數字的 token（色號）有在原文也算照抄。
    """
    o = _norm_for_grounding(order_text)
    c = _norm_for_grounding(color_text)
    if not c:
        return False
    if c in o:
        return True
    for tok in str(color_text or "").split():
        t = _norm_for_grounding(tok)
        if t and any(ch.isdigit() for ch in t) and t in o:
            return True
    return False


def _composition_flaws(png_path: str, artwork_path: str = "") -> list:
    """成品驗證：單隻鞋、正側面；有完稿時另驗「有沒有完稿上不存在的大面積顏色」。

    2026-09-02 PSS 6 案：nano-banana 第一發就把單隻側視完稿畫成一雙斜角產品照
    ——數字 OCR 與 ΔE 都攔不到構圖跑掉，補一通 flash 視覺數鞋子/看角度，不合
    就進既有重生迴圈。任何失敗（API 掛、格式歪）都放行，驗證器自己壞不擋交件。

    外來色（2026-09-02 5004 案）：完稿棕褐獨角獸印花被渲染成深綠鞋筒——ΔE 驗證
    只驗「該有的色在不在」不驗「不該有的色混沒混進來」，而壞圖外來色叢集對完稿
    主色的最小 ΔE 僅 15.7~22.8、全在容差 28 內，像素門檻分不開好壞（實測 margin
    太薄）。正解＝同一通 flash 連完稿一起看，問有沒有外來大面積顏色，零額外成本。
    """
    try:
        from PIL import Image
        from agent_core.gemini_client import GEMINI_MODEL, generate_content_tracked
        contents: list = [Image.open(png_path)]
        if artwork_path:
            contents.append(Image.open(artwork_path))
            contents.append(
                "第一張是渲染成品，第二張是客人上色完稿。回答三題："
                "1) 成品裡有幾隻鞋？2) 視角是否為正側面（水平視線、整隻鞋側輪廓朝左或朝右）？"
                "3) 成品是否出現完稿上不存在的大面積顏色（例如完稿無綠色但成品鞋筒是深綠）？"
                '只回 JSON（不要 code fence）：{"shoes": 1, "side_view": true, "foreign_color": ""}'
                "——沒有外來色 foreign_color 就是空字串，有就用中文短語描述（如「鞋筒變成深綠色」）。")
        else:
            contents.append(
                "這張鞋類產品圖裡有幾隻鞋？視角是否為正側面（水平視線、整隻鞋"
                "側輪廓朝左或朝右）？"
                '只回 JSON（不要 code fence）：{"shoes": 1, "side_view": true}')
        resp = generate_content_tracked(
            model=GEMINI_MODEL, contents=contents,
            caller="sample_order.generate_from_order")
        raw = (resp.text or "").strip()
        data = json.loads(raw.removeprefix("```json").removeprefix("```")
                          .removesuffix("```").strip())
        flaws = []
        if int(data.get("shoes", 1)) != 1:
            flaws.append(f"畫了 {int(data.get('shoes'))} 隻鞋（要求單隻）")
        if data.get("side_view") is False:
            flaws.append("視角不是正側面")
        fc = str(data.get("foreign_color") or "").strip()
        if fc and fc.lower() not in ("無", "沒有", "none", "no", "null"):
            flaws.append(f"出現完稿上沒有的顏色：{fc[:40]}")
        return flaws
    except Exception:  # noqa: BLE001 — 驗證器自己壞掉不擋生圖
        return []


def _first_inline_image(resp) -> bytes:
    """從 genai 回應抽第一張圖的 bytes；沒有 → b""。"""
    import base64
    for cand in (resp.candidates or []):
        for part in (getattr(cand.content, "parts", None) or []):
            if getattr(part, "inline_data", None) and part.inline_data.data:
                data = part.inline_data.data
                return base64.b64decode(data) if isinstance(data, str) else data
    return b""


def _render_prompt(spec: str, retry: bool, swatch_parts: list | None = None,
                   color_misses: list | None = None,
                   material_refs: list | None = None,
                   base_style: str = "",
                   colored_sketch: bool = False,
                   extra_note: str = "",
                   comp_flaws: list | None = None,
                   edit_previous: bool = False) -> str:
    """圖生圖 prompt。retry=True 補「上一版出現了數字」；color_misses 補偏色部位糾正。

    edit_previous=True＝修圖模式（2026-09-07 大王條件二）：第一張輸入圖是上一版
    渲染成品，只修使用者指正的部分、其餘與上一版完全一致——整張重生會把已經對的
    部位重新擲骰子，修好 A 弄壞 B（「以免犯第二次錯誤」）。
    """
    multi = bool(swatch_parts or material_refs or base_style)
    kind = "已上色的鞋款設計完稿" if colored_sketch else "鞋款設計線稿"
    lead = (f"第一張圖是{kind}。以它" if multi else f"以這張{kind}")
    if edit_previous:
        base = (
            "第一張圖是上一版的渲染成品。這次是局部修正、不是重新生成："
            f"只修改這段指正提到的部分——{extra_note}——"
            "除此之外的所有部位材質、顏色、形狀細節、單隻構圖、視角、光線與純白背景，"
            "必須與第一張圖保持完全一致，不得自行變動或重新發揮。"
            f"第二張圖是{kind}：修正處的形狀、位置、顏色與材質分佈以它為準。"
        )
    else:
        base = (
            # 「單隻、同視角」放在第一句定調——埋在句尾時 nano-banana 會漂去畫
            # 四視角產品集錦（2026-09-03 PSS 2 案：三次 attempt 全是 4 隻鞋斜角圖，
            # 構圖驗證打回也拉不回來）。
            f"{lead}為準，完全保留：鞋型與比例、幫面結構分割線、"
            "配件（魔鬼氈/鞋帶/裝飾）位置、鞋筒高度、鞋底樣式。"
            "渲染成一張真實產品攝影照：畫面裡只有這一隻鞋，"
            "維持與原圖完全相同的正側面視角，不要畫成多視角組圖。"
        )
    if colored_sketch and not edit_previous:
        # 客人自己上好的色是最權威的配色來源——完稿的顏色分佈整套沿用。
        # ⚠️ 修圖模式跳過（2026-09-08 review V1）：基準是上一版成品、只改指正
        # 部位，這裡的「每個部位用完稿的顏色」會叫模型整張重上色。
        base += "整體配色以完稿本身的顏色為準：每個部位用完稿上該部位的顏色，不得自行換色。"
        # 紋理分佈＝材質分佈（2026-09-01 PSS 8 二輪實測：織紋蔓延到整個鞋面，
        # 但完稿上平滑的部位是皮料）——邊界跟著結構分割線走，不得互相蔓延。
        base += ("材質紋理的分佈也嚴格依照完稿：完稿上畫有織紋/紋理的區域才是織物布料，"
                 "完稿上平滑無紋理的區域必須渲染成平滑皮料（PU/皮革）質感，"
                 "兩種材質的邊界沿著結構分割線，紋理不得蔓延到平滑部位。")
        # 成對部件一致性（2026-09-03 PSS 2 案：完稿兩條魔鬼氈畫法相同，成品卻
        # 上條平滑、下條織帶紋——規格文字裡的「織帶」被猜到黏扣帶上）。
        base += ("完稿上畫法相同的重複部件（例如上下兩條魔鬼氈黏扣帶）是同一種材質與"
                 "同一個顏色，成品必須渲染成完全一致——不得一條平滑、一條織紋；"
                 "規格文字提到的材質只能落在完稿畫出該質感的部位，"
                 "完稿沒有畫出的部位不得自行分配材質。")
    base += f"各部位材質與顏色：{spec}。"
    if base_style:
        # 形體也用「圖」錨定（Phase C）：同款既有成品實照——楦型/大底齒紋抄實物，
        # 不再靠模型想像（Richter 這類「形體＝大底編號」的客戶效果最大）。
        base += (f"第二張圖是同款（款號 {base_style}）既有成品的真實照片：整體形體、"
                 "楦型與大底齒紋以這張實照為準；但配色與材質不要照抄實照，"
                 "一律依本描述與後面的參考圖。")
    if material_refs:
        # 紋理也用「圖」錨定：附真實材質特寫，模型抄實物紋理而不是想像。
        # ⚠️ 明講「只取紋理不取顏色」——特寫圖的顏色不一定是這張單要的色。
        seq = "、".join(f"{m['label'] or m['key']}（用於{'/'.join(m['parts'][:3])}）"
                        for m in material_refs)
        base += (f"接著的 {len(material_refs)} 張圖是材質特寫參考，依序為：{seq}。"
                 "對應部位的表面紋理與質感必須依照材質特寫圖——只取紋理，"
                 + ("顏色一律以色票為準。" if swatch_parts else "顏色仍依上面的規格描述。"))
    if swatch_parts:
        names = "、".join(str(p.get("part", "?")) for p in swatch_parts)
        # 顏色以「圖」錨定不以「字」錨定：模型抄眼前的色塊，遠比理解色號文字準
        # （2026-09-01 大王要求 Pantone/文字色都要準——同線稿鎖形狀的思路）。
        # ⚠️ 客戶專有色名的字面意義常是錯的（2026-09-03 PSS 2 案：Richter 的
        # lagoon 是寶藍，按「環礁湖」字面聯想就成了藍綠）——明令以色塊為準。
        base += (f"最後一張圖是各部位的指定色票（{names}），每格已標部位名——"
                 + ("只有這次被修改的部位需要按對應色票上色，"
                    "未被點名的部位一律維持第一張圖原樣；" if edit_previous else
                    "各部位顏色必須精確使用對應色票的顏色，不得自行調亮調暗或換色；")
                 + "規格文字裡的色名（如 lagoon）以對應色票色塊為準，"
                 "不得按色名的字面意義自行聯想顏色；"
                 "色票圖本身不得出現在成品畫面裡。")
    base += (
        # 線稿已先洗掉彩色標號（_scrub_color_annotations），這句是雙保險：
        # 模型偶爾會自己「補」標註上去。
        "畫面上只有鞋子本身：除了鞋款原有的品牌壓印／繡標，不得加上任何數字、部位標號、"
        "指示引線、尺寸標註或文字浮水印。"
        # 修圖模式的視角/背景以上一版為準（使用者已接受，可能非正側面）。
        + ("維持第一張圖的視角、構圖與背景，單隻鞋、專業產品攝影、高解析。"
           if edit_previous else
           "專業產品攝影、純白背景、單隻鞋、正側面視角、真實皮革布料橡膠質感、高解析。")
    )
    if retry:
        base += "（上一版在鞋面上畫出了不該有的數字標號，這次務必完全乾淨。）"
    if color_misses:
        miss = "、".join(f"{m['part']}（應為 #{m['hex']}）" for m in color_misses[:4])
        base += (f"（這些部位的顏色仍偏離指定色票：{miss}——只修這些部位，"
                 "其他保持第一張圖原樣。）" if edit_previous else
                 f"（上一版這些部位的顏色偏離指定色票：{miss}——這次嚴格按色票上色。）")
    if comp_flaws:
        flaw_txt = "、".join(str(f) for f in comp_flaws[:2])
        base += (f"（上一版驗出問題：{flaw_txt}——只畫一隻鞋，其餘維持第一張圖不變。）"
                 if edit_previous else
                 f"（上一版驗出問題：{flaw_txt}"
                 "——這次構圖與配色必須跟第一張圖一致：只畫一隻鞋、正側面視角，"
                 "顏色只用完稿上存在的顏色。）")
    if extra_note and not edit_previous:
        # 視角護欄（2026-09-07 大王條件一）：修正指示只管它點名的內容；聊天層
        # LLM 曾自行加「45 度立體展示視角」進 extra_note，把正側面鎖定蓋掉、
        # 生出四視角組圖——除非指示明確要求別的視角，一律跟樣品單同視角。
        base += (f"另外務必遵守這些指示（使用者對前一版的修正）：{extra_note}。"
                 "（除非上面這段修正指示明確指定了別的視角，"
                 "視角一律維持與樣品單設計稿相同的單隻正側面。）")
    return base


def _resolve_parts(items: list, customer: str = "") -> list:
    """逐部位落色：Pantone 查表/客戶色彙/內建色名（確定性）優先；都查不到才收 LLM hex 推估。"""
    from agent_core import color_anchor
    parts = []
    for it in items[:8]:
        if not isinstance(it, dict):
            continue
        part = str(it.get("部位", "")).strip()[:12]
        if not part:
            continue
        color_text = str(it.get("顏色", "")).strip()
        got = color_anchor.resolve_color(color_text, customer=customer)
        if got:
            hexv, label, source = got["hex"], got["label"], got["source"]
        else:
            hexv = color_anchor.parse_hex(str(it.get("hex", "")))
            label, source = (color_text[:20] or "?", "llm") if hexv else ("", "")
        parts.append({"part": part, "material": str(it.get("材質", "")).strip()[:20],
                      "color_text": color_text[:40], "hex": hexv, "label": label,
                      "source": source})
    return parts


def _order_customer(order_file_path: str, order_text: str) -> str:
    """客戶名（給客戶色彙表定 scope 用）；辨識不了回 ""。零 LLM：純品牌字樣 regex。"""
    try:
        from agent_core.customer_identify import identify_customer
        ident = identify_customer(os.path.basename(order_file_path) + "\n" + order_text[:9000], [])
        return str(ident.get("customer") or "")
    except Exception:  # noqa: BLE001
        return ""


def _spec_for_render(order_text: str, colorway: str, customer: str = "") -> tuple[str, list, str]:
    """LLM 讀樣品單 →（渲染描述, 逐部位色錨清單, 款號）。結構化解析失敗退純文字（parts=[]）。

    一通 LLM 同時要三件事：渲染描述（進 prompt）＋逐部位 JSON（進色錨/色票/ΔE 驗證）
    ＋款號（查產品照索引附同款實照當形體參考）。LLM 對每個顏色也給 hex 推估，但只當
    最後備援——Pantone/色彙/內建色名解得出來的一律蓋掉它（_resolve_parts），文字色
    敘述才不會隨模型心情漂。
    """
    from agent_core.gemini_client import GEMINI_MODEL, generate_content_tracked
    cw = f"（指定配色 {colorway}）" if colorway.strip() else "（用第一個配色）"
    prompt = (
        f"這是製鞋廠樣品單。把數字標記部位 1-5 對應的材質與顏色{cw}整理成 JSON（不要 code fence）："
        "{\"描述\":\"一段繁體中文渲染用描述\",\"款號\":\"此配色所屬的型號/款號（沒有就空字串）\","
        "\"部位\":[{\"部位\":\"鞋身\",\"材質\":\"麂皮\","
        "\"顏色\":\"照抄樣品單原文（含 Pantone 色號務必完整照抄）\",\"hex\":\"#RRGGBB 你對該顏色的最佳推估\"}]}。"
        # 這段會原封不動進圖生圖 prompt：留著「部位 1」「1.」這種編號，模型就會
        # 把數字畫到鞋面上（2026-08-04 UserC 回報的第二個來源）。一律翻成部位名稱。
        "⚠️ 描述與部位名稱不可出現編號（不要寫「部位 1」「1.」「(2)」），一律改用部位名稱"
        "（鞋身/鞋面拼接/包邊/魔鬼氈/鞋帶/鞋墊/大底/扣件…）；認不出是哪個部位就用相對位置描述。"
        # 2026-09-03 PSS 2 案兩個「猜」的來源都出在這通 LLM 的自由發揮：
        # ①描述欄把 lagoon 意譯成「環礁湖藍綠」——客戶專有色名字面常是錯的，意譯直接把
        # 渲染帶偏；②「Web.tape」沒點名部位，被分配到魔鬼氈上→下條魔鬼氈變織帶紋。
        "⚠️ 描述裡的顏色一律照抄樣品單原文色名（lagoon、atlantic、Pantone 色號等），"
        "不得意譯成中文色系（不要把 lagoon 寫成「環礁湖」「藍綠」）——專有色名的字面"
        "意義常與實際用色不同。"
        "⚠️ 材質只能寫給樣品單明確對應的部位；規格沒點名部位的材質（如織帶 Web.tape）"
        "不得自行猜給魔鬼氈或鞋身——不確定就寫在描述裡帶過、不要落到任何部位上。"
        "\n\n樣品單：\n" + order_text[:8000]
    )
    resp = generate_content_tracked(model=GEMINI_MODEL, contents=[prompt],
                                    caller="sample_order.generate_from_order")
    raw = (resp.text or "").strip()
    cleaned = raw.removeprefix("```json").removeprefix("```").removesuffix("```").strip()
    try:
        data = json.loads(cleaned)
        desc = str(data.get("描述", "")).strip()
        items = data.get("部位")
        if not desc or not isinstance(items, list):
            raise ValueError("缺描述或部位")
    except Exception:  # noqa: BLE001 — LLM 沒照格式 → 舊行為：整段當描述、無色錨
        return raw[:400], [], ""
    style = str(data.get("款號", "")).strip()[:24]
    return desc[:400], _resolve_parts(items, customer=customer), style


def _find_base_photo(style: str):
    """款號 → 產品照索引裡的同款實照 meta（檔名前綴符合優先）；索引沒建/查無回 None。

    Phase C（2026-09-01）：Richter 這類客戶形體以大底編號區分，同款實照在索引的
    「客戶鞋照片」樹裡（index 腳本同批納入）。刻意只查本機索引、不打 Drive 搜尋
    ——employee 也走這條路徑，不多開網路查詢面；索引每週日自動重建。
    比對雙方都先去除空白：客人單上寫「PSS 7」、寄樣照檔名是「PSS7.jpg」
    （2026-09-02 PSS 7 案），帶空白比對必落空。
    """
    def norm(text) -> str:
        return "".join(str(text or "").split()).lower()

    s = norm(style)
    if len(s) < 4:
        return None
    try:
        from agent_core.logging_and_paths import DATA_DIR
        with open(os.path.join(DATA_DIR, "product_photos", "meta.json"), encoding="utf-8") as f:
            meta = json.load(f)
    except (OSError, ValueError):
        return None
    cands = [m for m in meta if s in norm(m.get("file", ""))]
    prefix = [m for m in cands if norm(m.get("file", "")).startswith(s)]
    pool = prefix or cands
    return pool[0] if pool else None


def _download_drive_image(file_id: str) -> bytes:
    """Drive 圖檔 bytes；任何失敗回 b""——形體參考是加分項，絕不擋生圖。"""
    try:
        from agent_core.google_auth import get_service
        from googleapiclient.http import MediaIoBaseDownload
        svc = get_service("drive", "v3")
        buf = io.BytesIO()
        d = MediaIoBaseDownload(buf, svc.files().get_media(fileId=file_id, supportsAllDrives=True))
        done = False
        while not done:
            _, done = d.next_chunk()
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return b""


# ────────────────────────────────────────────────────────────────────
# 部門員工（Telegram 自由對話）生圖：分艙 + 交付 + 花費上限
# ────────────────────────────────────────────────────────────────────
def _dept_scope_color() -> str:
    """這次呼叫是否在部門色 AgentRequest context 底下？回部門色或 ""。

    比照 doc_export._dept_scope_color（2026-08-03 UserC 產 Excel 案）：員工
    自由對話的工具都被 dept_tool_scope.wrap_tools_with_agent_caller 包進
    caller=<色> 的 context；大王 / REPL / daemon 路徑沒有 context → 回 ""。

    這顆色決定兩件事：
      1. 生圖存到 generated_images/dept/<色> —— 跟大王自己的生成圖分艙，
         回覆附圖白名單才綁得住（見 daemon_telegram._reply_photo_allowed_roots）。
      2. 交付改走 [[TG_PHOTO:]] 回覆附圖，不走 telegram_send_photo —— 後者的
         chat_id 閘只認大王 keyring chat，員工要的鞋圖會傳到大王手機。
    """
    try:
        from agent_core.agents.middleware import current_agent_request
        req = current_agent_request()
    except Exception:  # noqa: BLE001
        return ""
    if req is None:
        return ""
    color = str(getattr(getattr(req, "caller", None), "value", "") or "").strip().lower()
    return "" if color in ("", "red") else color


def _consume_render_quota(scope: str) -> tuple[bool, int, int]:
    """扣一次每日生圖配額。回 (放行, 今日已用, 上限)；上限 <=0 = 不限。

    為什麼要有這顆（2026-08-03）：generate_from_order 從 CONFIRM 降成 SAFE 才
    進得了員工白名單（員工 session 不收 CONFIRM —— +確認 token 不綁身分、放進去
    等於自己確認自己），但 CONFIRM 原本擋的是「意外花費」。改用硬上限接手：
    無論誰在呼叫、迴圈跑幾次，一天就是這麼多張。scope 分色計數（各部門互不
    佔用），大王路徑記在 "owner"。
    """
    from agent_core.env_utils import env_int
    limit = env_int("RED_ORDER_RENDER_DAILY_MAX", 20, min_value=0, max_value=100000)
    if limit <= 0:
        return True, 0, 0
    import fcntl
    from datetime import datetime
    from agent_core.logging_and_paths import STATE_DIR
    today = datetime.now().strftime("%Y-%m-%d")
    path = os.path.join(STATE_DIR, "order_render_quota.json")
    try:
        os.makedirs(STATE_DIR, exist_ok=True)
        # 10 隻 bot 是 10 個 process，read-modify-write 一定要鎖，不然併發漏算。
        with open(path + ".lock", "a+") as lf:
            fcntl.flock(lf, fcntl.LOCK_EX)
            try:
                try:
                    with open(path, encoding="utf-8") as f:
                        data = json.load(f)
                except (OSError, ValueError):
                    data = {}
                day = dict(data.get(today) or {})   # 只留今天 → 跨日自然歸零
                used = int(day.get(scope, 0) or 0)
                if used >= limit:
                    return False, used, limit
                day[scope] = used + 1
                tmp = path + ".tmp"
                with open(tmp, "w", encoding="utf-8") as f:
                    json.dump({today: day}, f)
                os.replace(tmp, path)
                return True, used + 1, limit
            finally:
                fcntl.flock(lf, fcntl.LOCK_UN)
    except OSError as e:  # noqa: BLE001
        # 配額檔壞掉不擋正事（best-effort，比照 repo 慣例）；只留 log。
        print(f"[sample_order] ⚠️ 生圖配額檔存取失敗，本次放行：{type(e).__name__}: {e}")
        return True, 0, limit


def _render_output_dir(color: str) -> str:
    from agent_core.logging_and_paths import GENERATED_IMAGES_DIR
    if not color:
        return GENERATED_IMAGES_DIR
    return os.path.join(GENERATED_IMAGES_DIR, "dept", color)


def _prune_dept_renders(out_dir: str) -> None:
    """清部門生圖目錄的過期舊檔（best-effort，失敗絕不擋生圖）。

    比照 doc_export._prune_dept_exports：這是新開的子目錄，沒有既有排程會掃到，
    不自己清就無限累積（一張圖 1-2MB × 每日上限）。
    """
    import time
    from agent_core.env_utils import env_int
    keep_days = env_int("RED_EXPORTS_KEEP_DAYS", 14, min_value=0, max_value=3650)
    if keep_days <= 0:
        return
    cutoff = time.time() - keep_days * 86400
    try:
        entries = os.listdir(out_dir)
    except OSError:
        return
    for fn in entries:
        if not fn.lower().endswith(".png"):
            continue
        p = os.path.join(out_dir, fn)
        try:
            if os.path.isfile(p) and os.path.getmtime(p) < cutoff:
                os.remove(p)
        except OSError:
            pass


def _mask_text_for_accents(path: str) -> str:
    """OCR 文字框塗白後的副本路徑（只餵點綴色抽取；OCR 不可用/失敗回原路徑）。

    2026-09-08 review：畫過色的完稿刻意不洗標號（#460 印花案，怕誤刪設計色），
    但 accent_hex_colors 的門檻（0.8%、chroma≥24）剛好看得見紅色部位標註——
    標註墨水一旦成為 ΔE 目標，正確渲染永遠驗不過、重試語還會叫模型把標註色
    畫上鞋面。文字遮罩只影響點綴色抽取這一路，渲染輸入圖與主色抽取不動。
    """
    try:
        from agent_core.vision_ocr import ocr_text_boxes
        boxes = ocr_text_boxes(path)
        if not boxes:
            return path
        from PIL import Image, ImageDraw
        with Image.open(path) as raw:
            im = raw.convert("RGB")
        w, h = im.size
        d = ImageDraw.Draw(im)
        pad = 2
        for _text, (bx, by, bw, bh) in boxes:
            # Vision 座標原點在左下 → PIL 的 top = (1 - y - h) * H
            d.rectangle((int(bx * w) - pad, int((1 - by - bh) * h) - pad,
                         int((bx + bw) * w) + pad, int((1 - by) * h) + pad),
                        fill=(255, 255, 255))
        out = os.path.join(os.path.dirname(path), "accent_mask.png")
        im.save(out)
        return out
    except Exception:  # noqa: BLE001 — 遮罩失敗退原圖，不擋生圖
        return path


def _art_color_targets(sketch: str, part_targets: list) -> list:
    """上色完稿 → ΔE 驗證目標：主色 top4 ＋ 高彩度點綴色（去重後、最多收 2 個）。

    2026-09-03 PSS 2 案：後套寶藍（Richter「lagoon」）只佔完稿 2.3% 進不了主色
    top4，色名當時又解不出 hex（source=llm 被 #458 濾網剔除）→ 這個部位全管線
    沒有任何色錨與驗證，被渲染成灰藍（實測 ΔE 36.7 遠超容差 28）照樣交件。
    點綴色補進驗證清單堵這個洞；已有同色目標（部位色票或完稿主色 ΔE ≤ 12）
    就不重複收，避免同一色塊報兩次 miss。
    2026-09-08 review 兩刀：①候選抓 6 個、去重後才截 2——彩色鞋身的明暗叢集
    排最前面，先截後去重會把名額吃光、真點綴色落空；②抽取源先過文字遮罩
    （_mask_text_for_accents），標註墨水不成為驗證目標。
    """
    from agent_core import color_anchor
    targets = [{"part": f"完稿主色{i + 1}", "hex": h}
               for i, (h, _f) in enumerate(color_anchor.dominant_hex_colors(sketch)[:4])]
    known = [color_anchor.rgb_to_lab(color_anchor.hex_to_rgb(t["hex"]))
             for t in targets + list(part_targets or []) if t.get("hex")]
    seq = 0
    for hexv, _frac in color_anchor.accent_hex_colors(
            _mask_text_for_accents(sketch), max_colors=6):
        lab = color_anchor.rgb_to_lab(color_anchor.hex_to_rgb(hexv))
        if any(color_anchor.delta_e(lab, k) <= 12.0 for k in known):
            continue
        seq += 1
        targets.append({"part": f"完稿點綴色{seq}", "hex": hexv})
        known.append(lab)
        if seq >= 2:
            break
    return targets


def generate_from_order(order_file_path: str, colorway: str = "",
                        extra_note: str = "",
                        previous_render_path: str = "") -> str:
    """客人樣品單(Excel/PDF) → 抽設計線稿/上色完稿 + 讀部位規格 → nano-banana 渲染成真實鞋圖。

    **「依這張樣品單自動生成鞋圖」就是用這顆**（開發部最常見的需求）：抽出客人的設計圖
    （xlsx 嵌入線稿、或 PDF 頁面上的線稿/上色完稿）鎖住形狀 + 讀部位材質/顏色 → 圖生圖渲染。
    比純文字生成準得多；PDF 樣品單（如 Richter PSS）有上色完稿時連配色也整套鎖定。
    視角依樣品單設計稿（通常＝單隻正側面），使用者明確要求別的角度才改。
    使用者看過成品後指出某部位錯了、其他都對 → **修圖模式**：previous_render_path
    帶上一版圖檔路徑（就是工具上次輸出的「圖檔路徑」）、extra_note 放使用者指正的
    **原話**——只重畫指正的部位，其餘與上一版完全一致，避免整張重生把已經對的
    部位改壞。
    ⚠️ 產出概念圖非生產精準圖。每日有生圖張數上限（撞到會明講）。
    材質術語：「羊巴戈」＝客人單上的 Pro Action（南亞南通廠合成皮），不是 Nubuck；
    Nubuck（牛巴戈）才是真牛皮磨砂。轉述/確認規格時材質名照樣品單原文，勿自行意譯。

    Args:
        order_file_path: 樣品單本機路徑（Telegram 上傳的 .xlsx 或 .pdf）。
        colorway: 指定配色（如「9110」），空=第一個配色。
        extra_note: 使用者的補充/修正指示——**照抄使用者原話**，不要自行改寫、
            添加視角/構圖/色系描述（使用者沒說的就不要出現）；空=無。
        previous_render_path: 修圖模式：上一版渲染圖的檔案路徑（工具上次輸出的
            「圖檔路徑」）。使用者只指正部分錯誤時必帶——只修 extra_note 點名的
            部位，其餘保留上一版原樣。空=整張重新生成。
    """
    from datetime import datetime

    if not os.path.exists(order_file_path):
        return f"❌ 找不到樣品單：{order_file_path}"
    prev_path = str(previous_render_path or "").strip()
    if prev_path:
        # 修圖模式（2026-09-07 大王條件二）：只修使用者指正的部分，其餘以上一版
        # 成品為準——整張重生會把已經對的部位重新擲骰子（修好 A 弄壞 B，
        # 「以免犯第二次錯誤」）。路徑限生圖輸出目錄（部門 context 再收窄到
        # 自己色的子目錄，比照 TG_PHOTO 白名單的分艙邏輯）。
        if not str(extra_note or "").strip():
            return ("❌ 修圖模式要同時帶 extra_note（使用者指正的原話）——"
                    "沒有指正內容就無從「只修那裡」")
        # 允許根目錄直接沿用寫入端的 _render_output_dir（review V5：先前逐字
        # 重刻一份，佈局一改驗證器就會拒絕工具自己剛輸出的路徑）。
        _root = os.path.realpath(_render_output_dir(_dept_scope_color()))
        prev_path = os.path.realpath(prev_path)
        if not prev_path.startswith(_root + os.sep) or not os.path.isfile(prev_path):
            return ("❌ previous_render_path 必須是先前生成的渲染圖檔"
                    "（generated_images 目錄下、部門僅限自己色的子目錄）")
        # 提早驗「解得開」（review V3）：Image.open 是惰性讀取，壞檔/半寫入的
        # png 會拖到扣完配額、付完規格 LLM 之後才以籠統「渲染失敗」爆掉。
        try:
            from PIL import Image as _probe_image
            with _probe_image.open(prev_path) as _probe:
                _probe.load()
        except Exception:  # noqa: BLE001
            return ("❌ previous_render_path 讀不開（檔案損壞或非圖片）——"
                    "請改用最新一次成功生成的圖檔路徑")
    pngs = _extract_order_images(order_file_path)
    if not pngs:
        return ("❌ 樣品單抽不到設計圖（支援 .xlsx 嵌入圖與 PDF 頁面線稿/完稿）；"
                "純文字單改用 parse_sample_order + generate_jf_shoe")
    sketch = _pick_shoe_sketch(pngs)
    if not sketch:
        return f"❌ 抽到 {len(pngs)} 張圖但認不出完整手繪鞋線稿（可能只有大底/logo/圖標）"
    sketch = _scrub_color_annotations(sketch)   # 標號進不了模型，就畫不出來
    # 配額先扣再開工（規格那通 LLM 也要錢）—— 撞上限就完全不打 Gemini。
    color = _dept_scope_color()
    allowed, used, limit = _consume_render_quota(color or "owner")
    if not allowed:
        return (f"⛔ 今天的樣品單生圖已達上限（{used}/{limit} 張"
                f"{'／' + color + ' 部門' if color else ''}）。明天會自動歸零；"
                "真的急需請找大王調 RED_ORDER_RENDER_DAILY_MAX。")
    try:
        order_text = _read_order_text(order_file_path)
        spec, parts, style = _spec_for_render(order_text, colorway,
                                              customer=_order_customer(order_file_path, order_text))
    except Exception as e:  # noqa: BLE001
        return f"❌ 讀規格失敗：{type(e).__name__}: {str(e)[:100]}"
    try:
        from PIL import Image
        from agent_core import color_anchor
        from agent_core.env_utils import env_float, env_int
        from agent_core.gemini_client import generate_content_tracked
        from google.genai import types as T
        sk = Image.open(sketch)
        prev_im = Image.open(prev_path) if prev_path else None
        # 上色完稿（PDF 樣品單常見）：客人自己上的色 = 最權威配色來源。
        # 全黑/深色款靠 _looks_painted 的深色路徑認（彩色佔比不夠，PSS 7 案）。
        sk_colored = _looks_painted(*_color_stats(sketch))
        # 色票條：解得出 hex 的部位畫成色塊圖，當額外輸入圖（顏色用圖錨定不用字）。
        # ⚠️ 完稿在場時 LLM 推估的 hex 不進色票（2026-09-01 PSS 8 實測：規格只寫
        # "natural color combination"，LLM 猜出一排淺米色、帶著「必須精確使用」
        # 的色票指令蓋過完稿的棕灰主色）——模型猜的永遠不能贏過客人自己畫的；
        # 確定性來源（Pantone/色彙/內建色名）才保留色票強調。
        targets = [p for p in parts if p.get("hex")]
        if sk_colored:
            # 完稿在場：LLM 推估不進色票（#458 PSS 8 案），且顏色敘述必須真的
            # 寫在樣品單原文（2026-09-02 PSS 6 案：LLM 腦補的「黑色」經內建色名表
            # 以 builtin 之姿穿過 source 濾網，黑大底蓋掉客人畫的白/灰大底）。
            targets = [p for p in targets if p.get("source") != "llm"
                       and _color_text_in_order(p.get("color_text", ""), order_text)]
        swatch_im = None
        if targets:
            sw = color_anchor.render_swatch_strip(
                targets, os.path.join(os.path.dirname(sketch), "swatch.png"))
            if sw:
                swatch_im = Image.open(sw)
        # 材質特寫：庫裡有的料附真實特寫圖（紋理也用圖錨定；空庫＝不附、零影響）。
        from agent_core import material_anchor
        mat_ims, mat_used = [], []
        for m in material_anchor.material_refs(parts):
            try:
                mat_ims.append(Image.open(m["path"]))
                mat_used.append(m)
            except Exception:  # noqa: BLE001 — 單張壞圖跳過，不擋生圖
                continue
        # 修正指示點名的料也要圖錨定（2026-09-02 PSS 6 案：「皮料改羊巴戈」只活在
        # extra_note，parts 的材質欄照單上仍是 PU → pro_action 特寫全程沒附，
        # 模型只能憑「羊巴戈」三個字想像質感）。庫裡有圖才附；沿用 3 張上限。
        if str(extra_note or "").strip():
            have = {m["key"] for m in mat_used}
            for em in material_anchor.materials_in_text(str(extra_note)):
                if em["key"] in have or len(mat_used) >= 3:
                    continue
                path = material_anchor.swatch_path(em["key"])
                if not path:
                    continue
                try:
                    mat_ims.append(Image.open(path))
                except Exception:  # noqa: BLE001
                    continue
                mat_used.append({"key": em["key"], "label": em["term"],
                                 "path": path, "parts": ["修正指示"]})
                have.add(em["key"])
        # 形體參考：款號在產品照索引有同款實照就附上（楦型/大底齒紋抄實物）。
        # ⚠️ 完稿在場時不附（2026-09-03 PSS 2 案）：索引命中的是 8/27 寄樣舊配色
        # （藍綠麂皮），與完稿新配色（深藍布面+寶藍飾片）衝突時 nano-banana 把實照
        # 整套照抄——材質配色版型全被舊樣蓋掉，正確色票＋ΔE 糾正語三連發都拉不回
        # （三 attempt 全數陣亡）。完稿本身已鎖形體＋配色＋材質分佈，實照只剩帶偏
        # 的風險；線稿路徑（無配色資訊）實照仍是加分項，照舊附。
        base_im, base_note = None, ""
        # 修圖模式也不附實照：基底＝上一版成品，任何額外照片參考都是帶偏源。
        hit = None if (sk_colored or prev_im is not None) else _find_base_photo(style)
        if hit:
            raw_photo = _download_drive_image(str(hit.get("id", "")))
            if raw_photo:
                try:
                    base_im = Image.open(io.BytesIO(raw_photo))
                    base_im.load()
                    base_note = f"{style}（{str(hit.get('file', ''))[:40]}）"
                except Exception:  # noqa: BLE001
                    base_im = None
        # 完稿主色＋高彩度點綴色直接進 ΔE 驗證——成品不得偏離客人畫的配色，
        # 小色塊（後套/織帶那類 1–4%）主色抽取看不見，另走點綴色補驗（PSS 2 案）。
        art_targets = _art_color_targets(sketch, targets) if sk_colored else []
        verify_targets = targets + art_targets
        # 修圖模式停用 ΔE（2026-09-08 review V1）：使用者的指正可能刻意偏離
        # 完稿/規格（「鞋帶改白色」），自動比對分不出「要求的改動」和「偏色」，
        # 只會把正確的修圖打回、燒光配額——修圖模式的色準由使用者驗收。
        color_check = (prev_im is None and verify_targets
                       and env_int("RED_ORDER_COLOR_CHECK", 1, min_value=0, max_value=1))
        comp_check = env_int("RED_ORDER_COMPOSITION_CHECK", 1, min_value=0, max_value=1)
        de_max = env_float("RED_ORDER_COLOR_DE_MAX", 28.0, min_value=5.0, max_value=100.0)
        out_dir = _render_output_dir(color)
        os.makedirs(out_dir, exist_ok=True)
        if color:
            _prune_dept_renders(out_dir)
        out = os.path.join(out_dir, f"order_{datetime.now().strftime('%Y%m%d_%H%M%S')}.png")
        # 洗完線稿還是有殘機率會自己補標號、上色也可能偏離色票 → OCR＋ΔE 驗過
        # 再交件，髒的/偏色的就重生一張。每次重生都真的再打一次 nano-banana，
        # 所以照樣扣配額（不然日花費上限會被重試偷偷放大 N 倍）；配額用完就交
        # 手上這張並說清楚。
        max_attempts = env_int("RED_ORDER_RENDER_MAX_ATTEMPTS", 3, min_value=1, max_value=5)
        digits: list = []
        color_misses: list = []
        comp_flaws: list = []
        digit_retry = False
        got_image = False
        attempt = 0
        while attempt < max_attempts:
            attempt += 1
            if attempt > 1:
                ok, _used, _limit = _consume_render_quota(color or "owner")
                if not ok:
                    break
            prompt = _render_prompt(spec, retry=digit_retry,
                                    swatch_parts=targets if swatch_im else None,
                                    color_misses=color_misses,
                                    material_refs=mat_used or None,
                                    base_style=style if base_im else "",
                                    colored_sketch=sk_colored,
                                    extra_note=str(extra_note or "").strip()[:1000],
                                    comp_flaws=comp_flaws,
                                    edit_previous=prev_im is not None)
            # 順序＝prompt 的描述順序：（修圖模式先上一版成品）→ 線稿/完稿 →
            # 同款實照 → 材質特寫們 → 色票（「最後一張」）。
            contents = (([prev_im] if prev_im else []) + [sk]
                        + ([base_im] if base_im else []) + mat_ims
                        + ([swatch_im] if swatch_im else []) + [prompt])
            resp = generate_content_tracked(
                model="nano-banana-pro-preview",
                contents=contents,
                config=T.GenerateContentConfig(response_modalities=["IMAGE", "TEXT"]),
                caller="sample_order.generate_from_order")
            data = _first_inline_image(resp)
            if not data:
                continue
            with open(out, "wb") as f:
                f.write(data)
            got_image = True
            digits = _rendered_digit_labels(out)
            color_misses = (color_anchor.find_color_misses(out, verify_targets, de_max)
                            if color_check else [])
            # 修圖模式不與完稿比對外來色、不硬驗正側面（review V1）：上一版的
            # 視角/配色已被使用者接受，比對只會打回使用者要的改動；留「單隻」。
            comp_flaws = (_composition_flaws(
                out, artwork_path="" if prev_im else (sketch if sk_colored else ""))
                if comp_check else [])
            if prev_im:
                comp_flaws = [f for f in comp_flaws if "視角" not in str(f)]
            if not digits and not color_misses and not comp_flaws:
                break
            digit_retry = bool(digits)
        # 輸入圖用完就關（2026-09-08 review：未關的 PIL handle 積到 GC 噴
        # ResourceWarning，還會把整張 raster 留在記憶體撐完整個函式）。
        for _im in [sk, prev_im, base_im, swatch_im] + mat_ims:
            try:
                if _im is not None:
                    _im.close()
            except Exception:  # noqa: BLE001
                pass
        if not got_image:
            return f"❌ nano-banana 沒回圖。規格：{spec[:100]}"
        # 員工一律、大王在 Telegram 對話發起時：圖跟著**回覆**走（daemon 抽
        # [[TG_PHOTO:]] 標記，用當前 bot token + 當前 chat_id 傳回；不經出站
        # 推送閘）。大王也走標記的原因（2026-09-02 green 聊天室 5004 案）：
        # 出站推送只認 keyring 預設 chat —— 在 green agent 聊天室要的圖會從
        # 紅 bot 主對話冒出來。大王的標記白名單本來就是整個 generated_images
        # （見 daemon_telegram._reply_photo_allowed_roots）。
        deliver_by_marker = bool(color)
        if not deliver_by_marker:
            try:
                from agent_core.channel_context import reply_consumes_tg_markers
                deliver_by_marker = reply_consumes_tg_markers()
            except Exception:  # noqa: BLE001
                deliver_by_marker = False
        if deliver_by_marker:
            delivery = (f"[[TG_PHOTO:{out}]]\n"
                        "⚠️ 回覆時請把上面 [[TG_PHOTO:...]] 標記行原樣保留在回覆最後，"
                        "系統會自動把圖傳給對方（標記本身不會顯示）。")
        else:
            # 大王在 REPL / 背景 daemon：沒有消費標記的回覆送出點，維持
            # best-effort 主動傳回 Telegram（chat_id 空=預設）。
            try:
                from agent_core.telegram import telegram_send_photo
                telegram_send_photo(out, caption="樣品單線稿驅動生成鞋圖", chat_id="")
            except Exception:  # noqa: BLE001
                pass
            delivery = ""
        # 色錨摘要：每部位落到哪個 hex、憑什麼（查表/色彙/內建/LLM 推估）——
        # 大王一眼就能看出哪個顏色是「有依據」哪個是「模型猜的」。
        src_zh = {"pantone": "Pantone 對照", "glossary": "客戶色彙", "builtin": "內建色名",
                  "llm": "LLM 推估⚠️"}
        anchor = ""
        if sk_colored:
            arts = " ".join("#" + t["hex"] for t in art_targets)
            anchor += ("   🖼️ 客人上色完稿：設計與配色整套鎖定"
                       + (f"（完稿鎖色 {arts}）" if arts else "") + "\n")
        if targets:
            rows = [f"{p['part']} #{p['hex']}（{p['label']}｜{src_zh.get(p['source'], '?')}）"
                    for p in targets[:6]]
            anchor += "   🎨 色錨：" + "；".join(rows) + "\n"
        if mat_used:
            anchor += ("   🧵 材質特寫已附："
                       + "；".join(f"{m['label'] or m['key']}→{'/'.join(m['parts'][:3])}"
                                   for m in mat_used) + "\n")
        if base_note:
            anchor += f"   📸 形體參考已附：同款實照 {base_note}\n"
        if prev_path:
            anchor += (f"   🖌️ 修圖模式：以上一版 {os.path.basename(prev_path)} 為基底，"
                       "只修指正部位、其餘保留（色準與視角由你驗收，自動驗證只查標號與單隻）\n")
        note_full = str(extra_note or "").strip()
        if note_full:
            anchor += f"   ✏️ 已套用修正指示：{note_full[:80]}\n"
            if len(note_full) > 1000:
                anchor += (f"   ⚠️ 修正指示 {len(note_full)} 字超過 1000 字上限、"
                           "尾段已截斷——過長的指正請分次修。\n")
        # 還是髒的/偏色的就照實講，別讓員工以為這張是乾淨的（UserC 要的就是無數字）。
        warn = ("" if not digits else
                f"   ⚠️ 這張仍殘留部位標號（OCR 讀到 {'/'.join(digits[:6])}），"
                f"已重生 {attempt} 次仍未乾淨——需要的話說「再生成一張」。\n")
        if color_check:
            warn += ("   ✅ 色準驗證：各部位主色皆在指定色票容差內（ΔE ≤ "
                     f"{de_max:.0f}）。\n" if not color_misses else
                     "   ⚠️ 色準未達標："
                     + "、".join(f"{m['part']}（ΔE {m['delta_e']}）" for m in color_misses[:4])
                     + f"——已重生 {attempt} 次仍偏色；螢幕近似色僅供確認方向，"
                       "正式對色請以實體色卡為準。\n")
        if comp_flaws:
            warn += ("   ⚠️ 構圖仍不符（" + "、".join(str(f) for f in comp_flaws[:2])
                     + f"）——已重生 {attempt} 次；需要的話說「再生成一張」。\n")
        return (f"✅ 樣品單「線稿驅動」生成鞋圖（鎖客人手繪形狀 + 規格上材質）：\n"
                f"   • {out}\n   規格：{spec[:130]}\n"
                + anchor +
                "   ⚠️ 概念圖、非生產精準；線稿抽自樣品單嵌入圖。\n"
                + warn + delivery)
    except Exception as e:  # noqa: BLE001
        return f"❌ 渲染失敗：{type(e).__name__}: {str(e)[:120]}"


SKILL_TOOLS = [parse_sample_order, generate_from_order]
