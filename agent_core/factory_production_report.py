"""結構化讀「X月份生產日報進度表.xlsx」(福群鞋廠生管日報)。

為什麼要這支工具（而不是 read_drive_file）：
  那張 Drive 進度表是「每天一個分頁(01..31)、每頁 ~990 列 PO × 71 欄」的超寬
  交叉表，全文匯成純文字約 158 萬字。read_drive_file 只給前 3 萬字(約 1.9%)，
  而且合併的中越文表頭一攤平就糊掉 → LLM 擠不出「每天每客戶幾雙」。小紅因此
  長期只能退回 email 拼湊(Decathlon 那幾天只剩『排程中/趕產』這種模糊值)。

這支工具直接用 openpyxl 解析，輸出精簡的：
  1. 每日 × 每客戶 各站日產量(日計)
  2. 截至最新一天 各客戶 包裝累計 / 未完(欠數)
     （本季已出完貨的客戶降成一行結案摘要，不佔進度表 —— 大王規則「已經出完貨的
       就不需要回報」，每日生產回報照這段念）
  3. 近期應出貨而尚未出貨的 PO

表頭以「文字錨點」偵測(客戶/指令/雙數 + 英文站名 Stitching/Injection/Packing)，
不寫死欄位位置；萬一下個月版型欄位位移，會在輸出標出 ⚠️ 提醒而非靜默算錯。
"""
import datetime as _dt
import io
import os
import re
import time as _time
import unicodedata

import openpyxl

# 英文站名(穩定錨點) → 中文顯示名。成型(Molding)在 PU 鞋走「灌注」故常為 0。
_STATIONS = [
    ("Stitching", "針車"),
    ("Molding", "成型"),
    ("Insock", "中底"),
    ("Injection", "灌注"),
    ("Packing", "包裝"),
]

# 廠內口語 → 英文錨點。「射出」是 PU 灌注成型那站的現場叫法（生管日報欄名寫「灌注」）。
_STATION_ALIASES = {"射出": "Injection", "射出成型": "Injection", "灌注成型": "Injection"}

# 產能折線圖的三站（生產管理部要看的：針車→射出→包裝 這條主線）。
_CAPACITY_STATIONS = ("Stitching", "Injection", "Packing")
_CAPACITY_LABELS = {"Stitching": "針車", "Injection": "射出(灌注)", "Packing": "包裝"}

# 站別英文錨點 → iter_production_rows() 對外 row dict 欄位名（小寫，對齊數據倉 schema）。
_STATION_FIELDS = {
    "Stitching": "stitching_day",
    "Molding": "molding_day",
    "Insock": "insock_day",
    "Injection": "injection_day",
    "Packing": "packing_day",
}

# 一張進度表內、被視為「某一天」的分頁名樣式：純 1–2 位數(01..31)。
_DAY_SHEET_RE = re.compile(r"^\s*(\d{1,2})\s*$")

_DEFAULT_PROGRESS_KEYWORD = "生產日報進度表"


def _to_num(v):
    """Excel 值 → float(取不出數字回 0.0)。"""
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _norm(s):
    """表頭比對用：去空白/換行。"""
    return re.sub(r"\s+", "", str(s or ""))


def _year_of(want_ship, work_order):
    """訂單年份：希望出貨日年份優先，取不到再退指令前綴(JFC26→2026)。回 int 或 None。

    用來把「往年下單、被生管原封複製到本月每天分頁」的舊單（NISSHIN/LURCHI 那種，
    多已出貨結案）跟本季訂單分開——**以年份判定，不以是否已出貨**（後者會誤殺本季
    已出貨的單，例 DECA 今年出掉的 10 萬雙）。生管日報本就是生產視圖，非訂單真相；
    訂單真實量以業務訂單夾為準。
    """
    if isinstance(want_ship, (_dt.datetime, _dt.date)):
        return want_ship.year
    m = re.search(r"(20\d\d)", str(want_ship or ""))
    if m:
        return int(m.group(1))
    m = re.match(r"\s*JFC(\d{2})\d", str(work_order or ""))
    if m:
        return 2000 + int(m.group(1))
    return None


def _order_year(row):
    """同 _year_of，但吃 _normalize_row 的 row dict。"""
    return _year_of(row.get("want_ship"), row.get("work_order"))


def _is_shipped(row):
    """這列是否已出貨：實際出貨欄有值（非空 / 非 NONE）。"""
    actual = row.get("actual_ship")
    return bool(actual) and str(actual).strip().upper() not in ("", "NONE")


def _detect_layout(rows):
    """從前幾列找表頭，回 (data_start_idx0, colmap) 或 (None, warnings)。

    rows: list[tuple]  (values_only，0-based)
    colmap: dict 欄名→欄索引(0-based)。站別日計鍵為 '<en>_day'，
            另有 pack_cum / pack_rem。
    """
    warnings = []
    header_i = None
    for i in range(min(8, len(rows))):
        cells = [_norm(c) for c in rows[i]]
        if "客戶" in cells and any("雙數" in c for c in cells):
            header_i = i
            break
    if header_i is None:
        return None, ["找不到表頭列(預期含『客戶』與『雙數』)"]

    hdr = [_norm(c) for c in rows[header_i]]
    sub = [_norm(c) for c in rows[header_i + 1]] if header_i + 1 < len(rows) else []

    def find(pred, start=0):
        for j in range(start, len(hdr)):
            if pred(hdr[j]):
                return j
        return None

    colmap = {}
    colmap["customer"] = find(lambda c: c == "客戶")
    colmap["work_order"] = find(lambda c: c == "指令")
    colmap["model"] = find(lambda c: c == "型體")
    # 客戶型體＝業務訂單上的 Supremo 型體號(74L/63L/36L…)，與「型體」(生管型體)同列；
    # 是業務訂單型體↔生管型體的現成對照（小紅按訂單型體查生產用）。
    colmap["cust_style"] = find(lambda c: c == "客戶型體")
    colmap["pairs"] = find(lambda c: "雙數" in c)
    colmap["want_ship"] = find(lambda c: "希望" in c and "出貨" in c)
    colmap["actual_ship"] = find(lambda c: "實際" in c and "出貨" in c)

    def sub_is(j, *needles):
        if j is None or j >= len(sub):
            return False
        return any(n in sub[j] for n in needles)

    # 站別日計：英文站名所在欄(站名合併在該站日計欄上)，並驗證 sub 列是「日計/day」。
    for en, _zh in _STATIONS:
        j = find(lambda c, e=en.lower(): e in c.lower())
        if j is not None and (not sub or sub_is(j, "日計", "day", "S/LTR", "SLTR") or not sub[j]):
            colmap[f"{en}_day"] = j

    # 包裝累計/欠數：以 Packing 日計欄為基準往左找最近的「累計(非不良NG)」與「欠數」。
    pack_day = colmap.get("Packing_day")
    if pack_day is not None:
        for j in range(pack_day - 1, max(pack_day - 6, -1), -1):
            if j >= len(sub):
                continue
            if "累計" in sub[j] and "NG" not in sub[j] and "不良" not in sub[j] and "pack_cum" not in colmap:
                colmap["pack_cum"] = j
            if ("欠數" in sub[j] or "remai" in sub[j].lower()) and "pack_rem" not in colmap:
                colmap["pack_rem"] = j
        # 解不到 → 整欄會被 _to_num 靜默算成 0（未完全綠假象）；必須標警告，不能靜默。
        if "pack_cum" not in colmap:
            warnings.append("找不到『包裝累計』欄(Packing 左側無累計子標——版型可能位移，累計會算成 0)")
        if "pack_rem" not in colmap:
            warnings.append("找不到『欠數』欄(Packing 左側無欠數子標——版型可能位移，未完會算成 0)")

    # 必要欄缺失 → 警告(仍盡量解析)。
    for key in ("customer", "work_order", "pairs"):
        if colmap.get(key) is None:
            warnings.append(f"表頭缺『{key}』欄")
    if not any(f"{en}_day" in colmap for en, _ in _STATIONS):
        warnings.append("找不到任何站別『日計』欄(版型可能變了)")

    return header_i + 2, (colmap, warnings)


def _iter_day_sheets(wb):
    """yield (day_int, sheet_name)；只挑分頁名是 1–2 位數者，依日排序。"""
    out = []
    for name in wb.sheetnames:
        m = _DAY_SHEET_RE.match(name)
        if m:
            out.append((int(m.group(1)), name))
    out.sort()
    return out


def _open_workbook(xlsx_bytes):
    """openpyxl 讀 workbook bytes（read_only 串流、data_only）。失敗讓呼叫端接例外。"""
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)


def _parse_day_filter(day):
    """'14' / '6/14' / '06-14' → '14'；空或無法解析 → ''。"""
    if not day:
        return ""
    m = re.search(r"(\d{1,2})\s*$", str(day).strip())
    return str(int(m.group(1))) if m else ""


def _normalize_row(r, colmap, day_int, latest_day):
    """單一資料列 + colmap → 正規化 dict；無客戶名（空列）回 None。

    站別日計用小寫欄位名（stitching_day…，對齊倉 schema）；want_ship/actual_ship
    保留原始 cell（datetime/str/None），交由呼叫端決定如何呈現／落庫。
    """
    def cell(key):
        j = colmap.get(key)
        return r[j] if (j is not None and j < len(r)) else None

    cust = cell("customer")
    if not cust or not str(cust).strip():
        return None
    row = {
        "prod_day": day_int,
        "is_latest_day": day_int == latest_day,
        "customer": str(cust).strip(),
        "work_order": str(cell("work_order") or "").strip(),
        "model": str(cell("model") or "").strip(),
        "cust_style": str(cell("cust_style") or "").strip(),
        "pairs": _to_num(cell("pairs")),
        "pack_cum": _to_num(cell("pack_cum")),
        "pack_rem": _to_num(cell("pack_rem")),
        "want_ship": cell("want_ship"),
        "actual_ship": cell("actual_ship"),
    }
    for en, field in _STATION_FIELDS.items():
        row[field] = _to_num(cell(f"{en}_day"))
    return row


def _iter_rows_for_workbook(wb, day_sheets=None, *, warnings=None, day_filter=""):
    """共用核心：對已開啟的 workbook，逐日分頁、逐列 yield 正規化 row dict。

    day_sheets: 預算好的 _iter_day_sheets(wb)（省重算）；None 則自行算。
    warnings:   選填 set，回填版型偵測警告（解不到表頭／缺站別欄）。
    day_filter: 選填，只跑該日分頁（'14'）—— 給單日查詢省去掃全部分頁；
                latest_day 仍以全部分頁的最後一天為準（is_latest_day 才正確）。
    """
    if day_sheets is None:
        day_sheets = _iter_day_sheets(wb)
    if not day_sheets:
        return
    latest_day = day_sheets[-1][0]
    for day_int, name in day_sheets:
        if day_filter and str(day_int) != day_filter:
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        data_start, layout = _detect_layout(rows)
        if data_start is None:
            if warnings is not None:
                warnings.update(layout)  # 此時 layout 是 warnings list
            continue
        colmap, warns = layout
        if warnings is not None:
            warnings.update(warns)
        for r in rows[data_start:]:
            row = _normalize_row(r, colmap, day_int, latest_day)
            if row is not None:
                yield row


def iter_production_rows(xlsx_bytes, *, warnings=None):
    """公開入口：解析生產日報 xlsx bytes，逐列 yield 正規化 row dict（給數據倉 loader）。

    row dict 欄位：prod_day, is_latest_day, customer, work_order, model, pairs,
    stitching_day, molding_day, insock_day, injection_day, packing_day,
    pack_cum, pack_rem, want_ship(raw cell), actual_ship(raw cell)。

    壞檔／非日分頁結構 → 不丟例外、yield 0 列（warnings 若給會記下原因）。
    這是「一份解析、多個消費者」的咽喉點：read_production_progress_sheet 的文字
    彙總與數據倉 loader 都走這條，版型偵測邏輯只有一份。
    """
    try:
        wb = _open_workbook(xlsx_bytes)
    except Exception as exc:  # noqa: BLE001
        if warnings is not None:
            warnings.add(f"解析 xlsx 失敗：{type(exc).__name__}: {exc}")
        return
    day_sheets = _iter_day_sheets(wb)
    if not day_sheets:
        if warnings is not None:
            warnings.add("找不到以日為分頁(01..31)的結構")
        return
    yield from _iter_rows_for_workbook(wb, day_sheets, warnings=warnings)


def _season_status(latest_rows, *, rem_col_broken=False):
    """最新日的列 → (本季在產, 本季已出完貨, 往年舊單) 三份 per-客戶彙總。

    「哪些客戶還要追」的**單一出處** —— _summarize 的文字彙總與 daily_production_report
    的每日回報共用，免得兩個消費者各判各的（同一份表念出兩套客戶清單是最難查的那種錯）。

    current: cust -> {pairs, shipped, pack_cum, pack_rem, pos:[(want, cust, wo, model, rem)]}
    done:    同上結構，但本季已全數出貨且欠數歸零 —— 大王規則「已經出完貨的就不需要回報」。
    prior:   cust -> {orders, pairs}，往年下單且已出貨的舊單。

    rem_col_broken: 欠數欄解析失敗（pack_rem 全 0 是假象）→ 停用 done 分流，寧可多念。
    """
    # current_year = 表上「雙數最多」的訂單年份（主導生產季），不是 max——表內常有
    # 0 雙的下一季(如 2027)佔位列，用 max 會把本季全打成往年。自我校準免寫死年份。
    year_pairs = {}
    for r in latest_rows:
        yy = _order_year(r)
        if yy is not None:
            year_pairs[yy] = year_pairs.get(yy, 0.0) + r["pairs"]
    current_year = max(year_pairs, key=year_pairs.get) if year_pairs else None

    current_status, prior_status = {}, {}
    for row in latest_rows:
        cust = row["customer"]
        y = _order_year(row)
        # 大王規則「去年出掉的就要排除」＝往年下單『且已出貨』→ 另列、不計本季。
        # 本季已出貨(DECA 今年出掉的)留在本季；往年但未出貨者也留(仍在產)。
        if current_year is not None and y is not None and y < current_year and _is_shipped(row):
            ps = prior_status.setdefault(cust, {"orders": 0, "pairs": 0.0})
            ps["orders"] += 1
            ps["pairs"] += row["pairs"]
            continue
        st = current_status.setdefault(
            cust, {"pairs": 0.0, "shipped": 0.0, "pack_cum": 0.0, "pack_rem": 0.0, "pos": []})
        st["pairs"] += row["pairs"]
        st["pack_cum"] += row["pack_cum"]
        rem = row["pack_rem"]
        st["pack_rem"] += rem
        if _is_shipped(row):
            st["shipped"] += row["pairs"]
        else:
            want = row["want_ship"]
            if want and rem > 0:
                st["pos"].append((str(want)[:10], cust, row["work_order"], row["model"], rem))

    # 大王規則（2026-08-07）「已經出完貨的就不需要回報」：本季在表全數出貨、且欠數
    # 歸零的客戶另列 —— 每日生產回報留著他們就是每天多一列全 0 的雜訊（8/6 那封的
    # NEW WAVE 就是：整季 1,440 雙早出完，三站全是「-」）。
    # 分流條件要「在產 0 **且** 未完 0」兩個都成立：包裝完但還沒出（在產>0）仍要追。
    # 這跟上面「往年已出貨另列」是同一族規則，差在本季／往年。
    done_status = {}
    if not rem_col_broken:
        for cust in [c for c, st in current_status.items()
                     if st["pairs"] > 0 and st["pack_rem"] <= 0 and st["shipped"] >= st["pairs"]]:
            done_status[cust] = current_status.pop(cust)
    return current_status, done_status, prior_status


def _summarize(xlsx_bytes, day="", customer="", month_hint=""):
    """解析 workbook bytes → 給 LLM 的精簡報告字串（跑在 _iter_rows_for_workbook 上）。"""
    cust_filter = (customer or "").strip().lower()
    day_filter = _parse_day_filter(day)

    try:
        wb = _open_workbook(xlsx_bytes)
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 解析 xlsx 失敗：{type(exc).__name__}: {exc}"

    day_sheets = _iter_day_sheets(wb)
    if not day_sheets:
        return ("⚠️ 這份檔找不到『以日為分頁(01..31)』的結構 — 可能不是生產日報進度表，"
                "或版型不同。分頁有：" + ", ".join(wb.sheetnames[:20]))

    latest_day = day_sheets[-1][0]
    all_warnings = set()
    per_day = {}            # day_int -> {cust: {station_zh: total}}
    latest_rows = []        # 最新日的列；後處理依「訂單年份」分本季 / 往年舊單

    for row in _iter_rows_for_workbook(wb, day_sheets, warnings=all_warnings, day_filter=day_filter):
        cust = row["customer"]
        if cust_filter and cust_filter not in cust.lower():
            continue
        agg = per_day.setdefault(row["prod_day"], {}).setdefault(cust, {})
        for en, zh in _STATIONS:
            d = row[_STATION_FIELDS[en]]
            if d:
                agg[zh] = agg.get(zh, 0.0) + d
        if row["is_latest_day"]:
            latest_rows.append(row)

    current_status, done_status, prior_status = _season_status(
        latest_rows, rem_col_broken=any("欠數" in w for w in all_warnings))

    # ---- 組裝輸出 ----
    title = month_hint or "生產日報進度表"
    days_present = [d for d, _ in day_sheets]
    lines = [f"📊 {title}"]
    if days_present:
        lo, hi = days_present[0], days_present[-1]
        missing = [d for d in range(lo, hi + 1) if d not in days_present]
        miss_txt = f"，缺 {'/'.join(str(m) for m in missing)}(多為週日休工)" if missing else ""
        lines.append(f"涵蓋生產日：{lo}–{hi} 號，共 {len(days_present)} 天{miss_txt}")
    lines.append("單位：雙。日計=當日各站產出；成型(Molding)此廠走 PU 灌注故多為 0。")
    if all_warnings:
        lines.append("⚠️ 解析提醒：" + "；".join(sorted(all_warnings)) + "（數字僅供參考，建議人工複核）")

    lines.append("")
    lines.append("【每日 × 各客戶 生產數量(日計)】")
    if not per_day or all(not dc for dc in per_day.values()):
        lines.append("(此條件下查無日產量資料)")
    for day_int in sorted(per_day):
        day_cust = per_day[day_int]
        if not day_cust:
            continue
        segs = []
        for cust, agg in sorted(day_cust.items(), key=lambda kv: -sum(kv[1].values())):
            inner = " ".join(f"{zh}{int(round(v))}" for zh, v in agg.items())
            if inner:
                segs.append(f"{cust}: {inner}")
        if segs:
            lines.append(f"{day_int}號  " + " ｜ ".join(segs))

    if current_status and not day_filter:
        lines.append("")
        lines.append(f"【截至 {latest_day} 號 本季各客戶 生產進度】（生管在表數，非業務訂單總額；訂單真實量請查業務訂單夾）")
        for cust, st in sorted(current_status.items(), key=lambda kv: -kv[1]["pack_rem"]):
            in_prod = st["pairs"] - st["shipped"]
            lines.append(
                f"{cust}: 本季在表 {int(st['pairs']):,} 雙（已出貨 {int(st['shipped']):,}｜在產 {int(in_prod):,}）"
                f"｜已包裝累計 {int(st['pack_cum']):,}｜未完 {int(st['pack_rem']):,}")

        # 近期應出貨而未出貨(依希望出貨日)
        ships = sorted(
            (p for st in current_status.values() for p in st["pos"]),
            key=lambda p: p[0],
        )[:15]
        if ships:
            lines.append("")
            lines.append("【近期應出貨而未出貨(客戶希望出貨日 / 未完雙數)】最多 15 筆")
            for want, cust, wo, model, rem in ships:
                lines.append(f"  {want}  {cust}  {wo} {model[:18]}  未完 {int(rem):,}")

    if done_status and not day_filter:
        lines.append("")
        tot_p = int(sum(st["pairs"] for st in done_status.values()))
        lines.append(f"【本季已出完貨（在產 0、未完 0）—— 每日回報不需列入】"
                     f"共 {len(done_status)} 個客戶 {tot_p:,} 雙")
        for cust, st in sorted(done_status.items(), key=lambda kv: -kv[1]["pairs"]):
            lines.append(f"  {cust}: 本季在表 {int(st['pairs']):,} 雙 已全數出貨"
                         f"（已包裝累計 {int(st['pack_cum']):,}）")

    if prior_status and not day_filter:
        lines.append("")
        tot_o = sum(c["orders"] for c in prior_status.values())
        tot_p = int(sum(c["pairs"] for c in prior_status.values()))
        lines.append(f"【往年舊單（往年下單且已出貨，不計本季）】共 {tot_o} 單 {tot_p:,} 雙")
        for cust, c in sorted(prior_status.items(), key=lambda kv: -kv[1]["pairs"]):
            lines.append(f"  {cust}: {c['orders']} 單 {int(c['pairs']):,} 雙")

    return "\n".join(lines)


def _find_latest_progress_file(keyword=_DEFAULT_PROGRESS_KEYWORD):
    """回 (file_id, name) — Drive 上檔名含 keyword、修改日最新的那份。找不到回 (None, None)。"""
    from agent_core.google_auth import get_service
    service = get_service("drive", "v3")
    safe = keyword.replace("\\", "\\\\").replace("'", "\\'")
    res = service.files().list(
        q=f"name contains '{safe}' and trashed = false",
        pageSize=25,
        fields="files(id, name, modifiedTime, mimeType)",
        includeItemsFromAllDrives=True,
        supportsAllDrives=True,
        corpora="allDrives",
    ).execute()
    items = [f for f in res.get("files", []) if "spreadsheet" in f.get("mimeType", "")
             or str(f.get("name", "")).lower().endswith((".xlsx", ".xls", ".xlsm"))]
    if not items:
        return None, None
    items.sort(key=lambda f: f.get("modifiedTime", ""), reverse=True)
    return items[0]["id"], items[0]["name"]


def _download_xlsx_bytes(file_id):
    """下載 Drive 檔成 bytes(記憶體，不落地)。"""
    from agent_core.google_auth import get_service
    from googleapiclient.http import MediaIoBaseDownload
    service = get_service("drive", "v3")
    meta = service.files().get(
        fileId=file_id, fields="id,name,modifiedTime,mimeType,size",
        supportsAllDrives=True,
    ).execute()
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, service.files().get_media(fileId=file_id, supportsAllDrives=True),
                             chunksize=8 * 1024 * 1024)
    done = False
    while not done:
        _status, done = dl.next_chunk()
    return buf.getvalue(), meta


def read_production_progress_sheet(file_id: str = "", day: str = "", customer: str = "") -> str:
    """讀「X月份生產日報進度表.xlsx」(福群鞋廠生管日報)，整理出每天每客戶的**生產進度**。

    這是查「每天/這個月 各客戶生產多少雙」「Decathlon/Jalas 進度」的**首選**工具 —
    比 read_drive_file 可靠(那張表 158 萬字會被截到只剩 1.9%、表頭攤平糊掉)。
    ⚠️ 這是**生產進度**（本季在表/已出貨/在產/累計/未完），**不是訂單總額**；查客戶
    『訂單下多少雙 / 有哪些 PO』請改用 read_customer_order_pos（讀業務訂單夾，權威來源）。

    Args:
        file_id: Drive 檔案 ID。**留空**會自動抓 Drive 上最新一份『X月份生產日報進度表』。
                 (檔名來源：search_drive_files('生產日報進度表') 的第一筆)
        day:     只看某一天，填日(如 '14'、'6/14'、'06-14')；留空=整月每天。
        customer: 只看某客戶，填關鍵字(如 'DECA'、'JALAS'、'LURCHI')；留空=全部。

    Returns:
        精簡報告：每日×各客戶 各站日產量(針車/灌注/包裝…) + 截至最新日的包裝累計/未完
        + 近期應出貨未出貨清單。引用時標明「檔名＋修改日」。
        本季**已出完貨**(在產 0 且未完 0)的客戶不在進度表內，另列一行結案摘要 —— 做
        每日回報時那一區整區略過（他們已經沒有進度可回報）。
    """
    fid = (file_id or "").strip()
    name = ""
    modified = ""
    try:
        if not fid:
            fid, name = _find_latest_progress_file()
            if not fid:
                return ("找不到任何『生產日報進度表』檔。請確認 Drive 上有"
                        "『X月份生產日報進度表XX-XX.xlsx』，或改用 search_drive_files 找正確檔名後傳 file_id。")
        xlsx_bytes, meta = _download_xlsx_bytes(fid)
        name = name or str(meta.get("name", ""))
        modified = str(meta.get("modifiedTime", ""))[:10]
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 讀取生產日報進度表失敗(file_id={fid})：{type(exc).__name__}: {exc}"

    month_hint = f"{name}（修改 {modified}）" if name else ""
    return _summarize(xlsx_bytes, day=day, customer=customer, month_hint=month_hint)


# ─────────────────────────────────────────────────────────────────────────────
# 每日生產數量回報（daily_production_8am 排程用的確定性版本）
#
# 為什麼要確定性：這封信本來是 Gemini 讀完 read_production_progress_sheet 自由發揮
# 產出的表格，同一份資料每輪產出都不一樣 —— 2026-08-07 那天早上四輪實測：10:09 只列
# 3 家（RICHTER 有 426 雙未完卻整家被漏掉）、10:40 列 4 家、11:16 才把 NEW WAVE 也列
# 進來。多列一家已出完貨的是雜訊，漏列一家還在產的是漏報。而且去重是拿整段文字的
# sha1 比對，LLM 每輪重寫 → 雜湊永遠不同 → 同一份資料一個上午寄四封。
# 表格改由程式排版後：客戶清單、欄位、數字每天固定，同一份資料只寄一封。
# ─────────────────────────────────────────────────────────────────────────────

_DAILY_FOCUS_N = 3        # ERP 交期核對抽查幾張（最急的前 N 張）
_DAILY_CHASE_HOUR = 11    # 過這個鐘點還沒涵蓋昨天才催收；之前安靜等下一輪
_DAILY_STATIONS = (("Injection", "灌注（成型）"), ("Packing", "包裝"), ("Stitching", "針車"))
# ERP 認定這單已經結束的狀態；日報卻還有未完 → 兩邊不一致，要標出來給人判斷
# （完成度以生管日報為權威，同 erp_delivery_risk_alert 的劃界，這裡只標不改數字）。
_ERP_CLOSED_STATES = ("完工", "出货", "出貨", "销货", "銷貨")


def _daily_report_rows(xlsx_bytes, yesterday_day):
    """解析進度表 → 每日回報要的料。

    回 dict：
      latest_day  截至第幾號（累計/未完 的口徑）
      day_totals  {cust: {站別中文: 當日雙數}}  昨天那一天
      current     還要追的客戶（_season_status 的本季在產）
      done_count/done_pairs  本季已出完貨的家數/雙數（只報數量，不列名）
      focus       最急的未完清單 [(want, cust, wo, model, rem)]，依希望出貨日升冪
      warnings    版型偵測警告
    """
    wb = _open_workbook(xlsx_bytes)
    day_sheets = _iter_day_sheets(wb)
    if not day_sheets:
        raise ValueError("這份檔找不到『以日為分頁(01..31)』的結構 —— 可能不是生產日報進度表")
    latest_day = day_sheets[-1][0]

    warns = set()
    day_totals = {}
    # 只解「昨天」與「最新日」兩個分頁：整月 30 頁全解要多花十幾秒，而這份報表只需要
    # 這兩天（日產量看昨天、累計/未完看最新日）。day_filter 不影響 is_latest_day 判定。
    for row in _iter_rows_for_workbook(wb, day_sheets, warnings=warns,
                                       day_filter=str(yesterday_day)):
        agg = day_totals.setdefault(row["customer"], {})
        for en, zh in _DAILY_STATIONS:
            d = row[_STATION_FIELDS[en]]
            if d:
                agg[zh] = agg.get(zh, 0.0) + d

    latest_rows = list(_iter_rows_for_workbook(wb, day_sheets, warnings=warns,
                                               day_filter=str(latest_day)))
    current, done, _prior = _season_status(
        latest_rows, rem_col_broken=any("欠數" in w for w in warns))

    focus = sorted((p for st in current.values() for p in st["pos"]),
                   key=lambda p: _ship_date_sort_key(p[0]))
    return {
        "latest_day": latest_day,
        "day_totals": day_totals,
        "current": current,
        "done_count": len(done),
        "done_pairs": int(sum(st["pairs"] for st in done.values())),
        "focus": focus,
        "warnings": sorted(warns),
    }


def _erp_order_facts(work_orders):
    """指令號 → ERP 鏡像的 {base單號: {交期, 狀態, 數量}}。查不到就少一筆，不丟例外。

    日報「指令」帶產線後綴（JFC26367-1-1），ERP 單號是 base 號 —— 剝法同
    erp_delivery_risk_alert。鏡像只有 CR_REQDATE（客戶交期），**沒有 FAC_DATE(交廠日)**，
    所以欄位就照實叫「ERP 客戶交期」；要交廠日得連線上 Oracle，背景報表不做。
    """
    bases = []
    for wo in work_orders:
        base = str(wo or "").strip().upper().split("-")[0]
        if base and base not in bases:
            bases.append(base)
    if not bases:
        return {}
    from agent_core.erp_stock_query import _db_ready, _run
    if not _db_ready():
        return {}
    placeholders = ", ".join("?" for _ in bases)
    rows = _run(
        # v_orders 是 (單號, 訂單序) 級：同一單多個 SEQ 各有交期/狀態，聚合成一列。
        f"SELECT 單號, MIN(substr(交期, 1, 10)), SUM(數量), "
        f"       string_agg(DISTINCT 狀態, '/') "
        f"FROM v_orders WHERE 單號 IN ({placeholders}) GROUP BY 單號",
        bases,
    )
    return {str(r[0]).strip().upper(): {"交期": str(r[1] or ""), "數量": int(r[2] or 0),
                                        "狀態": str(r[3] or "")}
            for r in rows}


def _erp_delivery_check_section(focus):
    """最急的前 N 張未完 → ERP 交期核對表（markdown）。ERP 讀不到就回說明、不讓整封信掛掉。"""
    picks = focus[:_DAILY_FOCUS_N]
    if not picks:
        return ""
    try:
        facts = _erp_order_facts([p[2] for p in picks])
    except Exception as exc:  # noqa: BLE001 - 這段掛掉不該讓昨日產量也寄不出去
        return f"### 🔍 ERP 交期核對\n⚠️ 查 ERP 鏡像失敗（{type(exc).__name__}: {str(exc)[:120]}），本段略過。"

    lines = [f"### 🔍 ERP 交期核對（最急的 {len(picks)} 張：日報上希望出貨日最早、且仍有未完）",
             "| 指令 | 客戶 | 型體 | 日報希望出貨日 | 日報未完 | ERP 客戶交期 | ERP 狀態 |",
             "| :--- | :--- | :--- | :---: | ---: | :---: | :--- |"]
    notes = []
    for want, cust, wo, model, rem in picks:
        base = str(wo or "").strip().upper().split("-")[0]
        f = facts.get(base)
        erp_date = f["交期"] if f else "查無此單"
        erp_state = f["狀態"] if f else "—"
        lines.append(f"| {wo} | {cust} | {str(model)[:18]} | {want} | {int(rem):,} 雙 "
                     f"| {erp_date} | {erp_state} |")
        if f and any(s in f["狀態"] for s in _ERP_CLOSED_STATES):
            notes.append(f"⚠️ {wo}：ERP 狀態已是「{f['狀態']}」，日報卻還有 {int(rem):,} 雙未完 —— 兩邊對不上。")
        elif f and f["交期"] and want and f["交期"] != want:
            notes.append(f"⚠️ {wo}：日報希望出貨日 {want}、ERP 客戶交期 {f['交期']} —— 兩邊對不上。")
    if not facts:
        notes.append("（ERP 本地鏡像讀不到，交期/狀態欄無資料。）")
    return "\n".join(lines + ([""] + notes if notes else []))


def _build_daily_production_report(now, xlsx_bytes, sheet_name, modified):
    """組整封每日生產回報（純函式：時間與檔案都由呼叫端給，測試不必碰 Drive）。"""
    yesterday = now.date() - _dt.timedelta(days=1)
    data = _daily_report_rows(xlsx_bytes, yesterday.day)
    latest_day = data["latest_day"]

    # 進度表還沒涵蓋昨天：早上先安靜等下一輪（生管通常上午才更新），過了 11 點
    # 且昨天是正常工作日才催收。週日工廠不生產，不催。
    if latest_day < yesterday.day:
        if now.hour < _DAILY_CHASE_HOUR or yesterday.weekday() == 6:
            return "(無新發現)"
        return (f"⚠️ 生管的生產日報進度表仍停在 {yesterday.month}/{latest_day}，"
                f"昨天（{yesterday.month}/{yesterday.day}）的還沒更新，建議催收。\n"
                f"（來源：{sheet_name}，修改 {modified}）")

    head = [f"📅 **{yesterday.isoformat()}（{'一二三四五六日'[yesterday.weekday()]}）每日生產數量回報**",
            f"資料來源：{sheet_name}（修改 {modified}）；"
            f"累計／未完為截至 {latest_day} 號。"]
    if data["warnings"]:
        head.append(f"⚠️ 解析提醒：{'；'.join(data['warnings'])}（數字建議人工複核）")

    # 列誰：昨天有產出的、或本季還有未完的。已出完貨的不列（大王規則），只報家數。
    day_totals, current = data["day_totals"], data["current"]
    custs = set(current) | {c for c, agg in day_totals.items() if any(agg.values())}
    # 排序：昨天做得多的在前，都沒做的按未完多寡 —— 讀的人先看到昨天真的在動的那幾家。
    def _key(c):
        return (-sum(day_totals.get(c, {}).values()), -current.get(c, {}).get("pack_rem", 0.0), c)

    def _cell(v):
        return f"{int(round(v)):,} 雙" if v else "－"

    table = [f"### 📊 昨天（{yesterday.month}/{yesterday.day}）各客戶生產動態與進度彙總",
             "| 客戶 | 灌注（成型） | 包裝 | 針車 | 本季已包裝累計 | 本季未完（剩餘） |",
             "| :--- | ---: | ---: | ---: | ---: | ---: |"]
    for c in sorted(custs, key=_key):
        agg = day_totals.get(c, {})
        st = current.get(c)
        cum = f"{int(st['pack_cum']):,} 雙" if st else "－"
        rem = f"{int(st['pack_rem']):,} 雙" if st else "－"
        table.append(f"| **{c}** | {_cell(agg.get('灌注（成型）', 0))} | {_cell(agg.get('包裝', 0))} "
                     f"| {_cell(agg.get('針車', 0))} | {cum} | {rem} |")
    if not custs:
        table.append("| （昨天各客戶都沒有產出，本季也沒有未完） | － | － | － | － | － |")
    if data["done_count"]:
        table.append("")
        table.append(f"（另有 {data['done_count']} 家本季已出完貨、共 {data['done_pairs']:,} 雙，"
                     f"照大王規則不列入回報。）")

    parts = ["\n".join(head), "\n".join(table)]
    erp = _erp_delivery_check_section(data["focus"])
    if erp:
        parts.append(erp)
    return "\n\n".join(parts)


def _short_error(exc):
    """例外 → 一行人看得懂的摘要。

    googleapiclient 的 HttpError 字串是整段 API URL＋Details 的 JSON（400 多字），
    原樣貼進通知只會洗版且看不出重點；只留狀態碼與 returned "..." 那句人話。
    """
    msg = str(exc).replace("\n", " ")
    said = re.search(r'returned "([^"]+)"', msg)
    code = re.search(r"HttpError (\d{3})", msg)
    if said:
        return f"{type(exc).__name__} {code.group(1) + ' ' if code else ''}{said.group(1)}"
    return f"{type(exc).__name__}: {msg[:120]}"


def _daily_production_report(now):
    """每日生產回報本體（now 由呼叫端給，測試不必動系統時鐘）。"""
    yesterday = now.date() - _dt.timedelta(days=1)
    try:
        xlsx_bytes, name, modified = _load_latest_sheet()
        # 跨月那幾天：最新的是新月份那份，昨天卻還在上個月 —— 直接指名昨天那個月的檔，
        # 否則會把「新月份的表只到 1 號」誤判成生管沒更新而發出假催收。
        if (_month_of_progress_name(name) or yesterday.month) != yesterday.month:
            floor = (now - _dt.timedelta(days=75)).strftime("%Y-%m-%dT%H:%M:%SZ")
            fid, nm = _find_month_progress_file(yesterday.month, min_modified=floor)
            if fid:
                xlsx_bytes, meta = _download_xlsx_bytes(fid)
                name = nm or str(meta.get("name", ""))
                modified = str(meta.get("modifiedTime", ""))[:10]
    except Exception as exc:
        # 讀不到進度表**不寄給收報表的四個人** —— 那是基礎設施問題，不是生產狀況。
        # 2026-08-10 踩到的：Drive 偶發 403 User rate limit exceeded，09:13 就把整段
        # HttpError（含 API URL）寄出去，而下一輪（30 分鐘後）其實就恢復了。
        # 不到催收時間 → 安靜等下一輪；到了還讀不到 → raise，讓 dispatcher 記
        # last_error，由排程健康告警（scheduler.task_health + dashboard_alerts
        # ._check_scheduled_tasks，#370）通知維運那條線。
        if now.hour < _DAILY_CHASE_HOUR:
            return "(無新發現)"
        raise RuntimeError(
            f"讀不到生管的生產日報進度表（Drive）：{_short_error(exc)}") from exc
    # 解析/組表失敗是程式的問題（版型變了、欄位對不上），同樣走告警那條線讓人來修，
    # 不要把例外字串寄給收報表的人 —— 他們對此無事可做。
    return _build_daily_production_report(now, xlsx_bytes, name, modified)


def daily_production_report() -> str:
    """每日生產數量回報（確定性、免 Gemini）：昨天各客戶各站雙數＋本季累計/未完＋ERP 交期核對。

    daily_production_8am 排程走 ``deterministic_tool`` 直接把這顆工具的回傳當信寄出、
    完全不經 LLM —— 客戶清單與數字每天固定，不會這輪列 5 家、下輪列 3 家（改成確定性
    之前實測同一份資料一個上午能產出四種表格）。互動問答也可直接叫。免 +確認。

    行為（沿用前身 prompt 的三選一）：
      A 進度表已涵蓋昨天 → 回報表格。
      B 還沒涵蓋、但現在不到 11:00 或昨天是週日 → 回「(無新發現)」，dispatcher 安靜。
      C 過了 11:00、昨天是工作日卻仍未涵蓋 → 回一則催收提醒。
    讀不到 Drive／解析失敗 → 不寄信，走排程健康告警（見 _daily_production_report）。

    Returns:
        markdown 報表（信件端會轉成 HTML 表格），或「(無新發現)」。
    """
    return _daily_production_report(_dt.datetime.now())


# 客戶 PO 編號樣式（業務開的 11 碼，如 20262602028）。
_CUST_PO_RE = re.compile(r"\b(20\d{9})\b")
_ORDER_ROOT_FOLDER = "4-客戶-訂單"


def _drive_children(service, folder_id, n=200):
    return service.files().list(
        q=f"'{folder_id}' in parents and trashed=false", pageSize=n,
        fields="files(id,name,mimeType)", orderBy="name",
        includeItemsFromAllDrives=True, supportsAllDrives=True,
        corpora="allDrives").execute().get("files", [])


def _find_folder_by_name(service, name, n=20):
    safe = name.replace("\\", "\\\\").replace("'", "\\'")
    return service.files().list(
        q=f"name = '{safe}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false",
        pageSize=n, fields="files(id,name)",
        includeItemsFromAllDrives=True, supportsAllDrives=True,
        corpora="allDrives").execute().get("files", [])


def _po_qty_from_analysis(service, file_id):
    """下載 PO 分析 xlsx 並解析 → {PO#: 雙數}（解不出回 {}）。"""
    import io as _io

    from googleapiclient.http import MediaIoBaseDownload
    buf = _io.BytesIO()
    dl = MediaIoBaseDownload(buf, service.files().get_media(fileId=file_id, supportsAllDrives=True),
                             chunksize=2 * 1024 * 1024)
    done = False
    while not done:
        _s, done = dl.next_chunk()
    return _po_qty_from_xlsx_bytes(buf.getvalue())


def _po_qty_from_xlsx_bytes(xlsx_bytes):
    """純解析 PO 分析 xlsx bytes → {PO#: 雙數}。找含 PO# 欄(值多為 20\\d{9})＋數量欄的
    工作表，group by PO# 加總。解不出回 {}（呼叫端會標明該 PO 量未解析，不靜默算錯）。"""
    import io as _io

    import pandas as _pd
    xls = _pd.ExcelFile(_io.BytesIO(xlsx_bytes))
    out = {}
    for sn in xls.sheet_names:
        df = xls.parse(sn, header=None)
        if df.empty:
            continue
        po_col = next((j for j in range(df.shape[1])
                       if df[j].astype(str).str.match(r"20\d{9}").sum() >= 3), None)
        if po_col is None:
            continue
        # 數量欄＝per-line 落在 1..5000、總和最大的非 PO 欄
        qcol, best = None, -1.0
        for j in range(df.shape[1]):
            if j == po_col:
                continue
            s = _pd.to_numeric(df[j], errors="coerce")
            tot = float(s[s.between(1, 5000)].sum())
            if tot > best:
                best, qcol = tot, j
        if qcol is None:
            continue
        for _i, r in df.iterrows():
            po = str(r[po_col]).strip()
            if not _CUST_PO_RE.fullmatch(po):
                continue
            q = _pd.to_numeric(r[qcol], errors="coerce")
            if not _pd.isna(q):
                out[po] = out.get(po, 0.0) + float(q)
    return out


def read_customer_order_pos(customer: str, fiscal_year: str = "") -> str:
    """查某客戶某年度的**訂單真實量**（客戶 PO）—— 來源=業務 Drive『4-客戶-訂單』夾。

    這是查「LURCHI/DECATHLON 這季**訂單**下多少雙」「客戶有哪些 PO」的**權威**工具。
    ⚠️ 生管日報（read_production_progress_sheet / chart_*）只給**生產進度**、不是訂單
    總額；訂單真實量一律以業務訂單夾為準（4-客戶-訂單/FY[年]/order [客戶]，內含客戶
    PO 確認）。可再比對業務 twsales@/owner@ 確認信。免 +確認。

    Args:
        customer: 客戶名（如 'LURCHI'、'DECATHLON'、'RICHTER'、'JALAS'）。
        fiscal_year: 會計年度（'FY26'/'26'/'2026'）；留空=最新的 FY 夾。

    Returns:
        PO 清單＋已解析數量（從 PO 分析表 Menge 加總）＋合計；標明哪些 PO 量未解析
        （多在 PDF / 後續 add-on）。引用時標明來源夾。
    """
    from agent_core.google_auth import get_service
    cust = (customer or "").strip()
    if not cust:
        return "請指定客戶名（如 LURCHI / DECATHLON / RICHTER / JALAS）。"
    fy = re.sub(r"^(FY|20)", "", (fiscal_year or "").strip().upper())  # FY26/2026 → 26
    try:
        service = get_service("drive", "v3")
        roots = _find_folder_by_name(service, _ORDER_ROOT_FOLDER)
        if not roots:
            return f"找不到業務訂單夾『{_ORDER_ROOT_FOLDER}』。請確認 Drive，或用 search_drive_files 找。"
        fy_folders = {}
        for root in roots:
            for ch in _drive_children(service, root["id"]):
                m = re.match(r"FY(\d{2})", ch["name"].upper())
                if ch["mimeType"].endswith("folder") and m:
                    fy_folders[m.group(1)] = ch
        if not fy_folders:
            return f"『{_ORDER_ROOT_FOLDER}』下找不到 FY 年度子夾。"
        def _order_folder_in(node):
            return next((c for c in _drive_children(service, node["id"])
                         if c["mimeType"].endswith("folder") and cust.lower() in c["name"].lower()), None)

        # 指定年度就用它；沒指定就從新到舊挑「第一個真的有該客戶訂單夾」的 FY
        # （不是 max——最新 FY 常是下一季、只有零星客戶；如 FY27 只有 Richter）。
        pick, order_node = None, None
        if fy and fy in fy_folders:
            pick, order_node = fy, _order_folder_in(fy_folders[fy])
        else:
            for yy in sorted(fy_folders, reverse=True):
                node = _order_folder_in(fy_folders[yy])
                if node:
                    pick, order_node = yy, node
                    break
            if pick is None:
                pick = max(fy_folders)  # 各年都沒這客戶 → 用最新 FY 報缺
        fy_node = fy_folders[pick]
        if not order_node:
            avail = [c["name"] for c in _drive_children(service, fy_node["id"]) if c["mimeType"].endswith("folder")]
            return f"FY{pick} 下找不到客戶『{cust}』訂單夾。現有：{avail}"
        files = _drive_children(service, order_node["id"])
        pos_in_names = {m.group(1) for f in files for m in _CUST_PO_RE.finditer(f["name"])}
        analysis = next((f for f in files
                         if f["name"].lower().endswith((".xls", ".xlsx"))
                         and ("分析" in f["name"] or "PO" in f["name"])), None)
        po_qty = {}
        if analysis:
            try:
                po_qty = _po_qty_from_analysis(service, analysis["id"])
            except Exception:  # noqa: BLE001
                po_qty = {}
        lines = [f"📑 {cust} FY{pick} 訂單（來源：業務 {_ORDER_ROOT_FOLDER}/{fy_node['name']}/{order_node['name']}）"]
        if po_qty:
            lines.append(f"已解析數量的 PO（{len(po_qty)} 張，合計 {int(sum(po_qty.values())):,} 雙）"
                         + (f"｜分析表：{analysis['name']}" if analysis else "") + "：")
            for po in sorted(po_qty, key=lambda p: -po_qty[p]):
                lines.append(f"  PO {po}: {int(po_qty[po]):,} 雙")
        missing = sorted(pos_in_names - set(po_qty))
        if missing:
            lines.append(f"另有 {len(missing)} 張 PO 出現在檔名但量未解析（多在 PDF / add-on）：{', '.join(missing)}")
        if not (po_qty or pos_in_names):
            lines.append(f"此夾共 {len(files)} 檔，未抓到 PO 編號；建議人工開夾或比對業務確認信。")
        lines.append("⚠️ 訂單常分多次 shipment(Main1/2/3)+add-on、部分只在 PDF；以上為已能解析的部分，"
                     "完整量請開夾或比對 twsales@/owner@ 確認信。生管日報是生產進度、非訂單。")
        return "\n".join(lines)
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 讀業務訂單夾失敗：{type(exc).__name__}: {exc}"


def _compute_completion(xlsx_bytes, customer=""):
    """從最新一天解析各客戶『本季在表量 / 已完包裝』→ list[dict]（依完工率升冪，落後的在前）。

    完工率 = 已完包裝(pack_cum) ÷ 本季在表量(pairs)。本季=排除往年已出貨舊單（同
    _summarize 的年份分流）。沿用 _detect_layout 的文字錨點，版型位移時自然解不到欄
    → 回 []（呼叫端會給可讀提示，不會靜默算錯）。生管在表量非業務訂單總額。
    """
    cust_filter = (customer or "").strip().lower()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return []
    day_sheets = _iter_day_sheets(wb)
    if not day_sheets:
        return []
    _latest_day, latest_name = day_sheets[-1]
    rows = list(wb[latest_name].iter_rows(values_only=True))
    data_start, layout = _detect_layout(rows)
    if data_start is None:
        return []
    colmap, _warns = layout
    if colmap.get("pairs") is None or colmap.get("pack_cum") is None:
        return []

    def cell(r, key):
        j = colmap.get(key)
        return r[j] if (j is not None and j < len(r)) else None

    data_rows = rows[data_start:]

    def _cell_shipped(r):
        a = cell(r, "actual_ship")
        return bool(a) and str(a).strip().upper() not in ("", "NONE")

    _year_pairs = {}
    for r in data_rows:
        yy = _year_of(cell(r, "want_ship"), cell(r, "work_order"))
        if yy is not None:
            _year_pairs[yy] = _year_pairs.get(yy, 0.0) + _to_num(cell(r, "pairs"))
    current_year = max(_year_pairs, key=_year_pairs.get) if _year_pairs else None

    agg = {}
    for r in data_rows:
        cust = cell(r, "customer")
        if not cust or not str(cust).strip():
            continue
        cust = str(cust).strip()
        if cust_filter and cust_filter not in cust.lower():
            continue
        y = _year_of(cell(r, "want_ship"), cell(r, "work_order"))
        if current_year is not None and y is not None and y < current_year and _cell_shipped(r):
            continue  # 往年已出貨舊單不計入本季達成率
        a = agg.setdefault(cust, {"ordered": 0.0, "done": 0.0})
        a["ordered"] += _to_num(cell(r, "pairs"))
        a["done"] += _to_num(cell(r, "pack_cum"))

    out = []
    for cust, a in agg.items():
        ordered = a["ordered"]
        if ordered <= 0:
            continue
        done = a["done"]
        pct = round(done / ordered * 100, 1)
        out.append({"customer": cust, "ordered": ordered, "done": done,
                    "remaining": max(0.0, ordered - done), "pct": pct})
    out.sort(key=lambda d: d["pct"])  # 落後的排最上面（畫圖反轉 y 軸後最顯眼）
    return out


def _completion_color(pct):
    """達成率分色：<50% 紅、50–80% 橘、≥80% 綠 —— 一眼看出誰落後。"""
    if pct >= 80:
        return "#2ECC71"
    if pct >= 50:
        return "#E67E22"
    return "#E74C3C"


def chart_production_completion(file_id: str = "", customer: str = "",
                                deliver: bool = True, chat_id: str = "") -> str:
    """畫『各客戶本季完工進度』橫條圖（已包裝 ÷ 本季在表量 %）並自動傳 Telegram。

    大王說「畫各客戶完工進度 / 達成率」「哪個客戶落後」「效率如何（以達成率看）」時用這個 ——
    直接解析 Drive 最新『X月份生產日報進度表』算出每客戶本季完工率：紅(<50%)/橘(50–80%)/
    綠(≥80%) 分級、條尾標「完工率 (已包裝/本季在表)＋欠數」、畫 100% 目標線。數字直接從生管表
    算，不靠模型口算，比叫小紅用 run_python_code 拼資料可靠；免 +確認。
    ⚠️ 這是**生產進度**（生管在表量、已排除往年已出貨舊單），不是業務訂單總額；
    要客戶訂單真實量請用 read_customer_order_pos（讀業務訂單夾）。

    Args:
        file_id: Drive 檔案 ID；留空自動抓最新一份『生產日報進度表』。
        customer: 只看某客戶關鍵字（如 'DECA'）；留空=全部。
        deliver: 是否自動傳到 Telegram（預設 True）。
        chat_id: 指定 Telegram chat；留空用大王預設。

    Returns:
        generate_chart 的結果（summary + PNG artifact）；找不到資料時回可讀說明字串。
    """
    import json
    from agent_core.chart_export import generate_chart
    try:
        xlsx_bytes, name, modified = _load_latest_sheet(file_id)
    except RuntimeError as e:
        return str(e)
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 讀取生產日報進度表失敗：{type(exc).__name__}: {exc}"

    rows = _compute_completion(xlsx_bytes, customer=customer)
    if not rows:
        return ("⚠️ 這份檔算不出各客戶完工達成率（可能版型不同、或缺『雙數/包裝累計』欄）。"
                "可先用 read_production_progress_sheet 看原始數字。")

    categories = [r["customer"] for r in rows]
    values = [r["pct"] for r in rows]
    colors = [_completion_color(r["pct"]) for r in rows]
    annotations = []
    for r in rows:
        done_i, ord_i, rem_i = (int(round(r["done"])), int(round(r["ordered"])),
                                int(round(r["remaining"])))
        tail = "　已結案" if r["pct"] >= 100 else (f"　欠 {rem_i:,}" if rem_i > 0 else "")
        annotations.append(f"{r['pct']:g}% ({done_i:,}/{ord_i:,}){tail}")

    subtitle = (f"資料：{name}（修改 {modified}）" if name else "") + "｜已包裝 ÷ 本季在表量（生管在表，非業務訂單）"
    spec = {
        "type": "barh",
        "title": "各客戶本季完工進度",
        "subtitle": subtitle,
        "xlabel": "完工率 (%)",
        "value_suffix": "%",
        "reference_line": 100,
        "reference_label": "100% 目標",
        "series": [{
            "name": "完工率", "values": values,
            "colors": colors, "annotations": annotations,
        }],
        "categories": categories,
    }
    fname = f"各客戶本季完工進度_{modified}" if modified else "各客戶本季完工進度"
    return generate_chart(json.dumps(spec, ensure_ascii=False),
                          filename=fname, deliver=deliver, chat_id=chat_id)


def _load_latest_sheet(file_id=""):
    """共用取檔：用 file_id 或自動抓最新『生產日報進度表』→ (xlsx_bytes, name, modified)。
    找不到檔丟 RuntimeError（訊息可直接回給使用者）。

    file_id 留空時走 production_schedule 的 5 分鐘 bytes 快取 —— 同一輪對話裡
    chart_* 與排程工具共用同一份下載，不重複抓 Drive。
    """
    fid = (file_id or "").strip()
    if not fid:
        from agent_core import production_schedule as _ps  # 延遲 import 免循環依賴
        return _ps._get_report()
    xlsx_bytes, meta = _download_xlsx_bytes(fid)
    return xlsx_bytes, str(meta.get("name", "")), str(meta.get("modifiedTime", ""))[:10]


def _resolve_station(station):
    """站別名稱(中/英/廠內口語) → 英文錨點；認不得回 None。

    「射出」是廠內對 PU 灌注成型那站的口語，生管日報的中文欄名寫「灌注」、英文錨點
    是 Injection —— 同一站兩個名字，兩邊都要認得（生管主管經理的需求單寫「射出」）。
    """
    zh_to_en = {zh: en for en, zh in _STATIONS}
    s = str(station or "").strip()
    if not s:
        return None
    if s in _STATION_ALIASES:
        return _STATION_ALIASES[s]
    if s in zh_to_en:
        return zh_to_en[s]
    for en, _zh in _STATIONS:
        if en.lower() == s.lower():
            return en
    return None


def _compute_daily_output(xlsx_bytes, station="包裝"):
    """整月每天某站(預設 包裝=完工雙數)的總產出 → [(day_int, total)]，依日排序。"""
    en = _resolve_station(station) or "Packing"
    return _compute_daily_output_multi(xlsx_bytes, [en]).get(en, [])


def _compute_daily_output_multi(xlsx_bytes, stations):
    """一次開檔算**多站**每日產出 → {英文站名: [(day_int, total)]}。

    單站版是這支的特例。分開呼叫 _compute_daily_output 三次會把整份 workbook
    (7MB、31 個分頁) 重解三遍(實測 6.4 秒/次)；三站折線圖一次解完即可。
    某站沒有日計欄 → 該站的 list 為空(不是 0，區分「這站沒欄位」與「這站沒產出」)。
    """
    wanted = [s for s in stations if s]
    out = {en: [] for en in wanted}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return out
    for day_int, sheet_name in _iter_day_sheets(wb):
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        data_start, layout = _detect_layout(rows)
        if data_start is None:
            continue
        colmap, _w = layout
        cust_j = colmap.get("customer")
        cols = {en: colmap.get(f"{en}_day") for en in wanted}
        if all(j is None for j in cols.values()):
            continue
        totals = {en: 0.0 for en in wanted if cols[en] is not None}
        for r in rows[data_start:]:
            cust = r[cust_j] if (cust_j is not None and cust_j < len(r)) else None
            if not cust or not str(cust).strip():
                continue
            for en in totals:
                j = cols[en]
                if j < len(r):
                    totals[en] += _to_num(r[j])
        for en, total in totals.items():
            out[en].append((day_int, total))
    return out


_MONTH_IN_NAME_RE = re.compile(r"^\s*(\d{1,2})\s*月份")


def _month_of_progress_name(name):
    """檔名 '08月份生產日報進度表08-04.xlsx' → 8；解不出回 None。

    折線圖的 X 軸要跨月（月初只有 3、4 天，光看本月那幾點看不出趨勢），標籤得是
    「7/29」而不是「29」，所以月份要從檔名拿 —— 分頁名只有日、沒有月。
    """
    m = _MONTH_IN_NAME_RE.match(str(name or ""))
    return int(m.group(1)) if m else None


def _find_month_progress_file(month, *, min_modified="", keyword=_DEFAULT_PROGRESS_KEYWORD):
    """指定月份那份進度表 → (file_id, name)；找不到回 (None, None)。

    同一個月份每年都有一份（'07月份…' 2025/2026 都在），所以除了比對檔名月份，
    還要求 modifiedTime ≥ min_modified（呼叫端傳「兩個月前」）——否則跨月那幾天
    會抓到去年同月的舊檔，畫出一條看似正常、其實差一年的線。
    """
    from agent_core.google_auth import get_service
    service = get_service("drive", "v3")
    safe = str(keyword).replace("\\", "\\\\").replace("'", "\\'")
    res = service.files().list(
        q=f"name contains '{safe}' and trashed = false",
        pageSize=50,
        fields="files(id, name, modifiedTime, mimeType)",
        includeItemsFromAllDrives=True,
        supportsAllDrives=True,
        corpora="allDrives",
    ).execute()
    items = []
    for f in res.get("files", []):
        name = str(f.get("name", ""))
        if not (name.lower().endswith((".xlsx", ".xls", ".xlsm"))
                or "spreadsheet" in f.get("mimeType", "")):
            continue
        if _month_of_progress_name(name) != int(month):
            continue
        if min_modified and str(f.get("modifiedTime", "")) < min_modified:
            continue
        items.append(f)
    if not items:
        return None, None
    items.sort(key=lambda f: f.get("modifiedTime", ""), reverse=True)
    return items[0]["id"], items[0]["name"]


def station_capacity_series(days=21, file_id="", today=None):
    """近 N 天 針車/射出/包裝 每日產量 → dict(points, sources, warnings)。

    points: [(label, {zh站名: 產量}), ...] 依日期升冪，label 形如 '7/29'。
    本月分頁不足 N 天時，往前補**上個月**那份進度表的尾巴（月初只有兩三個點的
    折線圖看不出任何趨勢）。找不到上個月的檔就只給本月、並在 warnings 說明。
    file_id 指定時只讀那一份（不跨月補），呼叫端要看哪份就是哪份。
    """
    today = today or _dt.date.today()
    warnings = []
    xlsx_bytes, name, modified = _load_latest_sheet(file_id)
    sources = [(name, modified)]
    cur_month = _month_of_progress_name(name) or today.month
    series = _compute_daily_output_multi(xlsx_bytes, _CAPACITY_STATIONS)
    days_present = sorted({d for pts in series.values() for d, _v in pts})
    if not days_present:
        # 空的有兩種意思，對使用者是兩件事：檔根本打不開（版型換了/檔壞了/根本不是
        # xlsx，要找人修檔），還是打得開但還沒有任何一天的數字（月初，等就好）。
        # _compute_daily_output_multi 兩種都吞成空 list，這裡補上區分。
        # 只有「空」這條路才多開一次 workbook：read_only 的 load_workbook 只讀壓縮檔
        # 索引不解列，而且真的打不開時是立刻拋錯，沒有 7MB 重解的成本。
        try:
            _open_workbook(xlsx_bytes)
        except Exception as exc:  # noqa: BLE001
            warnings.append(f"「{name}」打不開（{type(exc).__name__}），本月的數字全部缺")
    points = [(f"{cur_month}/{d}", {_CAPACITY_LABELS[en]: dict(series.get(en, [])).get(d, 0.0)
                                    for en in _CAPACITY_STATIONS})
              for d in days_present]

    need = max(int(days or 0), 1) - len(points)
    if need > 0 and not str(file_id or "").strip():
        prev_month = 12 if cur_month == 1 else cur_month - 1
        prev_bytes, pname, pmod, warn = _prev_month_sheet(prev_month, today)
        if warn:
            warnings.append(warn)
        if prev_bytes is not None:
            prev_series = _compute_daily_output_multi(prev_bytes, _CAPACITY_STATIONS)
            prev_days = sorted({d for pts in prev_series.values() for d, _v in pts})
            tail = prev_days[-need:] if need < len(prev_days) else prev_days
            points = [(f"{prev_month}/{d}",
                       {_CAPACITY_LABELS[en]: dict(prev_series.get(en, [])).get(d, 0.0)
                        for en in _CAPACITY_STATIONS})
                      for d in tail] + points
            sources.insert(0, (pname, pmod))

    return {"points": points[-max(int(days or 0), 1):], "sources": sources,
            "warnings": warnings}


# 上個月那份進度表的 bytes 快取（7MB、下載 ~11 秒）。本月的快取在
# production_schedule._CACHE；這裡補上上個月的，讓同一輪簡報裡「折線圖」與
# 「各型體產出速度」共用同一次下載，而不是各抓一次。
_PREV_CACHE: dict = {"month": None, "bytes": None, "name": "", "modified": "", "ts": 0.0}
_PREV_CACHE_TTL_S = 300.0


def _prev_month_sheet(prev_month, today):
    """上個月那份進度表 → (bytes, name, modified, warning)；取不到時 bytes 為 None。

    取不到一律回 warning 字串而不是丟例外 —— 補資料失敗不該讓本月的圖/速率整個沒有。
    """
    now = _time.time()
    if (_PREV_CACHE["bytes"] is not None and _PREV_CACHE["month"] == prev_month
            and now - _PREV_CACHE["ts"] < _PREV_CACHE_TTL_S):
        return _PREV_CACHE["bytes"], _PREV_CACHE["name"], _PREV_CACHE["modified"], ""
    # 「上個月的檔」至少要在 ~2 個月內改過，否則就是去年同月的舊檔。
    floor = (today - _dt.timedelta(days=75)).isoformat()
    try:
        pid, pname = _find_month_progress_file(prev_month, min_modified=floor)
    except Exception as exc:  # noqa: BLE001
        return None, "", "", f"找上個月進度表時出錯（{type(exc).__name__}），只用本月資料"
    if not pid:
        return None, "", "", f"Drive 上找不到 {prev_month} 月份的進度表，只用本月資料"
    try:
        prev_bytes, meta = _download_xlsx_bytes(pid)
    except Exception as exc:  # noqa: BLE001
        return None, "", "", f"讀上個月進度表失敗（{type(exc).__name__}），只用本月資料"
    modified = str(meta.get("modifiedTime", ""))[:10]
    _PREV_CACHE.update(month=prev_month, bytes=prev_bytes, name=pname,
                       modified=modified, ts=now)
    return prev_bytes, pname, modified, ""


def _pack_by_model(xlsx_bytes, day_filter_set=None):
    """該份進度表逐型體的包裝產出 → ({型體: 雙數}, 有資料的日數)。

    day_filter_set 給定時只算那幾天（跨月補窗用）。用「包裝」＝完工雙數，與系統其他
    地方對「產出」的定義一致（針車/射出是在製，拿來當產能會高估真正做得出來的量）。
    """
    out: dict = {}
    days = set()
    try:
        wb = _open_workbook(xlsx_bytes)
    except Exception:  # noqa: BLE001
        return out, 0
    for row in _iter_rows_for_workbook(wb):
        if day_filter_set is not None and row["prod_day"] not in day_filter_set:
            continue
        days.add(row["prod_day"])
        model = (row.get("model") or "").strip()
        if not model:
            continue
        out[model] = out.get(model, 0.0) + row.get("packing_day", 0.0)
    return out, len(days)


def style_output_rate(days=28, file_id="", today=None):
    """近 N 天各型體的實際包裝產出 → dict(pairs={型體: 雙數}, days=實際涵蓋日數, ...)。

    型體用生管日報的寫法（帶顏色尾碼，如 `DJS336189-02BEIGE`）；要對回 ERP 鞋款走
    weekly_rates_for_styles()。窗不足時同樣往前接上個月（與折線圖共用快取）。
    """
    today = today or _dt.date.today()
    warnings = []
    xlsx_bytes, name, _modified = _load_latest_sheet(file_id)
    cur_month = _month_of_progress_name(name) or today.month
    pairs, n_days = _pack_by_model(xlsx_bytes)

    need = max(int(days or 0), 1) - n_days
    if need > 0 and not str(file_id or "").strip():
        prev_month = 12 if cur_month == 1 else cur_month - 1
        prev_bytes, _pname, _pmod, warn = _prev_month_sheet(prev_month, today)
        if warn:
            warnings.append(warn)
        if prev_bytes is not None:
            try:
                wb = _open_workbook(prev_bytes)
                prev_days = [d for d, _n in _iter_day_sheets(wb)]
            except Exception:  # noqa: BLE001
                prev_days = []
            tail = set(prev_days[-need:]) if need < len(prev_days) else set(prev_days)
            prev_pairs, prev_n = _pack_by_model(prev_bytes, day_filter_set=tail)
            for model, qty in prev_pairs.items():
                pairs[model] = pairs.get(model, 0.0) + qty
            n_days += prev_n
    return {"pairs": pairs, "days": n_days, "warnings": warnings}


def weekly_rates_for_styles(rate, erp_styles):
    """ERP 鞋款 → 近期週產雙數（依 style_output_rate 的結果換算）。

    生管日報的型體帶顏色尾碼（`DJS336189-02BEIGE`），ERP 鞋款是 `DJS336189-02`，所以
    用前綴歸屬；同時符合多個 ERP 鞋款時取**最長**的那個 —— 否則舊款 `DJS336195` 會
    把 `DJS336195-01` 的產出整碗端走。
    無資料（近期沒生產、或沒讀到表）該款就不在回傳裡，呼叫端要顯示「無法換算」。
    """
    styles = sorted({str(s).strip() for s in erp_styles if str(s or "").strip()},
                    key=len, reverse=True)
    n_days = int(rate.get("days") or 0)
    if not styles or n_days <= 0:
        return {}
    totals: dict = {}
    for model, qty in (rate.get("pairs") or {}).items():
        owner = next((s for s in styles if model.startswith(s)), None)
        if owner is None:
            continue
        totals[owner] = totals.get(owner, 0.0) + qty
    return {s: qty / n_days * 7.0 for s, qty in totals.items() if qty > 0}


def _pad_cell(s, width):
    """以顯示寬度（CJK 算 2）右補空白，讓 code block 內的表格對齊。"""
    s = str(s)
    pad = width - sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in s)
    return s + " " * max(pad, 0)


def _capacity_table(points):
    """[(label, {站: 值})] → 對齊的文字表（含平均列）。"""
    stations = [_CAPACITY_LABELS[en] for en in _CAPACITY_STATIONS]
    header = ["日期"] + stations
    rows = [[label] + [f"{int(round(vals.get(st, 0.0))):,}" for st in stations]
            for label, vals in points]
    if points:
        avg = ["平均"] + [f"{int(round(sum(v.get(st, 0.0) for _l, v in points) / len(points))):,}"
                          for st in stations]
        rows.append(avg)
    widths = [max(sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in str(c))
                  for c in [header[i]] + [r[i] for r in rows])
              for i in range(len(header))]
    out = ["  ".join(_pad_cell(h, widths[i]) for i, h in enumerate(header))]
    out.append("  ".join("-" * w for w in widths))
    for i, r in enumerate(rows):
        if points and i == len(rows) - 1:
            out.append("  ".join("=" * w for w in widths))
        out.append("  ".join(_pad_cell(c, widths[j]) for j, c in enumerate(r)))
    return "\n".join(out)


def station_capacity_report(days=21, file_id="", deliver=False, chat_id="", today=None):
    """針車/射出/包裝 每日產能：折線圖 + 文字表（含 [[MAIL_FILE:]] 附件標記）。

    給 chart_station_capacity 工具與生管每日簡報共用。deliver=True 才推 Telegram；
    排程寄信走附件標記，不需要也不應該同時推大王的 Telegram。
    """
    import json
    from agent_core.chart_export import generate_chart
    try:
        data = station_capacity_series(days=days, file_id=file_id, today=today)
    except RuntimeError as e:      # 找不到進度表（訊息可直接回給使用者）
        return str(e)
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 讀取生產日報進度表失敗：{type(exc).__name__}: {exc}"

    points = data["points"]
    if not points:
        # warnings 帶著「為什麼一個點都沒有」（檔打不開／找不到上個月的檔）——
        # 以前這條路直接早退，那些線索一個都不會被看到，只剩一句通用的猜測。
        lines = ["⚠️ 這份進度表算不出每日產量（可能版型不同、或找不到站別日計欄）。"
                 "可先用 read_production_progress_sheet 看原始數字。"]
        lines += [f"ℹ️ {w}" for w in data["warnings"]]
        return "\n".join(lines)

    labels = [lab for lab, _v in points]
    series = [{"name": _CAPACITY_LABELS[en],
               "values": [v.get(_CAPACITY_LABELS[en], 0.0) for _l, v in points]}
              for en in _CAPACITY_STATIONS]
    src = "、".join(f"{n}（修改 {m}）" for n, m in data["sources"] if n)
    spec = {
        "type": "line",
        "title": "針車／射出／包裝 每日產能",
        "subtitle": f"{labels[0]} – {labels[-1]}｜資料：{src}" if src else f"{labels[0]} – {labels[-1]}",
        "xlabel": "日期", "ylabel": "產量（雙）",
        "categories": labels,
        "series": series,
    }
    chart = generate_chart(json.dumps(spec, ensure_ascii=False),
                           filename=f"三站每日產能_{labels[-1].replace('/', '-')}",
                           deliver=deliver, chat_id=chat_id)

    lines = [f"📈 針車／射出／包裝 每日產能（{labels[0]} – {labels[-1]}，共 {len(points)} 天）"]
    recent = points[-7:]
    tail = "、".join(
        f"{_CAPACITY_LABELS[en]} {int(round(sum(v.get(_CAPACITY_LABELS[en], 0.0) for _l, v in recent) / len(recent))):,}"
        for en in _CAPACITY_STATIONS)
    lines.append(f"近 {len(recent)} 天平均日產：{tail}（雙/天）")
    lines.append("```\n" + _capacity_table(points) + "\n```")
    for w in data["warnings"]:
        lines.append(f"ℹ️ {w}")
    if src:
        lines.append(f"資料來源：{src}")
    # 「日計」= 生管日報當天登記的產出；沒開工的日子是 0，不是漏抓。
    lines.append("（數字＝生管日報各站『日計』欄全客戶加總；停工日為 0。"
                 "射出即生管日報的『灌注』站。）")

    path = ""
    artifacts = getattr(chart, "artifacts", None) or []
    if artifacts:
        path = str(artifacts[0])
    if path and os.path.isfile(path):
        # 側通道：LLM 沒把標記抄進回覆時，dispatcher 仍取得到這個檔
        # （見 agent_core/deliverables 的說明）。標記本身保留不動。
        from agent_core.deliverables import register as _register_deliverable
        _register_deliverable(path)
        lines.append(f"[[MAIL_FILE:{path}]]")
    else:
        lines.append(f"⚠️ 折線圖產檔失敗（數字同上表）：{str(getattr(chart, 'summary', chart))[:160]}")
    return "\n".join(lines)


def chart_station_capacity(days: int = 21, deliver: bool = False, chat_id: str = "") -> str:
    """畫『針車／射出／包裝 每日產能』折線圖（三條線同圖）並回文字表，看產能趨勢用這個。

    大王或生產管理部問「三站每天做多少」「針車跟包裝跟不跟得上」「產能趨勢」時用這個。
    月初本月天數不足時會自動往前接上個月的進度表，湊滿 days 天才畫。數字直接從生管
    日報『日計』欄算，不經 LLM；SAFE tier 免確認。

    Args:
        days: 看最近幾天（預設 21，會跨月往前補）。
        deliver: 是否把圖推到 Telegram（預設 False —— 排程寄信是走附件，不推 TG）。
        chat_id: 指定 Telegram chat；留空用大王預設。
    Returns:
        文字表（近 N 天三站日產 + 平均）＋ 圖檔附件標記。
    """
    return station_capacity_report(days=days, deliver=deliver, chat_id=chat_id)


def _compute_delivery_risk(xlsx_bytes, top_n=None):
    """最新一天 未出貨且有欠數 的 PO → [{want, customer, wo, remaining}]，依希望出貨日升冪。
    top_n=None 回全部（達交風險圖會自行依客戶彙總，需要完整資料才不會少算）。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return []
    day_sheets = _iter_day_sheets(wb)
    if not day_sheets:
        return []
    rows = list(wb[day_sheets[-1][1]].iter_rows(values_only=True))
    data_start, layout = _detect_layout(rows)
    if data_start is None:
        return []
    colmap, _w = layout

    def cell(r, key):
        j = colmap.get(key)
        return r[j] if (j is not None and j < len(r)) else None

    pos = []
    for r in rows[data_start:]:
        cust = cell(r, "customer")
        if not cust or not str(cust).strip():
            continue
        rem = _to_num(cell(r, "pack_rem"))
        want = cell(r, "want_ship")
        actual = cell(r, "actual_ship")
        shipped = bool(actual) and str(actual).strip().upper() not in ("", "NONE")
        if want and not shipped and rem > 0:
            pos.append({"want": str(want)[:10], "customer": str(cust).strip(),
                        "wo": str(cell(r, "work_order") or "").strip(), "remaining": rem})
    pos.sort(key=lambda p: _ship_date_sort_key(p["want"]))
    return pos if top_n is None else pos[:top_n]


def _parse_ship_date(s, today=None):
    """'2026-06-20' / '6/20' / '06-20' → datetime.date 或 None。

    無年份時取「離今天最近」的年份（仿 email_timeline._closest_year）——
    一律補今年會讓 1 月讀到的『12/28』被推成今年 12 月而不再逾期。
    today 參數只給測試注入。
    """
    s = str(s or "").strip()
    if not s:
        return None
    today = today or _dt.date.today()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%m/%d", "%m-%d", "%m.%d"):
        try:
            d = _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
        if d.year != 1900:
            return d
        cands = []
        for y in (today.year - 1, today.year, today.year + 1):
            try:
                cands.append(d.replace(year=y))
            except ValueError:  # 2/29 非閏年
                continue
        if not cands:
            return None
        return min(cands, key=lambda x: abs((x - today).days))
    return None


def _ship_date_sort_key(want):
    """達交清單排序鍵：解析成真日期再排（want 可能混 ISO 字串與 '6/20' 這類無年格式，
    直接字串排序會排錯）。解析不出的放最後、以原字串穩定排序。"""
    d = _parse_ship_date(want)
    return (d is None, d or _dt.date.max, str(want))


def chart_daily_output(file_id: str = "", station: str = "包裝",
                       deliver: bool = True, chat_id: str = "") -> str:
    """畫『每日產量趨勢』長條圖（整月每天完工雙數）並標出產能高峰那天，自動傳 Telegram。

    大王問「哪天產能最高」「這個月每天做多少」「產能趨勢」時用這個。預設看『包裝』站
    (完工雙數)；可改 station='針車'/'灌注' 等看別站。數字直接從生管表算；SAFE tier 免確認。

    Args:
        file_id: Drive 檔案 ID；留空自動抓最新『生產日報進度表』。
        station: 看哪一站日產量（包裝/針車/灌注/成型/中底），預設 包裝(完工)。
        deliver: 是否自動傳 Telegram（預設 True）。
        chat_id: 指定 chat；留空用大王預設。
    """
    import json
    from agent_core.chart_export import generate_chart
    try:
        xlsx_bytes, name, modified = _load_latest_sheet(file_id)
    except RuntimeError as e:
        return str(e)
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 讀取生產日報進度表失敗：{type(exc).__name__}: {exc}"

    daily = _compute_daily_output(xlsx_bytes, station=station)
    if not daily:
        return ("⚠️ 這份檔算不出每日產量（可能版型不同、或找不到該站日計欄）。"
                "可先用 read_production_progress_sheet 看原始數字。")

    peak_idx = max(range(len(daily)), key=lambda i: daily[i][1])
    peak_day, peak_val = daily[peak_idx]
    categories = [str(d) for d, _v in daily]
    values = [v for _d, v in daily]
    colors = ["#E74C3C" if i == peak_idx else "#3498DB" for i in range(len(daily))]
    subtitle = (f"資料：{name}（修改 {modified}）｜" if name else "") \
        + f"產能高峰：{peak_day} 號 {int(round(peak_val)):,} 雙"
    spec = {
        "type": "bar",
        "title": f"每日{station}產量趨勢",
        "subtitle": subtitle,
        "xlabel": "日期（號）", "ylabel": f"{station}產量（雙）",
        "series": [{"name": f"{station}日產量", "values": values, "colors": colors}],
        "categories": categories,
    }
    fname = f"每日{station}產量_{modified}" if modified else f"每日{station}產量"
    return generate_chart(json.dumps(spec, ensure_ascii=False),
                          filename=fname, deliver=deliver, chat_id=chat_id)


def chart_delivery_risk(file_id: str = "", deliver: bool = True, chat_id: str = "",
                        top_n: int = 12) -> str:
    """畫『達交風險』橫條圖（各客戶未出貨欠數），有逾期/將到期單的客戶標紅橘，自動傳 Telegram。

    大王問「誰快趕不上交期」「哪個客戶欠最多沒出」「達交風險」時用這個。**依客戶彙總**未出貨
    且有欠數的 PO（單量大時 PO 級會洗版又少算，故彙總到客戶）：條長=該客戶總欠數、由多到少排；
    該客戶只要有逾期單→紅、有 7 日內到期單→橘、其餘→綠；條尾標單數/逾期數/最早希望出貨日。
    數字直接從生管表算；SAFE tier 免確認。

    Args:
        file_id: Drive 檔案 ID；留空自動抓最新『生產日報進度表』。
        deliver: 是否自動傳 Telegram（預設 True）。
        chat_id: 指定 chat；留空用大王預設。
        top_n: 最多顯示幾個客戶（依欠數排序，預設 12）。
    """
    import datetime as _dt
    import json
    from agent_core.chart_export import generate_chart
    try:
        xlsx_bytes, name, modified = _load_latest_sheet(file_id)
    except RuntimeError as e:
        return str(e)
    except Exception as exc:  # noqa: BLE001
        return f"⚠️ 讀取生產日報進度表失敗：{type(exc).__name__}: {exc}"

    pos = _compute_delivery_risk(xlsx_bytes)  # 全部未出貨 PO（不截斷，才能正確彙總/計數）
    if not pos:
        return "✅ 目前沒有『未出貨且有欠數』的 PO（或這份檔算不出希望出貨日/欠數欄）。"

    today = _dt.date.today()
    agg = {}
    for p in pos:
        a = agg.setdefault(p["customer"],
                           {"remaining": 0.0, "po": 0, "overdue": 0, "soon": 0, "earliest": None})
        a["remaining"] += p["remaining"]
        a["po"] += 1
        d = _parse_ship_date(p["want"])
        if d is not None:
            if a["earliest"] is None or d < a["earliest"]:
                a["earliest"] = d
            if d < today:
                a["overdue"] += 1
            elif (d - today).days <= 7:
                a["soon"] += 1

    ranked = sorted(agg.items(), key=lambda kv: -kv[1]["remaining"])[:top_n]
    categories, values, colors, annotations = [], [], [], []
    for cust, a in ranked:
        if a["overdue"]:
            color = "#E74C3C"
        elif a["soon"]:
            color = "#E67E22"
        else:
            color = "#2ECC71"
        ed = a["earliest"].isoformat() if a["earliest"] else "?"
        extra = (f"，{a['overdue']} 逾期" if a["overdue"]
                 else (f"，{a['soon']} 近7日" if a["soon"] else ""))
        categories.append(cust)
        values.append(a["remaining"])
        colors.append(color)
        annotations.append(f"欠 {int(round(a['remaining'])):,}（{a['po']} 單{extra}）最早 {ed}")

    grand_owed = int(round(sum(p["remaining"] for p in pos)))
    subtitle = (f"資料：{name}（修改 {modified}）｜" if name else "") \
        + f"{len(agg)} 客戶共 {len(pos)} 筆未出貨、總欠 {grand_owed:,} 雙"
    spec = {
        "type": "barh",
        "title": "達交風險：各客戶未出貨欠數",
        "subtitle": subtitle,
        "xlabel": "未完欠數（雙）",
        "value_labels": False,
        "series": [{"name": "欠數", "values": values,
                    "colors": colors, "annotations": annotations}],
        "categories": categories,
    }
    fname = f"達交風險_{modified}" if modified else "達交風險"
    return generate_chart(json.dumps(spec, ensure_ascii=False),
                          filename=fname, deliver=deliver, chat_id=chat_id)
