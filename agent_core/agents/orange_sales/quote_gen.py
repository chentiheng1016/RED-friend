"""Quote generation: openpyxl-driven Excel quote + optional email send.

Split off from agent_core/quote.py (phase 65 — quote.py was 669 lines
mixing two subsystems). quote.py now holds only history extraction
(CSV pipeline + Gemini prompts); this module consumes the CSV via
`_suggest_price_from_history` and produces the Excel file a salesperson
actually sends to the customer.

Imports `_QUOTE_CSV` from quote so both sides share the same CSV path —
the history writer and the generator price-suggester never drift.
"""
import os
import re
from datetime import datetime, timedelta

from agent_core.file_ops import _clean_path, _sanitize_filename
from agent_core.gmail import send_gmail_internal
# Phase 3c：orange_sales 內部跨檔 import 走真實路徑，避免循環依賴。
from agent_core.agents.orange_sales.quote import _QUOTE_CSV


_QUOTE_OUTPUT_DIR = os.path.expanduser("~/Documents/報價單")


def _suggest_price_from_history(customer: str, shoe_model: str, currency: str = "USD") -> tuple:
    """從 quote_history 推薦單價 —— **只取與請求 currency 同幣別的歷史列**。

    混幣平均（USD 跟 TWD 一起 mean）會給出無意義的建議價。無同幣別歷史時
    明說（note 說明、price=0 讓 Excel 留白），不退回混幣平均。
    回傳 (price, note)。找不到任何可參考歷史回 (0.0, '') 或 (0.0, 無同幣別說明)。
    """
    if not os.path.exists(_QUOTE_CSV):
        return 0.0, ""
    try:
        import pandas as pd
        df = pd.read_csv(_QUOTE_CSV)
    except Exception:
        return 0.0, ""
    if df.empty:
        return 0.0, ""

    cur = (currency or "USD").strip().upper()
    cur_series = df["currency"].astype(str).str.strip().str.upper()
    df = df[cur_series == cur]
    if df.empty:
        return 0.0, f"歷史庫沒有 {cur} 幣別的報價可參考（不採混幣平均）"

    mask = df["customer"].astype(str).str.contains(customer, case=False, na=False, regex=False)
    if shoe_model:
        mask &= df["sku"].astype(str).str.contains(shoe_model, case=False, na=False, regex=False)
    sub = df[mask & (df["direction"] == "out")]
    match_level = "同客戶 + 同鞋型"

    if sub.empty:
        sub = df[df["customer"].astype(str).str.contains(customer, case=False, na=False, regex=False)
                 & (df["direction"] == "out")]
        match_level = "同客戶（跨鞋型）"

    if sub.empty and shoe_model:
        sub = df[df["sku"].astype(str).str.contains(shoe_model, case=False, na=False, regex=False)
                 & (df["direction"] == "out")]
        match_level = "同鞋型（跨客戶）"

    prices = pd.to_numeric(sub["unit_price"], errors="coerce").dropna()
    if len(prices) == 0:
        return 0.0, f"歷史庫的 {cur} 報價裡沒有相符的客戶/鞋型可參考"

    avg = float(prices.mean())
    note = (f"建議：{avg:.2f} {cur}（{match_level}，n={len(prices)}，"
            f"範圍 {prices.min():.2f}–{prices.max():.2f} {cur}）")
    return avg, note


def generate_quote(
    customer: str,
    shoe_model: str,
    quantity: int,
    unit_price_usd: float = 0.0,
    currency: str = "USD",
    materials: str = "",
    delivery_days: int = 60,
    incoterm: str = "FOB Keelung",
    payment_terms: str = "T/T 30% deposit, 70% against B/L copy",
    notes: str = "",
    output_xlsx: str = "",
    use_history_suggest: bool = True,
    email_to: str = "",
):
    """產生報價單 Excel 檔。
    - customer：客戶名（必填）
    - shoe_model：鞋型/SKU 代號（必填）
    - quantity：數量（必填）
    - unit_price_usd：單價；=0 時會從 quote_history 找類似案例推薦（可手動覆寫）
    - currency：貨幣，預設 USD
    - materials：材料描述（自由文字，如 "Upper: PU / Outsole: Rubber / Lining: Mesh"）
    - delivery_days：交期天數（預設 60）
    - incoterm：貿易條件（預設 FOB Keelung）
    - payment_terms：付款條件
    - notes：備註
    - output_xlsx：輸出路徑，空則自動 ~/Documents/報價單/Quote_<客戶>_<鞋型>_<時間>.xlsx
    - use_history_suggest：unit_price_usd=0 時是否查歷史推薦（預設 True）
    - email_to：可選，產出後直接寄（附件）給這些 email（逗號分隔）
    """
    if not customer.strip() or not shoe_model.strip() or int(quantity) <= 0:
        return "錯誤：customer / shoe_model / quantity 都要填，且 quantity > 0"

    quantity = int(quantity)
    currency = (currency or "USD").upper()

    suggested = 0.0
    history_note = ""
    if unit_price_usd <= 0 and use_history_suggest:
        suggested, history_note = _suggest_price_from_history(customer, shoe_model, currency)
    price_to_use = unit_price_usd if unit_price_usd > 0 else suggested
    total = price_to_use * quantity if price_to_use else 0.0
    # Excel 數字格式：只有 USD 才配 '$' 符號；其他幣別由表頭/條款標示幣別，避免 '$' 誤導。
    money_fmt = '"$"#,##0.00' if currency == "USD" else '#,##0.00'

    if not output_xlsx:
        os.makedirs(_QUOTE_OUTPUT_DIR, exist_ok=True)
        ts = datetime.now().strftime("%Y-%m-%d_%H%M")
        safe_c = _sanitize_filename(customer, max_len=30)
        safe_m = _sanitize_filename(shoe_model, max_len=30)
        output_xlsx = os.path.join(_QUOTE_OUTPUT_DIR, f"Quote_{safe_c}_{safe_m}_{ts}.xlsx")
    else:
        output_xlsx = _clean_path(output_xlsx)
        os.makedirs(os.path.dirname(output_xlsx) or ".", exist_ok=True)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    except ImportError:
        return "缺 openpyxl：pip install openpyxl"

    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"

    font_title = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=11, bold=True)
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_normal = Font(name="Arial", size=11)
    fill_dark = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    fill_light = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")
    thin = Side(border_style="thin", color="888888")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:F1")
    ws["A1"] = "JAIFUNG CORPORATION — QUOTATION 報價單"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_dark
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    ws["A3"] = "Owner Name · General Manager"
    ws["A3"].font = font_bold
    ws["A4"] = "No.1 Example Rd., Taipei, Taiwan"
    ws["A5"] = "Tel: +886-2-00000000"
    ws["A6"] = "Email: owner@company.example"
    for r in (4, 5, 6):
        ws.cell(row=r, column=1).font = font_normal

    ts_full = datetime.now().strftime("%Y-%m-%d")
    quote_no = datetime.now().strftime("Q-%Y%m%d-%H%M")
    meta = [("Quote No.", quote_no), ("Date", ts_full),
            ("Valid Until", (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d"))]
    for i, (k, v) in enumerate(meta):
        ws.cell(row=3 + i, column=5, value=k).font = font_bold
        ws.cell(row=3 + i, column=6, value=v).font = font_normal

    ws["A8"] = "To (Customer):"
    ws["A8"].font = font_bold
    ws.merge_cells("B8:F8")
    ws["B8"] = customer
    ws["B8"].font = font_normal

    header_row = 10
    headers = ["Item", "Model / SKU", "Description / Materials", "Qty", f"Unit Price ({currency})", "Amount"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = font_header
        c.fill = fill_dark
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = box
    ws.row_dimensions[header_row].height = 24

    r = header_row + 1
    items = [(1, shoe_model, materials or "Per customer spec", quantity,
              price_to_use if price_to_use else None,
              total if price_to_use else None)]
    for it, sku, desc, qty, price, amt in items:
        ws.cell(row=r, column=1, value=it).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=2, value=sku)
        ws.cell(row=r, column=3, value=desc)
        ws.cell(row=r, column=4, value=qty).number_format = "#,##0"
        p_cell = ws.cell(row=r, column=5, value=price if price else "(please fill)")
        if price:
            p_cell.number_format = money_fmt
        a_cell = ws.cell(row=r, column=6, value=amt if amt else "")
        if amt:
            a_cell.number_format = money_fmt
        for col in range(1, 7):
            cell = ws.cell(row=r, column=col)
            cell.font = font_normal
            cell.border = box
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal=cell.alignment.horizontal or "left")
        ws.row_dimensions[r].height = 46

    r += 1
    ws.cell(row=r, column=5, value="Subtotal").font = font_bold
    ws.cell(row=r, column=5).alignment = Alignment(horizontal="right")
    ws.cell(row=r, column=5).border = box
    sc = ws.cell(row=r, column=6, value=total if price_to_use else "")
    if price_to_use:
        sc.number_format = money_fmt
    sc.font = font_bold
    sc.border = box
    sc.fill = fill_light

    r += 1
    ws.cell(row=r, column=5, value="GRAND TOTAL").font = font_bold
    ws.cell(row=r, column=5).alignment = Alignment(horizontal="right")
    ws.cell(row=r, column=5).border = box
    ws.cell(row=r, column=5).fill = fill_dark
    ws.cell(row=r, column=5).font = font_header
    gc = ws.cell(row=r, column=6, value=total if price_to_use else "")
    if price_to_use:
        gc.number_format = money_fmt
    gc.font = font_header
    gc.border = box
    gc.fill = fill_dark

    r += 3
    ws.cell(row=r, column=1, value="Terms & Conditions 交易條件").font = font_bold
    ws.cell(row=r, column=1).fill = fill_light
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    r += 1
    terms_lines = [
        f"Incoterm: {incoterm}",
        f"Delivery: {delivery_days} days after PO confirmation & deposit received",
        f"Payment: {payment_terms}",
        f"Currency: {currency}",
        "Quotation validity: 30 days",
    ]
    for line in terms_lines:
        ws.cell(row=r, column=1, value="• " + line).font = font_normal
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1

    if notes:
        r += 1
        ws.cell(row=r, column=1, value="Notes 備註：").font = font_bold
        r += 1
        ws.cell(row=r, column=1, value=notes).font = font_normal
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=6)
        r += 4

    r += 2
    ws.cell(row=r, column=1, value="Prepared by: Owner Name").font = font_normal
    ws.cell(row=r, column=5, value="Customer Confirmation:").font = font_bold
    r += 1
    ws.cell(row=r, column=1, value="Signature: ____________________").font = font_normal
    ws.cell(row=r, column=5, value="Signature: ____________________").font = font_normal

    widths = [6, 22, 42, 10, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    try:
        wb.save(output_xlsx)
    except Exception as e:
        return f"❌ 存 Excel 失敗：{e}"

    email_line = ""
    if email_to.strip():
        subj = f"Quotation {quote_no} – {customer} / {shoe_model}"
        body = (f"Dear {customer},\n\n"
                f"Please find attached our quotation {quote_no} for {shoe_model} "
                f"(qty {quantity:,}).\n\n"
                f"Summary:\n"
                f"• Unit price: {f'{price_to_use:.2f} {currency}' if price_to_use else '(TBD)'}\n"
                f"• Delivery: {delivery_days} days\n"
                f"• Incoterm: {incoterm}\n"
                f"• Payment: {payment_terms}\n"
                f"• Validity: 30 days\n\n"
                f"Looking forward to your feedback.\n\n"
                f"Best regards,\nOwner")
        try:
            for addr in [a.strip() for a in re.split(r"[,\s;，、]+", email_to) if a.strip()]:
                # generated_by：報價單是小紅算出來的衍生品（價格數字！），
                # 寄給內部地址時會被 RAG 吃回去 —— 一定要帶出處標記。
                r_send = send_gmail_internal(
                    addr, subj, body, attachments=output_xlsx,
                    generated_by=f"quote:{quote_no}",
                )
                email_line += f"\n   ✉️ {addr}: {r_send}"
        except Exception as _e:
            email_line = f"\n   ⚠️ 寄信失敗：{_e}"

    return (
        f"✅ 報價單產出：{quote_no}\n"
        f"   客戶：{customer}　鞋型：{shoe_model}　數量：{quantity:,}\n"
        f"   單價：{f'{price_to_use:.2f} {currency}' if price_to_use else '(未填，Excel 中留白)'}"
        + (f"　總額：{total:,.2f} {currency}" if total else "") + "\n"
        + (f"   💡 {history_note}\n" if history_note else "   💡 無歷史參考（首次報價）\n")
        + f"   交期 {delivery_days} 天、{incoterm}、{currency}\n"
        + f"   📂 檔案：{output_xlsx}"
        + email_line
    )
