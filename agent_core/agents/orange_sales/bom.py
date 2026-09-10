"""BOM (Bill of Materials) extraction from Pricing BOM xlsx files.

Reads pricing BOM Excel files (one sheet per SKU, header row at R15:
Material Description | Material Code | Vendor Name | Vendor Code |
Consumption | Waste % | Total Consumption | Unit | Unit Price € | EURO),
flattens every material line into a structured CSV, then exposes a query
tool the agent can hit instead of hand-synthesizing answers from RAG
snippets.

Driving incident: 小紅 once claimed "Jalas uses waterproof membrane" by
conflating Jalas + Huafon PU email co-occurrence. The actual BOM tells
the precise story (CASPER 005 GA IDROREPELLEN is **撥水**, not 防水膜),
so this tool keeps the agent honest by surfacing the row that decides
the question.
"""
from __future__ import annotations

import csv
import json
import os
import re
from datetime import datetime
from typing import Any, Iterable

from agent_core.logging_and_paths import BOM_HISTORY_DIR, logger


_BOM_CSV = os.path.join(BOM_HISTORY_DIR, "auto_extracted.csv")
_BOM_STATE = os.path.join(BOM_HISTORY_DIR, "ingested_files.json")
_BOM_DRIVE_CACHE = os.path.join(BOM_HISTORY_DIR, "_drive_cache")
_BOM_HEADER_ROW = 15

# Default Shared Drive that holds the Pricing BOM xlsx files. The 業務部門
# drive carries the canonical BOMs across all customers; override with
# RED_BOM_DRIVE_ID for staging / a different org.
_DEFAULT_BOM_DRIVE_ID = "0ABwagOW8-2UjUk9PVA"
_BOM_CSV_FIELDS = [
    "extracted_at",
    "source_file",
    "bom_date",
    "customer",
    "factory_code",
    "sku",
    "material_description",
    "material_code",
    "vendor_name",
    "vendor_code",
    "consumption",
    "waste_pct",
    "total_consumption",
    "unit",
    "unit_price_eur",
    "total_eur",
    "category",
]

# Material category classifiers — keyword → category, ordered by specificity
# so we hit "防水膜" before generic "撥水" / "襯裡". Keywords match
# case-insensitively against the material description.
#
# Important shoe-industry distinction (the whole reason this tool exists):
#   - 防水膜 (Waterproof Membrane): separate hydrophilic/microporous film
#     layer like Sympatex / Gore-Tex / eVent / OutDry. Truly waterproof.
#   - 撥水 (Water-Repellent / DWR / Idrorepellent): surface coating that
#     beads water but does NOT make the boot waterproof on its own.
#   - 襯裡 (Lining): inner fabric for comfort/moisture management; some
#     linings have a membrane backing but most do not.
_CATEGORY_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("防水膜", ("sympatex", "gore-tex", "goretex", "gtx", "event", "outdry",
                "porelle", "permatex", "hipora")),
    ("撥水", ("idrorepellen", "idroreplent", "dwr", "water repellent",
              "water-repellent", "撥水", "水撥", "repellente")),
    ("微纖維", ("microfiber", "microfibre", "huafon")),
    ("皮料", ("mast matriz", "mastrotto", "leather", "nappa", "suede",
              "split l")),
    ("襯裡", ("dri-lex", "drilex", "lining", "jersey", "casper")),
    ("緩衝補強", ("foam density", "foam", "biag ibisafe", "padding",
                  "reinforcement", "tape", "fiber tape")),
    ("扣件", ("eyelet", "shank", "hook", "buckle", "stud")),
    ("縫線", ("gral", "thread", "polyamide", " pa ", "yarn")),
    ("標籤", ("label", "size label", "tag")),
    ("鞋頭/鞋尾", ("avantgarde", "toe puff", "counter")),
    ("成型件", ("tpu", "tebox")),
)


def _classify_material(description: str) -> str:
    """Classify a material row into a high-level category.

    Returns "其他" when no keyword matches; the underlying description is
    always preserved in the CSV row so the LLM can still inspect it.
    """
    text = (description or "").lower()
    if not text.strip():
        return "其他"
    for category, keywords in _CATEGORY_RULES:
        for kw in keywords:
            if kw in text:
                return category
    return "其他"


def _ensure_dir() -> None:
    os.makedirs(BOM_HISTORY_DIR, exist_ok=True)


def _load_ingested() -> dict[str, str]:
    if not os.path.exists(_BOM_STATE):
        return {}
    try:
        with open(_BOM_STATE, "r", encoding="utf-8") as f:
            data = json.load(f) or {}
        return {str(k): str(v) for k, v in data.items() if isinstance(data, dict)}
    except Exception:
        return {}


def _save_ingested(state: dict[str, str]) -> None:
    _ensure_dir()
    tmp = _BOM_STATE + ".tmp"
    try:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False, indent=2, sort_keys=True)
        os.replace(tmp, _BOM_STATE)
    except Exception as exc:
        logger.warning("ingested_files 寫入失敗：%s", exc)


def _coerce_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        try:
            return float(str(value).strip().replace(",", ""))
        except (TypeError, ValueError):
            return None


def _coerce_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text:
        return ""
    # Already ISO-ish?
    m = re.match(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", text)
    if m:
        y, mo, d = m.groups()
        return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}"
    return text[:30]


def _infer_customer_from_filename(filename: str) -> str:
    """Heuristic customer detection from the BOM filename.

    Known factory↔customer aliases live in entity_aliases.json; we hit a
    small handful inline so the parser stays self-contained for offline
    tests. Anything not in the table returns "" so query_bom() can fall
    back to filename-based matching.
    """
    name = os.path.basename(filename).lower()
    aliases = {
        "2peak": "Jalas",
        "jalas": "Jalas",
        "ejendals": "Jalas",
        "lurchi": "Lurchi",
        "richter": "Richter",
        "blaklader": "Blaklader",
        "blåkläder": "Blaklader",
        "decathlon": "Decathlon",
        "isco": "Isco",
    }
    for token, customer in aliases.items():
        if token in name:
            return customer
    return ""


def _extract_factory_code(filename: str, header_text: str) -> str:
    """Pull the factory model code (e.g. "2PEAK") from filename or sheet header.

    Helps disambiguate "Jalas 1055" (customer SKU) from "2PEAK 1055"
    (the factory-internal name on the BOM cover). Both should be queryable.
    """
    for blob in (header_text, filename):
        text = (blob or "").upper()
        m = re.search(r"\b(2PEAK|1PEAK|RAPTOR|TPEX|JALAS)\b", text)
        if m:
            return m.group(1)
    return ""


def _iter_sheet_rows(ws) -> Iterable[tuple]:
    return ws.iter_rows(values_only=True)


def parse_bom_xlsx(
    file_path: str,
    *,
    customer_hint: str = "",
) -> list[dict[str, Any]]:
    """Parse one Pricing BOM xlsx file into structured material rows.

    Returns a list of dicts, one per material line, with the schema
    listed in _BOM_CSV_FIELDS. Empty / header-row sheets are skipped.
    Caller supplies customer_hint when the filename does not betray the
    customer (the function still tries _infer_customer_from_filename).
    """
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("需要 openpyxl：pip install openpyxl")

    if not os.path.isfile(file_path):
        raise FileNotFoundError(file_path)

    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    customer = (customer_hint or "").strip() or _infer_customer_from_filename(file_path)
    extracted_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows: list[dict[str, Any]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # openpyxl read_only sometimes returns ws.max_row=None until the
        # generator has been walked, so don't trust it as an upfront guard.
        iter_rows = list(_iter_sheet_rows(ws))
        if len(iter_rows) < _BOM_HEADER_ROW:
            continue
        # Pricing summary may carry the BOM date in R1 cell B; description in R3 A.
        bom_date = _coerce_date(iter_rows[0][1] if len(iter_rows[0]) > 1 else "")
        header_desc = str(iter_rows[2][0] if len(iter_rows) >= 3 and iter_rows[2] else "")
        factory_code = _extract_factory_code(file_path, header_desc)
        sku = sheet_name.strip()

        header = iter_rows[_BOM_HEADER_ROW - 1]
        if not header or not any(str(c or "").strip() for c in header):
            continue
        if "Material Description" not in str(header[0] or ""):
            # Not a recognizable BOM material table; skip rather than guess.
            continue

        for raw in iter_rows[_BOM_HEADER_ROW:]:
            if not raw or all(c in (None, "") for c in raw):
                continue
            cells = list(raw) + [None] * max(0, 10 - len(raw))
            description = str(cells[0] or "").strip()
            if not description:
                continue
            row = {
                "extracted_at": extracted_at,
                "source_file": os.path.basename(file_path),
                "bom_date": bom_date,
                "customer": customer,
                "factory_code": factory_code,
                "sku": sku,
                "material_description": description,
                "material_code": str(cells[1] or "").strip(),
                "vendor_name": str(cells[2] or "").strip(),
                "vendor_code": str(cells[3] or "").strip(),
                "consumption": _coerce_number(cells[4]),
                "waste_pct": _coerce_number(cells[5]),
                "total_consumption": _coerce_number(cells[6]),
                "unit": str(cells[7] or "").strip(),
                "unit_price_eur": _coerce_number(cells[8]),
                "total_eur": _coerce_number(cells[9]),
                "category": _classify_material(description),
            }
            rows.append(row)
    return rows


def _write_csv(rows: list[dict[str, Any]]) -> None:
    _ensure_dir()
    tmp = _BOM_CSV + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=_BOM_CSV_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for r in rows:
            writer.writerow({k: ("" if v is None else v) for k, v in r.items()})
    os.replace(tmp, _BOM_CSV)


def _read_csv_rows() -> list[dict[str, Any]]:
    if not os.path.exists(_BOM_CSV):
        return []
    with open(_BOM_CSV, "r", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def build_bom_history(folder: str = "", *, customer_hint: str = "") -> str:
    """掃描資料夾下所有 Pricing BOM xlsx，解析後合併寫入 auto_extracted.csv。

    - folder: BOM 來源資料夾。空字串時依序試 ~/Downloads → var/data/bom_history。
    - customer_hint: 若整批檔案都屬於同一個客戶，可一次指定（覆蓋檔名推斷）。

    重複跑會「全量重建 CSV」（每個檔解析後合併）。state 檔記錄每個檔的
    bom_date 摘要，方便人工確認。"""
    candidates: list[str] = []
    if folder.strip():
        candidates.append(os.path.expanduser(folder.strip()))
    else:
        candidates.append(os.path.expanduser("~/Downloads"))
        candidates.append(BOM_HISTORY_DIR)

    seen: set[str] = set()
    targets: list[str] = []
    for d in candidates:
        if not os.path.isdir(d):
            continue
        for fname in sorted(os.listdir(d)):
            lower = fname.lower()
            if not lower.endswith(".xlsx"):
                continue
            if "bom" not in lower and "pricing bom" not in lower:
                continue
            if fname.startswith("~$"):  # Excel lock files
                continue
            full = os.path.join(d, fname)
            if full in seen:
                continue
            seen.add(full)
            targets.append(full)

    if not targets:
        searched = ", ".join(candidates)
        return f"找不到 BOM xlsx。已搜尋：{searched}"

    all_rows: list[dict[str, Any]] = []
    state: dict[str, str] = {}
    errors: list[str] = []
    for fp in targets:
        try:
            rows = parse_bom_xlsx(fp, customer_hint=customer_hint)
        except Exception as exc:
            errors.append(f"{os.path.basename(fp)}: {type(exc).__name__}: {exc}")
            continue
        all_rows.extend(rows)
        # Use the freshest bom_date from the sheets as fingerprint.
        dates = sorted({r["bom_date"] for r in rows if r.get("bom_date")})
        state[fp] = dates[-1] if dates else ""

    _write_csv(all_rows)
    _save_ingested(state)

    summary = [
        f"📊 BOM history 重建完成：{len(targets)} 個檔，{len(all_rows)} 筆材料",
        f"  → {_BOM_CSV}",
    ]
    if errors:
        summary.append(f"⚠️ 失敗 {len(errors)} 個：")
        summary.extend(f"  - {e}" for e in errors[:5])
    return "\n".join(summary)


def _parse_drive_modified(ts: str) -> float:
    """RFC3339 → Unix epoch, swallowing parse errors so we never block sync."""
    if not ts:
        return 0.0
    try:
        # Drive returns "2026-05-19T14:23:01.500Z" — strip the trailing Z.
        from datetime import datetime as _dt
        clean = ts.rstrip("Z").split(".")[0]
        return _dt.fromisoformat(clean).timestamp()
    except Exception:
        return 0.0


def sync_bom_from_drive(
    drive_id: str = "",
    *,
    filename_filter: str = "BOM",
    customer_hint: str = "",
    max_files: int = 50,
) -> str:
    """從 Google Shared Drive 下載 BOM xlsx → 本機快取 → 重建 CSV。

    - drive_id: Shared Drive ID。空字串時用環境變數 RED_BOM_DRIVE_ID
      或內建預設（業務部門 drive）。
    - filename_filter: 檔名 contains 關鍵字，預設 "BOM"。
    - customer_hint: 整批檔案統一指定客戶（覆蓋檔名推斷）。
    - max_files: 單次同步上限，避免 Drive 配額爆掉。

    增量規則：本機快取的 mtime ≥ Drive modifiedTime 就跳過下載；
    其餘走 get_media 拉 bytes。最後呼叫 build_bom_history(folder=cache)
    重建 auto_extracted.csv。
    """
    target_drive = (
        drive_id.strip()
        or os.environ.get("RED_BOM_DRIVE_ID", "").strip()
        or _DEFAULT_BOM_DRIVE_ID
    )
    try:
        from agent_core.google_auth import get_service
        svc = get_service("drive", "v3")
    except Exception as exc:
        return f"Google Drive 服務初始化失敗：{type(exc).__name__}: {exc}"

    XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    try:
        result = svc.files().list(
            corpora="drive",
            driveId=target_drive,
            includeItemsFromAllDrives=True,
            supportsAllDrives=True,
            q=f"name contains '{filename_filter}' and mimeType='{XLSX_MIME}' and trashed=false",
            fields="files(id, name, mimeType, modifiedTime)",
            pageSize=min(int(max_files), 100),
            orderBy="modifiedTime desc",
        ).execute()
    except Exception as exc:
        return f"Drive 搜尋失敗（drive_id={target_drive}）：{type(exc).__name__}: {exc}"

    files = result.get("files", [])
    if not files:
        return f"Drive 沒找到任何 *{filename_filter}*.xlsx（drive_id={target_drive}）"

    os.makedirs(_BOM_DRIVE_CACHE, exist_ok=True)
    downloaded = 0
    skipped = 0
    errors: list[str] = []
    for f in files[:max_files]:
        # Drive 檔名可能含 / 等不適合本機路徑的字元 — 轉成底線。
        safe_name = re.sub(r'[/\\:*?"<>|]', "_", f["name"])
        if not safe_name.lower().endswith(".xlsx"):
            safe_name += ".xlsx"
        dest = os.path.join(_BOM_DRIVE_CACHE, safe_name)
        drive_mtime = _parse_drive_modified(f.get("modifiedTime", ""))
        if os.path.exists(dest):
            local_mtime = os.path.getmtime(dest)
            if drive_mtime and local_mtime >= drive_mtime - 1:
                skipped += 1
                continue
        try:
            data = svc.files().get_media(
                fileId=f["id"], supportsAllDrives=True,
            ).execute()
            if not isinstance(data, bytes):
                raise RuntimeError(f"unexpected type {type(data).__name__}")
            tmp = dest + ".tmp"
            with open(tmp, "wb") as fh:
                fh.write(data)
            os.replace(tmp, dest)
            if drive_mtime:
                os.utime(dest, (drive_mtime, drive_mtime))
            downloaded += 1
        except Exception as exc:
            errors.append(f"{f['name']}: {type(exc).__name__}: {exc}")

    summary = [
        f"☁️  Drive 同步：{len(files)} 個檔（下載 {downloaded} / 跳過 {skipped} / 失敗 {len(errors)}）",
        f"  快取：{_BOM_DRIVE_CACHE}",
    ]
    if errors:
        summary.append(f"⚠️ 下載失敗 {len(errors)} 個（顯示前 5）：")
        summary.extend(f"  - {e}" for e in errors[:5])

    summary.append("")
    summary.append(build_bom_history(folder=_BOM_DRIVE_CACHE, customer_hint=customer_hint))
    return "\n".join(summary)


def query_bom(
    customer: str = "",
    sku: str = "",
    material_keyword: str = "",
    category: str = "",
    *,
    limit: int = 30,
    expand_aliases: bool = True,
) -> str:
    """查 BOM 結構化資料 — 客戶/SKU/材料關鍵字/分類四維過濾。

    參數全部 optional，空字串=不過濾：
    - customer: 客戶名稱（模糊比對；expand_aliases=True 時自動展開 alias）
    - sku: 料號（模糊比對 sheet 名稱，例如 "1155"）
    - material_keyword: 材料描述/料號/供應商關鍵字（模糊）
    - category: 高階分類，可填「防水膜 / 撥水 / 微纖維 / 皮料 / 襯裡 /
      緩衝補強 / 扣件 / 縫線 / 標籤 / 鞋頭/鞋尾 / 成型件 / 其他」
    - limit: 最多回傳幾筆明細（預設 30）

    用途：被問「X 客戶用不用 Y 材料」時，用 category="防水膜"+customer=X
    直接拿結構化證據，不要再從 RAG snippets 猜。

    回傳：篩選後筆數 + 分類統計 + 明細（含 source_file，可當證據）。
    """
    rows = _read_csv_rows()
    if not rows:
        return (f"BOM history CSV 尚未建立。請先呼叫 "
                f"build_bom_history()。位置：{_BOM_CSV}")

    customer_aliases: list[str] = []
    resolved_info = ""
    if customer.strip() and expand_aliases:
        try:
            from agent_core.entity_resolver import resolve_entity
            resolved = resolve_entity(customer.strip(), entity_type="customers")
            if resolved.get("found"):
                customer_aliases = [a.lower() for a in resolved["aliases"]]
                resolved_info = (
                    f"\n🔗 alias 展開：「{customer}」→ canonical『{resolved['canonical']}』，"
                    f"{len(customer_aliases)} 個 variant 一起比對"
                )
        except Exception as exc:
            logger.debug("entity_resolver 失敗：%s", exc)

    cust_q = customer.strip().lower()
    sku_q = sku.strip().lower()
    mk_q = material_keyword.strip().lower()
    cat_q = category.strip()

    filtered: list[dict[str, Any]] = []
    for r in rows:
        if cust_q:
            blob = " ".join([
                str(r.get("customer", "")).lower(),
                str(r.get("factory_code", "")).lower(),
                str(r.get("source_file", "")).lower(),
            ])
            needles = customer_aliases or [cust_q]
            if not any(n in blob for n in needles):
                continue
        if sku_q and sku_q not in str(r.get("sku", "")).lower():
            continue
        if mk_q:
            blob = " ".join([
                str(r.get("material_description", "")).lower(),
                str(r.get("material_code", "")).lower(),
                str(r.get("vendor_name", "")).lower(),
            ])
            if mk_q not in blob:
                continue
        if cat_q and r.get("category") != cat_q:
            continue
        filtered.append(r)

    header = (
        f"🔍 查到 {len(filtered)} 筆材料（過濾：customer={customer!r} sku={sku!r} "
        f"material={material_keyword!r} category={category!r}；總庫 {len(rows)} 筆）"
    )
    lines = [header]
    if resolved_info:
        lines.append(resolved_info)
    if not filtered:
        # Special case the "客戶用不用 Y 材料" question: when filtered for a
        # category like "防水膜" and customer hit, surface the LACK of any
        # rows as the answer.
        if cat_q and cust_q:
            lines.append(
                f"\n💡 結論：BOM 庫裡沒有任何屬於『{cat_q}』類的材料登記給"
                f"「{customer}」。如果這份 BOM 是完整的，這就是『沒有用』的直接證據。"
            )
        return "\n".join(lines)

    # Category histogram for the filtered set.
    from collections import Counter
    cat_counts = Counter(r.get("category", "其他") for r in filtered)
    if len(cat_counts) > 1 or not cat_q:
        lines.append("\n📊 分類統計：")
        for cat, n in cat_counts.most_common():
            lines.append(f"  {cat:8s} {n}")

    # Distinct source files = evidence trail.
    sources = sorted({r["source_file"] for r in filtered if r.get("source_file")})
    if sources:
        lines.append(f"\n📁 來源檔案：{', '.join(sources)}")

    lines.append("\n📋 材料明細（前 {} 筆）：".format(min(limit, len(filtered))))
    for r in filtered[:limit]:
        desc = (r.get("material_description") or "")[:40]
        vendor = (r.get("vendor_name") or "")[:20]
        unit_p = r.get("unit_price_eur") or ""
        unit = r.get("unit") or ""
        sku_v = r.get("sku") or ""
        cat = r.get("category") or "其他"
        lines.append(
            f"  [{cat}] sku={sku_v} | {desc} | vendor={vendor} | "
            f"price={unit_p} EUR/{unit}"
        )
    if len(filtered) > limit:
        lines.append(f"  …還有 {len(filtered) - limit} 筆（提高 limit 看更多）")
    return "\n".join(lines)
